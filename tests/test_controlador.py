"""Edición manual, validación, presentación de modelos y sincronización (Controlador sin UI)."""
from pathlib import Path

import openpyxl
import pytest

import controlador as ctl
import fixtures_pdf as f
from main import Controlador


@pytest.fixture
def ctrl(tmp_path):
    c = Controlador(tmp_path / "t.db")
    c.cambiar_carpeta_clientes(tmp_path / "Clientes")
    yield c
    c.cerrar()


@pytest.fixture
def pepito(ctrl, tmp_path):
    cliente = ctrl.crear_cliente("Pepito", "Pérez", "X1234567A", pais="Francia", ejercicio=2026)
    cid = cliente["id"]
    ctrl.procesar_pdf(cid, f.escritura_compra(tmp_path / "esc.pdf"), 2026)
    ctrl.procesar_pdf(cid, f.recibo_ibi(tmp_path / "ibi.pdf"), 2026)
    return cid


def test_etiquetas_salen_de_la_plantilla():
    et = ctl.etiquetas("inmuebles")
    assert et["precio_escritura"] == "Precio escritura"
    assert et["gananciales_forzado"] == "Gananciales 50 % forzado"
    assert "valor_adquisicion_total" not in et  # columna fórmula, no editable


def test_validar_formatos(ctrl):
    assert ctrl.validar("inmuebles", {"precio_escritura": "180.000,50", "fecha_adquisicion": "10/05/2019"}) == {}
    errores = ctrl.validar("inmuebles", {"precio_escritura": "mucho", "fecha_adquisicion": "31/02/2019"})
    assert set(errores) == {"precio_escritura", "fecha_adquisicion"}
    assert "dd/mm/aaaa" in errores["fecha_adquisicion"]


def test_corregir_dato_de_la_ia_actualiza_excel(ctrl, pepito):
    res = ctrl.editar(pepito, "inmuebles", "INM-001", {"precio_escritura": "200.000", "direccion": "Calle Real 7"}, 2026)
    assert ctrl.db.get("inmuebles", pepito, "INM-001")["precio_escritura"] == 200_000
    wb = openpyxl.load_workbook(res["excel"])
    assert wb["02_INMUEBLES"]["J2"].value == 200_000 and wb["02_INMUEBLES"]["B2"].value == "Calle Real 7"
    wb.close()
    assert "200.000,00" in res["informe"].mensaje_patrimonio


def test_vaciar_campo_a_mano(ctrl, pepito):
    ctrl.editar(pepito, "inmuebles", "INM-001", {"notario": "Pérez Ruiz"}, 2026)
    ctrl.editar(pepito, "inmuebles", "INM-001", {"notario": ""}, 2026)
    assert ctrl.db.get("inmuebles", pepito, "INM-001")["notario"] is None


def test_editar_rechaza_valores_invalidos(ctrl, pepito):
    with pytest.raises(ValueError, match="importe"):
        ctrl.editar(pepito, "gastos", None, {"total": "abc"}, 2026)


def test_alta_manual_de_contrato_genera_210(ctrl, pepito):
    res = ctrl.editar(pepito, "contratos", None, {"inmueble": "INM-001", "fecha_inicio": "01/01/2026",
                                                  "renta_inicial": "750", "estado": "Vigente"}, 2026)
    assert res["codigo"] == "CTR-001"
    assert len(res["informe"].filas_210_alquiler) == 4


def test_marcar_presentado_congela_y_actualiza_obligacion(ctrl, pepito):
    inf = ctrl.sincronizar(pepito, 2026)["informe"]
    imp = inf.filas_210_imputacion[0]
    res = ctrl.marcar_presentado(pepito, "m210_imputacion", imp, 2026, "15/12/2027", "123,45", "NRC-1")
    fila = ctrl.db.get("m210_imputacion", pepito, imp)
    assert fila["estado"] == "Presentado" and fila["pagado"] == 123.45 and fila["nrc"] == "NRC-1"
    obl = [o for o in ctrl.db.listar("obligaciones", pepito, id_modelo=imp)][0]
    assert obl["estado"] == "Presentado"
    assert not any("Imputación" in p.titulo for p in res["informe"].pendientes)


def test_marcar_presentada_desde_obligacion_marca_el_modelo(ctrl, pepito):
    inf = ctrl.sincronizar(pepito, 2026)["informe"]
    imp = inf.filas_210_imputacion[0]
    obl = ctrl.db.listar("obligaciones", pepito, id_modelo=imp)[0]
    ctrl.marcar_presentado(pepito, "obligaciones", obl["codigo"], 2026, None, "100")
    assert ctrl.db.get("m210_imputacion", pepito, imp)["estado"] == "Presentado"


def test_borrar_registro(ctrl, pepito):
    gasto = ctrl.alta_gasto(pepito, 2026, "INM-001", "Comunidad", "2026-03-01", 300)["gasto"]
    ctrl.borrar_registro(pepito, "gastos", gasto, 2026)
    assert ctrl.db.get("gastos", pepito, gasto) is None


def test_abrir_cliente_sincroniza_si_el_excel_cambio(ctrl, pepito):
    ruta = ctrl.db.ruta_excel(ctrl.db.get_cliente(pepito))
    assert not ctl.excel_generator.excel_modificado(ctrl.db, pepito)
    ctrl.abrir_cliente(pepito, 2026)  # sin cambios: no reexporta
    wb = openpyxl.load_workbook(ruta)
    wb["02_INMUEBLES"]["L2"] = 950  # el usuario añade notaría a mano en Excel
    wb.save(ruta)
    wb.close()
    assert ctl.excel_generator.excel_modificado(ctrl.db, pepito)
    ctrl.abrir_cliente(pepito, 2026)
    assert ctrl.db.get("inmuebles", pepito, "INM-001")["notaria"] == 950
    assert not ctl.excel_generator.excel_modificado(ctrl.db, pepito)


def test_pendientes_en_informe(ctrl, pepito):
    inf = ctrl.sincronizar(pepito, 2026)["informe"]
    titulos = [p.titulo for p in inf.pendientes]
    assert "Falta: Certificado de residencia fiscal" in titulos
    assert not any(t == "Falta: Escritura de compra" for t in titulos)  # la escritura ya está archivada
    falta = [d for d in ctrl.db.listar("documentos", pepito) if d["estado"] == "Falta"]
    wb = openpyxl.load_workbook(ctrl.db.ruta_excel(ctrl.db.get_cliente(pepito)))
    estados = [wb["13_DOCUMENTACION"][f"G{r}"].value for r in range(2, 12)]
    wb.close()
    assert estados.count("Falta") == len(falta) >= 1  # 00_RESUMEN!H10 los cuenta


def test_lector_y_escritor_en_hilos(ctrl, pepito):
    import threading
    errores = []

    def leer():
        try:
            for _ in range(50):
                ctrl.lector.listar("inmuebles", pepito)
        except Exception as exc:  # pragma: no cover - solo si falla
            errores.append(exc)

    hilos = [threading.Thread(target=leer) for _ in range(3)]
    for h in hilos:
        h.start()
    for _ in range(5):
        ctrl.editar(pepito, "inmuebles", "INM-001", {"notario": f"N{_}"}, 2026)
    for h in hilos:
        h.join()
    assert errores == []

"""Casos operativos del cliente, reproducidos con expedientes ficticios y sin API."""
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
import pytest

import database
import tax_calculator as tc
from controlador import Controlador, CalculoConfirmadoSinExcel


UE = tc.Residencia('Francia', 'UE/EEE', .19, True)
EXTRA = tc.Residencia('Rusia', 'Extra-UE', .24, False)
CAT = dict(valor_catastral=100000, valor_construccion=60000, valor_suelo=40000,
           anio_revision=2020, revision_pendiente=0)


def inmueble(compra='2020-01-01'):
    return dict(codigo='INM-001', fecha_adquisicion=compra, pct_titularidad=1,
                precio_escritura=100000, estado='Activo')


def contrato(inicio, fin, renta):
    return dict(codigo='CTR-001', inmueble='INM-001', fecha_inicio=inicio,
                fecha_fin_prevista=fin, fecha_fin_real=fin, estado='Finalizado',
                tipo='Temporada', renta_inicial=renta, periodicidad='Mensual')


@pytest.mark.parametrize('compra,inicio,fin,mensual,propiedad,alquilado,vacio,ingresos', [
    ('2025-07-01', '2025-09-01', '2025-12-31', 1000, 184, 122, 62, 4000),
    ('2020-01-01', '2025-01-01', '2025-03-31', 1000, 365, 90, 275, 3000),
    ('2020-01-01', '2025-01-01', '2025-06-30', 850, 365, 181, 184, 5100),
])
def test_fechas_e_ingresos_exactos_del_cliente(compra, inicio, fin, mensual,
                                              propiedad, alquilado, vacio, ingresos):
    inm, ctr = inmueble(compra), contrato(inicio, fin, mensual)
    tenencia = tc.periodo_tenencia(inm, 2025)
    assert tc.dias(*tenencia) == propiedad
    assert tc.dias_totales(tc.intervalos_alquilados([ctr], tenencia)) == alquilado
    imp = tc.calcular_210_imputacion(inm, [ctr], CAT, UE, 2025)
    assert imp.dias == vacio
    liquidacion = tc.calcular_210_alquiler(inm, [ctr], {}, [], CAT, UE, 2025, 'Anual')
    assert len(liquidacion) == 1 and liquidacion[0].ingresos == ingresos
    esperado = {62: 186.85, 275: 828.77, 184: 554.52}[vacio]
    assert imp.renta_imputada == esperado


def test_gasto_anual_90_dias_y_extra_ue_sin_deduccion():
    inm, ctr = inmueble(), contrato('2025-01-01', '2025-03-31', 1000)
    gasto = dict(codigo='GAS-001', categoria='IBI', ejercicio=2025, total=365,
                 periodo_desde='2025-01-01', periodo_hasta='2025-12-31')
    ue = tc.calcular_210_alquiler(inm, [ctr], {}, [gasto], CAT, UE, 2025, 'Anual')[0]
    ru = tc.calcular_210_alquiler(inm, [ctr], {}, [gasto], CAT, EXTRA, 2025, 'Anual')[0]
    assert ue.gastos == 533.84  # IBI90 + amortización443,84, prorrateados sobre el año.
    assert ru.gastos == 0 and ru.base == 3000 and ru.cuota == 720


@pytest.fixture
def expediente(tmp_path):
    c = Controlador(tmp_path/'ficticio.db')
    c.cambiar_carpeta_clientes(tmp_path/'Clientes')
    cid = c.crear_cliente('Prueba', 'Contratos', 'X0000001T', pais='Francia', ejercicio=2025)['id']
    inm = c.db.upsert('inmuebles', cid, inmueble('2025-07-01'))
    c.db.upsert('catastro', cid, {'inmueble': inm, 'ejercicio': 2025, **CAT})
    yield c, cid, inm
    c.cerrar()


def test_guardar_contrato_actualiza_vista_excel_e_imputacion(expediente):
    c, cid, inm = expediente
    datos = contrato('2025-09-01', '2025-12-31', 1000)
    datos.pop('codigo');datos['inmueble'] = inm
    resultado = c.editar(cid, 'contratos', None, datos, 2025)
    assert not resultado.get('error_excel')
    vista = c.vista_inmuebles(cid, 2025)[0]
    assert (vista['_dias_propiedad'], vista['_dias_alquilado'], vista['_dias_vacio']) == (184, 122, 62)
    assert vista['_ingresos_devengados'] == 4000
    assert c.db.listar('m210_imputacion', cid, ejercicio=2025)[0]['dias'] == 62
    ruta = c.db.ruta_excel(c.db.get_cliente(cid))
    wb = load_workbook(ruta)
    t = database.TABLAS['contratos'];ws = wb[t.hoja]
    cod_col = next(x.letra for x in t.columnas_excel if x.campo=='codigo')
    renta_col = next(x.letra for x in t.columnas_excel if x.campo=='renta_inicial')
    filas = [i for i in range(t.fila_inicio, ws.max_row+1) if ws[f'{cod_col}{i}'].value==resultado['codigo']]
    assert len(filas)==1 and ws[f'{renta_col}{filas[0]}'].value==1000
    assert wb['00_RESUMEN']['E9'].value.startswith('=SUMIFS(')
    wb.close()


def test_default_anual_y_referencia_por_ejercicio(expediente):
    c,cid,inm = expediente
    c.db.upsert('catastro',cid,{'inmueble':inm,'ejercicio':2022,'valor_referencia':180000})
    c.db.upsert('catastro',cid,{'inmueble':inm,'ejercicio':2025,'valor_referencia':210000})
    c.db.upsert('inmuebles',cid,{'codigo':inm,'fecha_adquisicion':'2022-06-01'})
    v = c.vista_inmuebles(cid,2025)[0]
    assert v['_periodicidad_210']=='Anual'
    assert v['_valor_referencia_compra']==180000 and v['_valor_referencia_actual']==210000
    c.db.upsert('inmuebles', cid, {'codigo': inm, 'valor_referencia': 190000})
    assert c.vista_inmuebles(cid,2025)[0]['_valor_referencia_compra'] == 190000


def test_sin_cambios_excel_no_recalcula_ni_reescribe(expediente,monkeypatch):
    c,cid,inm=expediente
    c.sincronizar(cid,2025)
    monkeypatch.setattr(c,'sincronizar',lambda *args:pytest.fail('No debe sincronizar un archivo sin cambios'))
    assert c.sincronizar_excel_si_cambio(cid,2025)=={}


def test_excel_editado_importa_contrato_y_cobro_y_actualiza_resumen(expediente):
    c, cid, inm = expediente
    c.db.upsert('inmuebles', cid, {'codigo': inm, 'fecha_adquisicion': '2020-01-01'})
    datos = contrato('2025-01-01', '2025-06-30', 850)
    datos['inmueble'] = inm
    c.editar(cid, 'contratos', None, datos, 2025)
    assert c.resumen_alquileres(cid, 2025)['ingresos_devengados'] == 5100
    ruta = c.db.ruta_excel(c.db.get_cliente(cid))
    wb = load_workbook(ruta)
    for tabla, campo, valor in [('contratos', 'renta_inicial', 900), ('cobros', 'importe_cobrado', 800)]:
        t = database.TABLAS[tabla]
        col = next(x.letra for x in t.columnas_excel if x.campo == campo)
        wb[t.hoja][f'{col}{t.fila_inicio}'] = valor
    wb.save(ruta)
    wb.close()
    resultado = c.sincronizar_excel_si_cambio(cid, 2025)
    assert resultado['excel_importado'] and not resultado.get('error_excel')
    resumen = c.resumen_alquileres(cid, 2025)
    assert resumen['ingresos_devengados'] == 5400
    assert resumen['ingresos_cobrados'] == 800
    assert resumen['contratos'][0]['_dias_alquilado'] == 181
    wb = load_workbook(ruta)
    ws = wb['00_RESUMEN']
    cabecera = next(row[0].row for row in ws.iter_rows() if row[0].value == 'OCUPACIÓN Y RENTAS · EJERCICIO 2025')
    assert [ws.cell(cabecera + 3, col).value for col in range(3, 8)] == [365, 181, 184, 5400, 800]
    assert ws['E9'].value.startswith('=SUMIFS(') and ws['E10'].value.startswith('=SUMIFS(')
    wb.close()
    assert c.sincronizar_excel_si_cambio(cid, 2025) == {}


def test_previa_normaliza_texto_y_propone_duracion_sin_escribir(expediente):
    c, cid, inm = expediente
    datos = {'inmueble': inm, 'fecha_inicio': '01/09/2025', 'renta_inicial': '1.000,00',
             'tipo': 'Temporada', 'duracion_meses': 4}
    previa = c.previsualizar_contrato(cid, datos, 2025)
    assert previa['fecha_fin_prevista'] == date(2025, 12, 31)
    assert previa['dias_alquilado'] == 122 and previa['ingresos_devengados'] == 4000
    assert c.db.listar('contratos', cid) == []


def test_previa_nueva_no_hereda_rentas_de_un_codigo_interno(expediente):
    c, cid, inm = expediente
    c.db.upsert('rentas', cid, {'contrato': 'PREVISUALIZACION', 'desde': '2025-09-01',
                               'renta': 9000, 'periodicidad': 'Mensual'})
    datos = {'inmueble': inm, 'fecha_inicio': '2025-09-01', 'fecha_fin_prevista': '2025-12-31',
             'renta_inicial': 1000, 'tipo': 'Temporada'}
    assert c.previsualizar_contrato(cid, datos, 2025)['ingresos_devengados'] == 4000


def test_previa_incluye_rentas_actualizadas_y_advierte_prorroga(expediente):
    c, cid, inm = expediente
    datos = {**contrato('2025-09-01', '2025-12-31', 1000), 'inmueble': inm}
    cod = c.db.upsert('contratos', cid, datos)
    c.db.upsert('rentas', cid, {'contrato': cod, 'desde': '2025-11-01', 'renta': 1100,
                               'periodicidad': 'Mensual'})
    assert c.previsualizar_contrato(cid, datos, 2025)['ingresos_devengados'] == 4200
    datos.update(tipo='Vivienda', estado='Vigente', fecha_fin_real=None)
    assert c.previsualizar_contrato(cid, datos, 2025)['avisos']


def test_catastro_referencia_no_oculta_datos_catastrales_del_mismo_anio(expediente):
    c, cid, inm = expediente
    previo = c.db.listar('catastro', cid)[0]
    c.db.upsert('catastro', cid, {'codigo': previo['codigo'], 'pct_imputacion': .011, 'revision_pendiente': 1})
    c.db.upsert('catastro', cid, {'inmueble': inm, 'ejercicio': 2025, 'valor_referencia': 210000})
    cat = c.db.catastro(cid, inm, 2025)
    assert cat['valor_catastral'] == 100000 and cat['valor_construccion'] == 60000
    assert cat['valor_referencia'] == 210000 and '_heredado' not in cat
    assert cat['pct_imputacion'] == .011 and cat['revision_pendiente'] == 1
    assert c.db.valor_referencia_ejercicio(cid, inm, 2024) is None


def test_periodicidad_por_inmueble_y_ejercicio_sin_falsos_pendientes(expediente):
    c, cid, inm = expediente
    datos = {**contrato('2025-09-01', '2025-12-31', 1000), 'inmueble': inm}
    c.editar(cid, 'contratos', None, datos, 2025)
    modelos = c.db.listar('m210_alquiler', cid, inmueble=inm, ejercicio=2025)
    assert [m['periodo'] for m in modelos] == ['Anual']
    with pytest.raises(ValueError, match='agrupación'):
        c.establecer_periodicidad_210(cid, inm, 2025, 'Trimestral')
    assert tc.periodicidad_210_alquiler(c.db, cid, inm, 2025) == 'Anual'
    assert tc.periodicidad_210_alquiler(c.db, cid, inm, 2026) == 'Anual'
    c.establecer_periodicidad_210(cid, inm, 2025, 'Anual')
    trimestrales = [m for m in c.db.listar('m210_alquiler', cid, ejercicio=2025) if m['periodo'] != 'Anual']
    assert all(m['estado'] == 'No procede' and m['cuota'] == 0 for m in trimestrales)
    with pytest.raises(ValueError, match='2024'):
        c.establecer_periodicidad_210(cid, inm, 2023, 'Anual')
    anual = c.db.get_por_clave('m210_alquiler', cid, f'210A|2025|{inm}|Anual')
    c.db.upsert('m210_alquiler', cid, {'codigo': anual['codigo'], 'estado': 'Presentado', 'cuota': 777})
    # Reflejar la presentación en Excel antes de otra acción que importa sus ediciones.
    c.commit_and_sync_excel(cid, 2025)
    with pytest.raises(ValueError, match='agrupación'):
        c.establecer_periodicidad_210(cid, inm, 2025, 'Trimestral')
    assert c.db.get('m210_alquiler', cid, anual['codigo'])['cuota'] == 777


@pytest.mark.parametrize('renta,periodo,tabla', [('01', '0a', 'm210_alquiler'),
    ('35', '0A', 'm210_alquiler'), ('02', '0A', 'm210_imputacion')])
def test_pdf_presentado_se_reconoce_sin_api(expediente, tmp_path, renta, periodo, tabla):
    import fixtures_pdf as f
    import pdf_processor as pp
    c, cid, inm = expediente
    c.db.upsert('inmuebles', cid, {'codigo': inm, 'ref_catastral': f.REF})
    pdf = f.modelo_210(tmp_path/'justificante.pdf', ejercicio=2025, tipo_renta=renta, periodo=periodo)
    pp.procesar_pdf(c.db, cid, pdf)
    filas = c.db.listar(tabla, cid, ejercicio=2025, inmueble=inm)
    assert len(filas) == 1 and filas[0]['estado'] == 'Presentado' and filas[0]['cuota'] == 285
    if renta != '02':
        assert filas[0]['periodo'] == 'Anual'


def test_propiedad_bisiesta_y_solapes_no_duplican_dias():
    inm = inmueble()
    a = contrato('2024-01-01', '2024-02-29', 1000)
    b = {**contrato('2024-02-01', '2024-03-31', 1000), 'codigo': 'CTR-002'}
    resumen = tc.resumen_arrendamiento(inm, [a, b], {}, 2024)
    assert (resumen['dias_propiedad'], resumen['dias_alquilado'], resumen['dias_vacio']) == (366, 91, 275)


def test_borrador_manual_conserva_calendario_y_exige_revision_antes_de_cambiar(expediente):
    c, cid, inm = expediente
    c.db.upsert('m210_alquiler', cid, {'ejercicio': 2025, 'inmueble': inm, 'periodo': '1T',
                                     'estado': 'En preparación', 'cuota': 123})
    assert tc.periodicidad_210_alquiler(c.db, cid, inm, 2025) == 'Anual'
    with pytest.raises(ValueError, match='preparación'):
        c.establecer_periodicidad_210(cid, inm, 2025, 'Anual')
    assert c.db.listar('m210_alquiler', cid)[0]['cuota'] == 123


def test_borrar_inmueble_limpia_preferencia_antes_de_reutilizar_codigo(expediente):
    c, cid, inm = expediente
    c.db.set_config(f'periodicidad_210|{cid}|{inm}|2025', 'Trimestral')
    c.db.borrar('inmuebles', cid, inm)
    nuevo = c.db.upsert('inmuebles', cid, inmueble())
    assert nuevo == inm and tc.periodicidad_210_alquiler(c.db, cid, nuevo, 2025) == 'Anual'


def test_contrato_fin_anterior_al_inicio_no_se_guarda(expediente):
    c, cid, inm = expediente
    datos = {**contrato('2025-09-01', '2025-08-31', 1000), 'inmueble': inm}
    with pytest.raises(ValueError, match='anterior al inicio'):
        c.editar(cid, 'contratos', None, datos, 2025)
    assert c.db.listar('contratos', cid) == []


def test_contratos_de_otros_ejercicios_siguen_visibles_sin_sumar_devengo(expediente):
    c, cid, inm = expediente
    for codigo, inicio, fin in [('CTR-ANT', '2023-01-01', '2023-12-31'),
                               ('CTR-FUT', '2026-01-01', '2026-12-31')]:
        c.db.upsert('contratos', cid, {**contrato(inicio, fin, 1000), 'codigo': codigo, 'inmueble': inm})
    r = c.resumen_alquileres(cid, 2025)
    assert {f['codigo'] for f in r['contratos']} == {'CTR-ANT', 'CTR-FUT'}
    assert r['ingresos_devengados'] == 0
    assert all(f['_dias_alquilado'] == 0 for f in r['contratos'])


def test_contrato_abierto_no_pisa_cambio_concurrente_del_mismo_campo(expediente):
    c, cid, inm = expediente
    datos = {**contrato('2025-09-01', '2025-12-31', 1000), 'inmueble': inm}
    codigo = c.editar(cid, 'contratos', None, datos, 2025)['codigo']
    original = c.lector.get('contratos', cid, codigo)
    c.editar(cid, 'contratos', codigo, {'renta_inicial': 900}, 2025)
    with pytest.raises(ValueError, match='cambió desde que abriste'):
        c.editar(cid, 'contratos', codigo, {'renta_inicial': 1100}, 2025, esperados=original)
    assert c.db.get('contratos', cid, codigo)['renta_inicial'] == 900
    # Un campo distinto no debe bloquearse por un cambio ajeno que no se va a sobrescribir.
    c.editar(cid, 'contratos', codigo, {'fianza': 1500}, 2025, esperados=original)
    assert c.db.get('contratos', cid, codigo)['renta_inicial'] == 900


def test_periodicidad_publica_desglose_confirmado_antes_de_exportar_excel(expediente, monkeypatch):
    c, cid, inm = expediente
    datos = {**contrato('2025-09-01', '2025-12-31', 1000), 'inmueble': inm}
    c.editar(cid, 'contratos', None, datos, 2025)
    pasos = []
    exportar = c._exportar_informe

    def al_recalcular(resultado):
        assert c.db._nivel_transaccion == 0
        assert c.lector.get_config(f'periodicidad_210|{cid}|{inm}|2025') == 'Anual'
        vista = c.vista_inmuebles(cid, 2025)[0]
        assert {f['periodo'] for f in vista['_desglose_210']} == {'Anual'}
        assert [p['ingresos_devengados'] for p in vista['_ocupacion_trimestral']] == [0, 0, 1000, 3000]
        assert all(f['estado'] != 'No procede' for f in vista['_desglose_210'])
        assert sum(f['ingresos'] for f in vista['_desglose_210']) == 4000
        assert resultado['informe']
        pasos.append('calculo')

    def guardar(*args):
        assert pasos == ['calculo']
        pasos.append('excel')
        return exportar(*args)

    monkeypatch.setattr(c, '_exportar_informe', guardar)
    resultado = c.establecer_periodicidad_210(cid, inm, 2025, 'Anual', al_recalcular=al_recalcular)
    assert pasos == ['calculo', 'excel'] and not resultado.get('error_excel')


def test_fallo_de_notificacion_visual_no_impide_guardar_excel(expediente, monkeypatch):
    c, cid, inm = expediente
    c.editar(cid, 'contratos', None, {**contrato('2025-09-01', '2025-12-31', 1000), 'inmueble': inm}, 2025)
    exportado = []
    exportar = c._exportar_informe
    def guardar(*args):
        exportado.append(True)
        return exportar(*args)
    def notificacion_fallida(_resultado):
        raise RuntimeError('fallo ficticio de UI')
    monkeypatch.setattr(c, '_exportar_informe', guardar)
    resultado = c.establecer_periodicidad_210(cid, inm, 2025, 'Anual', al_recalcular=notificacion_fallida)
    assert exportado == [True] and not resultado.get('error_excel')


def test_vista_no_espera_al_lock_del_exportador_para_mostrar_nuevo_desglose(expediente, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    import threading
    c, cid, inm = expediente
    c.editar(cid, 'contratos', None, {**contrato('2025-09-01', '2025-12-31', 1000), 'inmueble': inm}, 2025)
    calculado, continuar = threading.Event(), threading.Event()
    exportar = c._exportar_informe
    def exportacion_lenta(*args):
        assert continuar.wait(5), 'La prueba debe liberar el guardado'
        return exportar(*args)
    monkeypatch.setattr(c, '_exportar_informe', exportacion_lenta)
    with ThreadPoolExecutor(max_workers=2) as pool:
        guardar = pool.submit(c.establecer_periodicidad_210, cid, inm, 2025, 'Anual',
                              lambda _res: calculado.set())
        try:
            assert calculado.wait(5)
            leer = pool.submit(c.vista_inmuebles, cid, 2025)
            vista = leer.result(timeout=3)[0]
            assert vista['_periodicidad_210'] == 'Anual'
            assert {m['periodo'] for m in vista['_desglose_210']} == {'Anual'}
            assert not guardar.done()  # los datos se leen mientras el worker sigue guardando
        finally:
            continuar.set()
        assert not guardar.result(timeout=15).get('error_excel')


def test_error_exportando_conserva_informe_del_calculo_ya_confirmado(expediente, monkeypatch):
    c, cid, inm = expediente
    c.editar(cid, 'contratos', None, {**contrato('2025-09-01', '2025-12-31', 1000), 'inmueble': inm}, 2025)
    def fallar_exportacion(*_args):
        raise PermissionError('fallo ficticio')
    monkeypatch.setattr(c, '_exportar_informe', fallar_exportacion)
    with pytest.raises(CalculoConfirmadoSinExcel) as error:
        c.establecer_periodicidad_210(cid, inm, 2025, 'Anual')
    assert (error.value.cliente_id, error.value.ejercicio) == (cid, 2025)
    assert error.value.resultado['informe']
    assert c.lector.get_config(f'periodicidad_210|{cid}|{inm}|2025') == 'Anual'
    assert 'fallo ficticio' not in str(error.value)

"""Casos de uso de la app, sin interfaz gráfica (testeable).

Autores: Antigravity (versión inicial en main.py), Claude (refactor, edición manual, IA, pendientes).
Orden obligatorio en toda escritura: importar Excel -> cambiar BD -> recalcular -> exportar Excel.
"""
from __future__ import annotations

import calendar
import gc
import hashlib
import json
import logging
import math
import os
import shutil
import threading
import uuid
from contextlib import nullcontext
from datetime import date, timedelta
from functools import lru_cache
from pathlib import Path
from typing import Any, Callable, Optional

import carteras
import database
import documentos
import excel_generator
import fiscal_residentes
import ia_extractor
import pdf_processor
import tax_calculator
from carteras import NR, RES, cartera_de
from database import TABLAS, Database, a_fecha

ConfirmarFn = Callable[[pdf_processor.Extraccion], pdf_processor.Extraccion]


def _carpeta_sin_enlaces(carpeta: Path) -> Path:
    """Rechaza una junction tanto en la carpeta como en cualquiera de sus padres."""
    carpeta = Path(os.path.abspath(carpeta))
    try:
        if carpeta.resolve() != carpeta:
            raise ValueError("La carpeta apunta a otra ubicación (enlace). No se usa.")
    except (OSError, RuntimeError):
        raise ValueError("No se puede verificar la ubicación de la carpeta.") from None
    return carpeta


class CanceladoPorUsuario(Exception):
    """El usuario descartó el documento en el diálogo de revisión."""


class CalculoConfirmadoSinExcel(RuntimeError):
    """La BD ya se confirmó: conservar el informe calculado aunque falle la exportación posterior."""

    def __init__(self, cid: int, ejercicio: int, resultado: dict, causa: Exception):
        self.cliente_id, self.ejercicio, self.resultado = cid, ejercicio, resultado
        registro = f"Registro {resultado['codigo']} guardado. " if resultado.get('codigo') else ''
        super().__init__(registro + "El cálculo se guardó en la base de datos, pero no se pudo completar la actualización del Excel "
                         f"({type(causa).__name__}). "
                         "Revisa permisos y espacio disponible y pulsa Sincronizar para reintentarlo.")


# Etiquetas de los campos que no tienen columna en la plantilla (solo BD)
ETIQUETAS_SOLO_BD = {
    "ejercicio": "Ejercicio", "gananciales_forzado": "Gananciales 50 % forzado",
    "pct_titularidad_escritura": "% titularidad según escritura", "valor_comprobado": "Valor comprobado Administración",
    "gastos_venta": "Gastos de venta", "valor_referencia": "Valor de referencia", "origen": "Origen",
    "renta_imputada": "Renta imputada", "cuota": "Cuota", "sha256": "Huella", "nombre_original": "Nombre original",
}

SI_NO = ("Sí", "No", "Revisar")
OPCIONES: dict[tuple[str, str], tuple[str, ...]] = {
    ("inmuebles", "estado"): database.ESTADOS_INMUEBLE,
    ("inmuebles", "uso_valor_referencia"): SI_NO,
    ("inmuebles", "tipo_adquisicion"): ("Compraventa", "Herencia", "Donación", "Obra nueva", "Otro"),
    ("contratos", "estado"): database.ESTADOS_CONTRATO,
    ("contratos", "periodicidad"): database.PERIODICIDADES,
    ("contratos", "tipo"): ("Vivienda", "Temporada", "Uso distinto de vivienda", "Habitación", "Otro"),
    ("rentas", "periodicidad"): database.PERIODICIDADES,
    ("inquilinos", "tipo_persona"): ("Persona física", "Persona jurídica"),
    ("catastro", "fuente"): database.FUENTES_CATASTRO,
    ("gastos", "categoria"): database.CATEGORIAS_GASTO,
    ("gastos", "incluido_210"): ("Sí", "No", "Pendiente"),
    ("cobros", "estado"): database.ESTADOS_COBRO,
    ("m210_alquiler", "estado"): database.ESTADOS_MODELO,
    ("m210_imputacion", "estado"): database.ESTADOS_MODELO,
    ("patrimonio_714", "estado"): database.ESTADOS_MODELO,
    ("its_718", "estado"): database.ESTADOS_MODELO,
    ("obligaciones", "estado"): database.ESTADOS_MODELO,
    ("obligaciones", "tipo"): database.TIPOS_OBLIGACION,
    ("documentos", "estado"): database.ESTADOS_DOCUMENTO,
    ("residencias", "certificado"): ("Sí", "No", "Pendiente", "No procede"),
    ("residencias", "gastos_deducibles"): SI_NO,
    # ---- residentes
    ("actividades", "regimen_iva"): carteras.REGIMENES_IVA,
    ("actividades", "regimen_irpf"): carteras.REGIMENES_IRPF,
    ("actividades", "obligado_130"): ("Sí", "No"),
    ("facturas_emitidas", "estado_cobro"): database.ESTADOS_COBRO_FACTURA,
    ("facturas_recibidas", "categoria"): database.CATEGORIAS_FACTURA_RECIBIDA,
    ("m303", "estado"): database.ESTADOS_MODELO,
    ("m130", "estado"): database.ESTADOS_MODELO,
    ("m111_115", "estado"): database.ESTADOS_MODELO,
    ("m111_115", "modelo"): ("111", "115"),
    ("modelos_presentados", "estado"): database.ESTADOS_MODELO,
}
OPCIONES[("obligaciones", "tipo")] = database.TIPOS_OBLIGACION + database.TIPOS_OBLIGACION_RES[:-1]
# Campos que la app calcula sola: se muestran pero se avisa de que se recalculan
CALCULADOS = {
    "m210_alquiler": {"ingresos", "gastos", "base", "tipo", "cuota", "desde", "hasta", "pais", "contrato"},
    "m210_imputacion": {"dias", "valor_catastral", "pct_imputacion", "pct_titularidad", "tipo", "renta_imputada", "cuota"},
    "patrimonio_714": {"valor_bienes", "base", "minimo_exento", "base_liquidable", "cuota_integra", "resultado"},
    "its_718": {"base", "cuota", "deducciones", "resultado"},
    "cobros": {"importe_contractual"},
    "facturas_emitidas": {"ejercicio", "trimestre"},
    "facturas_recibidas": {"ejercicio", "trimestre"},
    "m303": {"base_devengada", "iva_devengado", "base_deducible", "iva_deducible", "compensacion_anterior",
             "resultado", "a_ingresar", "a_compensar"},
    "m130": {"ingresos", "gastos", "rendimiento", "pago_20", "retenciones", "pagos_anteriores", "resultado",
             "a_ingresar"},
    "m111_115": {"perceptores", "base", "retenciones"},
}
OCULTOS = {"codigo", "sha256", "nombre_original", "origen", "pct_titularidad_escritura", "renta_imputada", "cuota",
           "detalle_limite", "revision_pendiente"}
CAMPOS_PRESENTACION = {
    "m210_alquiler": ("fecha_presentacion", "pagado", "nrc"),
    "m210_imputacion": ("fecha_presentacion", "pagado", "nrc"),
    "patrimonio_714": ("fecha_presentacion", "pagado", "csv"),
    "its_718": ("fecha_presentacion", "pagado", "csv"),
    "obligaciones": ("fecha_presentacion", "pagado"),
    "m303": ("fecha_presentacion", "nrc"),
    "m130": ("fecha_presentacion", "nrc"),
    "m111_115": ("fecha_presentacion", "nrc"),
}
# Declaraciones de residentes: el recálculo los congela al estar presentados
_DECLARACIONES_RES = ("m303", "m130", "m111_115")
_PREFIJO_A_TABLA = {t.prefijo: t.nombre for t in TABLAS.values() if t.prefijo}


@lru_cache(maxsize=1)
def _cabeceras_plantilla() -> dict[str, dict[str, str]]:
    """{hoja: {letra: texto de la fila 1}} leído de la plantilla (modo solo lectura, rápido)."""
    import openpyxl
    from openpyxl.utils import get_column_letter
    out = {}
    for plantilla in database.PLANTILLAS.values():  # las hojas no se repiten entre las dos plantillas
        if not Path(plantilla).exists():
            continue
        wb = openpyxl.load_workbook(plantilla, read_only=True)
        try:
            for t in TABLAS.values():
                for hoja in {t.hoja, t.hoja_res} - {""}:
                    if hoja in wb.sheetnames and hoja not in out:
                        fila = next(wb[hoja].iter_rows(min_row=1, max_row=1, values_only=True), ())
                        out[hoja] = {get_column_letter(i): str(v) for i, v in enumerate(fila, 1) if v}
        finally:
            wb.close()
    return out


def etiquetas(tabla: str) -> dict[str, str]:
    t = TABLAS[tabla]
    cab = _cabeceras_plantilla().get(t.hoja, {})
    return {c.campo: (cab.get(c.letra) if c.letra else None) or ETIQUETAS_SOLO_BD.get(c.campo)
            or c.campo.replace("_", " ").capitalize() for c in t.columnas_bd}


def campos_editables(tabla: str) -> list[database.Columna]:
    return [c for c in TABLAS[tabla].columnas_bd if c.campo not in OCULTOS]


class Controlador:
    """Orquestador de casos de uso. Seguro entre hilos: escrituras bajo ``_lock``; la UI lee con ``lector``."""

    def __init__(self, db_path: Path | None = None):
        ruta = Path(db_path) if db_path is not None else database.DB_PATH_DEFECTO
        from migracion_actualizacion import preparar_actualizacion, validar_almacen_existente
        validar_almacen_existente(ruta)
        self.db = Database(ruta, multihilo=True)
        try:
            self.copia_actualizacion = preparar_actualizacion(self.db)
            self.lector = Database(ruta, multihilo=True)  # conexión aparte para que la UI lea sin bloquear
        except BaseException:
            self.db.cerrar()
            raise
        self._lock = threading.RLock()
        self._firmas: dict[int, Optional[str]] = {}  # firma del Excel leída en la última importación
        self.aviso_migracion: Optional[str] = None
        try:  # una sola vez: cada expediente a la raíz de su cartera (todo o nada, con copia de la BD)
            self.migracion = self.db.migrar_carteras()
        except Exception as exc:  # la app sigue con las rutas antiguas y se reintenta al reabrir
            self.migracion = None
            detalle = ("La recuperación quedó incompleta: revisa la copia de la BD y el manifiesto de "
                       "recuperación antes de continuar." if getattr(exc, "requiere_recuperacion", False)
                       else "No se ha movido nada; se reintentará al volver a abrir la app.")
            self.aviso_migracion = ("No se pudieron ordenar las carpetas en No residentes / Residentes "
                                    f"({exc.__class__.__name__}: {exc}). {detalle}")

    def cerrar(self) -> None:
        with self._lock:
            self.db.cerrar()
            self.lector.cerrar()

    # ------------------------------------------------------------------ configuración
    def cambiar_carpeta_clientes(self, ruta: Path) -> None:
        with self._lock:
            self.db.carpeta_clientes = Path(ruta)

    def ajustes_ia(self) -> ia_extractor.AjustesIA:
        return ia_extractor.AjustesIA.desde_db(self.lector)

    def guardar_ajustes_ia(self, ajustes: ia_extractor.AjustesIA, clave: Optional[str] = None) -> None:
        with self._lock:
            ajustes.guardar(self.db)
            if clave:
                ia_extractor.guardar_clave(clave)

    def borrar_clave_ia(self) -> None:
        ia_extractor.borrar_clave()

    def estado_ia(self) -> tuple[bool, str]:
        aj = self.ajustes_ia()
        if not aj.activada:
            return False, "Lectura con IA desactivada · lector local"
        if not ia_extractor.hay_clave():
            return False, "Falta la clave API de Claude · lector local"
        return True, f"Lectura con IA activa · {aj.modelo}"

    # ------------------------------------------------------------------ clientes
    def crear_cliente(self, nombre: str, apellidos: str, nif: str, tipo_cliente: str = "No residente",
                      pais: str | None = None, ejercicio: int | None = None, regimen_iva: str | None = None,
                      regimen_irpf: str | None = None) -> dict:
        """Alta en la carpeta de su cartera. NR: país de residencia. Residente: régimen de IVA e IRPF."""
        with self._lock:
            if self.db.get_cliente_por_nif((nif or "").strip().upper()):
                raise ValueError("Ya existe un cliente con ese NIF/NIE. Abre su ficha para corregir sus datos; no se ha creado ni modificado ningún expediente.")
            carteras.validar_tipo(tipo_cliente)
            if pais and ejercicio is not None:
                documentos._validar_ejercicio(ejercicio)
            cliente = self.db.crear_cliente(nombre, apellidos, nif, tipo_cliente=tipo_cliente)
            if cartera_de(cliente.get("tipo_cliente")) == RES:
                if (regimen_iva or regimen_irpf) and not self.db.listar("actividades", cliente["id"]):
                    self.db.upsert("actividades", cliente["id"], {
                        "descripcion": "Actividad principal", "regimen_iva": regimen_iva,
                        "regimen_irpf": regimen_irpf if cliente["tipo_cliente"] == carteras.TIPO_AUTONOMO else None,
                        "observaciones": "Datos indicados al crear el cliente; sube el 036/037 para completarlos."})
                return cliente
            if pais:
                ej = int(ejercicio or date.today().year)
                grupo = tax_calculator.grupo_territorio(pais)
                self.db.guardar_residencia(cliente["id"], ej, pais=pais, grupo=grupo,
                                           gastos_deducibles="Sí" if grupo == tax_calculator.GRUPO_UE else "No")
            return cliente

    def eliminar_cliente(self, cid: int) -> Path:
        """Quita al cliente de la app. Su carpeta (PDFs y Excel) NO se borra: se mueve a
        ``<carpeta de clientes>/_Eliminados/<carpeta>_<fecha-hora>`` para poder recuperarla.

        Orden seguro: primero se renombra la carpeta, sin copiar ni borrar archivos parcialmente.
        Si el sistema rechaza el traslado (p. ej. entre volúmenes), no se toca la BD.
        Después se borra de la BD; si la BD fallara, la carpeta vuelve a su sitio."""
        from datetime import datetime
        with self._lock:
            cliente = self.db.get_cliente(cid)
            if not cliente:
                raise ValueError("El cliente ya no existe.")
            origen = self.db.verificar_cartera(cliente)
            # La papelera tampoco puede redirigir los datos a otra cartera mediante una junction.
            papelera = _carpeta_sin_enlaces(origen.parent / "_Eliminados")
            destino = papelera / f"{origen.name}_{datetime.now():%Y%m%d_%H%M%S}"
            n = 2
            while destino.exists():
                destino = papelera / f"{origen.name}_{datetime.now():%Y%m%d_%H%M%S}_{n}"
                n += 1
            movida = False
            if origen.exists():
                papelera.mkdir(parents=True, exist_ok=True)
                try:
                    origen.rename(destino)
                except (PermissionError, OSError) as exc:
                    raise ValueError("No se pudo mover la carpeta del cliente: cierra su Excel o sus PDFs "
                                     "abiertos y vuelve a intentarlo. Comprueba también los permisos y que "
                                     "_Eliminados esté en la misma unidad. "
                                     f"({exc.__class__.__name__})") from None
                movida = True
            try:
                # Las tablas del cliente tienen ON DELETE CASCADE hacia clientes
                self.db.con.execute("DELETE FROM clientes WHERE id=?", (cid,))
                self.db.con.execute("DELETE FROM config WHERE clave IN (?,?,?,?,?,?,?)",
                                    tuple(f'{prefijo}|{cid}' for prefijo in ('excel_firma', 'excel_pendiente',
                                        'excel_resumen', 'excel_version', 'excel_nombre', 'excel_filas', 'excel_revision')))
                for nombre in ("periodicidad_210", "conciliacion_210", "revision_calendario_210", "amortizacion_venta", "saldo303"):
                    prefijo = f"{nombre}|{cid}|"
                    self.db.con.execute("DELETE FROM config WHERE substr(clave,1,?)=?", (len(prefijo), prefijo))
                self.db.con.commit()
            except Exception:
                self.db.con.rollback()
                if movida:
                    destino.rename(origen)
                raise
            self._firmas.pop(cid, None)
            return destino

    def abrir_cliente(self, cid: int, ejercicio: int) -> dict:
        """Carga el expediente. Si el usuario tocó el Excel fuera de la app, sincroniza primero (A3-P3)."""
        with self._lock:
            cliente = self.db.get_cliente(cid)
            if cliente is None:
                raise ValueError('El cliente ya no existe. Actualiza la lista de expedientes.')
            if not self.db.verificar_cartera(cliente).is_dir():
                raise FileNotFoundError('No se encuentra la carpeta del expediente; no se ha recreado. '
                                        'Comprueba su ubicación y la unidad antes de continuar.')
            documentos.recuperar_archivos(self.db, cid)
            if (excel_generator.excel_modificado(self.db, cid)
                    or excel_generator.necesita_actualizar_formato(self.db, cid)):
                return self.sincronizar(cid, ejercicio)
            if self._es_residente(cid):
                ej = int(ejercicio)
                todos = fiscal_residentes.ejercicios_encadenados(self.db, cid, {ej}, ej)
                with self.db.transaccion():
                    informes = {e: fiscal_residentes.recalcular_residente(self.db, cid, e) for e in todos}
                return {"informe": informes[ej], "estados_autonomo": informes[ej].estados, "excel": None,
                        "error_excel": None}
            with self.db.transaccion():
                informe = tax_calculator.recalcular_cliente(self.db, cid, int(ejercicio))
                estados_aut = self._autonomo(cid, int(ejercicio))
                if estados_aut:
                    informe.pendientes = tax_calculator.revisar_pendientes(self.db, cid, int(ejercicio))
            return {"informe": informe, "estados_autonomo": estados_aut,
                    "excel": None, "error_excel": None}

    def _es_residente(self, cid: int) -> bool:
        cliente = self.db.get_cliente(cid)
        return bool(cliente) and cartera_de(cliente.get("tipo_cliente")) == RES

    def _autonomo(self, cid: int, ej: int) -> list[str]:
        cliente = self.db.get_cliente(cid)
        if cliente and cliente.get("tipo_cliente") == "Autónomo":
            return tax_calculator.recalcular_autonomo(self.db, cid, ej)
        return []

    # ------------------------------------------------------------------ flujo central
    def _recalcular_bd(self, cid: int, ejercicios: set[int], ejercicio_activo: int) -> dict:
        ej = int(ejercicio_activo)
        # Ventana automática de cuatro ejercicios anteriores y el actual,
        # además del activo y la historia explícita; no decide la prescripción legal.
        hoy = date.today().year
        desde = min({int(e) for e in ejercicios} | {ej})
        tablas_modelo = (_DECLARACIONES_RES if self._es_residente(cid) else
                         ("m210_alquiler", "m210_imputacion", "patrimonio_714", "its_718"))
        # El libro contiene varios ejercicios. Al renovar formato o editar desde
        # un año posterior, actualizar también los modelos ya guardados dentro
        # de la ventana de revisión; no crear años sin datos ni dejar la BD con
        # importes anteriores mientras sus fórmulas de Excel ya son nuevas.
        desde_guardados = min(desde, hoy - 4)
        # Un modelo ya calculado debe seguir coherente aunque sea histórico o
        # una previsión futura. No crear años nuevos por una etiqueta del PDF.
        existentes = {int(r[0]) for tabla in (*tablas_modelo, "obligaciones", "modelos_presentados")
            for r in self.db.con.execute(f"SELECT DISTINCT ejercicio FROM {tabla} WHERE cliente_id=? AND ejercicio>=?",
                                         (cid, desde_guardados)) if r[0] is not None}
        if self._es_residente(cid):
            # El 303 encadena lo «a compensar» de un año al siguiente: se recalcula en orden
            encadenados = set(fiscal_residentes.ejercicios_encadenados(self.db, cid, ejercicios, ej))
            # Completar los años intermedios necesarios para arrastrar saldos
            # hasta una previsión que ya existe, sin crear previsiones por etiquetas.
            hasta = max(encadenados | existentes)
            todos = sorted(encadenados | existentes | set(range(max(encadenados) + 1, hasta + 1)))
            with self.db.transaccion():
                informes = {e: fiscal_residentes.recalcular_residente(self.db, cid, e) for e in todos}
            estados_aut = informes[ej].estados
        else:
            todos = sorted({int(e) for e in ejercicios if hoy - 4 <= int(e) <= hoy}
                           | existentes | {ej})
            with self.db.transaccion():
                informes = tax_calculator.recalcular_ejercicios(self.db, cid, todos)
                estados_aut: list[str] = []
                for e in todos:
                    est = self._autonomo(cid, e)
                    if est:
                        informes[e].pendientes = tax_calculator.revisar_pendientes(self.db, cid, e)
                    if e == ej:
                        estados_aut = est
        return {"informe": informes[ej], "estados_autonomo": estados_aut}

    def _exportar_informe(self, cid: int, ejercicio_activo: int, resultado: dict) -> dict:
        ej = int(ejercicio_activo)
        firma = self._firmas.get(cid, excel_generator.FIRMA_NO_ESPECIFICADA)
        try:
            self.db.set_config(f"excel_pendiente|{cid}", "1")
            if self._excel_cambiado_durante(cid):
                # El usuario guardó el Excel mientras trabajábamos (A4-P01): no se pisa su versión
                cliente = self.db.get_cliente(cid)
                base = self.db.ruta_excel(cliente)
                # Cada conflicto conserva su propia copia, incluso dos el mismo día.
                pendiente = base.with_name(f"{base.stem}_PENDIENTE_{uuid.uuid4().hex[:12]}{base.suffix}")
                if len(str(pendiente)) > 247:
                    pendiente = base.with_name(f"PENDIENTE_{uuid.uuid4().hex[:12]}{base.suffix}")
                ruta_excel = excel_generator.exportar_cliente(self.db, cid, ej, destino=pendiente)
                error_excel = ("Guardaste el Excel mientras se procesaba: tu versión se ha conservado. La versión "
                               f"calculada está en {pendiente.name}. La copia pendiente no se importa "
                               "automáticamente: conserva cualquier edición que hagas en ella. "
                               "Al sincronizar se leerá tu Excel principal y se volverá a generar desde los datos guardados.")
            else:
                opciones = {} if firma is excel_generator.FIRMA_NO_ESPECIFICADA else {'firma_esperada': firma}
                ruta_excel, error_excel = excel_generator.exportar_cliente(self.db, cid, ej, **opciones), None
        except excel_generator.ExcelBloqueadoError as exc:
            ruta_excel, error_excel = exc.ruta_pendiente, str(exc)
        except Exception as exc:
            raise CalculoConfirmadoSinExcel(cid, ej, resultado, exc) from exc
        if error_excel:
            self.db.set_config(f"excel_pendiente|{cid}", str(ruta_excel))
        return {**resultado, "excel": ruta_excel, "error_excel": error_excel}

    def _recalcular_y_exportar(self, cid: int, ejercicios: set[int], ejercicio_activo: int) -> dict:
        return self.commit_and_sync_excel(cid, ejercicio_activo, ejercicios)

    def commit_and_sync_excel(self, cid: int, ejercicio: int, ejercicios: Optional[set[int]] = None) -> dict:
        """Punto común de recálculo y exportación tras importar/guardar una modificación validada.

        El ejercicio es explícito para no escribir accidentalmente el año actual al editar un histórico.
        Los guardados concurrentes del asesor y los Excel bloqueados conservan sus protecciones.
        """
        with self._lock:
            ej = int(ejercicio)
            return self._exportar_informe(cid, ej, self._recalcular_bd(cid, set(ejercicios or {ej}) | {ej}, ej))

    def sincronizar_excel_si_cambio(self, cid: int, ejercicio: int) -> dict:
        """Importa al regresar a la app; si la firma no cambió no recalcula ni reescribe el archivo."""
        with self._lock:
            if self.db.get_cliente(cid) is None:
                raise ValueError("El cliente no existe.")
            if not excel_generator.excel_modificado(self.db, cid):
                return {}
            return {**self.sincronizar(cid, int(ejercicio)), "excel_importado": True}

    def conciliar_historial_210(self, cid: int, inmueble: str, ejercicio_historico: int,
                               ejercicio_activo: int, saldos: dict, observacion: str,
                               firma_esperada: str, confirmado: bool = False) -> dict:
        """El asesor verifica un cierre legado para continuar años posteriores.

        No cambia declaraciones ni arrastres históricos. La firma evita aceptar
        una conciliación sobre filas modificadas después de abrir el diálogo.
        """
        if confirmado is not True:
            raise ValueError("Confirma que has contrastado el histórico y sus saldos con los justificantes.")
        origen, activo = int(ejercicio_historico), int(ejercicio_activo)
        if not 2024 <= origen < activo:
            raise ValueError("La conciliación solo permite continuar ejercicios posteriores al histórico.")
        if not isinstance(observacion, str) or len(observacion.strip()) < 8:
            raise ValueError("Indica el justificante o la comprobación realizada por el asesor.")
        if not isinstance(saldos, dict):
            raise ValueError("Revisa los saldos por año de origen.")
        verificados = {}
        for año, importe in saldos.items():
            a = int(año)
            n = database.a_euros(importe)
            if (not origen - 3 <= a <= origen or isinstance(importe, bool) or n is None
                    or not math.isfinite(n) or not 0 <= n <= 1e12):
                raise ValueError("Los saldos deben ser importes no negativos de los cuatro orígenes vigentes.")
            verificados[a] = tax_calculator.r2(n)
        with self._lock:
            if not self.db.get("inmuebles", cid, inmueble):
                raise ValueError("El inmueble no pertenece a este cliente.")
            self._importar(cid)
            modelos = self.db.listar("m210_alquiler", cid, inmueble=inmueble)
            firma = tax_calculator.firma_historial_210(modelos, inmueble, origen)
            if firma != firma_esperada:
                raise ValueError("El histórico cambió desde que abriste la revisión. Vuelve a contrastarlo.")
            filas = [f for f in modelos if f.get("ejercicio") == origen and f.get("estado") != "No procede"]
            periodos = [tax_calculator.periodo_210_o_none(f.get("periodo")) for f in filas]
            if (not filas or any(f.get("estado") != "Presentado" for f in filas)
                    or not set(periodos) <= {"1T", "2T", "3T", "4T"}
                    or len(periodos) != len(set(periodos))):
                raise ValueError("Primero revisa los periodos y acredita las declaraciones presentadas, sin duplicados.")
            with self.db.transaccion():
                # Pago, firma y cierre forman una sola operación: un fallo del
                # recálculo no debe dejar un pago completado sin conciliación.
                # Completar un pago vacío por NRC sigue la regla ya autorizada;
                # no se alteran cuotas. La firma se calcula después de ese paso.
                for fila in filas:
                    tax_calculator.completar_pago_210_nrc(self.db, cid, "m210_alquiler", fila['codigo'])
                firma = tax_calculator.firma_historial_210(
                    self.db.listar("m210_alquiler", cid, inmueble=inmueble), inmueble, origen)
                self.db.set_config(f"conciliacion_210|{cid}|{inmueble}|{origen}", json.dumps({
                    "version": 1, "firma": firma, "saldos": verificados,
                    "observacion": observacion.strip(), "fecha": date.today().isoformat()}, sort_keys=True))
                resultado = self._recalcular_bd(cid, {activo}, activo)
            return self._exportar_informe(cid, activo, resultado)

    def confirmar_amortizacion_venta(self, cid: int, inmueble: str, ejercicio: int, importe: Any,
                                    justificacion: str, firma_esperada: str, confirmado: bool = False) -> dict:
        """Registra la cifra acreditada por el asesor sin reinterpretar su histórico."""
        if confirmado is not True:
            raise ValueError("Confirma que has contrastado la amortización acumulada con sus justificantes.")
        n = database.a_euros(importe)
        if isinstance(importe, bool) or n is None or not math.isfinite(n) or not 0 <= n <= 1e12:
            raise ValueError("Indica un importe de amortización acumulada no negativo y finito.")
        if not isinstance(justificacion, str) or len(justificacion.strip()) < 8:
            raise ValueError("Indica el historial o justificante utilizado para acreditar la amortización.")
        with self._lock:
            self._importar(cid)
            inm = self.db.get("inmuebles", cid, inmueble)
            fv = a_fecha((inm or {}).get("fecha_venta"))
            if not inm or not fv or not inm.get("precio_venta"):
                raise ValueError("Primero registra la fecha y el precio de venta de este inmueble.")
            previa = self.db.get_por_clave("obligaciones", cid, f"VENTA|{inmueble}")
            if tax_calculator._presentado(previa):
                raise ValueError("La venta ya está presentada: conserva el justificante y revisa su rectificación.")
            firma = tax_calculator.firma_amortizacion_venta(self.db, cid, inmueble)
            if firma != firma_esperada:
                raise ValueError("El expediente cambió desde que abriste la revisión. Vuelve a contrastarlo.")
            if n > tax_calculator.valor_adquisicion_total(inm) * tax_calculator.pct_titularidad(inm) + .005:
                raise ValueError("La amortización del cliente no puede superar su valor de adquisición registrado.")
            with self.db.transaccion():
                self.db.set_config(f"amortizacion_venta|{cid}|{inmueble}", json.dumps({
                    "version": 1, "importe": tax_calculator.r2(n), "justificacion": justificacion.strip(),
                    "firma": firma, "fecha": date.today().isoformat()}, sort_keys=True))
                resultado = self._recalcular_bd(cid, {fv.year, int(ejercicio)}, int(ejercicio))
            return self._exportar_informe(cid, int(ejercicio), resultado)

    def establecer_periodicidad_210(self, cid: int, inmueble: str, ejercicio: int, periodicidad: str,
                                    al_recalcular: Optional[Callable[[dict], None]] = None) -> dict:
        if periodicidad not in ("Anual", "Trimestral"):
            raise ValueError("Elige Anual o Trimestral.")
        ej = int(ejercicio)
        if ej < 2024:
            raise ValueError("El selector anual/trimestral solo está disponible desde el ejercicio 2024.")
        if periodicidad == "Trimestral":
            raise ValueError("Desde 2024 la agrupación del 210 de alquiler es anual. "
                             "El desglose trimestral de ocupación no es una declaración agrupada. "
                             "La declaración separada por devengo requiere revisión del asesor.")
        with self._lock:
            if not self.db.get("inmuebles", cid, inmueble):
                raise ValueError("El inmueble no pertenece a este cliente.")
            self._importar(cid)
            borradores = [f for f in self.db.listar("m210_alquiler", cid, inmueble=inmueble, ejercicio=ej)
                          if (f.get("estado") or "Pendiente") not in
                          (*tax_calculator.ESTADOS_AUTOMATICOS, "Presentado")
                          # un periodo no reconocible cuenta como incompatible: que lo revise el asesor
                          and (tax_calculator.periodo_210_o_none(f.get("periodo")) is None
                               or (tax_calculator.periodo_210_o_none(f.get("periodo")) == "Anual")
                               != (periodicidad == "Anual"))]
            if borradores:
                raise ValueError("Hay modelos en preparación o revisión con otra periodicidad. "
                                 "Revisa sus datos y déjalos en Pendiente antes de cambiar el calendario.")
            presentados = {tax_calculator.periodo_210_o_none(f.get("periodo"))
                           for f in self.db.listar("m210_alquiler", cid, inmueble=inmueble, ejercicio=ej)
                           if f.get("estado") == "Presentado"} - {None}
            if presentados and ((periodicidad == "Anual" and presentados != {"Anual"})
                                or (periodicidad == "Trimestral" and "Anual" in presentados)):
                raise ValueError("La periodicidad difiere de un modelo ya presentado. Revisa el histórico antes de cambiarla.")
            with self.db.transaccion():
                self.db.set_config(f"periodicidad_210|{cid}|{inmueble}|{ej}", periodicidad)
                resultado = self._recalcular_bd(cid, {ej}, ej)
            # El cálculo ya está confirmado: la UI puede mostrarlo sin esperar a serializar Excel.
            # Un fallo de presentación nunca impide guardar el libro ni libera el turno de escritura.
            if al_recalcular is not None:
                try:
                    al_recalcular(dict(resultado))
                except Exception as exc:
                    logging.getLogger(__name__).warning(
                        "No se pudo notificar el avance del recálculo (%s).", type(exc).__name__)
            try:
                return self._exportar_informe(cid, ej, resultado)
            except Exception as exc:
                raise CalculoConfirmadoSinExcel(cid, ej, resultado, exc) from exc

    def previsualizar_contrato(self, cid: int, datos: dict, ejercicio: int) -> dict:
        """Previsualización sin escrituras; una duración propone fin previsto, nunca una rescisión real."""
        columnas = {c.campo: c for c in TABLAS["contratos"].columnas_bd}
        limpios = {k: database._normalizar(columnas[k], v)
                   for k, v in datos.items() if k in columnas}
        inicio = a_fecha(limpios.get("fecha_inicio"))
        if inicio is None:
            raise ValueError("Indica la fecha de inicio del contrato.")
        meses = datos.get("duracion_meses", datos.get("meses"))
        if meses not in (None, "") and not limpios.get("fecha_fin_prevista"):
            if isinstance(meses, bool) or not str(meses).strip().isdigit() or not 1 <= int(meses) <= 1200:
                raise ValueError("La duración debe ser de 1 a 1200 meses enteros.")
            indice = inicio.year * 12 + inicio.month - 1 + int(meses)
            anio, mes = divmod(indice, 12)
            mes += 1
            limite = date(anio, mes, min(inicio.day, calendar.monthrange(anio, mes)[1]))
            limpios["fecha_fin_prevista"] = limite - timedelta(days=1)
        for nombre in ("fecha_fin_prevista", "fecha_fin_real"):
            if limpios.get(nombre) and a_fecha(limpios[nombre]) < inicio:
                raise ValueError("La fecha de fin no puede ser anterior al inicio.")
        inm = self.lector.get("inmuebles", cid, limpios.get("inmueble"))
        if not inm:
            raise ValueError("Selecciona un inmueble de este cliente.")
        limpios.setdefault("codigo", "PREVISUALIZACION")
        limpios.setdefault("estado", "Vigente")
        limpios.setdefault("periodicidad", "Mensual")
        rentas = (self.lector.listar("rentas", cid, contrato=limpios["codigo"])
                  if datos.get("codigo") else [])
        resumen = tax_calculator.resumen_arrendamiento(
            inm, [limpios], {limpios["codigo"]: rentas}, int(ejercicio))
        avisos = []
        fin_legal = tax_calculator.fin_con_prorroga(limpios)
        if fin_legal and limpios.get("fecha_fin_prevista") and fin_legal > a_fecha(limpios["fecha_fin_prevista"]):
            avisos.append("Vivienda vigente: se aplica la prórroga legal. La fecha prevista no acredita el fin real del alquiler.")
        return {**resumen, "fecha_fin_prevista": (a_fecha(limpios.get("fecha_fin_prevista")) or None),
                "avisos": avisos}

    def resumen_alquileres(self, cid: int, ejercicio: int) -> dict:
        """Todos los contratos para poder corregirlos; sus métricas corresponden solo al ejercicio."""
        ej = int(ejercicio)
        inmuebles = {i["codigo"]: i for i in self.lector.listar("inmuebles", cid)}
        contratos = []
        for fila in self.lector.listar("contratos", cid):
            inm = inmuebles.get(fila.get("inmueble"))
            resumen = (tax_calculator.resumen_arrendamiento(inm, [fila],
                       {fila["codigo"]: self.lector.listar("rentas", cid, contrato=fila["codigo"])}, ej)
                       if inm else {"dias_alquilado": 0, "ingresos_devengados": 0})
            contratos.append({**fila, "_dias_alquilado": resumen["dias_alquilado"],
                              "_ingresos_devengados": resumen["ingresos_devengados"]})
        cobros = self.lector.listar("cobros", cid, ejercicio=ej)
        return {"contratos": contratos, "cobros": cobros,
                "ingresos_devengados": tax_calculator.r2(sum(c["_ingresos_devengados"] for c in contratos)),
                "ingresos_cobrados": tax_calculator.r2(sum(float(c.get("importe_cobrado") or 0) for c in cobros))}

    def _importar(self, cid: int) -> None:
        """Importa lo manual del Excel y recuerda qué versión del archivo se leyó."""
        cliente = self.db.get_cliente(cid)
        ruta = self.db.ruta_excel(cliente) if cliente else None
        # Captura ANTES: un guardado externo durante la importación debe detectarse al exportar.
        firma = excel_generator._firma_excel(ruta) if ruta and ruta.exists() else None
        excel_generator.import_excel_to_db(self.db, cid)
        if ruta and ruta.exists() and self.db.get_config(f'excel_nombre|{cid}') != ruta.name:
            self.db.set_config(f'excel_nombre|{cid}', ruta.name)
        self._firmas[cid] = firma

    def _excel_cambiado_durante(self, cid: int) -> bool:
        if cid not in self._firmas:
            return False
        cliente = self.db.get_cliente(cid)
        ruta = self.db.ruta_excel(cliente) if cliente else None
        actual = excel_generator._firma_excel(ruta) if ruta and ruta.exists() else None
        return self._firmas.pop(cid) != actual

    # --------------------------------------------------------------- gestión de documentos
    def listar_documentos(self, cid: int, papelera: bool = False, limite: int = 51, desplazamiento: int = 0) -> list[dict]:
        return documentos.listar(self.lector, cid, papelera, limite, desplazamiento)

    def obtener_documento(self, cid: int, codigo: str) -> dict:
        return documentos._documento(self.lector, cid, codigo)

    def _version_documental(self, cid: int, codigo: str) -> str:
        doc = self.db.get("documentos", cid, codigo)
        if not doc:
            raise ValueError("El documento no pertenece a este cliente.")
        maximo = self.db.con.execute("SELECT COALESCE(MAX(id),0) FROM documentos_cambios WHERE cliente_id=?", (cid,)).fetchone()[0]
        cliente = self.db.get_cliente(cid)
        ruta = self.db.ruta_excel(cliente)
        meta = documentos._documento(self.db, cid, codigo)
        pdf = documentos._ruta(self.db, cid, meta['ruta_papelera'] if meta['en_papelera'] else doc['ruta'], papelera=meta['en_papelera'])
        datos = (doc, documentos._huella(pdf) if pdf.is_file() else None, documentos.revision(self.db, cid), maximo,
                 excel_generator._firma_excel(ruta) if ruta.exists() else None)
        return hashlib.sha256(json.dumps(datos, sort_keys=True, default=str).encode("utf-8")).hexdigest()

    def _preparar_documento(self, cid: int, codigo: str, version_esperada=None) -> None:
        cliente = self.db.get_cliente(cid)
        if not cliente:
            raise ValueError("El cliente ya no existe.")
        self.db.verificar_cartera(cliente)
        documentos.recuperar_archivos(self.db, cid)
        self._importar(cid)
        if version_esperada is not None and self._version_documental(cid, codigo) != version_esperada:
            raise ValueError("Los datos o el Excel han cambiado desde la previsualización. Revisa de nuevo el documento antes de confirmar.")
        if not self.db.get("documentos", cid, codigo):
            raise ValueError("El documento no pertenece a este cliente.")

    def previsualizar_documento(self, cid: int, codigo: str, accion: str, ejercicio_destino=None) -> dict:
        with self._lock:
            self._preparar_documento(cid, codigo)
            accion_backend = "cambiar_ejercicio" if accion == "corregir_ejercicio" else accion
            vista = documentos.previsualizar(self.db, cid, codigo, accion_backend, ejercicio_destino)
            doc = self.db.get("documentos", cid, codigo)
            cliente = self.db.get_cliente(cid)
            return {**vista, "version": self._version_documental(cid, codigo), "documento": doc,
                    "cliente": {k: cliente[k] for k in ("id", "nombre", "apellidos", "nif")}}

    def _terminar_documentos(self, cid: int, ejercicio: int, calculo: dict, resultado=None) -> dict:
        avisos = []
        try:
            documentos.recuperar_archivos(self.db, cid)
        except Exception as exc:
            avisos.append("La operación está guardada. El movimiento del archivo queda pendiente de recuperación (" + type(exc).__name__ + ").")
        try:
            res = self._exportar_informe(cid, ejercicio, calculo)
        except Exception as exc:
            res = {**calculo, "excel": None, "error_excel": "Los datos se guardaron y recalcularon, pero no se pudo actualizar el Excel. Conserva el archivo y pulsa Sincronizar cuando esté disponible (" + type(exc).__name__ + ")."}
        return {**res, "resultado": resultado, "documentos_actualizados": True, "cliente_documentos": cid,
                "avisos_documentales": avisos}

    def _operar_documento(self, cid, codigo, ejercicio, accion, destino=None, version_esperada=None):
        with self._lock:
            self._preparar_documento(cid, codigo, version_esperada)
            proteccion = (documentos.proteger_movimiento(self.db, cid, codigo, accion == "retirar")
                          if accion in ("retirar", "restaurar") else nullcontext())
            with proteccion as copias, self.db.transaccion():
                if accion == "retirar":
                    afectados = documentos.retirar(self.db, cid, codigo, copias=copias)
                elif accion == "restaurar":
                    afectados = documentos.restaurar(self.db, cid, codigo, copias=copias)
                else:
                    afectados = documentos.cambiar_ejercicio(self.db, cid, codigo, destino)
                calculo = self._recalcular_bd(cid, set(afectados) | {int(ejercicio)}, int(ejercicio))
            return self._terminar_documentos(cid, int(ejercicio), calculo)

    def retirar_documento(self, cid, codigo, ejercicio_activo, version_esperada=None):
        return self._operar_documento(cid, codigo, ejercicio_activo, "retirar", version_esperada=version_esperada)

    def restaurar_documento(self, cid, codigo, ejercicio_activo, version_esperada=None):
        return self._operar_documento(cid, codigo, ejercicio_activo, "restaurar", version_esperada=version_esperada)

    def corregir_ejercicio_documento(self, cid, codigo, nuevo_ejercicio, ejercicio_activo, version_esperada=None):
        return self._operar_documento(cid, codigo, ejercicio_activo, "cambiar_ejercicio", nuevo_ejercicio, version_esperada)

    def preparar_accion_documental(self, cid, plan):
        from acciones_documentales import validar_plan
        import gestion_archivos
        plan = validar_plan(plan)
        with self._lock:
            self._preparar_documento(cid, plan['codigo'])
            doc = self.db.get('documentos', cid, plan['codigo'])
            if plan['accion'] == 'cambiar_ejercicio':
                vista = documentos.previsualizar(self.db, cid, plan['codigo'], 'cambiar_ejercicio', plan['ejercicio_destino'])
                documentos._exigir(vista)
                nombre = f"{plan['ejercicio_destino']}_{plan['codigo']}{Path(doc['ruta']).suffix.lower()}"
            else:
                vista = {'advertencias': [], 'ejercicio_origen': doc.get('ejercicio')}
                nombre = plan['nombre']
            _, _, relativa, _, _ = gestion_archivos.preparar(self.db, cid, plan['codigo'], nombre)
            return {**vista, 'version': self._version_documental(cid, plan['codigo']),
                    'ruta_origen': doc['ruta'], 'ruta_destino': relativa}

    def ejecutar_accion_documental(self, propuesta):
        from acciones_documentales import Propuesta, validar_plan
        import gestion_archivos
        if type(propuesta) is not Propuesta:
            raise ValueError('Se requiere una propuesta documental confirmada.')
        cid, ejercicio = propuesta.cliente_id, propuesta.ejercicio
        plan = validar_plan(dict(propuesta.plan))
        codigo = plan['codigo']
        with self._lock:
            self._preparar_documento(cid, codigo, propuesta.version)
            nombre = (f"{plan['ejercicio_destino']}_{codigo}{Path(self.db.get('documentos', cid, codigo)['ruta']).suffix.lower()}"
                      if plan['accion'] == 'cambiar_ejercicio' else plan['nombre'])
            with gestion_archivos.proteger(self.db, cid, codigo, nombre) as copias, self.db.transaccion():
                afectados = set()
                if plan['accion'] == 'cambiar_ejercicio':
                    afectados |= documentos.cambiar_ejercicio(self.db, cid, codigo, plan['ejercicio_destino'])
                afectados |= gestion_archivos.renombrar(self.db, cid, codigo, nombre, plan.get('descripcion'), copias=copias)
                self.db.con.execute('INSERT INTO documentos_operaciones(cliente_id,documento,accion,detalle) VALUES(?,?,?,?)',
                                    (cid, codigo, plan['accion'], documentos._json(plan)))
                calculo = self._recalcular_bd(cid, afectados | {ejercicio}, ejercicio)
            return self._terminar_documentos(cid, ejercicio, calculo)

    def sustituir_documento(self, cid, codigo, ruta_pdf, ejercicio_activo, confirmar=None, version_esperada=None):
        with self._lock:
            self._preparar_documento(cid, codigo, version_esperada)
            doc = self.db.get("documentos", cid, codigo)
            ruta_pdf = Path(ruta_pdf)
            # Revisar el nuevo PDF antes de modificar el soporte anterior.
            huella_nueva = pdf_processor._hash(ruta_pdf)
            ext = self.extraer(ruta_pdf)
            if confirmar is not None:
                ext = confirmar(ext)
                if not isinstance(ext, pdf_processor.Extraccion):
                    raise TypeError("La revisión debe devolver una extracción válida.")
            if pdf_processor._hash(ruta_pdf) != huella_nueva:
                raise ValueError("El PDF cambió durante la revisión. Vuelve a seleccionarlo antes de sustituir el documento.")
            self._preparar_documento(cid, codigo, version_esperada)
            if huella_nueva == doc.get("sha256"):
                raise ValueError("Has elegido el mismo PDF. El documento anterior sigue activo.")
            cliente = self.db.get_cliente(cid)
            soportes = database.subcarpeta_cliente(Path(cliente["carpeta"]))
            anteriores = set(soportes.glob("*"))
            try:
                with documentos.proteger_movimiento(self.db, cid, codigo, True) as copias, self.db.transaccion():
                    with documentos.conservar_correcciones(self.db, cid, codigo):
                        afectados = documentos.retirar(self.db, cid, codigo, copias=copias)
                        resultado = pdf_processor.aplicar(self.db, cid, ext, ruta_pdf, ejercicio=int(ejercicio_activo),
                                                           huella_esperada=huella_nueva)
                        if resultado.duplicado:
                            raise ValueError("El PDF elegido ya está registrado. El documento anterior se conserva activo.")
                    calculo = self._recalcular_bd(cid, set(afectados) | set(resultado.ejercicios_afectados) | {int(ejercicio_activo)}, int(ejercicio_activo))
            except BaseException:
                # Solo retirar copias nuevas creadas por esta importación fallida y verificadas por contenido.
                for archivo in set(soportes.glob("*")) - anteriores:
                    relativo = archivo.relative_to(Path(cliente["carpeta"])).as_posix()
                    registrado = self.db.con.execute("SELECT 1 FROM documentos WHERE cliente_id=? AND ruta=?", (cid, relativo)).fetchone()
                    if not registrado and archivo.resolve().parent == soportes and pdf_processor._hash(archivo) == huella_nueva:
                        archivo.unlink()
                raise
            return self._terminar_documentos(cid, int(ejercicio_activo), calculo, resultado)

    def extraer(self, ruta_pdf: Path) -> pdf_processor.Extraccion:
        return pdf_processor.extraer(Path(ruta_pdf), self.ajustes_ia())

    def procesar_pdf(self, cid: int, ruta_pdf: Path | str, ejercicio: int, confirmar: ConfirmarFn | None = None,
                     inmueble_destino: str | None = None, comentario_guia: str = "") -> dict:
        with self._lock:
            ej = int(ejercicio)
            cliente = self.db.get_cliente(cid)
            if cliente is None:
                raise ValueError(f"Cliente inexistente: {cid}")
            self.db.verificar_cartera(cliente)  # nunca se archiva un soporte fuera de la cartera del cliente
            self._importar(cid)
            ext, huella = pdf_processor.preparar_pdf(self.db, Path(ruta_pdf), confirmar=confirmar,
                                                   comentario_guia=comentario_guia)
            # Confirmar importación y cálculo juntos; exportar solo después del commit.
            with pdf_processor.archivos_importacion(), self.db.transaccion():
                resultado = pdf_processor.aplicar(self.db, cid, ext, Path(ruta_pdf),
                                                  inmueble_destino=inmueble_destino, ejercicio=ej,
                                                  huella_esperada=huella)
                informe = self._recalcular_bd(cid, set(resultado.ejercicios_afectados) | {ej}, ej)
            return {"resultado": resultado, **self._exportar_informe(cid, ej, informe)}

    def sincronizar(self, cid: int, ejercicio: int) -> dict:
        with self._lock:
            self._importar(cid)
            existentes = {int(f["ejercicio"]) for f in self.db.listar("m210_alquiler", cid) if f.get("ejercicio")}
            return self._recalcular_y_exportar(cid, existentes | {int(ejercicio)}, int(ejercicio))

    def procesar_bandeja(self, cid: int, ejercicio: int, confirmar: ConfirmarFn | None = None,
                         comentario_guia: str = "", rutas: list[Path] | None = None) -> list[dict]:
        with self._lock:
            cliente = self.db.get_cliente(cid)
            if cliente is None:
                raise ValueError(f"Cliente inexistente: {cid}")
            bandeja = database.subcarpeta_cliente(Path(cliente["carpeta"]), database.SUBCARPETA_ENTRADA)
            bandeja.mkdir(parents=True, exist_ok=True)
            resultados: list[dict] = []
            elegidas = None if rutas is None else {Path(r).absolute() for r in rutas}
            for pdf_path in sorted(p for p in bandeja.iterdir() if p.is_file()
                    and p.suffix.lower() in {".pdf", ".xlsx", ".xls", ".docx", ".csv"} and p.resolve().parent == bandeja
                    and (elegidas is None or p.absolute() in elegidas)):
                error_msg: Optional[str] = None
                try:
                    res = self.procesar_pdf(cid, pdf_path, ejercicio, confirmar=confirmar,
                                            **({"comentario_guia": comentario_guia} if comentario_guia else {}))
                    pdf_path.unlink(missing_ok=True)
                    resultados.append({**res, "archivo": pdf_path.name, "ok": True, "error": None})
                except CanceladoPorUsuario:
                    # Se aparta: si se queda, el vigilante vuelve a abrir el diálogo cada 2 s (A3-01)
                    destino = self._mover_unico(pdf_path, bandeja / "_descartados")
                    resultados.append({"archivo": pdf_path.name, "ok": False, "cancelado": True,
                                       "error": "Descartado por el usuario", "ruta_descartado": destino})
                except Exception as exc:
                    error_msg = str(exc) or exc.__class__.__name__
                if error_msg is not None:
                    # Fuera del except: la traza ya no retiene el PDF abierto por PyMuPDF (WinError 32)
                    gc.collect()
                    destino = self._mover_unico(pdf_path, bandeja / "_errores")
                    txt = destino.with_suffix(".txt")
                    txt.write_text(f"Archivo: {pdf_path.name}\nError al procesar PDF: {error_msg}\n", encoding="utf-8")
                    resultados.append({"archivo": pdf_path.name, "ok": False, "error": error_msg,
                                       "ruta_error": destino, "txt_error": txt})
            return resultados

    @staticmethod
    def _mover_unico(origen: Path, carpeta: Path) -> Path:
        if carpeta.resolve() != origen.parent.resolve() / carpeta.name:
            raise ValueError("La carpeta de destino sale de la bandeja del cliente.")
        carpeta.mkdir(parents=True, exist_ok=True)
        destino, n = carpeta / origen.name, 2
        while destino.exists():
            destino = carpeta / f"{origen.stem}_{n}{origen.suffix}"
            n += 1
        shutil.move(str(origen), str(destino))
        return destino

    # ------------------------------------------------------------------ edición manual
    def validar(self, tabla: str, datos: dict) -> dict[str, str]:
        """{campo: mensaje} con los valores que no se pueden interpretar. {} = todo correcto."""
        columnas = {c.campo: c for c in TABLAS[tabla].columnas_bd}
        nombres = etiquetas(tabla)
        errores = {}
        for campo, valor in datos.items():
            col = columnas.get(campo)
            if col is None or valor in (None, ""):
                continue
            try:
                database._normalizar(col, valor)
            except (ValueError, TypeError):
                ayuda = {"date": "usa dd/mm/aaaa", "eur": "usa un importe como 1.234,56",
                         "pct": "usa un porcentaje entre 0 % y 100 % (por ejemplo 50 % o 0,5)",
                         "int": "usa un número entero"}.get(col.tipo, "valor no válido")
                errores[campo] = f"{nombres.get(campo, campo)}: {ayuda}"
        return errores

    def editar(self, cid: int, tabla: str, codigo: Optional[str], datos: dict, ejercicio: int,
               esperados: Optional[dict] = None, version_esperada: Optional[str] = None) -> dict:
        """Alta (codigo=None) o corrección manual de una fila. Un campo enviado vacío se borra."""
        if tabla == "documentos" and codigo and "ejercicio" in datos:
            previo = self.db.get(tabla, cid, codigo)
            if previo and str(datos["ejercicio"]) != str(previo.get("ejercicio")):
                if set(datos) != {"ejercicio"}:
                    raise ValueError("Corrige el ejercicio con su acción específica antes de editar otros datos del documento.")
                return self.corregir_ejercicio_documento(cid, codigo, int(datos["ejercicio"]), ejercicio)
        errores = self.validar(tabla, datos)
        if errores:
            raise ValueError("; ".join(errores.values()))
        with self._lock:
            self._importar(cid)
            validos = {c.campo for c in TABLAS[tabla].columnas_bd}
            limpios = {k: v for k, v in datos.items() if k in validos and k not in {"codigo", "revision_pendiente"}}
            if tabla in ("gastos", "facturas_emitidas", "facturas_recibidas") and not codigo:
                limpios.setdefault("origen", "app")
            if tabla == "residencias" and limpios.get("ejercicio"):
                codigo = codigo or str(int(limpios["ejercicio"]))
            previa = self.db.get(tabla, cid, codigo) if codigo else None
            if version_esperada is not None:
                version = hashlib.sha256(json.dumps(previa, sort_keys=True, default=str).encode('utf-8')).hexdigest()
                if version != version_esperada:
                    raise ValueError("El registro cambió durante la confirmación. Solicita una nueva propuesta.")
            if codigo and esperados is not None:
                if previa is None:
                    raise ValueError("El registro ya no existe. Reabre la ficha antes de guardar.")
                columnas = {c.campo: c for c in TABLAS[tabla].columnas_bd}
                for campo in limpios.keys() & esperados.keys():
                    anterior = database._normalizar(columnas[campo], esperados[campo])
                    propuesto = database._normalizar(columnas[campo], limpios[campo])
                    if previa.get(campo) != anterior and previa.get(campo) != propuesto:
                        sujeto = "El contrato" if tabla == "contratos" else "El registro"
                        raise ValueError(sujeto + " cambió desde que abriste el diálogo (por ejemplo, en Excel). "
                                         "No se han sobrescrito tus datos. Reabre la ficha y revisa los cambios.")
            if (tabla not in ('facturas_emitidas', 'facturas_recibidas')
                    and limpios.get('estado', (previa or {}).get('estado')) != 'Presentado'):
                columnas = {c.campo: c for c in TABLAS[tabla].columnas_bd}
                cambiados = [c for c in CALCULADOS.get(tabla, set()) & limpios.keys()
                             if database._normalizar(columnas[c], limpios[c]) != (previa or {}).get(c)]
                if cambiados:
                    raise ValueError("Estos campos se recalculan automáticamente: " + ', '.join(sorted(cambiados))
                                     + ". Corrige los datos de origen (contrato, renta, gastos o catastro).")
            if tabla == "contratos" and (not codigo or any(
                    k in limpios for k in ("fecha_inicio", "fecha_fin_prevista", "fecha_fin_real"))):
                propuesta = {**(previa or {}), **limpios}
                inicio = a_fecha(propuesta.get("fecha_inicio"))
                for campo in ("fecha_fin_prevista", "fecha_fin_real"):
                    fin = a_fecha(propuesta.get(campo))
                    if inicio and fin and fin < inicio:
                        raise ValueError("La fecha de fin del contrato no puede ser anterior al inicio.")
            if tabla == "catastro" and limpios.get("anio_revision"):
                anio_valor = limpios.get("ejercicio") or (previa or {}).get("ejercicio") or ejercicio
                if int(limpios["anio_revision"]) > int(anio_valor):
                    raise ValueError("La revisión catastral no puede ser posterior al ejercicio del valor.")
            with self.db.transaccion():
                if codigo:
                    nuevo = self.db.upsert(tabla, cid, {"codigo": codigo, **limpios}, respetar_manuales=False)
                    actual = self.db.get(tabla, cid, nuevo) or {}
                    for campo, valor in limpios.items():
                        if valor in (None, "") and actual.get(campo) not in (None, ""):
                            self.db.vaciar_campo(tabla, cid, nuevo, campo)
                else:
                    nuevo = self.db.upsert(tabla, cid, limpios, respetar_manuales=False)
                fila = self.db.get(tabla, cid, nuevo) or {}
                # Los ajustes administrativos de una presentada conservan su justificante.
                # Cambiar su identidad fiscal exige validar el nuevo destino, aun omitiendo estado.
                cambia_destino = any((previa or {}).get(k) != fila.get(k)
                                     for k in ("ejercicio", "inmueble", "id_modelo"))
                if fila.get("estado") == "Presentado" and (
                        (previa or {}).get("estado") != "Presentado" or cambia_destino):
                    self._validar_imputacion_presentable(cid, tabla, nuevo, fila_propuesta=fila,
                        presentada_previamente=tax_calculator._presentado(previa) and not cambia_destino)
                if tabla == "inmuebles":
                    self.db.proteger_campos_inmueble(cid, nuevo, (
                        k for k in limpios if (previa or {}).get(k) != fila.get(k)))
                afectados = {int(ejercicio)}
                if tabla == "gastos":
                    afectados |= tax_calculator.ejercicios_gasto(previa or {}) | tax_calculator.ejercicios_gasto(fila)
                elif tabla in ("inmuebles", "contratos", "rentas", "residencias"):
                    afectados |= {int(f["ejercicio"]) for f in self.db.listar("m210_alquiler", cid) if f.get("ejercicio")}
                if isinstance((previa or {}).get("ejercicio"), int):
                    afectados.add(previa["ejercicio"])
                if isinstance(fila.get("ejercicio"), int):
                    afectados.add(fila["ejercicio"])
                resultado = self._recalcular_bd(cid, afectados, int(ejercicio))
            return self._exportar_informe(cid, int(ejercicio), {"codigo": nuevo, **resultado})

    def aplicar_correccion_chat(self, propuesta, cliente_activo: int, ejercicio_activo: int, al_guardar=None):
        """Solo comandos confirmados del usuario; nunca se invoca desde una respuesta del modelo."""
        import chat_correcciones as cc
        if type(propuesta) is not cc.Propuesta or (propuesta.cliente_id, propuesta.ejercicio) != (cliente_activo, ejercicio_activo):
            raise cc.ErrorCorreccion("La propuesta no corresponde al expediente activo.")
        if propuesta.campo != "valor_adquisicion":
            raise cc.ErrorCorreccion("Campo de corrección no permitido.")
        valor = cc.normalizar_importe(propuesta.valor)
        cid, ej = propuesta.cliente_id, propuesta.ejercicio
        with self._lock:
            cliente = self.db.get_cliente(cid)
            if not cliente or cartera_de(cliente.get("tipo_cliente")) != NR:
                raise cc.ErrorCorreccion("La corrección de inmuebles requiere un expediente de No residentes.")
            self.db.verificar_cartera(cliente)
            self._importar(cid)
            if self.db.get("inmuebles", cid, propuesta.codigo) is None:
                raise cc.ErrorCorreccion("El inmueble no pertenece a este cliente.")
            # Un precio es interanual: no reinterpretar ningún histórico presentado.
            for tabla in ("m210_alquiler", "m210_imputacion", "patrimonio_714", "its_718", "obligaciones"):
                if any(tax_calculator._presentado(f) for f in self.db.iterar(tabla, cid)):
                    raise cc.ErrorCorreccion("El expediente tiene declaraciones presentadas. Revisa la corrección manualmente.")
            resultado = self.editar(cid, "inmuebles", propuesta.codigo, {"precio_escritura": float(valor)}, ej)
            if resultado.get("error_excel"):
                raise cc.GuardadoNoVerificado("El dato se guardó en la BD y se recalculó, pero el Excel está pendiente. "
                                             + resultado["error_excel"])
            fila = self.db.get("inmuebles", cid, propuesta.codigo)
            if not fila or self.db.get_por_clave("patrimonio_714", cid, f"714|{ej}") is None:
                raise cc.GuardadoNoVerificado("Revisa la ficha: no se ha podido verificar el recálculo de Patrimonio.")
            if al_guardar is not None:
                al_guardar(resultado)
            return cc.ReciboGuardado(propuesta.id, cid, ej, propuesta.codigo, propuesta.campo,
                                    cc.importe_leido(fila["precio_escritura"]), 1)

    def preparar_correccion_expediente(self, cid: int, ejercicio: int, plan: dict) -> dict:
        from correcciones_expediente import validar_plan
        plan = validar_plan(plan)
        with self._lock:
            cliente = self.db.get_cliente(cid)
            if cliente is None:
                raise ValueError("El cliente no existe.")
            self.db.verificar_cartera(cliente)
            permitidas = {t.nombre for t, _ in database.hojas_cartera(cartera_de(cliente.get('tipo_cliente')))}
            if plan['tabla'] not in permitidas:
                raise ValueError("El registro no pertenece a la cartera activa.")
            self._importar(cid)
            fila = self.db.get(plan['tabla'], cid, plan['codigo'])
            if fila is None:
                raise ValueError("El registro no pertenece a este cliente o ya no existe.")
            errores = self.validar(plan['tabla'], {plan['campo']: plan['valor']})
            if errores:
                raise ValueError('; '.join(errores.values()))
            columna = next(c for c in TABLAS[plan['tabla']].columnas_bd if c.campo == plan['campo'])
            despues = database._normalizar(columna, plan['valor'])
            if plan['tabla'] == 'contratos':
                self.previsualizar_contrato(cid, {**fila, plan['campo']: despues}, int(ejercicio))
            version = hashlib.sha256(json.dumps(fila, sort_keys=True, default=str).encode('utf-8')).hexdigest()
            tablas_presentadas = _DECLARACIONES_RES if self._es_residente(cid) else ('m210_alquiler', 'm210_imputacion', 'patrimonio_714', 'its_718')
            advertencias = (['Las declaraciones presentadas conservarán sus importes. Si la corrección les afecta, revisa su rectificación.']
                            if any(tax_calculator._presentado(f) for tabla in tablas_presentadas for f in self.db.iterar(tabla, cid)) else [])
            if plan['tabla'] in ('facturas_emitidas', 'facturas_recibidas'):
                advertencias.append('Comprueba también los tipos y cuotas de IVA y retención de la factura: corregir un dato no autoriza cambiar los demás importes documentados.')
            return {'version': version, 'antes': fila.get(plan['campo']), 'despues': despues,
                    'cliente': ' '.join(str(cliente.get(c) or '') for c in ('nombre', 'apellidos')).strip(),
                    'etiqueta': etiquetas(plan['tabla']).get(plan['campo'], plan['campo']),
                    'advertencias': advertencias}

    def ejecutar_correccion_expediente(self, propuesta) -> dict:
        from acciones_documentales import Propuesta
        from correcciones_expediente import validar_plan
        if type(propuesta) is not Propuesta:
            raise ValueError("Se requiere una propuesta confirmada.")
        plan = validar_plan(dict(propuesta.plan))
        cid, ej = propuesta.cliente_id, propuesta.ejercicio
        with self._lock:
            vista = self.preparar_correccion_expediente(cid, ej, plan)
            if vista['version'] != propuesta.version:
                raise ValueError("El registro cambió desde la propuesta. Revisa los datos y solicita otra corrección.")
            try:
                resultado = self.editar(cid, plan['tabla'], plan['codigo'], {plan['campo']: plan['valor']}, ej,
                                        esperados={plan['campo']: vista['antes']}, version_esperada=propuesta.version)
            except CalculoConfirmadoSinExcel as exc:
                resultado = {**exc.resultado, 'error_excel': str(exc), 'excel': None}
            actual = self.db.get(plan['tabla'], cid, plan['codigo'])
            if actual is None or actual.get(plan['campo']) != vista['despues']:
                raise ValueError("No se pudo verificar el dato guardado. Revisa la ficha antes de repetirlo.")
            return {**resultado, 'registro_actualizado': True}

    # Dependientes que se borran con su inmueble o contrato (las referencias son texto, sin claves foráneas)
    _DEPENDIENTES = {
        "inmuebles": ("inmueble", ("catastro", "contratos", "inquilinos", "rentas", "cobros", "gastos",
                                   "m210_alquiler", "m210_imputacion", "obligaciones", "arrastres_gastos")),
        "contratos": ("contrato", ("inquilinos", "rentas", "cobros")),
    }
    _DECLARACIONES = ("m210_alquiler", "m210_imputacion", "obligaciones")

    def editar_datos_cliente(self, cid: int, ejercicio: int, datos: dict, residencia: dict,
                            nuevo_tipo: str) -> dict:
        """Confirma ficha, residencia y cálculo juntos; exporta después del commit."""
        permitidos = {'nombre', 'apellidos', 'pasaporte', 'nacionalidad', 'email', 'telefono', 'responsable'}
        if set(datos) - permitidos or set(residencia) - {'pais', 'certificado'}:
            raise ValueError('La ficha contiene campos no editables.')
        if 'nombre' in datos and not str(datos['nombre'] or '').strip():
            raise ValueError('El nombre del cliente no puede quedar vacío.')
        nuevo_tipo = carteras.validar_tipo(nuevo_tipo)
        ej = int(ejercicio)
        with self._lock:
            self._importar(cid)
            cliente = self.db.get_cliente(cid)
            if not cliente:
                raise ValueError('El cliente ya no existe.')
            if cartera_de(cliente.get('tipo_cliente')) != cartera_de(nuevo_tipo):
                raise ValueError('El cambio de cartera requiere su confirmación específica.')
            with self.db.transaccion():
                self.db.actualizar_cliente(cid, tipo_cliente=nuevo_tipo, **datos)
                if residencia:
                    cambios = dict(residencia)
                    if cambios.get('pais'):
                        cambios['grupo'] = tax_calculator.grupo_territorio(cambios['pais'])
                    codigo = self.db.guardar_residencia(cid, ej, **cambios)
                    for campo, valor in residencia.items():
                        if valor == '':
                            self.db.vaciar_campo('residencias', cid, codigo, campo)
                resultado = self._recalcular_bd(cid, {ej}, ej)
            return self._exportar_informe(cid, ej, resultado)

    def cambiar_tipo_cliente(self, cid: int, nuevo_tipo: str, ejercicio: int) -> dict:
        """Cambia el tipo de cliente. Si cambia de cartera, mueve su carpeta a la otra raíz (todo o nada).

        El Excel anterior pasa a _backups (la otra cartera usa otra plantilla) y se genera el nuevo. Se rechaza
        si el cliente tiene declaraciones presentadas: el histórico fiscal de una cartera no se reinterpreta."""
        nuevo_tipo = carteras.validar_tipo(nuevo_tipo)
        with self._lock:
            cliente = self.db.get_cliente(cid)
            if not cliente:
                raise ValueError("El cliente ya no existe.")
            origen = self.db.verificar_cartera(cliente)
            self._importar(cid)
            antes, despues = cartera_de(cliente.get("tipo_cliente")), cartera_de(nuevo_tipo)
            if antes == despues:
                return self.editar_datos_cliente(cid, ejercicio, {}, {}, nuevo_tipo)
            tablas = (self._DECLARACIONES + ("patrimonio_714", "its_718")) if antes == NR else (*_DECLARACIONES_RES, "obligaciones")
            presentadas = [f"{t} {r['codigo']}" for t in tablas for r in self.db.listar(t, cid)
                           if r.get("estado") == "Presentado"]
            if presentadas:
                raise ValueError(f"Tiene declaraciones presentadas ({', '.join(presentadas[:3])}): no se cambia de "
                                 "cartera para conservar su histórico. Crea un expediente nuevo en la otra cartera.")
            raiz = _carpeta_sin_enlaces(self.db.raiz(despues))
            raiz.mkdir(parents=True, exist_ok=True)
            destino, n = raiz / origen.name, 2
            while destino.exists():
                destino, n = raiz / f"{origen.name}_{n}", n + 1
            destino = _carpeta_sin_enlaces(destino)
            excel_viejo = self.db.ruta_excel(cliente)
            movida, respaldo = False, None
            viejo = destino / excel_viejo.name
            try:
                with self.db.transaccion():
                    if origen.exists():
                        try:
                            origen.rename(destino)
                        except OSError as exc:
                            raise ValueError("No se pudo mover la carpeta del cliente: cierra su Excel o sus "
                                             f"PDFs y vuelve a intentarlo ({exc.__class__.__name__}).") from None
                        movida = True
                    self.db.actualizar_cliente(cid, tipo_cliente=nuevo_tipo, carpeta=str(destino))
                    self.db.asegurar_carpetas(self.db.get_cliente(cid))
                    if viejo.exists():
                        copias = _carpeta_sin_enlaces(destino / "_backups")
                        copias.mkdir(exist_ok=True)
                        copia = copias / f"{viejo.stem}_antes_cambio_cartera{viejo.suffix}"
                        numero = 2
                        while copia.exists():
                            copia = copias / f"{viejo.stem}_antes_cambio_cartera_{numero}{viejo.suffix}"
                            numero += 1
                        _carpeta_sin_enlaces(copia)
                        viejo.rename(copia)
                        respaldo = copia
                    self.db.con.execute("DELETE FROM config WHERE clave IN (?,?,?,?,?)",
                        tuple(f'{p}|{cid}' for p in ('excel_firma', 'excel_nombre', 'excel_version', 'excel_filas', 'excel_revision')))
            except Exception as exc:
                fallos = []
                if respaldo is not None:
                    try:
                        respaldo.rename(viejo)
                    except OSError as error:
                        fallos.append(f"Excel en {respaldo}: {error.__class__.__name__}")
                if movida:
                    try:
                        destino.rename(origen)
                    except OSError as error:
                        fallos.append(f"carpeta en {destino}: {error.__class__.__name__}")
                if fallos:
                    raise RuntimeError("No se pudo restaurar por completo el cambio de cartera; requiere "
                                       "recuperación manual. " + "; ".join(fallos)) from exc
                raise
            self._firmas.pop(cid, None)
            return self._recalcular_y_exportar(cid, {int(ejercicio)}, int(ejercicio))

    def borrar_registro(self, cid: int, tabla: str, codigo: str, ejercicio: int) -> dict:
        if tabla == "documentos":
            return self.retirar_documento(cid, codigo, ejercicio)
        """Borra una fila. Inmuebles y contratos arrastran sus datos dependientes (A4-P04), salvo que el
        inmueble tenga declaraciones ya presentadas: entonces se conserva el histórico."""
        with self._lock:
            self._importar(cid)
            previa = self.db.get(tabla, cid, codigo)
            if tax_calculator._presentado(previa) and tabla in (*self._DECLARACIONES, 'patrimonio_714', 'its_718', *_DECLARACIONES_RES):
                raise ValueError("La declaración está presentada: conserva el histórico y revisa su rectificación.")
            if tabla == 'contratos' and previa:
                inmueble = previa.get('inmueble')
                if inmueble and any(tax_calculator._presentado(f) for t in self._DECLARACIONES
                        for f in self.db.listar(t, cid, inmueble=inmueble)):
                    raise ValueError("El inmueble tiene declaraciones presentadas: conserva el contrato como soporte "
                                     "histórico y corrige sus fechas o estado, sin borrarlo.")
            afectados = {int(ejercicio)}
            if (previa or {}).get('ejercicio'):
                afectados.add(int(previa['ejercicio']))
            if tabla == 'gastos':
                afectados |= tax_calculator.ejercicios_gasto(previa or {})
            if tabla in ('inmuebles', 'contratos', 'rentas', 'residencias'):
                modelos = _DECLARACIONES_RES if self._es_residente(cid) else ('m210_alquiler', 'm210_imputacion', 'patrimonio_714', 'its_718')
                for modelo in modelos:
                    afectados.update(int(r[0]) for r in self.db.con.execute(
                        f'SELECT DISTINCT ejercicio FROM {modelo} WHERE cliente_id=? AND ejercicio IS NOT NULL', (cid,)))
            with self.db.transaccion():
                self._borrar_registro_dependientes(cid, tabla, codigo)
                resultado = self._recalcular_bd(cid, afectados, int(ejercicio))
            return self._exportar_informe(cid, int(ejercicio), resultado)

    def _borrar_registro_dependientes(self, cid: int, tabla: str, codigo: str) -> None:
        """Solo escrituras SQLite: el llamador garantiza transacción y protecciones."""
        # Se conserva la comprobación del inmueble y su cascada explícita.
        if tabla in self._DEPENDIENTES:
            campo, tablas = self._DEPENDIENTES[tabla]
            presentadas = [f"{t} {r['codigo']}" for t in tablas if t in self._DECLARACIONES
                           for r in self.db.listar(t, cid, **{campo: codigo}) if r.get("estado") == "Presentado"]
            if presentadas:
                raise ValueError(f"{codigo} tiene declaraciones presentadas ({', '.join(presentadas[:3])}). "
                                 "No se borra para conservar el histórico: revisa sus fechas y estado sin eliminarlo.")
            contratos = ([r["codigo"] for r in self.db.listar("contratos", cid, inmueble=codigo)]
                         if tabla == "inmuebles" else [codigo])
            for t in tablas:
                for r in self.db.listar(t, cid, **{campo: codigo}):
                    self.db.borrar(t, cid, r["codigo"])
            for ctr in contratos:  # inquilinos/rentas/cobros enlazados por contrato
                for t in ("inquilinos", "rentas", "cobros"):
                    for r in self.db.listar(t, cid, contrato=ctr):
                        self.db.borrar(t, cid, r["codigo"])
        self.db.borrar(tabla, cid, codigo)

    def _validar_imputacion_presentable(self, cid: int, tabla: str, codigo: str,
                                          fila_propuesta: Optional[dict] = None, *,
                                          presentada_previamente: Optional[bool] = None) -> None:
        from validacion_catastro import revision_pendiente
        fila = fila_propuesta if fila_propuesta is not None else (self.db.get(tabla, cid, codigo) or {})
        previa = self.db.get(tabla, cid, codigo) or {}
        ya_presentada = (tax_calculator._presentado(previa) if presentada_previamente is None
                         else presentada_previamente)
        if tabla == 'm303' and not ya_presentada:
            anio303 = fila.get('ejercicio') or previa.get('ejercicio')
            periodo303 = fila.get('periodo') or previa.get('periodo')
            if not anio303 or periodo303 not in fiscal_residentes.TRIMESTRES:
                raise ValueError('El modelo 303 vinculado no existe o no tiene un ejercicio y periodo válidos. Revisa el registro antes de presentarlo.')
            pendientes = fiscal_residentes.pendientes_saldo_303(self.db, cid,
                int(anio303), periodo303)
            if pendientes:
                raise ValueError('El cálculo 303 es provisional: confirma primero con el asesor el destino de los saldos negativos anteriores.')
        if (tabla == "obligaciones" and str(previa.get("clave") or "").startswith("VENTA|")
                and not ya_presentada
                and tax_calculator.amortizacion_venta_confirmada(self.db, cid, previa.get("inmueble")) is None):
            raise ValueError("La venta es provisional: confirma la amortización acumulada acreditada en la ficha del inmueble.")
        if tabla == "obligaciones":
            modelo = fila.get("id_modelo")
            destino_modelo = _PREFIJO_A_TABLA.get(modelo.rsplit("-", 1)[0]) if modelo else None
            if destino_modelo in ('m210_imputacion', 'm303'):
                return self._validar_imputacion_presentable(cid, destino_modelo, modelo,
                    fila_propuesta=self.db.get(destino_modelo, cid, modelo))
        if tabla == "m210_imputacion" and (fila_propuesta is not None or fila.get("estado") != "Presentado"):
            if int(fila.get("ejercicio") or 0) >= 2027 and not ya_presentada:
                raise ValueError("La previsión de imputación desde 2027 es provisional: la nueva normativa "
                                 "requiere validar el motor agregado antes de presentar desde este cálculo. "
                                 "Las declaraciones ya presentadas y sus justificantes se conservan.")
            cat = self.db.catastro(cid, fila.get("inmueble"), fila.get("ejercicio"))
            if revision_pendiente(cat):
                raise ValueError("La imputación es provisional: acredita la revisión catastral o confirma el porcentaje manual en Catastro antes de marcarla presentada.")

    def marcar_presentado(self, cid: int, tabla: str, codigo: str, ejercicio: int,
                          fecha_presentacion: Any = None, pagado: Any = None, justificante: Optional[str] = None) -> dict:
        """Marca una declaración como presentada (y su obligación/modelo enlazado)."""
        if tabla not in CAMPOS_PRESENTACION:
            raise ValueError(f"{tabla} no es una declaración")
        datos: dict[str, Any] = {"estado": "Presentado", "fecha_presentacion": fecha_presentacion or date.today(),
                                 "pagado": pagado}
        campo_just = next((c for c in CAMPOS_PRESENTACION[tabla] if c in ("nrc", "csv")), None)
        if campo_just and justificante:
            datos[campo_just] = justificante
        errores = self.validar(tabla, datos)
        if errores:
            raise ValueError("; ".join(errores.values()))
        with self._lock:
            self._importar(cid)
            self._validar_imputacion_presentable(cid, tabla, codigo)
            with self.db.transaccion():
                self.db.upsert(tabla, cid, {"codigo": codigo, **datos})
                enlazados = []
                if tabla == "obligaciones":
                    obl = self.db.get("obligaciones", cid, codigo) or {}
                    modelo = obl.get("id_modelo")
                    destino = _PREFIJO_A_TABLA.get((modelo or "").rsplit("-", 1)[0]) if modelo else None
                    if destino in CAMPOS_PRESENTACION:
                        enlazados.append((destino, modelo))
                else:
                    enlazados += [("obligaciones", o["codigo"]) for o in self.db.listar("obligaciones", cid, id_modelo=codigo)]
                for t, cod in enlazados:
                    self.db.upsert(t, cid, {"codigo": cod, "estado": "Presentado",
                                            "fecha_presentacion": datos["fecha_presentacion"], "pagado": pagado})
                resultado = self._recalcular_bd(cid, {int(ejercicio)}, int(ejercicio))
            return self._exportar_informe(cid, int(ejercicio), resultado)

    def confirmar_saldo_303(self, cid: int, codigo: str, ejercicio: int, destino: str,
                            saldo: Any, justificacion: str, firma_esperada: str, *, confirmado: bool = False) -> dict:
        """Confirma el arrastre, sin reescribir el justificante presentado ni solicitar devoluciones a Hacienda."""
        if destino not in ('compensar', 'devolver', 'pendiente'):
            raise ValueError('Elige compensar, devolución solicitada o pendiente de revisar.')
        if not confirmado:
            raise ValueError('El asesor debe confirmar la elección y el saldo con su justificante.')
        nota = str(justificacion or '').strip()
        if destino != 'pendiente' and len(nota) < 8:
            raise ValueError('Indica el justificante contrastado por el asesor.')
        valor = database.a_euros(saldo) if destino == 'compensar' else 0.0
        if valor is None or not math.isfinite(valor) or not 0 <= valor <= 1e12:
            raise ValueError('Indica el saldo final a compensar acreditado, no negativo.')
        with self._lock:
            if not self._es_residente(cid):
                raise ValueError('El modelo 303 pertenece a la cartera de residentes.')
            self._importar(cid)
            fila = self.db.get('m303', cid, codigo)
            if not fila or not fiscal_residentes.saldo_303(self.db, cid, fila)['requiere']:
                raise ValueError('No hay un saldo presentado pendiente de decisión para este registro.')
            if fiscal_residentes.firma_saldo_303(fila) != firma_esperada:
                raise ValueError('El justificante cambió. Vuelve a abrir la revisión antes de confirmar.')
            if destino == 'devolver' and fila.get('periodo') != '4T':
                raise ValueError('Este motor trimestral general solo contempla solicitar devolución en 4T. '
                                 'Los supuestos especiales requieren revisión fuera de este flujo.')
            minimo = fiscal_residentes.minimo_saldo_303(fila)
            if destino == 'compensar' and valor + .005 < minimo:
                raise ValueError('El saldo final no puede ser inferior al resultado negativo declarado. Revisa el justificante.')
            dato = {'version':1, 'firma':firma_esperada, 'destino':destino, 'saldo':round(valor,2),
                    'justificacion':nota, 'fecha_confirmacion':date.today().isoformat()}
            with self.db.transaccion():
                self.db.set_config(f'saldo303|{cid}|{codigo}', json.dumps(dato, ensure_ascii=False))
                resultado = self._recalcular_bd(cid, {int(fila['ejercicio']),int(ejercicio)}, int(ejercicio))
            return self._exportar_informe(cid, int(ejercicio), resultado)

    # ------------------------------------------------------------------ atajos usados por la UI y los tests
    def _datos_gananciales(self, cid: int, inmueble: str, activo: bool) -> dict:
        """Al forzar el 50 % se guarda el % de la escritura (H muestra 50 % en Excel); al quitarlo se restaura."""
        fila = self.db.get("inmuebles", cid, inmueble) or {}
        if activo:
            datos = {"codigo": inmueble, "gananciales_forzado": 1, "pct_titularidad": 0.5}
            if not fila.get("gananciales_forzado"):
                datos["pct_titularidad_escritura"] = fila.get("pct_titularidad") if fila.get("pct_titularidad") is not None else 1.0
            return datos
        original = fila.get("pct_titularidad_escritura")
        return {"codigo": inmueble, "gananciales_forzado": 0,
                "pct_titularidad": original if original is not None else fila.get("pct_titularidad")}

    def set_gananciales(self, cid: int, inmueble: str, activo: bool, ejercicio: int) -> dict:
        with self._lock:
            self._importar(cid)
            datos = self._datos_gananciales(cid, inmueble, activo)
            with self.db.transaccion():
                codigo = self.db.upsert("inmuebles", cid, datos, respetar_manuales=False)
                self.db.proteger_campos_inmueble(cid, codigo, datos.keys())
            existentes = {int(f["ejercicio"]) for f in self.db.listar("m210_alquiler", cid) if f.get("ejercicio")}
            return {"inmueble": codigo, **self._recalcular_y_exportar(cid, existentes | {int(ejercicio)}, int(ejercicio))}

    def marcar_vendido(self, cid: int, inmueble: str, fecha_venta: str | date, precio_venta: float | str,
                       ejercicio: int, gastos_venta: float | None = None) -> dict:
        with self._lock:
            self._importar(cid)
            datos: dict[str, Any] = {"codigo": inmueble, "estado": "Vendido", "fecha_venta": fecha_venta,
                                     "precio_venta": precio_venta}
            if gastos_venta is not None:
                datos["gastos_venta"] = gastos_venta
            with self.db.transaccion():
                codigo = self.db.upsert("inmuebles", cid, datos, respetar_manuales=False)
                self.db.proteger_campos_inmueble(cid, codigo, datos.keys())
            fv = a_fecha(fecha_venta)
            afectados = {int(ejercicio), fv.year} if fv else {int(ejercicio)}
            return {"inmueble": codigo, **self._recalcular_y_exportar(cid, afectados, int(ejercicio))}

    def alta_gasto(self, cid: int, ejercicio: int, inmueble: str, categoria: str, fecha: str | date,
                   total: float | str, **extra: Any) -> dict:
        datos = {"ejercicio": int(ejercicio), "inmueble": inmueble, "categoria": categoria, "fecha": fecha,
                 "total": total, "origen": "app", "incluido_210": extra.pop("incluido_210", "Pendiente"), **extra}
        res = self.editar(cid, "gastos", None, datos, ejercicio)
        return {"gasto": res["codigo"], **res}

    # ------------------------------------------------------------------ vistas de lectura para la UI
    def vista_inmuebles(self, cid: int, ejercicio: int, limite: Optional[int] = None,
                        desplazamiento: int = 0) -> list[dict]:
        out = []
        # El detector ya necesita el histórico del cliente. Reutilizar esa misma
        # lectura evita repetir consultas por inmueble para el desglose y la firma.
        modelos_210 = self.lector.listar("m210_alquiler", cid, orden="periodo")
        conciliadas = tax_calculator.conciliaciones_210(self.lector, cid, modelos_210)
        conflictos = tax_calculator._conflictos_modelos_210(modelos_210, int(ejercicio), set(conciliadas))
        historias = {}
        for modelo in modelos_210:
            historias.setdefault(modelo.get("inmueble"), []).append(modelo)
        inmuebles = (self.lector.listar_pagina("inmuebles", cid, limite, desplazamiento)
                     if limite is not None else self.lector.iterar("inmuebles", cid))
        for inm in inmuebles:
            cat = self.lector.catastro(cid, inm["codigo"], int(ejercicio))
            vp = tax_calculator.valor_patrimonial(inm, cat)
            tit = tax_calculator.pct_titularidad(inm)
            contratos = self.lector.listar("contratos", cid, inmueble=inm["codigo"])
            rentas = {c["codigo"]: self.lector.listar("rentas", cid, contrato=c["codigo"]) for c in contratos}
            ocupacion = tax_calculator.resumen_arrendamiento(inm, contratos, rentas, int(ejercicio))
            compra = a_fecha(inm.get("fecha_adquisicion"))
            ref_compra = inm.get("valor_referencia")
            if ref_compra in (None, "") and compra:
                ref_compra = self.lector.valor_referencia_ejercicio(cid, inm["codigo"], compra.year)
            cobros = self.lector.listar("cobros", cid, inmueble=inm["codigo"], ejercicio=int(ejercicio))
            historia = historias.get(inm["codigo"], [])
            desglose = [m for m in historia if m.get("ejercicio") == int(ejercicio)
                        and (m.get("estado") != "No procede" or inm["codigo"] in conflictos)]
            ejercicios_historicos = sorted({int(f["ejercicio"]) for f in historia
                if f.get("ejercicio") and 2024 <= int(f["ejercicio"]) < int(ejercicio)
                and tax_calculator.periodo_210_o_none(f.get("periodo")) in ("1T", "2T", "3T", "4T")
                and f.get("estado") == "Presentado"})
            out.append({**inm, "_vp": vp, "_tit": tit, "_computado": round(vp.valor * tit, 2), "_catastro": cat,
                        "_venta": ({"firma": tax_calculator.firma_amortizacion_venta(self.lector, cid, inm["codigo"]),
                                    "acreditada": tax_calculator.amortizacion_venta_confirmada(self.lector, cid, inm["codigo"]),
                                    "presentada": tax_calculator._presentado(self.lector.get_por_clave(
                                        "obligaciones", cid, f"VENTA|{inm['codigo']}"))}
                                   if inm.get("fecha_venta") and inm.get("precio_venta") else None),
                        **{"_" + k: v for k, v in ocupacion.items()},
                        "_ingresos_cobrados": tax_calculator.r2(sum(float(c.get("importe_cobrado") or 0) for c in cobros)),
                        "_valor_referencia_compra": ref_compra,
                        "_valor_referencia_actual": self.lector.valor_referencia_ejercicio(cid, inm["codigo"], int(ejercicio)),
                        "_anio_compra": compra.year if compra else None,
                        "_desglose_210": desglose,
                        "_ocupacion_trimestral": [
                            {"periodo": periodo, **tax_calculator.resumen_arrendamiento(
                                inm, contratos, rentas, int(ejercicio), (desde, hasta))}
                            for periodo, desde, hasta in tax_calculator.trimestres(int(ejercicio))],
                        "_revision_210": conflictos.get(inm["codigo"]),
                        "_conciliaciones_pendientes": [
                            {"ejercicio": e, "firma": tax_calculator.firma_historial_210(historia, inm["codigo"], e)}
                            for e in ejercicios_historicos if (inm["codigo"], e) not in conciliadas],
                        "_periodicidad_210": tax_calculator.periodicidad_210_alquiler(self.lector, cid, inm["codigo"], int(ejercicio))})
        return out

"""Regresiones con dos inmuebles, confirmaciones locales y archivos Excel ficticios."""
import re

import openpyxl
import pytest

import acciones_documentales as ad
import correcciones_expediente as ce
import tax_calculator as tc
from controlador import CalculoConfirmadoSinExcel, Controlador
from database import TABLAS, Database


@pytest.fixture
def caso(tmp_path):
    ctrl = Controlador(tmp_path / 'ficticio.sqlite')
    ctrl.db.carpeta_clientes = tmp_path / 'Clientes'
    cid = ctrl.db.crear_cliente('Prueba', 'Aislamiento', 'X0000077T')['id']
    ctrl.db.guardar_residencia(cid, 2025, pais='Alemania')
    inmuebles, contratos = [], []
    for _ in range(2):
        inm = ctrl.db.upsert('inmuebles', cid, {'fecha_adquisicion': '2020-01-01',
            'precio_escritura': 100000, 'pct_titularidad': 1, 'estado': 'Activo'})
        ctr = ctrl.db.upsert('contratos', cid, {'inmueble': inm, 'fecha_inicio': '2025-01-01',
            'fecha_fin_real': '2025-12-31', 'renta_inicial': 1000, 'periodicidad': 'Mensual',
            'tipo': 'Temporada', 'estado': 'Finalizado'})
        ctrl.db.upsert('catastro', cid, {'inmueble': inm, 'ejercicio': 2025,
            'valor_catastral': 90000, 'valor_suelo': 30000, 'valor_construccion': 60000,
            'pct_imputacion': .011}, clave=f'CAT|{inm}|2025')
        inmuebles.append(inm)
        contratos.append(ctr)
    try:
        yield ctrl, cid, inmuebles, contratos
    finally:
        ctrl.cerrar()


def duplicar(db, cid, inm, estado='Presentado', ejercicio=2025, claves=('OA', 'Anual')):
    return [db.upsert('m210_alquiler', cid, {'inmueble': inm, 'ejercicio': ejercicio,
        'periodo': periodo, 'estado': estado, 'base': base, 'cuota': base * .19,
        'pagado': base * .19 if estado == 'Presentado' else None},
        clave=f'210A|{ejercicio}|{inm}|{clave}')
        for periodo, clave, base in zip(('OA', 'Anual'), claves, (10000, 12000))]


@pytest.mark.parametrize('estado,claves', [('Presentado', ('OA', 'Anual')),
    ('En preparación', ('viejo1', 'viejo2')), ('En preparación', ('OA', 'Anual'))])
def test_conflicto_no_bloquea_otro_inmueble_ni_imputacion(caso, estado, claves):
    ctrl, cid, (malo, bueno), _ = caso
    contrato = ctrl.db.listar('contratos', cid, inmueble=malo)[0]
    ctrl.db.upsert('contratos', cid, {'codigo': contrato['codigo'], 'fecha_fin_real': '2025-09-30'})
    codigos = duplicar(ctrl.db, cid, malo, estado, claves=claves)
    originales = [ctrl.db.get('m210_alquiler', cid, c) for c in codigos]
    informe = tc.recalcular_cliente(ctrl.db, cid, 2025)
    assert set(informe.alquileres_bloqueados) == {malo}
    assert [ctrl.db.get('m210_alquiler', cid, c) for c in codigos] == originales
    assert not ctrl.db.listar('arrastres_gastos', cid, inmueble=malo)
    assert any(ctrl.db.get('m210_alquiler', cid, c)['inmueble'] == bueno for c in informe.filas_210_alquiler)
    assert all(ctrl.db.get('m210_alquiler', cid, c)['inmueble'] != malo for c in informe.filas_210_alquiler)
    assert any('PROVISIONAL' in e and malo in e for e in informe.estados)
    assert any(p.nivel == 'grave' and malo in p.titulo for p in informe.pendientes)
    assert ctrl.db.listar('m210_imputacion', cid, inmueble=malo)  # cálculo independiente sigue


def test_historico_ambiguo_no_reutiliza_arrastre_y_se_recupera_al_resolver(caso):
    ctrl, cid, (malo, bueno), _ = caso
    codigos = duplicar(ctrl.db, cid, malo, ejercicio=2024)
    arr = ctrl.db.upsert('arrastres_gastos', cid, {'inmueble': malo, 'ejercicio': 2024,
        'generado': 777, 'pendiente': 777, 'saldos': '{"2024":777}', 'estado': 'Calculado'},
        clave=f'ARR|{malo}|2024')
    original = ctrl.db.get('arrastres_gastos', cid, arr)
    for ejercicio in (2025, 2026):
        informe = tc.recalcular_cliente(ctrl.db, cid, ejercicio)
        assert malo in informe.alquileres_bloqueados
        assert '2024' in informe.alquileres_bloqueados[malo]
        assert not ctrl.db.listar('arrastres_gastos', cid, inmueble=malo, ejercicio=ejercicio)
        assert ctrl.db.get('arrastres_gastos', cid, arr) == original
    ctrl.db.borrar('m210_alquiler', cid, codigos[1])  # conciliación explícita simulada
    informe = tc.recalcular_cliente(ctrl.db, cid, 2025)
    assert not informe.alquileres_bloqueados
    assert ctrl.db.listar('arrastres_gastos', cid, inmueble=malo, ejercicio=2025)


def test_errores_inesperados_no_se_ocultan_como_conflicto(caso, monkeypatch):
    ctrl, cid, _, _ = caso
    monkeypatch.setattr(tc, 'liquidar_210_alquiler', lambda *a, **k: (_ for _ in ()).throw(ValueError('fallo inesperado')))
    with pytest.raises(ValueError, match='fallo inesperado'):
        tc.recalcular_cliente(ctrl.db, cid, 2025)


def propuesta(campo='fecha_fin_real', valor='2025-09-30', tabla='contratos', codigo='CTR-001'):
    return {'accion': 'editar_registro', 'tabla': tabla, 'codigo': codigo, 'campo': campo, 'valor': valor}


def confirmar(ctrl, cid, plan):
    gestor = ad.GestorAcciones(lambda: (cid, 2025))
    mensaje = gestor.proponer(plan, cid, 2025, lambda c, p: ctrl.preparar_correccion_expediente(c, 2025, p))
    token = re.search(r'/confirmar_cambio ([a-f0-9]{24})', mensaje)[1]
    texto = gestor.procesar('/confirmar_cambio ' + token, cid, 2025, ctrl.ejecutar_correccion_expediente)
    return gestor, token, texto


def valor_excel(ruta, tabla, codigo, campo):
    wb = openpyxl.load_workbook(ruta)
    try:
        t = TABLAS[tabla]
        col = next(c for c in t.columnas_excel if c.campo == campo)
        cod = next(c for c in t.columnas_excel if c.campo == 'codigo')
        for fila in range(t.fila_inicio, wb[t.hoja].max_row + 1):
            if wb[t.hoja][f'{cod.letra}{fila}'].value == codigo:
                return wb[t.hoja][f'{col.letra}{fila}'].value
        pytest.fail('No se exportó el registro')
    finally:
        wb.close()


def test_chat_corrige_ultimo_trimestre_y_excel_persiste_al_reabrir(caso):
    ctrl, cid, (inm, _), (ctr, _) = caso
    gestor, token, texto = confirmar(ctrl, cid, propuesta(codigo=ctr))
    assert 'Excel actualizado' in texto
    fila = ctrl.db.listar('m210_alquiler', cid, inmueble=inm, ejercicio=2025)[0]
    assert fila['ingresos'] == 9000
    assert ctrl.db.listar('m210_imputacion', cid, inmueble=inm, ejercicio=2025)[0]['dias'] == 92
    ruta = ctrl.db.ruta_excel(ctrl.db.get_cliente(cid))
    assert valor_excel(ruta, 'm210_alquiler', fila['codigo'], 'ingresos') == 9000
    assert valor_excel(ruta, 'contratos', ctr, 'fecha_fin_real').date().isoformat() == '2025-09-30'
    ctrl.sincronizar(cid, 2025)
    assert ctrl.db.get('contratos', cid, ctr)['fecha_fin_real'] == '2025-09-30'
    ruta_bd = ctrl.db.db_path
    ctrl.cerrar()
    nuevo = Controlador(ruta_bd)
    try:
        nuevo.abrir_cliente(cid, 2025)
        assert nuevo.db.get('contratos', cid, ctr)['fecha_fin_real'] == '2025-09-30'
        assert nuevo.db.get('m210_alquiler', cid, fila['codigo'])['ingresos'] == 9000
    finally:
        nuevo.cerrar()
    with pytest.raises(ValueError, match='pendiente'):
        gestor.procesar('/confirmar_cambio ' + token, cid, 2025, ctrl.ejecutar_correccion_expediente)


@pytest.mark.parametrize('tabla,campo,valor', [('m210_alquiler', 'cuota', '0'), ('contratos', 'ruta', '../x'),
    ('contratos', 'inmueble', 'INM-002'), ('documentos', 'ruta', '../x'), ('inmuebles', 'manual_override', '1')])
def test_chat_rechaza_calculos_rutas_y_enlaces(tabla, campo, valor):
    with pytest.raises(ValueError):
        ce.validar_plan(propuesta(campo, valor, tabla))


def test_confirmacion_rechaza_datos_cambiados_y_fecha_imposible(caso):
    ctrl, cid, _, (ctr, _) = caso
    with pytest.raises(ValueError):
        ctrl.preparar_correccion_expediente(cid, 2025, propuesta(valor='2019-01-01', codigo=ctr))
    gestor = ad.GestorAcciones(lambda: (cid, 2025))
    texto = gestor.proponer(propuesta(codigo=ctr), cid, 2025,
        lambda c, p: ctrl.preparar_correccion_expediente(c, 2025, p))
    token = re.search(r'/confirmar_cambio ([a-f0-9]{24})', texto)[1]
    ctrl.db.upsert('contratos', cid, {'codigo': ctr, 'renta_inicial': 1500})
    with pytest.raises(ValueError, match='cambió'):
        gestor.procesar('/confirmar_cambio ' + token, cid, 2025, ctrl.ejecutar_correccion_expediente)
    assert ctrl.db.get('contratos', cid, ctr)['fecha_fin_real'] == '2025-12-31'


def test_error_exportacion_conserva_dato_y_no_afirma_excel_actualizado(caso, monkeypatch):
    ctrl, cid, _, (ctr, _) = caso
    import excel_generator
    monkeypatch.setattr(excel_generator, 'exportar_cliente', lambda *a, **k: (_ for _ in ()).throw(PermissionError()))
    _, _, texto = confirmar(ctrl, cid, propuesta(codigo=ctr))
    assert 'Excel actualizado' not in texto and 'no se pudo' in texto
    assert ctrl.db.get('contratos', cid, ctr)['fecha_fin_real'] == '2025-09-30'
    assert ctrl.db.get_config(f'excel_pendiente|{cid}')


@pytest.mark.parametrize('tabla,campo,valor,esperado', [
    ('inmuebles', 'precio_escritura', '123.456,78', 123456.78),
    ('contratos', 'renta_inicial', '800', 800),
    ('catastro', 'valor_catastral', '85000', 85000),
    ('gastos', 'total', '222,50', 222.5),
    ('rentas', 'renta', '875', 875),
    ('cobros', 'importe_cobrado', '999,99', 999.99),
])
def test_correcciones_de_origen_llegan_a_excel_y_no_revierten(caso, tabla, campo, valor, esperado):
    ctrl, cid, (inm, _), (ctr, _) = caso
    if tabla == 'gastos':
        codigo = ctrl.db.upsert(tabla, cid, {'inmueble': inm, 'ejercicio': 2025,
            'fecha': '2025-03-01', 'categoria': 'IBI', 'total': 100, 'pct_inmueble': 1, 'pct_alquiler': 1})
    elif tabla == 'rentas':
        codigo = ctrl.db.upsert(tabla, cid, {'inmueble': inm, 'contrato': ctr, 'desde': '2025-01-01',
            'hasta': '2025-12-31', 'renta': 700, 'periodicidad': 'Mensual'})
    elif tabla == 'cobros':
        tc.recalcular_cliente(ctrl.db, cid, 2025)
        codigo = ctrl.db.listar(tabla, cid, inmueble=inm)[0]['codigo']
    else:
        codigo = inm if tabla == 'inmuebles' else ctr if tabla == 'contratos' else ctrl.db.catastro(cid, inm, 2025)['codigo']
    resultado = ctrl.editar(cid, tabla, codigo, {campo: valor}, 2025)
    assert valor_excel(resultado['excel'], tabla, codigo, campo) == esperado
    ctrl.sincronizar(cid, 2025)
    assert ctrl.db.get(tabla, cid, codigo)[campo] == esperado


def test_calculo_no_admite_cambio_silenciosamente_descartado(caso):
    ctrl, cid, _, _ = caso
    informe = tc.recalcular_cliente(ctrl.db, cid, 2025)
    codigo = informe.filas_210_alquiler[0]
    anterior = ctrl.db.get('m210_alquiler', cid, codigo)
    with pytest.raises(ValueError, match='datos de origen'):
        ctrl.editar(cid, 'm210_alquiler', codigo, {'base': '0'}, 2025)
    assert ctrl.db.get('m210_alquiler', cid, codigo)['base'] == anterior['base']


def test_chat_no_cambia_importes_de_presentadas(caso):
    ctrl, cid, (inm, _), (ctr, _) = caso
    tc.recalcular_cliente(ctrl.db, cid, 2025)
    fila = ctrl.db.listar('m210_alquiler', cid, inmueble=inm)[0]
    ctrl.db.upsert('m210_alquiler', cid, {'codigo': fila['codigo'], 'estado': 'Presentado', 'pagado': 2280})
    confirmar(ctrl, cid, propuesta(codigo=ctr))
    actual = ctrl.db.get('m210_alquiler', cid, fila['codigo'])
    assert (actual['base'], actual['cuota'], actual['pagado'], actual['estado']) == (fila['base'], fila['cuota'], 2280, 'Presentado')


def test_residente_corrige_factura_por_chat_y_excel_conserva_formulas(caso):
    ctrl, _, _, _ = caso
    cid = ctrl.db.crear_cliente('Ficticio', 'Residente', '00000000T', tipo_cliente='Autónomo')['id']
    codigo = ctrl.db.upsert('facturas_recibidas', cid, {'fecha': '2025-03-01', 'base': 100,
        'cuota_iva': 21, 'tipo_iva': .21, 'pct_iva_deducible': 1, 'pct_deducible_irpf': 1})
    _, _, texto = confirmar(ctrl, cid, propuesta('base', '200', 'facturas_recibidas', codigo))
    assert 'Excel actualizado' in texto
    ruta = ctrl.db.ruta_excel(ctrl.db.get_cliente(cid))
    assert valor_excel(ruta, 'facturas_recibidas', codigo, 'base') == 200
    ctrl.sincronizar(cid, 2025)
    assert ctrl.db.get('facturas_recibidas', cid, codigo)['base'] == 200


def test_excel_y_tarjeta_advierten_importes_historicos_no_actualizados(caso):
    import excel_generator
    import ui_cartera as uc
    import ui_estilo as ui
    from test_ui_cartera import textos
    ctrl, cid, (malo, _), _ = caso
    duplicar(ctrl.db, cid, malo)
    ctrl._recalcular_bd(cid, {2025}, 2025)
    vista = next(v for v in ctrl.vista_inmuebles(cid, 2025) if v['codigo'] == malo)
    assert len(uc.filas_desglose_210(vista, 'Anual')) == 2  # también muestra el histórico OA
    assert 'PENDIENTE DE REVISIÓN' in textos(uc.desglose_210(vista, 'Anual', 2025, ui.CLARO)[0])
    ruta = excel_generator.exportar_cliente(ctrl.db, cid, 2025)
    wb = openpyxl.load_workbook(ruta)
    try:
        assert 'PENDIENTE DE REVISIÓN' in wb['00_RESUMEN']['H7'].comment.text
        assert 'PENDIENTE DE REVISIÓN' in wb['00_RESUMEN']['E13'].comment.text
    finally:
        wb.close()


def test_confirmacion_se_invalida_por_cambio_de_cliente_y_caducidad():
    activo, reloj = [7, 2025], [0]
    gestor = ad.GestorAcciones(lambda: tuple(activo), reloj=lambda: reloj[0])
    plan = propuesta()
    preparar = lambda *a: {'version': 'v1', 'antes': '2025-12-31', 'despues': '2025-09-30'}
    for cambiar in ('cliente', 'caducidad'):
        activo[:], reloj[0] = [7, 2025], 0
        texto = gestor.proponer(plan, 7, 2025, preparar)
        token = re.search(r'/confirmar_cambio ([a-f0-9]{24})', texto)[1]
        if cambiar == 'cliente':
            activo[0] = 8
        else:
            reloj[0] = 121
        with pytest.raises(ValueError, match='caducó|cambió'):
            gestor.procesar('/confirmar_cambio ' + token, 7, 2025, lambda *a: pytest.fail('No debe escribir'))


def test_edicion_y_recalculo_son_atomicos(caso, monkeypatch):
    ctrl, cid, (inm, _), _ = caso
    monkeypatch.setattr(ctrl, '_recalcular_bd', lambda *a: (_ for _ in ()).throw(ValueError('Fallo de cálculo')))
    with pytest.raises(ValueError, match='Fallo de cálculo'):
        ctrl.editar(cid, 'inmuebles', inm, {'precio_escritura': '250000'}, 2025)
    assert ctrl.db.get('inmuebles', cid, inm)['precio_escritura'] == 100000


def test_segunda_importacion_no_aplica_una_propuesta_obsoleta(caso, monkeypatch):
    ctrl, cid, _, (ctr, _) = caso
    original, llamadas = ctrl._importar, [0]
    def importar(cliente):
        llamadas[0] += 1
        original(cliente)
        if llamadas[0] == 3:
            ctrl.db.upsert('contratos', cid, {'codigo': ctr, 'renta_inicial': 1500})
    monkeypatch.setattr(ctrl, '_importar', importar)
    with pytest.raises(ValueError, match='cambió'):
        confirmar(ctrl, cid, propuesta(codigo=ctr))
    actual = ctrl.db.get('contratos', cid, ctr)
    assert actual['fecha_fin_real'] == '2025-12-31' and actual['renta_inicial'] == 1500


@pytest.mark.parametrize('periodos', [('Anual', '1T'), ('Enero',)])
def test_historico_invalido_se_aisla_sin_ocultar_otros_errores(caso, periodos):
    ctrl, cid, (malo, bueno), _ = caso
    for i, periodo in enumerate(periodos):
        ctrl.db.upsert('m210_alquiler', cid, {'inmueble': malo, 'ejercicio': 2024,
            'periodo': periodo, 'estado': 'Presentado', 'base': 1000, 'cuota': 190}, clave=f'viejo-{i}')
    informe = tc.recalcular_cliente(ctrl.db, cid, 2025)
    assert malo in informe.alquileres_bloqueados
    assert any(ctrl.db.get('m210_alquiler', cid, c)['inmueble'] == bueno for c in informe.filas_210_alquiler)
    assert not ctrl.db.listar('arrastres_gastos', cid, inmueble=malo, ejercicio=2025)


def test_alta_con_excel_fallido_informa_codigo_y_no_invita_a_repetir(caso, monkeypatch):
    import excel_generator
    ctrl, cid, (inm, _), _ = caso
    monkeypatch.setattr(excel_generator, 'exportar_cliente', lambda *a, **k: (_ for _ in ()).throw(PermissionError()))
    with pytest.raises(CalculoConfirmadoSinExcel) as error:
        ctrl.editar(cid, 'gastos', None, {'inmueble': inm, 'ejercicio': 2025,
            'fecha': '2025-05-01', 'categoria': 'IBI', 'total': 100}, 2025)
    codigo = error.value.resultado['codigo']
    assert codigo in str(error.value)
    assert len(ctrl.db.listar('gastos', cid)) == 1
    assert ctrl.db.get('gastos', cid, codigo)['total'] == 100

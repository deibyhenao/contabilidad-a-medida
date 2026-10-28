"""Regresiones de error humano con BD, PDFs y Excel ficticios; API siempre simulada."""
import json
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

import asistente
import documentos
import fixtures_pdf as f
import ia_extractor
import pdf_processor as pp
from controlador import Controlador, CanceladoPorUsuario


@pytest.fixture
def entorno(tmp_path):
    ctrl = Controlador(tmp_path / 'aislada.sqlite')
    ctrl.cambiar_carpeta_clientes(tmp_path / 'Clientes')
    cliente = ctrl.crear_cliente('Cliente', 'Ficticio', 'X1111111A')
    yield ctrl, cliente['id'], tmp_path
    ctrl.cerrar()


def importar_certificado(entorno, nombre='certificado.pdf', pais='Francia', ejercicio=2025):
    ctrl, cid, tmp = entorno
    ruta = f.certificado_residencia(tmp / nombre, pais=pais, ejercicio=ejercicio)
    resultado = ctrl.procesar_pdf(cid, ruta, 2026)['resultado']
    return resultado.documento, ruta


def test_nif_duplicado_no_modifica_pais_ni_identidad(entorno):
    ctrl, cid, _ = entorno
    ctrl.db.guardar_residencia(cid, 2025, pais='Francia')
    with pytest.raises(ValueError, match='Ya existe'):
        ctrl.crear_cliente('Otro', 'Erroneo', 'x1111111a', pais='Rusia', ejercicio=2025)
    assert ctrl.db.get_cliente(cid)['nombre'] == 'Cliente'
    assert ctrl.db.residencia(cid, 2025)['pais'] == 'Francia'
    assert ctrl.db.contar_clientes() == 1


def test_prevision_existente_encadena_saldos_sin_crear_anios_por_etiqueta(entorno):
    ctrl, _, _ = entorno
    cid = ctrl.crear_cliente('Residente', 'Previsión', 'RES-PREV', tipo_cliente='Autónomo', ejercicio=2025,
                            regimen_iva='General', regimen_irpf='Directa normal')['id']
    ctrl.db.upsert('facturas_recibidas', cid, {'fecha':'2025-02-10','base':2000,'cuota_iva':420})
    ctrl.db.upsert('m303', cid, {'ejercicio':2028,'periodo':'1T','estado':'Pendiente'}, clave='303|2028|1T')
    ctrl.db.upsert('documentos', cid, {'ejercicio':2029,'tipo':'Documento sin modelo'})
    ctrl._recalcular_bd(cid, {2025,2029}, 2025)
    assert ctrl.db.get_por_clave('m303', cid, '303|2028|1T')['compensacion_anterior'] == 420
    assert ctrl.db.get_por_clave('m303', cid, '303|2026|4T')['a_compensar'] == 420
    assert ctrl.db.get_por_clave('m303', cid, '303|2027|4T')['a_compensar'] == 420
    assert not ctrl.db.listar('m303', cid, ejercicio=2029)


@pytest.mark.parametrize('presentado', [False, True])
def test_fecha_manual_otro_anio_se_recalcula_o_bloquea_si_presentado(entorno, presentado):
    ctrl, _, tmp = entorno
    cid = ctrl.crear_cliente('Residente', 'Fecha manual', 'RES-FECHA', tipo_cliente='Autónomo', ejercicio=2025,
                            regimen_iva='General', regimen_irpf='Directa normal')['id']
    ruta = f._crear(tmp/'fecha_manual.pdf', 'FACTURA FICTICIA\nDocumento de prueba de cambio de año')
    ext = pp.Extraccion(tipo=pp.TipoDocumento.FACTURA_GASTO, confianza=1, nombre_archivo=ruta.name,
        campos={'fecha':'2025-02-10','base':100,'iva':21,'total':121,'n_factura':'FECHA-1'})
    res = pp.aplicar(ctrl.db, cid, ext, ruta)
    ctrl.sincronizar(cid, 2025)
    ctrl.editar(cid, 'facturas_recibidas', res.codigos['factura'], {'fecha':'2027-04-10'}, 2027)
    modelo = ctrl.db.get_por_clave('m303', cid, '303|2027|2T')
    assert modelo['base_deducible'] == 100
    if presentado:
        ctrl.db.upsert('m303', cid, {'codigo':modelo['codigo'],'estado':'Presentado'})
        vista = ctrl.previsualizar_documento(cid, res.documento, 'retirar')
        assert not vista['confirmable']
        with pytest.raises(ValueError, match='presentadas'):
            ctrl.retirar_documento(cid, res.documento, 2025)
        assert not ctrl.obtener_documento(cid, res.documento)['en_papelera']
    else:
        ctrl.retirar_documento(cid, res.documento, 2025)
        assert ctrl.db.get_por_clave('m303', cid, '303|2027|2T')['base_deducible'] == 0
        ctrl.restaurar_documento(cid, res.documento, 2025)
        assert ctrl.db.get_por_clave('m303', cid, '303|2027|2T')['base_deducible'] == 100


@pytest.mark.parametrize('accion', ['ejercicio', 'sustituir'])
@pytest.mark.parametrize('legado', [False, True])
def test_correcciones_manuales_factura_se_conservan_antes_recalcular(entorno, accion, legado):
    ctrl, _, tmp = entorno
    cid = ctrl.crear_cliente('Residente', 'Ficticio', 'RES-DOC', tipo_cliente='Autónomo', ejercicio=2025,
                            regimen_iva='General', regimen_irpf='Directa normal')['id']
    ruta = f._crear(tmp / 'factura_original.pdf', 'FACTURA DE GASTO FICTICIA\nNúmero de factura: DOC-FACT-1\nFecha factura: 10/02/2025\nBase imponible: 100,00 euros\nIVA: 21,00 euros\nTotal: 121,00 euros')
    campos = {'fecha': '2025-02-10', 'base': 100, 'iva': 21, 'total': 121, 'n_factura': 'DOC-FACT-1'}
    ext = pp.Extraccion(tipo=pp.TipoDocumento.FACTURA_GASTO, confianza=1, campos=campos, nombre_archivo=ruta.name)
    res = pp.aplicar(ctrl.db, cid, ext, ruta)
    codigo = res.codigos['factura']
    ctrl.sincronizar(cid, 2025)
    ctrl.editar(cid, 'facturas_recibidas', codigo, {'base': 77, 'fecha': '2025-04-10'}, 2025)
    if legado:
        ctrl.db.con.execute('DELETE FROM documentos_cambios WHERE cliente_id=?', (cid,))
        ctrl.db.con.execute('DELETE FROM documentos_gestion WHERE cliente_id=?', (cid,))
        ctrl.db.con.commit()
    if accion == 'ejercicio':
        ctrl.corregir_ejercicio_documento(cid, res.documento, 2026, 2026)
        ejercicio = 2026
        assert ctrl.db.get('documentos', cid, res.documento)['fecha_documento'] == '2026-04-10'
    else:
        nueva = f._crear(tmp / 'factura_nueva.pdf', 'FACTURA FICTICIA\nVersión revisada')
        ctrl.sustituir_documento(cid, res.documento, nueva, 2025, confirmar=lambda e: ext)
        ejercicio = 2025
    factura = ctrl.db.get('facturas_recibidas', cid, codigo)
    assert len(ctrl.db.listar('facturas_recibidas', cid)) == 1
    assert factura['base'] == 77
    assert factura['fecha'] == f'{ejercicio}-04-10' and factura['ejercicio'] == ejercicio
    modelo = ctrl.db.get_por_clave('m303', cid, f'303|{ejercicio}|2T')
    assert modelo['base_deducible'] == 77
    if accion == 'ejercicio':
        ctrl.retirar_documento(cid, res.documento, ejercicio)
        factura = ctrl.db.get('facturas_recibidas', cid, codigo)
        assert factura['base'] == 77 and factura['fecha'] == '2026-04-10' and factura['ejercicio'] == 2026
        ctrl.restaurar_documento(cid, res.documento, ejercicio)
        assert ctrl.db.get('facturas_recibidas', cid, codigo)['fecha'] == '2026-04-10'


def test_retira_exporta_sin_reactivar_y_restaura_pdf_y_residencia(entorno):
    ctrl, cid, _ = entorno
    codigo, _ = importar_certificado(entorno)
    doc = ctrl.db.get('documentos', cid, codigo)
    carpeta = Path(ctrl.db.get_cliente(cid)['carpeta'])
    original = carpeta / doc['ruta']
    contenido = original.read_bytes()
    vista = ctrl.previsualizar_documento(cid, codigo, 'retirar')
    resultado = ctrl.retirar_documento(cid, codigo, 2026, version_esperada=vista['version'])
    assert resultado['documentos_actualizados'] and not resultado['avisos_documentales']
    assert not original.exists() and not ctrl.db.residencia(cid, 2025)
    papelera = ctrl.listar_documentos(cid, papelera=True)
    assert any(d['codigo'] == codigo for d in papelera)
    wb = openpyxl.load_workbook(resultado['excel'])
    hoja = wb['13_DOCUMENTACION']
    fila = next(r for r in range(2, hoja.max_row + 1) if hoja.cell(r, 1).value == codigo)
    estados = [hoja.cell(fila, col).value for col in range(1, hoja.max_column + 1)]
    assert 'En papelera' in estados
    assert any(c.hyperlink and '.Papelera_Documentos/' in c.hyperlink.target for c in hoja[fila])
    wb['00_RESUMEN']['B30'] = 'Edición manual ficticia'
    wb.save(resultado['excel'])
    wb.close()
    ctrl.sincronizar(cid, 2026)
    assert ctrl.obtener_documento(cid, codigo)['en_papelera']
    assert ctrl.db.get('documentos', cid, codigo)['ruta'] == doc['ruta']
    # Excel conserva una fila vacía identificada por año; no debe recuperar sus datos.
    residencia = ctrl.db.residencia(cid, 2025) or {}
    assert not any(residencia.get(k) for k in ('pais', 'certificado', 'grupo', 'ruta'))
    ctrl.restaurar_documento(cid, codigo, 2026)
    assert original.read_bytes() == contenido
    assert ctrl.db.residencia(cid, 2025)['pais'] == 'Francia'


def test_correccion_traslada_residencia_y_no_resucita_soporte_anterior(entorno):
    ctrl, cid, _ = entorno
    a, _ = importar_certificado(entorno)
    b, _ = importar_certificado(entorno, 'rusia.pdf', 'Rusia')
    ctrl.corregir_ejercicio_documento(cid, a, 2026, 2026)
    ctrl.retirar_documento(cid, b, 2026)
    assert not ctrl.db.residencia(cid, 2025)
    assert ctrl.db.residencia(cid, 2026)['pais'] == 'Francia'


@pytest.mark.parametrize('fallo', ['cancelacion', 'extraccion', 'recalculo'])
def test_sustitucion_fallida_conserva_original_y_no_deja_nuevo_soporte(entorno, monkeypatch, fallo):
    ctrl, cid, _ = entorno
    codigo, _ = importar_certificado(entorno)
    carpeta = Path(ctrl.db.get_cliente(cid)['carpeta'])
    anteriores = {p.name: p.read_bytes() for p in (carpeta / 'Soportes_PDF').glob('*.pdf')}
    nueva = f.certificado_residencia(entorno[2] / 'nuevo.pdf', pais='Rusia', ejercicio=2025)
    def lanzar(*_args, **_kwargs):
        raise RuntimeError('fallo simulado')
    confirmar = None
    error = RuntimeError
    if fallo == 'cancelacion':
        def confirmar(_ext):
            raise CanceladoPorUsuario()
        error = CanceladoPorUsuario
    elif fallo == 'extraccion':
        monkeypatch.setattr(ctrl, 'extraer', lanzar)
    else:
        monkeypatch.setattr(ctrl, '_recalcular_bd', lanzar)
    with pytest.raises(error):
        ctrl.sustituir_documento(cid, codigo, nueva, 2026, confirmar=confirmar)
    assert not ctrl.obtener_documento(cid, codigo)['en_papelera']
    assert ctrl.db.residencia(cid, 2025)['pais'] == 'Francia'
    assert {p.name: p.read_bytes() for p in (carpeta / 'Soportes_PDF').glob('*.pdf')} == anteriores
    assert not ctrl.db.con.execute('SELECT 1 FROM documentos_archivos').fetchone()
    assert not list((carpeta / '.Papelera_Documentos').glob('*.pdf'))


def test_pdf_cambiado_justo_antes_copia_no_sustituye_ni_deja_copias(entorno, monkeypatch):
    ctrl, cid, tmp = entorno
    codigo, _ = importar_certificado(entorno)
    ruta = f.certificado_residencia(tmp / 'mutable_tarde.pdf', pais='Rusia', ejercicio=2025)
    archivar_original = pp.archivar_pdf
    def cambiar_y_archivar(*args, **kwargs):
        ruta.write_bytes(b'%PDF-cambiado-despues-de-revisar')
        return archivar_original(*args, **kwargs)
    monkeypatch.setattr(pp, 'archivar_pdf', cambiar_y_archivar)
    with pytest.raises(ValueError, match='cambi'):
        ctrl.sustituir_documento(cid, codigo, ruta, 2026)
    assert not ctrl.obtener_documento(cid, codigo)['en_papelera']
    assert ctrl.db.residencia(cid, 2025)['pais'] == 'Francia'
    carpeta = Path(ctrl.db.get_cliente(cid)['carpeta'])
    assert len(list((carpeta / 'Soportes_PDF').glob('*.pdf'))) == 1
    assert not list((carpeta / '.Papelera_Documentos').glob('*.pdf'))


def test_sustituye_solo_tras_revision_y_conserva_anterior_recuperable(entorno):
    ctrl, cid, tmp = entorno
    codigo, _ = importar_certificado(entorno)
    ruta = f.certificado_residencia(tmp / 'correcto.pdf', pais='Alemania', ejercicio=2025)
    revisiones = []
    def revisar(ext):
        assert not ctrl.obtener_documento(cid, codigo)['en_papelera']
        revisiones.append(ext.tipo)
        return ext
    res = ctrl.sustituir_documento(cid, codigo, ruta, 2026, confirmar=revisar)
    assert revisiones == [pp.TipoDocumento.CERTIFICADO_RESIDENCIA]
    assert ctrl.obtener_documento(cid, codigo)['en_papelera']
    assert not ctrl.obtener_documento(cid, res['resultado'].documento)['en_papelera']
    assert ctrl.db.residencia(cid, 2025)['pais'] == 'Alemania'
    ctrl.restaurar_documento(cid, codigo, 2026)
    assert ctrl.db.residencia(cid, 2025)['pais'] == 'Alemania'


def test_confirmacion_obsoleta_detecta_edicion_manual_excel(entorno):
    ctrl, cid, _ = entorno
    codigo, _ = importar_certificado(entorno)
    vista = ctrl.previsualizar_documento(cid, codigo, 'retirar')
    ruta = ctrl.db.ruta_excel(ctrl.db.get_cliente(cid))
    wb = openpyxl.load_workbook(ruta)
    wb['00_RESUMEN']['B30'] = 'Cambio posterior a la previsualización'
    wb.save(ruta)
    wb.close()
    with pytest.raises(ValueError, match='han cambiado'):
        ctrl.retirar_documento(cid, codigo, 2026, version_esperada=vista['version'])
    assert not ctrl.obtener_documento(cid, codigo)['en_papelera']


@pytest.mark.parametrize('sustitucion', [False, True])
def test_pdf_cambiado_durante_revision_no_se_archiva(entorno, sustitucion):
    ctrl, cid, tmp = entorno
    codigo, _ = importar_certificado(entorno)
    ruta = f.certificado_residencia(tmp / 'mutable.pdf', pais='Rusia', ejercicio=2025)
    def revisar(ext):
        ruta.write_bytes(b'%PDF-contenido-distinto')
        return ext
    with pytest.raises(ValueError, match='cambi'):
        if sustitucion:
            ctrl.sustituir_documento(cid, codigo, ruta, 2026, confirmar=revisar)
        else:
            ctrl.procesar_pdf(cid, ruta, 2026, confirmar=revisar)
    assert not ctrl.obtener_documento(cid, codigo)['en_papelera']
    assert ctrl.db.residencia(cid, 2025)['pais'] == 'Francia'


@pytest.mark.parametrize('caso', ['otro_cliente', 'ejercicio_contradictorio'])
def test_factura_incorrecta_rechazada_antes_de_archivar(entorno, caso):
    ctrl, cid, tmp = entorno
    ruta = f._crear(tmp / 'factura.pdf', 'FACTURA FICTICIA')
    campos = {'fecha': '2025-05-10', 'base': 100, 'iva': 21, 'total': 121}
    if caso == 'otro_cliente':
        campos.update(nif_proveedor='B12345678', nif_receptor='X2222222B')
    else:
        campos['ejercicio'] = 2026
    ext = pp.Extraccion(tipo=pp.TipoDocumento.FACTURA_GASTO, confianza=1, campos=campos, nombre_archivo=ruta.name)
    with pytest.raises(ValueError, match='NIF|ejercicio'):
        pp.aplicar(ctrl.db, cid, ext, ruta, ejercicio=2026)
    carpeta = Path(ctrl.db.get_cliente(cid)['carpeta']) / 'Soportes_PDF'
    assert not list(carpeta.glob('*.pdf'))


def test_ia_excluye_pdf_retirado_incluso_con_original_pendiente_y_limpia_historial(entorno, monkeypatch):
    ctrl, cid, _ = entorno
    codigo, _ = importar_certificado(entorno)
    llamadas = []
    monkeypatch.setattr(ia_extractor, 'hay_clave', lambda: True)
    def respuesta(ctx, mensajes, *_args):
        llamadas.append(json.dumps({'ctx': ctx, 'mensajes': mensajes}, default=str))
        return SimpleNamespace(model='modelo-simulado', usage=None, stop_reason='end_turn',
                               content=[SimpleNamespace(type='text', text='Respuesta ficticia')])
    monkeypatch.setattr(asistente, '_llamar', respuesta)
    conv = asistente.Conversacion(cid)
    antes, _ = asistente.construir_contexto(ctrl.db, cid, 2026)
    conv.preguntar('Analiza el certificado', antes, ia_extractor.AjustesIA())
    assert 'CERTIFICADO' in llamadas[-1]
    # Simular commit seguido de caída antes del movimiento definitivo del archivo.
    with ctrl.db.transaccion():
        documentos.retirar(ctrl.db, cid, codigo)
    doc = ctrl.db.get('documentos', cid, codigo)
    assert (Path(ctrl.db.get_cliente(cid)['carpeta']) / doc['ruta']).exists()
    despues, _ = asistente.construir_contexto(ctrl.db, cid, 2026)
    assert doc['ruta'] not in despues and 'CERTIFICADO DE RESIDENCIA FISCAL' not in despues
    conv.preguntar('Revisa los soportes activos', despues, ia_extractor.AjustesIA())
    assert 'Analiza el certificado' not in llamadas[-1]
    assert 'CERTIFICADO DE RESIDENCIA FISCAL' not in llamadas[-1]
    assert len(conv.turnos) == 2


@pytest.mark.parametrize('accion', ['retirar', 'restaurar', 'ejercicio'])
def test_excel_antiguo_no_revierte_correccion_y_permite_conciliar(entorno, monkeypatch, accion):
    import excel_generator as eg
    ctrl, cid, _ = entorno
    codigo, _ = importar_certificado(entorno)
    ruta = ctrl.db.ruta_excel(ctrl.db.get_cliente(cid))
    if accion == 'restaurar':
        ctrl.retirar_documento(cid, codigo, 2026)
    guardar = eg._guardar_atomico
    pendiente = ruta.with_stem(ruta.stem + '_PENDIENTE_PRUEBA')
    def bloqueado(wb, destino, carpeta, **opciones):
        wb.save(pendiente)
        raise eg.ExcelBloqueadoError(pendiente)
    monkeypatch.setattr(eg, '_guardar_atomico', bloqueado)
    if accion == 'ejercicio':
        res = ctrl.corregir_ejercicio_documento(cid, codigo, 2026, 2026)
    else:
        res = getattr(ctrl, accion + '_documento')(cid, codigo, 2026)
    assert res['error_excel'] and pendiente.exists()
    assert ctrl.db.get_config(f'excel_pendiente|{cid}')
    monkeypatch.setattr(eg, '_guardar_atomico', guardar)
    residencias = ctrl.db.listar('residencias', cid)
    documento = ctrl.obtener_documento(cid, codigo)
    celda_email = eg.database.CELDAS_RESUMEN[eg.NR]['email']
    wb = openpyxl.load_workbook(ruta)
    wb['00_RESUMEN'][celda_email] = 'manual@example.invalid'
    wb.save(ruta)
    wb.close()
    original_manual = ruta.read_bytes()
    with pytest.raises(eg.ExcelConflictoImportacion, match='No se ha importado'):
        ctrl.sincronizar(cid, 2026)
    assert ctrl.db.listar('residencias', cid) == residencias
    assert ctrl.obtener_documento(cid, codigo) == documento
    assert ctrl.db.get_cliente(cid)['email'] != 'manual@example.invalid'
    assert ruta.read_bytes() == original_manual
    # Conciliación indicada en el mensaje: trasladar la edición al libro actualizado.
    wb = openpyxl.load_workbook(pendiente)
    wb['00_RESUMEN'][celda_email] = 'manual@example.invalid'
    wb.save(ruta)
    wb.close()
    assert not ctrl.sincronizar(cid, 2026)['error_excel']
    assert not ctrl.db.get_config(f'excel_pendiente|{cid}')
    assert ctrl.db.get_cliente(cid)['email'] == 'manual@example.invalid'
    assert ctrl.obtener_documento(cid, codigo)['en_papelera'] == (accion == 'retirar')
    residencia = ctrl.db.residencia(cid, 2026 if accion == 'ejercicio' else 2025) or {}
    assert residencia.get('pais') == (None if accion == 'retirar' else 'Francia')


def test_excel_bloqueado_sin_edicion_reintenta_exportar(entorno, monkeypatch):
    import excel_generator as eg
    ctrl, cid, _ = entorno
    codigo, _ = importar_certificado(entorno)
    ruta = ctrl.db.ruta_excel(ctrl.db.get_cliente(cid))
    guardar = eg._guardar_atomico
    def bloqueado(wb, destino, carpeta, **opciones):
        raise eg.ExcelBloqueadoError(destino.with_stem(destino.stem + '_PENDIENTE'))
    monkeypatch.setattr(eg, '_guardar_atomico', bloqueado)
    assert ctrl.retirar_documento(cid, codigo, 2026)['error_excel']
    monkeypatch.setattr(eg, '_guardar_atomico', guardar)
    assert not ctrl.sincronizar(cid, 2026)['error_excel']
    assert not (ctrl.db.residencia(cid, 2025) or {}).get('pais')
    assert not ctrl.db.get_config(f'excel_pendiente|{cid}')


def test_edicion_manual_excel_normal_sigue_permitida(entorno):
    import excel_generator as eg
    ctrl, cid, _ = entorno
    importar_certificado(entorno)
    ruta = ctrl.db.ruta_excel(ctrl.db.get_cliente(cid))
    assert not ctrl.db.get_config(f'excel_pendiente|{cid}')
    wb = openpyxl.load_workbook(ruta)
    ws = wb['01_RESIDENCIA_FISCAL']
    # Usar el mapa de columnas real evita depender de su orden.
    pais = next(c for c in eg.database.TABLAS['residencias'].columnas_excel if c.campo == 'pais')
    ws[f'{pais.letra}2'] = 'Alemania'
    wb.save(ruta)
    wb.close()
    ctrl.sincronizar(cid, 2026)
    assert ctrl.db.residencia(cid, 2025)['pais'] == 'Alemania'


def test_excel_pendiente_preserva_fila_ausente_y_correccion_de_ficha(entorno):
    import excel_generator as eg
    ctrl, cid, _ = entorno
    ctrl.db.actualizar_cliente(cid, pasaporte='ANTERIOR-FICTICIO')
    ruta = eg.exportar_cliente(ctrl.db, cid, 2026)
    ctrl.db.actualizar_cliente(cid, pasaporte='CORREGIDO-FICTICIO')
    codigo = ctrl.db.upsert('gastos', cid, {'origen':'excel', 'ejercicio':2026, 'base':77, 'total':77})
    ctrl.db.set_config(f'excel_pendiente|{cid}', '1')
    celdas = eg.database.CELDAS_RESUMEN[eg.NR]
    wb = openpyxl.load_workbook(ruta)
    wb['00_RESUMEN'][celdas['email']] = 'manual@example.invalid'
    wb.save(ruta)
    wb.close()
    contenido = ruta.read_bytes()
    with pytest.raises(eg.ExcelConflictoImportacion, match='00_RESUMEN'):
        eg.import_excel_to_db(ctrl.db, cid)
    assert ctrl.db.get_cliente(cid)['pasaporte'] == 'CORREGIDO-FICTICIO'
    assert ctrl.db.get('gastos', cid, codigo)['base'] == 77
    assert ruta.read_bytes() == contenido
    wb = openpyxl.load_workbook(ruta)
    wb['00_RESUMEN'][celdas['pasaporte']] = 'CORREGIDO-FICTICIO'
    wb.save(ruta)
    wb.close()
    eg.import_excel_to_db(ctrl.db, cid)
    assert ctrl.db.get_cliente(cid)['email'] == 'manual@example.invalid'
    assert ctrl.db.get('gastos', cid, codigo)['base'] == 77
    # La instantánea local auxiliar tampoco queda huérfana tras eliminar el cliente.
    ctrl.eliminar_cliente(cid)
    for clave in ('excel_firma', 'excel_pendiente', 'excel_resumen'):
        assert ctrl.db.get_config(f'{clave}|{cid}') is None

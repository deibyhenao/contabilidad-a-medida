"""Regresiones A11: operaciones de carpetas con expedientes exclusivamente ficticios."""
import subprocess
from pathlib import Path

import pytest

import database
from carteras import NR, RES, RAIZ
from controlador import Controlador
from database import Database


@pytest.fixture
def ctrl(tmp_path, monkeypatch):
    monkeypatch.setattr(database, "PROYECTO_DIR", tmp_path)
    ruta = tmp_path / "controlador.db"
    with Database(ruta) as db:
        db.carpeta_base = tmp_path
        db.set_config("version_carteras", "1")
    c = Controlador(ruta)
    yield c
    c.cerrar()


def _junction(enlace, destino):
    resultado = subprocess.run(["cmd", "/c", "mklink", "/J", str(enlace), str(destino)],
                              capture_output=True, text=True)
    if resultado.returncode:
        pytest.skip("Este entorno no permite crear junctions temporales")


def test_eliminar_rechaza_papelera_junction_hacia_otro_cliente(ctrl, tmp_path):
    nr = ctrl.crear_cliente("NR", "Ficticio", "A11NR")
    res = ctrl.crear_cliente("RES", "Ficticio", "A11RES", tipo_cliente="Autónomo")
    origen = Path(res["carpeta"])
    _junction(tmp_path / RAIZ[RES] / "_Eliminados", Path(nr["carpeta"]))
    with pytest.raises(ValueError, match="enlace|ubicación"):
        ctrl.eliminar_cliente(res["id"])
    assert ctrl.db.get_cliente(res["id"]) is not None and origen.is_dir()
    assert not list(Path(nr["carpeta"]).glob("RES_*"))


@pytest.mark.parametrize("operacion", ["eliminar", "cambiar"])
@pytest.mark.parametrize("modo", ["raiz_equivocada", "junction"])
def test_rechaza_origen_ajeno_sin_tocar_bd(ctrl, tmp_path, monkeypatch, operacion, modo):
    cliente = ctrl.crear_cliente("Origen", "Ficticio", "A11ORIGEN", tipo_cliente="Autónomo")
    original = Path(cliente["carpeta"])
    if modo == "raiz_equivocada":
        ajeno = tmp_path / RAIZ[NR] / original.name
        original.rename(ajeno)
    else:
        otro = ctrl.crear_cliente("Otro", "Ficticio", "A11OTRO")
        ajeno = original.parent / "Junction"
        _junction(ajeno, Path(otro["carpeta"]))
    ctrl.db.actualizar_cliente(cliente["id"], carpeta=str(ajeno))
    monkeypatch.setattr(ctrl, "_importar", lambda cid: None)
    monkeypatch.setattr(ctrl, "_recalcular_y_exportar", lambda *a: {})
    with pytest.raises(ValueError, match="enlace|mezclar"):
        if operacion == "eliminar":
            ctrl.eliminar_cliente(cliente["id"])
        else:
            ctrl.cambiar_tipo_cliente(cliente["id"], "No residente", 2025)
    actual = ctrl.db.get_cliente(cliente["id"])
    assert actual["tipo_cliente"] == "Autónomo" and Path(actual["carpeta"]) == ajeno
    assert ajeno.exists()


def test_cambio_restaurado_si_falla_crear_subcarpetas(ctrl, monkeypatch):
    cliente = ctrl.crear_cliente("Cambio", "Ficticio", "A11CAMBIO")
    origen = Path(cliente["carpeta"])
    (origen / "ficticio.txt").write_text("conservar", encoding="utf-8")
    monkeypatch.setattr(ctrl, "_importar", lambda cid: None)
    def fallo(_cliente):
        raise PermissionError("subcarpeta bloqueada ficticia")
    monkeypatch.setattr(ctrl.db, "asegurar_carpetas", fallo)
    with pytest.raises(PermissionError, match="ficticia"):
        ctrl.cambiar_tipo_cliente(cliente["id"], "Autónomo", 2025)
    actual = ctrl.db.get_cliente(cliente["id"])
    assert actual["tipo_cliente"] == "No residente" and actual["cartera"] == NR
    assert Path(actual["carpeta"]) == origen
    assert (origen / "ficticio.txt").read_text(encoding="utf-8") == "conservar"
    assert not (ctrl.db.raiz(RES) / origen.name).exists()


def test_cambio_rechaza_raiz_destino_junction(ctrl, tmp_path, monkeypatch):
    cliente = ctrl.crear_cliente("Raiz", "Ficticio", "A11RAIZ")
    raiz_res = tmp_path / RAIZ[RES]
    raiz_res.rmdir()  # raíz temporal vacía creada por la fixture
    _junction(raiz_res, tmp_path / RAIZ[NR])
    monkeypatch.setattr(ctrl, "_importar", lambda cid: None)
    monkeypatch.setattr(ctrl, "_recalcular_y_exportar", lambda *a: {})
    with pytest.raises(ValueError, match="enlace|ubicación"):
        ctrl.cambiar_tipo_cliente(cliente["id"], "Autónomo", 2025)
    assert ctrl.db.get_cliente(cliente["id"])["tipo_cliente"] == "No residente"
    assert Path(cliente["carpeta"]).is_dir()


def test_cambio_informa_rollback_incompleto(ctrl, monkeypatch):
    cliente = ctrl.crear_cliente("Rollback", "Ficticio", "A11ROLLBACK")
    origen = Path(cliente["carpeta"])
    destino = ctrl.db.raiz(RES) / origen.name
    monkeypatch.setattr(ctrl, "_importar", lambda cid: None)
    def fallo(_cliente):
        raise PermissionError("asegurar ficticio")
    monkeypatch.setattr(ctrl.db, "asegurar_carpetas", fallo)
    rename = Path.rename
    def fallo_rollback(self, target):
        if self == destino and Path(target) == origen:
            raise PermissionError("rollback ficticio")
        return rename(self, target)
    monkeypatch.setattr(Path, "rename", fallo_rollback)
    with pytest.raises(RuntimeError, match="recuperación|restaurar") as error:
        ctrl.cambiar_tipo_cliente(cliente["id"], "Autónomo", 2025)
    assert str(destino) in str(error.value) and destino.is_dir()
    assert ctrl.db.get_cliente(cliente["id"])["tipo_cliente"] == "No residente"


def test_cambio_restaurado_si_excel_abierto_bloquea_respaldo(ctrl, monkeypatch):
    cliente = ctrl.crear_cliente("Excel", "Ficticio", "A11EXCEL")
    origen = Path(cliente["carpeta"])
    excel = ctrl.db.ruta_excel(cliente)
    excel.write_bytes(b"Excel ficticio conservado")
    monkeypatch.setattr(ctrl, "_importar", lambda cid: None)
    rename, replace = Path.rename, Path.replace
    def bloqueado(self, destino, original):
        if Path(destino).parent.name == "_backups":
            raise PermissionError("Excel ficticio abierto")
        return original(self, destino)
    monkeypatch.setattr(Path, "rename", lambda self, destino: bloqueado(self, destino, rename))
    monkeypatch.setattr(Path, "replace", lambda self, destino: bloqueado(self, destino, replace))
    with pytest.raises(PermissionError, match="abierto"):
        ctrl.cambiar_tipo_cliente(cliente["id"], "Autónomo", 2025)
    assert ctrl.db.get_cliente(cliente["id"])["tipo_cliente"] == "No residente"
    assert Path(ctrl.db.get_cliente(cliente["id"])["carpeta"]) == origen
    assert excel.read_bytes() == b"Excel ficticio conservado"


def test_cambio_conserva_respaldo_anterior_y_funciona_con_base_cambiada(ctrl, tmp_path, monkeypatch):
    cliente = ctrl.crear_cliente("Base", "Ficticio", "A11BASE")
    origen = Path(cliente["carpeta"])
    excel = ctrl.db.ruta_excel(cliente)
    excel.write_bytes(b"Excel actual ficticio")
    copias = origen / "_backups"
    copias.mkdir()
    nombre_copia = f"{excel.stem}_antes_cambio_cartera{excel.suffix}"
    (copias / nombre_copia).write_bytes(b"Excel anterior ficticio")
    nueva_base = tmp_path / "NuevaBase"
    ctrl.cambiar_carpeta_clientes(nueva_base)
    monkeypatch.setattr(ctrl, "_importar", lambda cid: None)
    monkeypatch.setattr(ctrl, "_recalcular_y_exportar", lambda *a: {})
    ctrl.cambiar_tipo_cliente(cliente["id"], "Autónomo", 2025)
    actual = ctrl.db.get_cliente(cliente["id"])
    destino = Path(actual["carpeta"])
    assert destino.parent == nueva_base / RAIZ[RES] and not origen.exists()
    assert (destino / "_backups" / nombre_copia).read_bytes() == b"Excel anterior ficticio"
    segunda = destino / "_backups" / f"{excel.stem}_antes_cambio_cartera_2{excel.suffix}"
    assert segunda.read_bytes() == b"Excel actual ficticio"


def test_aviso_migracion_no_afirma_rollback_si_requiere_recuperacion(tmp_path, monkeypatch):
    monkeypatch.setattr(database, "PROYECTO_DIR", tmp_path)
    ruta = tmp_path / "aviso.db"
    with Database(ruta) as db:
        db.carpeta_base = tmp_path
    error = RuntimeError("Ver manifiesto ficticio")
    error.requiere_recuperacion = True
    def fallo(self):
        raise error
    monkeypatch.setattr(Database, "migrar_carteras", fallo)
    c = Controlador(ruta)
    try:
        assert "No se ha movido nada" not in c.aviso_migracion
        assert "recuperación" in c.aviso_migracion
        assert "se reintentará" not in c.aviso_migracion
    finally:
        c.cerrar()


def test_cambio_cartera_importa_presentado_manual_antes_de_comprobar(ctrl):
    import openpyxl

    cliente = ctrl.crear_cliente("Presentado", "Ficticio", "A11PRESENTADO", tipo_cliente="Autónomo")
    cid = cliente["id"]
    ctrl.sincronizar(cid, 2025)
    origen = Path(cliente["carpeta"])
    ruta = ctrl.db.ruta_excel(cliente)
    wb = openpyxl.load_workbook(ruta)
    try:
        ws = wb["04_MODELO_303"]
        assert ws["C2"].value == "1T"
        codigo = ws["A2"].value
        ws["L2"] = "Presentado"
        ws["N2"] = "NRC-ficticio-manual"
        wb.save(ruta)
    finally:
        wb.close()
    assert ctrl.db.get("m303", cid, codigo)["estado"] != "Presentado"
    with pytest.raises(ValueError, match="declaraciones presentadas"):
        ctrl.cambiar_tipo_cliente(cid, "No residente", 2025)
    actual = ctrl.db.get_cliente(cid)
    assert actual["tipo_cliente"] == "Autónomo" and Path(actual["carpeta"]) == origen
    assert origen.is_dir() and not (ctrl.db.raiz(NR) / origen.name).exists()
    presentada = ctrl.db.get("m303", cid, codigo)
    assert presentada["estado"] == "Presentado" and presentada["nrc"] == "NRC-ficticio-manual"


def test_cambio_subtipo_importa_factura_manual_sin_perder_excel(ctrl):
    import openpyxl

    cliente = ctrl.crear_cliente("Factura", "Ficticia", "A11FACTURA", tipo_cliente="Autónomo")
    cid = cliente["id"]
    ctrl.sincronizar(cid, 2025)
    ruta = ctrl.db.ruta_excel(cliente)
    wb = openpyxl.load_workbook(ruta)
    try:
        ws = wb["02_FACTURAS_EMITIDAS"]
        for columna, valor in zip("BDEFGIJK", (2025, "2025-03-01", "MANUAL-11", "Cliente ficticio",
                                              "A11RECEPTOR", 500, 0.21, 105)):
            ws[f"{columna}2"] = valor
        wb.save(ruta)
    finally:
        wb.close()
    assert ctrl.db.listar("facturas_emitidas", cid) == []
    ctrl.cambiar_tipo_cliente(cid, "Pyme", 2025)
    actual = ctrl.db.get_cliente(cid)
    assert actual["tipo_cliente"] == "Pyme" and actual["carpeta"] == cliente["carpeta"]
    filas = ctrl.db.listar("facturas_emitidas", cid, ejercicio=2025)
    assert len(filas) == 1 and filas[0]["numero"] == "MANUAL-11"
    assert filas[0]["base"] == 500 and filas[0]["cuota_iva"] == 105
    wb = openpyxl.load_workbook(ruta)
    try:
        assert wb["02_FACTURAS_EMITIDAS"]["E2"].value == "MANUAL-11"
        assert wb["02_FACTURAS_EMITIDAS"]["I2"].value == 500
    finally:
        wb.close()


def test_cadena_303_interanual_y_obligaciones_no_procede(ctrl):
    import fiscal_residentes as fr

    with pytest.raises(ValueError, match="porcentaje"):
        fr.iva_deducible({"base": 100, "cuota_iva": 21, "pct_iva_deducible": 1.5})

    cliente = ctrl.crear_cliente("Cadena", "IVA", "A11CADENA303", tipo_cliente="Autónomo")
    cid = cliente["id"]
    ctrl.db.upsert("actividades", cid, {"epigrafe_iae": "699", "regimen_iva": "General",
                                        "regimen_irpf": "Directa simplificada", "modelos": "303,130,111,347",
                                        "ruta": "Soportes_PDF/036.pdf"})
    # Factura recibida en 4T 2025 sin recalcular 2025 explícitamente antes de abrir 2026
    ctrl.db.upsert("facturas_recibidas", cid, {"fecha": "2025-11-15", "numero": "G-4T", "proveedor": "Prof",
                                               "nif_proveedor": "A11PROF01", "categoria": "Servicios profesionales",
                                               "base": 1000, "tipo_iva": 0.21, "cuota_iva": 210, "retencion": 150})
    ctrl.db.upsert("facturas_emitidas", cid, {"fecha": "2026-02-10", "numero": "E-1T", "receptor": "Cli",
                                              "nif_receptor": "A11CLI01", "base": 2000, "tipo_iva": 0.21,
                                              "cuota_iva": 420})
    res_2026 = ctrl.abrir_cliente(cid, 2026)
    t1_2026 = res_2026["informe"].trimestres[0]
    assert t1_2026.a_ingresar_303 == 210.0  # 420 devengado - 210 a compensar del 4T 2025

    # Editar factura de 2025 estando en 2025 actualiza también el 303 ya existente de 2026
    f_2025 = ctrl.db.listar("facturas_recibidas", cid, ejercicio=2025)[0]
    ctrl.editar(cid, "facturas_recibidas", f_2025["codigo"], {"base": 500, "cuota_iva": 105, "retencion": 0}, 2025)
    m303_2026_1t = ctrl.db.get_por_clave("m303", cid, "303|2026|1T")
    assert m303_2026_1t["compensacion_anterior"] == 105.0
    assert m303_2026_1t["a_ingresar"] == 315.0
    # Al quitar la retención de 2025, la obligación del 111 pasa a No procede
    obl_111 = ctrl.db.get_por_clave("obligaciones", cid, "OBL|111|2025|4T")
    assert obl_111["estado"] == "No procede"

    # Al cambiar a Pyme y quitar 347 de actividades, las obligaciones de 130 y 347 pasan a No procede
    act = ctrl.db.listar("actividades", cid)[0]
    ctrl.db.upsert("actividades", cid, {"codigo": act["codigo"], "modelos": "303"})
    ctrl.cambiar_tipo_cliente(cid, "Pyme", 2025)
    for p in ("1T", "2T", "3T", "4T"):
        assert ctrl.db.get_por_clave("obligaciones", cid, f"OBL|130|2025|{p}")["estado"] == "No procede"
    assert ctrl.db.get_por_clave("obligaciones", cid, "AUT|347|2025|Anual")["estado"] == "No procede"


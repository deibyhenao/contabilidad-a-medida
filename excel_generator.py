"""Intercambio bidireccional SQLite/plantilla, sin alterar fórmulas ni estilos."""
from __future__ import annotations

import os
import math
import json
import hashlib
import logging
import shutil
import tempfile
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path, PureWindowsPath

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.workbook.properties import CalcProperties
from openpyxl.comments import Comment

import database
import valores_inmuebles
from version import VERSION
from carteras import NR, RES, cartera_de
from database import Database, a_fecha
from gastos import ejercicio_registrado, fraccion
from rutas_seguras import sanitize_safe_path

FIRMA_NO_ESPECIFICADA = object()
REVISION_FORMATO = "1.0.7-fiscal-20261029"


def _formula_imputacion_redondeada(formula):
    """Excepción autorizada: base y cuota no presentadas con dos decimales."""
    if not isinstance(formula, str):
        return None
    original = re.fullmatch(r'=IF\((A\d+)="","",(.+)\)', formula)
    if not original:
        return None
    expresion = original[2]
    # Conserva los días naturales de D; limita únicamente la proporción anual.
    expresion = re.sub(r'\b(D\d+)/365\b', r'MIN(\1,365)/365', expresion)
    return f'=IF({original[1]}="","",ROUND({expresion},2))'

class ExcelBloqueadoError(Exception):
    """El libro anterior se conserva; la nueva versión queda pendiente."""

    ruta_pendiente: Path

    def __init__(self, ruta_pendiente: Path):
        self.ruta_pendiente = Path(ruta_pendiente)
        super().__init__(
            "Cierra el Excel abierto, comprueba los permisos de su carpeta y de _backups, y pulsa Sincronizar. "
            f"Se guardó una copia en {self.ruta_pendiente}"
        )


class ExcelConflictoImportacion(ValueError):
    """Un libro pendiente contiene datos anteriores a una corrección documental."""


class ExcelCambioConcurrenteError(ExcelBloqueadoError):
    """El usuario guardó durante la serialización; el principal se conserva."""

    def __init__(self, ruta_pendiente: Path):
        self.ruta_pendiente = Path(ruta_pendiente)
        Exception.__init__(self, 'El Excel cambió mientras se preparaba la exportación. '
                           'Se ha conservado tu archivo principal. La versión calculada quedó en '
                           + self.ruta_pendiente.name + '. Revisa ambas versiones antes de sincronizar.')

def _conflicto_excel(db, cid, ruta, celda, campo, candidato, actual):
    pendiente = db.get_config(f"excel_pendiente|{cid}")
    copia = Path(pendiente).name if pendiente and pendiente != "1" else "No se pudo crear una copia actualizada"
    raise ExcelConflictoImportacion(
        f"La sincronización se ha detenido en {celda} ({campo}).\n"
        f"Valor del Excel: {candidato if candidato is not None else '(vacío)'}.\n"
        f"Valor actual en la ficha: {actual if actual is not None else '(vacío)'}.\n\n"
        "No se ha importado ni sobrescrito ningún dato.\n"
        f"Carpeta: {ruta.parent}\nExcel original: {ruta.name}\nCopia actualizada: {copia}\n\n"
        "1. Haz una copia de seguridad del Excel original antes de continuar.\n"
        "2. Traslada tus cambios manuales a la copia actualizada indicada, conservando los datos corregidos de la ficha. "
        "Si no existe esa copia, ajusta la celda indicada en el original al valor actual de la ficha.\n"
        "3. Guarda el libro conciliado con el nombre del Excel original y pulsa Sincronizar. "
        "Si aparece otro conflicto, revísalo del mismo modo.\n\n"
        "Después de sincronizar podrás hacer nuevas correcciones en la ficha o el Excel.")


def _comprobar_datos_pendientes(db, cid, tabla, columnas, datos, previo, hoja, fila, ruta):
    # Solo durante una exportación pendiente: una edición normal de Excel sigue
    # siendo una corrección manual válida. La consulta usa el índice por fila.
    campos = {r[0] for r in db.con.execute(
        "SELECT DISTINCT campo FROM documentos_cambios WHERE cliente_id=? AND tabla=? AND codigo=?",
        (cid, tabla, datos.get("codigo")))}
    if tabla == "documentos" and previo:
        campos |= {"ejercicio", "ruta", "tipo", "fecha_documento", "fecha_caducidad"}
    if tabla == "residencias":
        campos.discard("ejercicio")  # La clave del año sigue visible aunque sus datos se retiren.
    for col in columnas:
        if col.campo not in campos:
            continue
        candidato = database._normalizar(col, datos.get(col.campo))
        actual = database._normalizar(col, (previo or {}).get(col.campo))
        if candidato != actual:
            _conflicto_excel(db, cid, ruta, f"{hoja}!{col.letra}{fila}",
                             col.campo.replace("_", " ").capitalize(), candidato, actual)


def _ruta(db: Database, cid: int, ruta: Path | None) -> Path:
    cliente = db.get_cliente(cid)
    if cliente is None:
        raise ValueError(f"Cliente inexistente: {cid}")
    base = Path(cliente["carpeta"])
    sanitize_safe_path(base, base.absolute())
    if not base.is_dir():
        raise FileNotFoundError('No se encuentra la carpeta del expediente. No se ha recreado; '
                                'comprueba su ubicación y que la unidad esté conectada.')
    destino = Path(ruta) if ruta is not None else db.ruta_excel(cliente)
    # La exportación/importación explícita permite una copia elegida por el usuario.
    # El destino automático siempre permanece dentro del expediente.
    raiz = destino.absolute().parent if ruta is not None else base
    return sanitize_safe_path(raiz, destino.absolute())


# Tablas en las que una fila escrita a mano en el Excel (sin ID) es un alta y su ausencia, una baja
MANUALES_EXCEL = ("gastos", "facturas_emitidas", "facturas_recibidas")


def _firma_fila_manual(tabla, fila) -> str:
    campos = {c.campo: fila.get(c.campo) for c in _columnas(database.TABLAS[tabla])}
    campos['origen'] = fila.get('origen')
    contenido = json.dumps(campos, sort_keys=True, ensure_ascii=False, default=str)
    return hashlib.sha256(contenido.encode('utf-8')).hexdigest()


def _columnas(tabla):
    return [c for c in tabla.columnas_excel if c.tipo != "formula"]


def _link(cell, valor, carpeta_cliente: Path | None = None) -> None:
    if valor in (None, ""):
        cell.value = None
        cell.hyperlink = None
        return
    relativo = str(valor).replace("\\", "/")
    cell.value = relativo.rstrip("/").rsplit("/", 1)[-1]
    cell.data_type = 's'  # El nombre es una etiqueta de archivo, nunca una fórmula.
    # El texto de una celda importada o extraída no concede permiso para
    # convertir URLs, rutas UNC o archivos externos en vínculos ejecutables.
    ruta = Path(relativo)
    permitido = (not ruta.is_absolute() and not PureWindowsPath(relativo).drive
                 and ".." not in ruta.parts and len(ruta.parts) == 2
                 and ruta.parts[0] in {"Soportes_PDF", ".Papelera_Documentos"}
                 and ruta.suffix.lower() in {".pdf", ".xlsx", ".xls", ".docx", ".csv"})
    if permitido:
        try:
            sanitize_safe_path(carpeta_cliente or Path.cwd(), relativo)
        except ValueError:
            permitido = False
    cell.hyperlink = relativo if permitido else None
    if not permitido:
        cell.value = str(valor)  # conserva el dato manual, sin convertirlo en vínculo
        cell.data_type = 's'
        return
    fuente = copy(cell.font)
    fuente.color = "0000FF"
    fuente.underline = "single"
    cell.font = fuente


def _valor_excel(col, valor):
    if valor in (None, ""):
        return None
    if col.tipo == "date":
        return a_fecha(valor)
    if col.tipo in ("eur", "pct"):
        return float(valor)
    if col.tipo == "int":
        return int(valor)
    return valor


def _ampliar_fila(ws, fila: int, tabla, limite: int) -> None:
    if fila <= limite:
        return
    ws.row_dimensions[fila].height = ws.row_dimensions[2].height
    formulas = {c.letra for c in tabla.columnas_excel if c.tipo == "formula"}
    for origen in ws[2]:
        celda = ws.cell(fila, origen.column)
        celda._style = copy(origen._style)
        if origen.column_letter in formulas and origen.data_type == "f":
            celda.value = Translator(origen.value, origin=origen.coordinate).translate_formula(celda.coordinate)


def _guardar_atomico(wb, destino: Path, carpeta_cliente: Path, *,
                     firma_esperada=FIRMA_NO_ESPECIFICADA, guardar_version: bool = False) -> None:
    destino = sanitize_safe_path(destino.parent, destino.absolute())
    carpeta_cliente = sanitize_safe_path(carpeta_cliente, carpeta_cliente.absolute())
    if not carpeta_cliente.is_dir():
        raise FileNotFoundError('No se encuentra la carpeta del expediente; no se ha recreado.')
    if destino.parent != carpeta_cliente:
        destino.parent.mkdir(parents=True, exist_ok=True)
    if firma_esperada is FIRMA_NO_ESPECIFICADA:
        firma_esperada = _firma_excel(destino) if destino.exists() else None

    def pendiente_de(temporal):
        sello = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        pendiente = destino.with_name(f"{destino.stem}_PENDIENTE_{sello}.xlsx")
        if len(str(pendiente)) > 247:
            descriptor, nombre = tempfile.mkstemp(prefix="PENDIENTE_", suffix=".xlsx", dir=destino.parent)
            os.close(descriptor)
            pendiente = Path(nombre)
        shutil.copyfile(temporal, pendiente)
        return pendiente

    def comprobar_firma(temporal):
        actual = _firma_excel(destino) if destino.exists() else None
        if actual != firma_esperada:
            raise ExcelCambioConcurrenteError(pendiente_de(temporal))
    descriptor, nombre_tmp = tempfile.mkstemp(prefix=".nr_", suffix=".xlsx", dir=destino.parent)
    os.close(descriptor)
    temporal = Path(nombre_tmp)
    try:
        wb.save(temporal)
        try:
            comprobar_firma(temporal)
            if destino.exists():
                backups = sanitize_safe_path(carpeta_cliente, "_backups")
                backups.mkdir(parents=True, exist_ok=True)
                if guardar_version:
                    digest = hashlib.sha256(destino.read_bytes()).hexdigest()
                    nombre_version = f"ANTES_{VERSION}_{digest[:12]}.xlsx"
                    if len(str(backups / nombre_version)) > 247:
                        nombre_version = f"V{VERSION.replace('.', '')}_{digest[:8]}.xlsx"
                    duradera = sanitize_safe_path(backups, backups / nombre_version)
                    if duradera.exists() and hashlib.sha256(duradera.read_bytes()).hexdigest() != digest:
                        raise ValueError('La copia previa a la versión fue modificada; se conserva sin sobrescribir.')
                    if not duradera.exists():
                        shutil.copy2(destino, duradera)
                sello = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
                corto = "CAM_" + hashlib.sha256(destino.name.encode("utf-8")).hexdigest()[:12]
                nombre = f"{destino.stem}_{sello}.xlsx"
                if len(str(backups / nombre)) > 247:
                    # La ruta del cliente puede ser larga aunque el Excel quepa.
                    # Evita superar MAX_PATH únicamente al añadir el sello de copia.
                    nombre = f"{corto}_{sello}.xlsx"
                copia = sanitize_safe_path(backups, backups / nombre)
                patron = re.compile("(?:" + re.escape(destino.stem) + "|" + corto + r")_(\d{8}_\d{6}_\d{6})\.xlsx")
                copias = []
                for candidata in backups.glob("*.xlsx"):
                    if not patron.fullmatch(candidata.name):
                        continue
                    try:
                        candidata = sanitize_safe_path(backups, candidata)
                    except ValueError:
                        continue
                    if candidata.is_file():
                        copias.append(candidata)
                # El sello del nombre es la fecha de copia; copy2 conserva la
                # fecha del original, que puede ser antigua al restaurar un libro.
                copias.sort(key=lambda p: patron.fullmatch(p.name).group(1), reverse=True)
                if copias and hashlib.sha256(copias[0].read_bytes()).hexdigest() == hashlib.sha256(destino.read_bytes()).hexdigest():
                    copia = copias[0]  # Reintentos bloqueados no duplican el mismo respaldo.
                else:
                    shutil.copy2(destino, copia)
                    copias.insert(0, copia)
            else:
                copias = []
            comprobar_firma(temporal)
            os.replace(temporal, destino)
            # No purgar el historial si Windows rechazó el reemplazo del principal.
            # Un fallo al depurar tampoco invalida un Excel ya guardado correctamente.
            for antigua in copias[10:]:
                try:
                    antigua.unlink()
                except OSError:
                    logging.getLogger(__name__).warning('No se pudieron depurar algunas copias antiguas; se conservan.')
        except PermissionError as exc:
            raise ExcelBloqueadoError(pendiente_de(temporal)) from exc
    finally:
        temporal.unlink(missing_ok=True)


def _firma_excel(ruta: Path) -> str:
    st = ruta.stat()
    h = hashlib.sha256()
    with ruta.open('rb') as entrada:
        for bloque in iter(lambda: entrada.read(1024 * 1024), b''):
            h.update(bloque)
    return f"{st.st_mtime_ns}|{st.st_size}|{h.hexdigest()}"


def excel_modificado(db: Database, cliente_id: int) -> bool:
    """True si el usuario ha tocado el Excel desde la última exportación de la app."""
    cliente = db.get_cliente(cliente_id)
    if not cliente:
        return False
    ruta = db.ruta_excel(cliente)
    return ruta.exists() and db.get_config(f"excel_firma|{cliente_id}") != _firma_excel(ruta)


def necesita_actualizar_formato(db: Database, cliente_id: int) -> bool:
    """Un libro de una versión anterior se importa y regenera al primer uso.

    La marca solo se confirma al guardar el libro canónico: un PENDIENTE nunca
    da por actualizado un Excel antiguo. Un fichero eliminado se reconstruye de BD.
    """
    cliente = db.get_cliente(cliente_id)
    if not cliente:
        return False
    return (not db.ruta_excel(cliente).exists()
            or db.get_config(f'excel_version|{cliente_id}') != VERSION
            or db.get_config(f'excel_revision|{cliente_id}') != REVISION_FORMATO
            or bool(db.get_config(f'excel_pendiente|{cliente_id}')))


def _resumen_gastos(ws, inmuebles, limite: int = 3000, *, sin_asignar: bool = False) -> None:
    """Añade el desglose en celdas vacías; las fórmulas siguen el ejercicio E4."""
    categorias = database.CATEGORIAS_GASTO
    fila_total = 31 + len(categorias)
    fila_nota = fila_total + 2
    for fila in ws.iter_rows(min_row=29, max_row=fila_nota, min_col=1, max_col=8):
        if any(c.value is not None for c in fila):
            raise ValueError("El área A29:H del desglose de gastos debe estar vacía en la plantilla")
    if any(r.min_row <= fila_nota and r.max_row >= 29 and r.min_col <= 8
           for r in ws.merged_cells.ranges):
        raise ValueError("El área del desglose de gastos contiene celdas combinadas de la plantilla")

    for col in range(1, 9):
        ws.cell(29, col)._style = copy(ws["A18"]._style)
        ws.cell(30, col)._style = copy(ws["A19"]._style)
    ws["A29"] = "GASTOS REGISTRADOS DEL EJERCICIO POR INMUEBLE"
    ws.merge_cells("A29:H29")
    ws.row_dimensions[29].height = ws.row_dimensions[18].height
    ws.row_dimensions[30].height = 45
    ws["A30"], ws["B30"] = "Categoría", "Total"
    resto = len(inmuebles) > 6 or sin_asignar
    visibles = inmuebles[:5 if resto else 6]
    for col, inmueble in enumerate(visibles, 3):
        direccion = " ".join(str(inmueble.get("direccion") or "").split())[:36]
        cell = ws.cell(30, col, f"{inmueble['codigo']}\n{direccion}")
        cell.data_type = "s"
    if resto:
        ws["H30"] = "Resto / sin asignar"

    def rango(col):
        return "'08_GASTOS'!$" + col + "$2:$" + col + "$" + str(limite)

    def suma(fila, inmueble=None):
        # Importe por gasto = R (deducible final) si está escrito; si no, N × (O o 100 %) × (P o 100 %).
        # No se usa Q: su fórmula de plantilla (N*O*P) da 0 cuando O/P están vacíos, como en los gastos de PDF.
        # (X="") vale 1 en celdas vacías, así un porcentaje vacío cuenta como 100 % sin fórmulas matriciales.
        filtro = f"({rango('B')}=$E$4)*({rango('H')}=$A{fila})"
        if inmueble:
            filtro += f"*({rango('C')}={inmueble})"
        total = f"({rango('N')}+({rango('N')}=\"\")*{rango('L')})"
        importe = (f"({rango('R')}+({rango('R')}=\"\")*{total}"
                   f"*({rango('O')}+({rango('O')}=\"\"))*({rango('P')}+({rango('P')}=\"\")))")
        return f"=SUMPRODUCT({filtro}*ROUND({importe},2))"

    for fila, categoria in enumerate(categorias, 31):
        ws.cell(fila, 1, categoria)._style = copy(ws["A20"]._style)
        for col in range(2, 9):
            cell = ws.cell(fila, col)
            cell._style = copy(ws["E12"]._style)
            cell.number_format = '#,##0.00 [$€-es-ES]'
            if col == 2:
                cell.value = suma(fila)
            elif col - 3 < len(visibles):
                letra = cell.column_letter
                # Código en la primera línea de la cabecera; la dirección queda visible.
                codigo = f'LEFT({letra}$30,FIND(CHAR(10),{letra}$30&CHAR(10))-1)'
                cell.value = suma(fila, codigo)
            elif col == 8 and resto:
                # Incluye los demás inmuebles y los gastos aún sin asignar.
                cell.value = f"=B{fila}-SUM(C{fila}:G{fila})"
    ws.cell(fila_total, 1, "TOTAL")._style = copy(ws["G7"]._style)
    for col in range(2, 9):
        cell = ws.cell(fila_total, col)
        cell._style = copy(ws["H7"]._style)
        cell.number_format = '#,##0.00 [$€-es-ES]'
        cell.value = f"=SUM({cell.column_letter}31:{cell.column_letter}{fila_total - 1})"
    nota = ws.cell(fila_nota, 1,
                   "Gastos registrados; la deducción fiscal depende del régimen y del cálculo del 210. "
                   "Detalle y PDFs en 08_GASTOS (columna U) y 13_DOCUMENTACION (columna J).")
    fuente = copy(ws["A20"].font)
    fuente.color = "808080"
    fuente.italic = True
    nota.font = fuente


def _resumen_valores(wb, db: Database, cid: int, ejercicio: int) -> None:
    """Vista de consulta en celdas vacías; las hojas de entrada siguen intactas."""
    ws = wb["00_RESUMEN"]
    inicio = max(50, 36 + len(database.CATEGORIAS_GASTO))
    cantidad = db.con.execute("SELECT COUNT(*) FROM inmuebles WHERE cliente_id=?", (cid,)).fetchone()[0]
    fin = inicio + 2 + cantidad
    if any(c.value is not None for row in ws.iter_rows(min_row=inicio, max_row=fin, max_col=8) for c in row):
        raise ValueError("El área del resumen de valores debe estar vacía en la plantilla.")
    if any(r.min_row <= fin and r.max_row >= inicio and r.min_col <= 8 for r in ws.merged_cells.ranges):
        raise ValueError("El resumen de valores coincide con celdas combinadas de la plantilla.")
    ws.cell(inicio, 1, f"VALORES DE LOS INMUEBLES · EJERCICIO {int(ejercicio)}")
    ws.cell(inicio + 1, 1, "Vista de consulta al sincronizar · importes antes de titularidad. Edita 02_INMUEBLES y 03_CATASTRO_IBI.")
    encabezados = ("Inmueble", "Valor catastral", "Valor de adquisición total", "Valor de referencia disponible",
                   "Precio de escritura", "Gastos de compra", "Referencia en adquisición", "Referencia Catastro")
    for c, texto in enumerate(encabezados, 1):
        celda = ws.cell(inicio + 2, c, texto)
        celda._style = copy(wb["02_INMUEBLES"].cell(1, c)._style)
        alineacion = copy(celda.alignment)
        alineacion.wrap_text = True
        celda.alignment = alineacion
    ws.row_dimensions[inicio + 2].height = 44
    for i, inmueble in enumerate(db.iterar("inmuebles", cid), 2):
        fila = inicio + 1 + i
        fin = max(fin, fila)  # la siguiente sección empieza después de la última fila realmente exportada
        cat = db.catastro(cid, inmueble["codigo"], int(ejercicio)) or {}
        compra = a_fecha(inmueble.get("fecha_adquisicion"))
        referencia_compra = (db.valor_referencia_ejercicio(cid, inmueble["codigo"], compra.year)
                             if compra else None)
        referencia_actual = db.valor_referencia_ejercicio(cid, inmueble["codigo"], ejercicio)
        referencia_alternativa = str(referencia_compra) if referencia_compra is not None else '""'
        vista = valores_inmuebles.resumen(inmueble, cat)
        origen = "'02_INMUEBLES'!"
        valores = [inmueble["codigo"] + "\n" + (inmueble.get("direccion") or "Sin dirección"), vista["catastral"],
            f'=IF(COUNT({origen}J{i}:P{i})=0,"",{origen}Q{i})',
            f'=IF(COUNT(G{fila}:H{fila})=0,"",IF(ISNUMBER(H{fila}),H{fila},G{fila}))',
            f'=IF(ISBLANK({origen}J{i}),"",{origen}J{i})',
            f'=IF(COUNT({origen}K{i}:P{i})=0,"",ROUND(SUM({origen}K{i}:P{i}),2))',
            f'=IF(ISBLANK({origen}R{i}),{referencia_alternativa},'
            f'{origen}R{i})', referencia_actual]
        for c, valor in enumerate(valores, 1):
            celda = ws.cell(fila, c, valor)
            if c == 1:
                celda.data_type = "s"
                celda.alignment = copy(wb["02_INMUEBLES"]["B2"].alignment)
            else:
                celda.number_format = '#,##0.00 [$€-es-ES]'
        ws.row_dimensions[fila].height = 32
        if cat.get("_heredado"):
            ws.cell(fila, 2).comment = Comment(f"Dato catastral heredado de {cat['_heredado']}. Se actualiza al sincronizar.", "Contabilidad a Medida")
        ws.cell(fila, 7).comment = Comment(
            f"Referencia en adquisición: {compra.year if compra else 'año sin dato'}. "
            "Se mantiene separada de la referencia del ejercicio consultado.", "Contabilidad a Medida")
        ws.cell(fila, 8).comment = Comment(
            f"Referencia acreditada para {ejercicio}; vacío si no consta en ese ejercicio.", "Contabilidad a Medida")
        if vista["parcial"]:
            ws.cell(fila, 3).comment = Comment("Al sincronizar faltaba registrar el precio de escritura; revisa el total tras completar los datos.", "Contabilidad a Medida")
    hoja = wb["02_INMUEBLES"]
    for campo, etiqueta in valores_inmuebles.CAMPOS_COMPRA:
        col = next(c for c in database.TABLAS["inmuebles"].columnas_excel if c.campo == campo)
        hoja[f"{col.letra}1"].comment = Comment(etiqueta + ". Importe individual; no introduzcas aquí el total de adquisición. Si no consta, deja vacío.", "Contabilidad a Medida")
    for celda, texto in {
        "Q1": "Total automático: precio de escritura + ITP/IVA-AJD + notaría + registro + agencia/inmobiliaria + abogado/gestoría + otros gastos registrados. No incluye valor catastral ni referencia.",
        "R1": "Valor de referencia en la adquisición. Se conserva separado del precio de escritura y sus gastos.",
        "S1": "Base usada en el impuesto de adquisición: dato distinto del valor de adquisición total. Las reglas fiscales existentes no se cambian.",
    }.items():
        hoja[celda].comment = Comment(texto, "Contabilidad a Medida")
    wb["03_CATASTRO_IBI"]["E1"].comment = Comment("Valor catastral del ejercicio. Se conserva separado del precio de escritura, los gastos y el valor de referencia.", "Contabilidad a Medida")
    _resumen_ocupacion(wb, db, cid, int(ejercicio), fin + 3)


def _resumen_ocupacion(wb, db: Database, cid: int, ejercicio: int, inicio: int) -> None:
    """Consulta del motor al sincronizar; no duplica las reglas fiscales en fórmulas de Excel."""
    import tax_calculator as tc
    ws = wb["00_RESUMEN"]
    inmuebles = db.listar("inmuebles", cid)
    fin = inicio + 2 + len(inmuebles)
    if any(c.value is not None for row in ws.iter_rows(min_row=inicio, max_row=fin, max_col=8) for c in row):
        raise ValueError("El área del resumen de ocupación debe estar vacía en la plantilla.")
    if any(r.min_row <= fin and r.max_row >= inicio and r.min_col <= 8 for r in ws.merged_cells.ranges):
        raise ValueError("El resumen de ocupación coincide con celdas combinadas de la plantilla.")
    ws.cell(inicio, 1, f"OCUPACIÓN Y RENTAS · EJERCICIO {ejercicio}")
    ws.cell(inicio + 1, 1, "Consulta al sincronizar con la app. Devengado y cobrado son conceptos distintos.")
    encabezados = ("Inmueble", "Fecha de compra", "Días propiedad", "Días alquilados", "Días vacíos",
                   "Renta devengada", "Cobros registrados", "Ejercicio")
    for col, texto in enumerate(encabezados, 1):
        celda = ws.cell(inicio + 2, col, texto)
        celda._style = copy(wb["02_INMUEBLES"].cell(1, col)._style)
        alineacion = copy(celda.alignment)
        alineacion.wrap_text = True
        celda.alignment = alineacion
    ws.row_dimensions[inicio + 2].height = 44
    for fila, inm in enumerate(inmuebles, inicio + 3):
        contratos = db.listar("contratos", cid, inmueble=inm["codigo"])
        rentas = {c["codigo"]: db.listar("rentas", cid, contrato=c["codigo"]) for c in contratos}
        resumen = tc.resumen_arrendamiento(inm, contratos, rentas, ejercicio)
        cobros = db.listar("cobros", cid, inmueble=inm["codigo"], ejercicio=ejercicio)
        valores = (inm["codigo"], a_fecha(inm.get("fecha_adquisicion")), resumen["dias_propiedad"],
                   resumen["dias_alquilado"], resumen["dias_vacio"], resumen["ingresos_devengados"],
                   tc.r2(sum(float(c.get("importe_cobrado") or 0) for c in cobros)), ejercicio)
        for col, valor in enumerate(valores, 1):
            celda = ws.cell(fila, col, valor)
            if col == 1:
                celda.data_type = "s"
            elif col == 2:
                celda.number_format = "dd/mm/yyyy"
            elif col in (6, 7):
                celda.number_format = '#,##0.00 [$€-es-ES]'


def _aclaraciones_resumen(wb, db: Database, cid: int, ejercicio: int) -> None:
    """Aclara el alcance de las fórmulas existentes sin cambiarlas ni alterar la plantilla."""
    from tax_calculator import fmt_eur
    ws = wb["00_RESUMEN"]
    notas = {
        "H7": "Total de impuestos pagados registrados para el ejercicio. No suma cuotas pendientes ni provisionales.",
        "E10": "Suma de movimientos cobrados registrados; puede incluir cobros pendientes de conciliar. No acredita su naturaleza ni su devengo fiscal.",
        "E12": "Suma de importes anotados en gastos. Contrasta la deducción fiscal con residencia, límites y modelos 210; no acredita por sí sola gastos efectivamente deducidos.",
        "H10": "Cuenta documentos faltantes de todo el expediente, no solo del ejercicio seleccionado. Filtra 13_DOCUMENTACION por ejercicio para revisarlos.",
    }
    from tax_calculator import conflictos_210_alquiler
    conflictos = conflictos_210_alquiler(db, cid, ejercicio)
    if conflictos:
        aviso = "PENDIENTE DE REVISIÓN: " + " ".join(conflictos.values())
        notas["H7"] += " " + aviso + " Los pagos registrados no constituyen un total conciliado."
        notas["E13"] = aviso + " Se conservan los importes históricos; no hay un cálculo actualizado del alquiler afectado."
    pendientes = [f for f in db.listar("m210_imputacion", cid, ejercicio=ejercicio)
                  if f.get("estado") not in ("Presentado", "No procede")]
    total = sum(float(f.get("cuota") or 0) for f in pendientes)
    notas["E14"] = ("Solo suma importes pagados, no la cuota pendiente. Cuota de modelos no presentados: "
                    + fmt_eur(total) + ". Consulta 10_210_IMPUTACION."
                    + " Una declaración presentada también puede tener pagos pendientes."
                    + (" Cálculo PROVISIONAL: falta documentación por acreditar." if any(
                        f.get("estado") == "Falta documentación" for f in pendientes) else ""))
    for coordenada, texto in notas.items():
        anterior = ws[coordenada].comment
        if anterior:
            texto = anterior.text if texto in anterior.text else anterior.text + "\n\n" + texto
            comentario = copy(anterior)
            comentario.text = texto
        else:
            comentario = Comment(texto, "Contabilidad a Medida")
        comentario.width = max(comentario.width, 360)
        comentario.height = max(comentario.height, 160)
        ws[coordenada].comment = comentario


def exportar_cliente(db: Database, cliente_id: int, ejercicio: int, destino: Path | None = None, *,
                     firma_esperada=FIRMA_NO_ESPECIFICADA) -> Path:
    """Exporta BD; el controlador debe importar el Excel ANTES de modificar BD.

    Importar aquí revertiría los datos que acaba de aplicar el PDF/calculadora.
    Para un ciclo completo, usar sincronizar o el flujo de 01-arquitectura.md.
    """
    destino = _ruta(db, cliente_id, destino)
    if any(destino.resolve() == Path(pl).resolve() for pl in database.PLANTILLAS.values()):
        raise ValueError("La plantilla no puede ser el destino de una exportación")
    cliente = db.get_cliente(cliente_id)
    cartera = cartera_de(cliente.get("tipo_cliente"))
    canonico = destino.resolve() == db.ruta_excel(cliente).resolve()
    if firma_esperada is FIRMA_NO_ESPECIFICADA or not canonico:
        firma_esperada = _firma_excel(destino) if destino.exists() else None
    guardar_version = canonico and (
        db.get_config(f'excel_version|{cliente_id}') != VERSION
        or db.get_config(f'excel_revision|{cliente_id}') != REVISION_FORMATO)
    if canonico:
        db.set_config(f"excel_pendiente|{cliente_id}", "1")
    wb = openpyxl.load_workbook(database.PLANTILLAS[cartera])
    firmas_filas = {t: {} for t in MANUALES_EXCEL}
    try:
        for tabla, hoja in database.hojas_cartera(cartera):
            nombre = tabla.nombre
            ws = wb[hoja]
            limite = ws.max_row
            # Guardar la fórmula de plantilla antes de que una primera fila
            # presentada la sustituya por un importe fijo.
            formulas_imputacion = ({letra: ws[f"{letra}{tabla.fila_inicio}"].value for letra in ("H", "J")}
                                   if nombre == "m210_imputacion" else {})
            for fila, datos in enumerate(db.iterar(nombre, cliente_id, incluir_retirados=True), tabla.fila_inicio):
                if nombre in MANUALES_EXCEL:
                    firmas_filas[nombre][datos['codigo']] = _firma_fila_manual(nombre, datos)
                if nombre == "documentos":
                    gestion = db.con.execute("SELECT en_papelera,ruta_papelera FROM documentos_gestion WHERE cliente_id=? AND codigo=?",
                                             (cliente_id, datos["codigo"])).fetchone()
                    if gestion and gestion["en_papelera"]:
                        datos = {**datos, "estado": "En papelera", "ruta": gestion["ruta_papelera"]}
                _ampliar_fila(ws, fila, tabla, limite)
                if nombre == 'm303':
                    import fiscal_residentes as fr
                    decision_303 = fr.saldo_303(db, cliente_id, datos)
                for col in _columnas(tabla):
                    celda = ws[f"{col.letra}{fila}"]
                    valor = datos.get(col.campo)
                    if nombre == "gastos" and col.campo == "ejercicio" and valor is None:
                        valor = ejercicio_registrado(datos)
                    if nombre == "residencias" and col.campo == "codigo":
                        valor = int(datos.get("ejercicio") or datos["codigo"])
                    if col.tipo == "link":
                        _link(celda, valor, Path(cliente["carpeta"]))
                    else:
                        celda.value = _valor_excel(col, valor)
                        if isinstance(celda.value, str):
                            celda.data_type = "s"
                    if nombre == 'm303' and col.campo == 'a_compensar' and decision_303['requiere']:
                        mensaje = ('PENDIENTE DE CONFIRMAR POR EL ASESOR. Este importe histórico no se aplica '
                                   'a los periodos siguientes; sus cálculos son provisionales.'
                                   if decision_303['pendiente'] else
                                   f"Destino confirmado: {decision_303['destino']}. Saldo aplicado en los cálculos posteriores: "
                                   f"{decision_303['saldo']:.2f} euros. {decision_303.get('justificacion', '')}")
                        celda.comment = Comment(mensaje + ' Se conserva aquí el dato histórico presentado.', 'Contabilidad a Medida')
                    if nombre == 'm303' and col.campo == 'a_ingresar' and valor is not None and float(valor) < 0:
                        # El resumen suma esta columna como deuda, no como resultado
                        # de liquidación. No restar devoluciones de otros impuestos.
                        celda.value = 0
                        celda.comment = Comment(
                            f'Resultado negativo del justificante conservado en la base de datos: {float(valor):.2f} euros. '
                            'A ingresar se muestra como cero porque no es una deuda negativa. '
                            'El destino y saldo para periodos posteriores requieren confirmación del asesor; '
                            'consulta el comentario de a compensar. No se modifica el justificante presentado.',
                            'Contabilidad a Medida')
                if nombre == "m210_imputacion" and datos.get("estado") == "Presentado":
                    # Congelada: renta y cuota son las del justificante (la fórmula usaría los datos de la app)
                    for formula, campo in (("renta_imputada_formula", "renta_imputada"), ("cuota_formula", "cuota")):
                        col = next((c for c in tabla.columnas_excel if c.campo == formula), None)
                        if col and datos.get(campo) not in (None, ""):
                            ws[f"{col.letra}{fila}"].value = float(datos[campo])
                            ws[f"{col.letra}{fila}"].comment = Comment(
                                "Importe de la declaración presentada (justificante). Sustituye a la fórmula para "
                                "conservar lo declarado; si deja de estar presentada, vuelve la fórmula.",
                                "Contabilidad a Medida")
                elif nombre == "m210_imputacion":
                    for campo in ("renta_imputada_formula", "cuota_formula"):
                        col = next((c for c in tabla.columnas_excel if c.campo == campo), None)
                        if col:
                            celda = ws[f"{col.letra}{fila}"]
                            modelo = formulas_imputacion.get(col.letra)
                            origen = f"{col.letra}{tabla.fila_inicio}"
                            formula = _formula_imputacion_redondeada(
                                Translator(modelo, origin=origen).translate_formula(celda.coordinate)
                                if isinstance(modelo, str) and modelo.startswith("=") else celda.value)
                            if formula:
                                celda.value = formula
                                celda.comment = Comment(
                                    "Importe no presentado: se redondea a dos decimales, igual que en la aplicación. "
                                    "La cuota se calcula sobre la renta imputada ya redondeada. "
                                    "La proporción no supera el año completo; D conserva los días naturales.",
                                    "Contabilidad a Medida")
        resumen = {
            **cliente,
            "nombre_completo": f"{cliente['nombre']} {cliente['apellidos']}".strip(),
            "fecha_revision": date.today(),
            "ejercicio": int(ejercicio),
        }
        ws = wb["00_RESUMEN"]
        celdas = database.CELDAS_RESUMEN[cartera]
        for campo, coordenada in celdas.items():
            ws[coordenada] = resumen.get(campo)
            if isinstance(ws[coordenada].value, str):
                ws[coordenada].data_type = "s"
        ws[celdas["carpeta"]].hyperlink = Path(cliente["carpeta"]).resolve().as_uri()
        if cartera == NR:  # el desglose dinámico por inmueble solo existe en la plantilla de no residentes
            sin_asignar = db.con.execute(
                "SELECT 1 FROM gastos g WHERE g.cliente_id=? AND NOT EXISTS "
                "(SELECT 1 FROM inmuebles i WHERE i.cliente_id=g.cliente_id AND i.codigo=g.inmueble) LIMIT 1",
                (cliente_id,)).fetchone() is not None
            _resumen_gastos(ws, db.listar("inmuebles", cliente_id), max(3000, wb["08_GASTOS"].max_row),
                            sin_asignar=sin_asignar)
            _resumen_valores(wb, db, cliente_id, ejercicio)
            _aclaraciones_resumen(wb, db, cliente_id, ejercicio)
        if wb.calculation is None:
            wb.calculation = CalcProperties()
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        _guardar_atomico(wb, destino, Path(cliente["carpeta"]), firma_esperada=firma_esperada,
                         guardar_version=guardar_version)
        if canonico:
            with db.transaccion():
                db.set_config(f"excel_firma|{cliente_id}", _firma_excel(destino))
                db.set_config(f"excel_version|{cliente_id}", VERSION)
                db.set_config(f"excel_revision|{cliente_id}", REVISION_FORMATO)
                db.set_config(f"excel_nombre|{cliente_id}", destino.name)
                db.set_config(f"excel_filas|{cliente_id}", json.dumps(firmas_filas, ensure_ascii=False))
                db.set_config(f"excel_resumen|{cliente_id}", json.dumps({c: cliente.get(c) for c in
                    ("pasaporte", "nacionalidad", "email", "telefono", "responsable")}, ensure_ascii=False))
                db.set_config(f"excel_pendiente|{cliente_id}", "")
    except ExcelBloqueadoError as exc:
        if canonico:
            db.set_config(f"excel_pendiente|{cliente_id}", str(exc.ruta_pendiente))
        raise
    finally:
        wb.close()
    return destino


def import_excel_to_db(db: Database, cliente_id: int, ruta: Path | None = None) -> dict[str, int]:
    """Importa celdas editables: un vacío explícito borra el campo en BD.

    No escribe el libro. Los códigos de filas manuales se verán en la próxima
    exportación. Una fila de gasto manual sin ID idéntica a otra ya importada
    se reconoce para permitir reintentos antes de exportar.
    """
    ruta = _ruta(db, cliente_id, ruta)
    if not ruta.exists():
        return {}
    cliente = db.get_cliente(cliente_id)
    if cliente and ruta.resolve() == db.ruta_excel(cliente).resolve():
        if db.get_config(f"excel_firma|{cliente_id}") == _firma_excel(ruta):
            return {}
    cartera = cartera_de((cliente or {}).get("tipo_cliente"))
    pendiente = bool(db.get_config(f"excel_pendiente|{cliente_id}"))
    try:
        filas_exportadas = json.loads(db.get_config(f'excel_filas|{cliente_id}', '{}'))
    except (ValueError, TypeError):
        filas_exportadas = {}
    hojas = database.hojas_cartera(cartera)
    wb = openpyxl.load_workbook(ruta, data_only=False)
    faltan = [hoja for hoja in ['00_RESUMEN', *(h for _, h in hojas)] if hoja not in wb.sheetnames]
    if faltan:
        wb.close()
        raise ValueError(f"El Excel no corresponde a la plantilla de {'residentes' if cartera == RES else 'no residentes'}"
                         f" (faltan {', '.join(faltan[:3])}).")
    conteos: dict[str, int] = {}
    try:
        # El nombre y la carpeta no prueban la identidad del contenido: el asesor
        # puede haber copiado por error el libro de otro expediente sobre este.
        celda_nif = database.CELDAS_RESUMEN[cartera]['nif']
        nif_libro = str(wb['00_RESUMEN'][celda_nif].value or '').strip().upper()
        if not nif_libro:
            raise ValueError('Falta el NIF/NIE en el resumen del Excel. No se ha importado: '
                             'comprueba que corresponde a este cliente y restaura su identificación.')
        if nif_libro != str((cliente or {}).get('nif') or '').strip().upper():
            raise ValueError('El Excel corresponde a otro cliente (NIF/NIE diferente). '
                             'No se ha importado ni sobrescrito ningún dato del expediente.')
        with db.transaccion():
            # Instantánea anterior a las notas generadas por la propia sincronización.
            anteriores_tablas = {t.nombre: {r["codigo"]: r for r in db.iterar(t.nombre, cliente_id, incluir_retirados=True)} for t, _ in hojas}
            for tabla, hoja in hojas:
                nombre = tabla.nombre
                ws = wb[hoja]
                columnas = _columnas(tabla)
                anteriores = anteriores_tablas[nombre]
                vistos = set()
                conteos[nombre] = 0
                for fila in range(tabla.fila_inicio, ws.max_row + 1):
                    datos = {}
                    for col in columnas:
                        cell = ws[f"{col.letra}{fila}"]
                        valor = (cell.hyperlink.target or cell.value) if col.tipo == "link" and cell.hyperlink else cell.value
                        if col.tipo == "link" and valor:
                            valor = str(valor).replace("\\", "/")
                        if isinstance(valor, datetime):
                            valor = valor.date()
                        if col.tipo == "pct" and isinstance(valor, (int, float)):
                            # Excel almacena porcentajes como fracción: 2 significa 200 %, nunca 2 %.
                            fraccion(valor)
                        if cell.data_type == "f":
                            raise ValueError(f"Fórmula en celda editable {hoja}!{cell.coordinate}")
                        datos[col.campo] = valor
                    if all(v in (None, "") for v in datos.values()):
                        continue
                    codigo = datos.get("codigo")
                    if nombre == "residencias":
                        if codigo in (None, ""):
                            raise ValueError(f"Falta ejercicio en {hoja}!A{fila}")
                        datos.update(codigo=str(int(codigo)), ejercicio=int(codigo))
                    elif codigo not in (None, ""):
                        datos["codigo"] = str(codigo).strip()
                    else:
                        datos.pop("codigo", None)
                        if nombre in MANUALES_EXCEL:
                            datos["origen"] = "excel"
                            # Comparar con la misma normalización que aplica upsert.
                            normalizados = {c.campo: database._normalizar(c, datos.get(c.campo))
                                            for c in columnas if c.campo != "codigo"}
                            for cod, r in anteriores.items():
                                if cod not in vistos and r.get("origen") == "excel" and all(
                                    (r.get(k) if r.get(k) != "" else None) == v for k, v in normalizados.items()
                                ):
                                    datos["codigo"] = cod
                                    break
                    codigo = datos.get('codigo')
                    if codigo is not None and codigo in vistos:
                        raise ValueError(f'Código duplicado {codigo} en {hoja}, fila {fila}. '
                                         'Revisa la fila copiada; no se ha importado el libro.')
                    previo = anteriores.get(codigo)
                    if nombre == "documentos" and previo:
                        gestion = db.con.execute("SELECT en_papelera FROM documentos_gestion WHERE cliente_id=? AND codigo=?",
                                                 (cliente_id, previo["codigo"])).fetchone()
                        if gestion and gestion["en_papelera"]:
                            vistos.add(previo["codigo"])
                            conteos[nombre] += 1
                            continue  # El Excel muestra la papelera, pero no cambia su ruta original ni la reactiva.
                    if previo and previo.get("estado") == "Presentado":
                        # La declaración ya presentada conserva todos sus datos, también ante vaciados.
                        vistos.add(previo["codigo"])
                        conteos[nombre] += 1
                        continue
                    if nombre == 'contratos':
                        campos_fecha = ('fecha_inicio', 'fecha_fin_prevista', 'fecha_fin_real')
                        fechas = {campo: a_fecha(datos.get(campo)) for campo in campos_fecha}
                        # No bloquear un histórico intacto por una anomalía anterior;
                        # sí impedir que una edición nueva introduzca un intervalo invertido.
                        cambio_fechas = not previo or any(
                            fechas[campo] != a_fecha(previo.get(campo)) for campo in campos_fecha)
                        if cambio_fechas and fechas['fecha_inicio'] and any(
                                fechas[campo] and fechas[campo] < fechas['fecha_inicio']
                                for campo in ('fecha_fin_prevista', 'fecha_fin_real')):
                            raise ValueError(f'{hoja}, fila {fila}: la fecha de fin del contrato '
                                             'no puede ser anterior al inicio. No se ha importado el libro.')
                    if nombre in ('m303', 'obligaciones') and datos.get('estado') == 'Presentado':
                        import fiscal_residentes as fr
                        modelo_303 = (datos if nombre == 'm303' else
                            db.get('m303', cliente_id, datos.get('id_modelo') or (previo or {}).get('id_modelo') or ''))
                        if modelo_303 and fr.pendientes_saldo_303(db, cliente_id,
                                int(modelo_303.get('ejercicio') or (previo or {}).get('ejercicio')),
                                modelo_303.get('periodo') or (previo or {}).get('periodo')):
                            raise ValueError(f'{hoja}, fila {fila}: confirma en la app el destino de los saldos 303 anteriores. '
                                             'No se ha importado el libro ni presentado nada.')
                    if (nombre == 'obligaciones' and previo
                            and str(previo.get('clave') or '').startswith('VENTA|')
                            and datos.get('estado') == 'Presentado'):
                        from tax_calculator import amortizacion_venta_confirmada
                        if amortizacion_venta_confirmada(db, cliente_id, previo.get('inmueble')) is None:
                            raise ValueError(f'{hoja}, fila {fila}: confirma primero en la app la amortización '
                                             'acumulada acreditada de la venta. El libro no se ha importado.')
                    if ((nombre == 'm210_imputacion' or
                         (nombre == 'obligaciones' and
                          (datos.get('tipo') or (previo or {}).get('tipo')) == '210 Imputación'))
                            and datos.get('estado') == 'Presentado'
                            and int(datos.get('ejercicio') or (previo or {}).get('ejercicio') or 0) >= 2027):
                        raise ValueError(f'{hoja}, fila {fila}: la previsión desde 2027 requiere validar '
                                         'la nueva normativa antes de presentarse desde este cálculo.')
                    if pendiente:
                        _comprobar_datos_pendientes(db, cliente_id, nombre, columnas, datos, previo, hoja, fila, ruta)
                    nuevo = db.upsert(nombre, cliente_id, datos, respetar_manuales=False)
                    vistos.add(nuevo)
                    if previo:
                        for col in columnas:
                            if col.campo != "codigo" and datos.get(col.campo) in (None, "") \
                                    and previo.get(col.campo) not in (None, ""):
                                db.vaciar_campo(nombre, cliente_id, nuevo, col.campo)
                    if nombre == "inmuebles":
                        actual = db.get(nombre, cliente_id, nuevo) or {}
                        db.proteger_campos_inmueble(cliente_id, nuevo, (
                            c.campo for c in columnas if (previo or {}).get(c.campo) != actual.get(c.campo)))
                    conteos[nombre] += 1
                for codigo, previo in anteriores.items():
                    if codigo in vistos or previo.get("estado") == "Presentado":
                        continue
                    if nombre in MANUALES_EXCEL and previo.get("origen") == "excel":
                        anterior = filas_exportadas.get(nombre, {}).get(codigo)
                        if pendiente and anterior is None and nombre in filas_exportadas:
                            # La fila nació en BD después del último libro confirmado:
                            # el usuario nunca la vio allí y no puede haberla borrado.
                            db.nota(cliente_id, "Sincronización", f"Fila nueva pendiente: {codigo}",
                                    f"{codigo} de {hoja} aún no estaba en el último Excel guardado; "
                                    "se conserva el registro nuevo de la aplicación.",
                                    clave=f"NUEVA_PENDIENTE|{nombre}|{codigo}")
                            continue
                        if pendiente and anterior != _firma_fila_manual(nombre, previo):
                            raise ExcelConflictoImportacion(
                                f'La fila {codigo} de {hoja} se eliminó del Excel, pero hay una exportación '
                                'pendiente y no se puede confirmar que la ficha siga igual. Se conservan ambas '
                                'versiones; no se ha borrado ni restaurado la fila automáticamente. '
                                'Revisa la copia pendiente y concilia el libro antes de sincronizar.')
                        db.borrar(nombre, cliente_id, codigo)
                    else:
                        db.nota(cliente_id, "Sincronización", f"Fila ausente: {codigo}",
                                f"{codigo} de {hoja} no aparece en el Excel; se conserva en la BD.",
                                clave=f"BORRADA|{nombre}|{codigo}")
            ws = wb["00_RESUMEN"]
            celdas = database.CELDAS_RESUMEN[cartera]
            campos = [c for c in ("pasaporte", "nacionalidad", "email", "telefono", "responsable") if c in celdas]
            if pendiente:
                resumen_anterior = json.loads(db.get_config(f"excel_resumen|{cliente_id}", "{}"))
                for campo in campos:
                    candidato, actual = ws[celdas[campo]].value, cliente.get(campo)
                    if campo in resumen_anterior and resumen_anterior[campo] != actual and candidato != actual:
                        _conflicto_excel(db, cliente_id, ruta, f"00_RESUMEN!{celdas[campo]}",
                                         campo.capitalize(), candidato, actual)
            db.actualizar_cliente(cliente_id, **{c: ws[celdas[c]].value for c in campos})
    finally:
        wb.close()
    return conteos


def sincronizar(db: Database, cliente_id: int, ejercicio: int) -> Path:
    from tax_calculator import recalcular_ejercicios

    import_excel_to_db(db, cliente_id)
    if cartera_de(db.get_cliente(cliente_id)["tipo_cliente"]) == RES:
        import fiscal_residentes
        for e in fiscal_residentes.ejercicios_encadenados(db, cliente_id, {int(ejercicio)}, int(ejercicio)):
            fiscal_residentes.recalcular_residente(db, cliente_id, e)
    else:
        existentes = {int(f["ejercicio"]) for f in db.listar("m210_alquiler", cliente_id) if f.get("ejercicio")}
        recalcular_ejercicios(db, cliente_id, existentes | {int(ejercicio)})
    return exportar_cliente(db, cliente_id, ejercicio)


def _celdas_imputacion_congeladas(ws) -> set[str]:
    """Renta y cuota de una imputación presentada: llevan el importe del justificante en lugar de la fórmula."""
    tabla = database.TABLAS["m210_imputacion"]
    if ws.title not in {tabla.hoja_en(c) for c in (NR, RES)}:
        return set()
    letra = {c.campo: c.letra for c in tabla.columnas_excel}
    if "estado" not in letra:
        return set()
    return {f"{letra[campo]}{fila}"
            for fila in range(tabla.fila_inicio, ws.max_row + 1) if ws[f"{letra['estado']}{fila}"].value == "Presentado"
            for campo in ("renta_imputada_formula", "cuota_formula")
            if campo in letra and type(ws[f"{letra[campo]}{fila}"].value) in (int, float)
            and math.isfinite(ws[f"{letra[campo]}{fila}"].value)}


def verificar_estructura(ruta: Path, plantilla_ruta: Path | None = None) -> list[str]:
    """Diferencias respecto a la plantilla (NR por defecto), sin evaluar fórmulas."""
    diferencias = []
    plantilla = openpyxl.load_workbook(plantilla_ruta or database.PLANTILLA_EXCEL)
    libro = openpyxl.load_workbook(ruta)
    try:
        if libro.sheetnames != plantilla.sheetnames:
            diferencias.append("Nombres u orden de las hojas diferentes")
        for original in plantilla:
            if original.title not in libro:
                continue
            ws = libro[original.title]
            if [c.value for c in original[1]] != [c.value for c in ws[1]]:
                diferencias.append(f"{ws.title}: fila 1 diferente")
            if {k: v.width for k, v in ws.column_dimensions.items()} != {
                k: v.width for k, v in original.column_dimensions.items()
            }:
                diferencias.append(f"{ws.title}: anchos de columna diferentes")
            def validaciones(hoja):
                return sorted((str(v.sqref), str(v.formula1))
                              for v in hoja.data_validations.dataValidation)
            if validaciones(ws) != validaciones(original):
                diferencias.append(f"{ws.title}: validaciones diferentes")
            combinadas = set(map(str, ws.merged_cells.ranges))
            esperadas = set(map(str, original.merged_cells.ranges))
            # Única extensión autorizada del resumen: el título del desglose de gastos.
            if ws.title == "00_RESUMEN" and ws["A29"].value in (
                    "GASTOS DEDUCIBLES DEL EJERCICIO POR INMUEBLE",
                    "GASTOS REGISTRADOS DEL EJERCICIO POR INMUEBLE"):
                esperadas.add("A29:H29")
            if combinadas != esperadas:
                diferencias.append(f"{ws.title}: celdas combinadas diferentes")
            if ws.max_column != original.max_column:
                diferencias.append(f"{ws.title}: número de columnas diferente")
            # También protege las fórmulas de resumen.
            congeladas = _celdas_imputacion_congeladas(ws)
            for fila in original:
                for celda in fila:
                    if celda.coordinate in congeladas:
                        continue
                    if (ws.title == "10_210_IMPUTACION" and celda.column_letter in ("H", "J")
                            and ws[f"A{celda.row}"].value not in (None, "")
                            and ws[f"O{celda.row}"].value != "Presentado"
                            and celda.data_type == "f"
                            and ws[celda.coordinate].value == _formula_imputacion_redondeada(celda.value)):
                        continue
                    if celda.data_type == "f" and ws[celda.coordinate].value != celda.value:
                        diferencias.append(f"{ws.title}!{celda.coordinate}: fórmula diferente o ausente")
    finally:
        plantilla.close()
        libro.close()
    return diferencias

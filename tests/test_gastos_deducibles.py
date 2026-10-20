"""Lector local de gastos y desglose del resumen, sin llamadas a IA."""
import hashlib
from copy import copy
from pathlib import Path

import openpyxl
import pytest

import database
import excel_generator as eg
import fixtures_pdf as f
import ia_extractor
import pdf_processor as pp


CASOS = [
    ("Suministros", "Luz"),
    ("Suministros", "Agua"),
    ("Suministros", "Gas"),
    ("Comunidad", "Cuota de comunidad"),
    ("Seguro", "Seguro del hogar"),
    ("Intereses/financiación", "Hipoteca"),
    ("Reparación/conservación", "Reparación de fontanería"),
]


@pytest.fixture
def db(tmp_path, monkeypatch):
    monkeypatch.setattr(ia_extractor, "disponible", lambda *args, **kwargs: False)
    d = database.Database(tmp_path / "gastos.db")
    d.carpeta_clientes = tmp_path / "Clientes"
    yield d
    d.cerrar()


@pytest.fixture
def cliente(db):
    return db.crear_cliente("Cliente", "Sintético", "X1234567A")


@pytest.mark.parametrize("categoria,concepto", CASOS)
def test_facturas_clasificacion_y_extraccion_local(tmp_path, categoria, concepto):
    hipoteca = categoria == "Intereses/financiación"
    ruta = f.factura_gasto(tmp_path / "entrada.pdf", categoria=categoria, concepto=concepto,
                           base=123.45, iva=25.92, total=700 if hipoteca else 149.37)
    anterior = ruta.read_bytes()
    ext = pp.extraer_local(ruta)
    assert pp.clasificar(ext.texto)[0] == pp.TipoDocumento.FACTURA_GASTO
    assert ext.tipo == pp.TipoDocumento.FACTURA_GASTO
    assert ext.motor == "local" and not ext.necesita_ocr
    assert ext.campos["categoria"] == categoria
    assert ext.campos["proveedor"] == "Servicios del Hogar SA"
    assert ext.campos["nif_proveedor"] == "A12345678"
    assert ext.campos["n_factura"] == "F-2026-0001"
    assert ext.campos["fecha"] == "2026-03-10"
    assert ext.campos["periodo_desde"] == "2026-02-01"
    assert ext.campos["periodo_hasta"] == "2026-02-28"
    assert ext.campos["ref_catastral"] == f.REF
    assert ext.campos["direccion"] == f.DIRECCION
    assert ext.campos["total"] == (123.45 if hipoteca else 149.37)
    if hipoteca:
        assert "intereses" in ext.campos["concepto"].lower()
        assert "base" not in ext.campos and "iva" not in ext.campos
    else:
        assert ext.campos["concepto"] == concepto
        assert ext.campos["base"] == 123.45
        assert ext.campos["iva"] == 25.92  # no el 21 % impreso en la etiqueta
    assert ruta.read_bytes() == anterior


@pytest.mark.parametrize("texto,tipo", [
    ("Recibo IBI. Tasa de basuras. Importe total: 60 euros", pp.TipoDocumento.FACTURA_GASTO),
    ("Factura. Recogida de residuos. Referencia catastral: " + f.REF, pp.TipoDocumento.FACTURA_GASTO),
    ("Recibo IBI. Impuesto sobre Bienes Inmuebles. Tasa de basuras. IVA. Importe total.",
     pp.TipoDocumento.CATASTRO_IBI),
    ("Factura. Valor catastral. Electricidad. Consumo. kWh. Importe total.", pp.TipoDocumento.CATASTRO_IBI),
    ("Recibo de comunidad de propietarios. Cuota de comunidad: 50 euros", pp.TipoDocumento.FACTURA_GASTO),
    ("Póliza. Prima anual de seguro del hogar: 200 euros", pp.TipoDocumento.FACTURA_GASTO),
    ("Préstamo hipotecario. Intereses: 40 euros", pp.TipoDocumento.FACTURA_GASTO),
    ("Presupuesto de reparación y mantenimiento: 90 euros", pp.TipoDocumento.FACTURA_GASTO),
    ("Honorarios de servicios profesionales: 120 euros", pp.TipoDocumento.FACTURA_GASTO),
])
def test_prioridad_de_clasificacion(texto, tipo):
    assert pp.clasificar(texto)[0] == tipo


@pytest.mark.parametrize("nif", ["A12345678", "B1234567J", "12345678Z", "X1234567A"])
def test_nif_y_preferencia_de_total(tmp_path, nif):
    ruta = f._crear(tmp_path / "entrada.pdf", f"""FACTURA DE SUMINISTRO DE ELECTRICIDAD
Proveedor: Empresa Sintética SL
CIF proveedor: {nif}
Factura nº: SUM-42
Fecha de factura: 10/03/2026
Periodo de facturación: 01/02/2026 - 28/02/2026
Base imponible: 1.000,00 euros
IVA 21%: 210,00 euros
Total: 999,00 euros
Total a pagar: 1.211,00 euros
Importe total: 1.212,00 euros
Total factura: 1.210,00 euros
""")
    campos = pp.extraer_local(ruta).campos
    assert campos["nif_proveedor"] == nif
    assert campos["n_factura"] == "SUM-42"
    assert campos["total"] == 1210
    assert campos["base"] == 1000 and campos["iva"] == 210
    assert campos["periodo_desde"] == "2026-02-01"
    assert campos["periodo_hasta"] == "2026-02-28"


@pytest.mark.parametrize("categoria,concepto", CASOS)
def test_aplicacion_local_a_bd(db, cliente, tmp_path, categoria, concepto):
    cid = cliente["id"]
    inmueble = db.upsert("inmuebles", cid, {"ref_catastral": f.REF, "direccion": f.DIRECCION})
    ruta = f.factura_gasto(tmp_path / "entrada.pdf", categoria=categoria, concepto=concepto,
                           base=100, total=700 if categoria == "Intereses/financiación" else 121)
    resultado = pp.procesar_pdf(db, cid, ruta, ejercicio=2026)
    gasto = db.get("gastos", cid, resultado.codigos["gasto"])
    assert gasto["inmueble"] == inmueble and gasto["ejercicio"] == 2026
    assert gasto["categoria"] == categoria and gasto["origen"] == "pdf"
    assert gasto["total"] == (100 if categoria == "Intereses/financiación" else 121)
    assert gasto["proveedor"] == "Servicios del Hogar SA"
    assert gasto["nif_proveedor"] == "A12345678" and gasto["n_factura"] == "F-2026-0001"
    assert gasto["fecha"] == "2026-03-10"
    assert gasto["periodo_desde"] == "2026-02-01" and gasto["periodo_hasta"] == "2026-02-28"
    assert gasto["observaciones"]
    assert gasto["ruta"] == resultado.ruta_relativa
    assert gasto["ruta"].startswith("Soportes_PDF/") and not Path(gasto["ruta"]).is_absolute()
    assert (Path(cliente["carpeta"]) / gasto["ruta"]).is_file()
    duplicado = pp.procesar_pdf(db, cid, ruta, ejercicio=2026)
    assert duplicado.duplicado and len(db.listar("gastos", cid)) == 1


@pytest.mark.parametrize("linea", ["", "Intereses:", "Intereses: no desglosados"])
def test_hipoteca_sin_intereses_no_aplica_cuota_completa(db, cliente, tmp_path, linea):
    cid = cliente["id"]
    db.upsert("inmuebles", cid, {"ref_catastral": f.REF})
    ruta = f._crear(tmp_path / "entrada.pdf", f"""RECIBO DE PRÉSTAMO HIPOTECARIO
Proveedor: Banco Sintético SA
Referencia catastral: {f.REF}
{linea}
Amortización de capital: 700,00 euros
Base imponible: 700,00 euros
Total a pagar: 700,00 euros
""")
    ext = pp.extraer_local(ruta)
    assert ext.tipo == pp.TipoDocumento.FACTURA_GASTO
    assert "total" not in ext.campos and "base" not in ext.campos
    assert any("intereses" in aviso.lower() for aviso in ext.avisos)
    pp.procesar_pdf(db, cid, ruta, ejercicio=2026)
    assert db.listar("gastos", cid) == []


def test_intereses_cero_son_un_importe_valido(tmp_path):
    ruta = f.factura_gasto(tmp_path / "entrada.pdf", categoria="Intereses/financiación",
                           concepto="Hipoteca", base=0, total=700)
    assert pp.extraer_local(ruta).campos["total"] == 0


@pytest.mark.parametrize("etiqueta", ["Tasa de basura", "Tasa de basuras", "Recogida de residuos"])
def test_tasa_basura_ibi_se_extrae_y_aplica(db, cliente, tmp_path, etiqueta):
    ruta = f.recibo_ibi(tmp_path / "entrada.pdf", tasa_basura=87.65)
    texto, _ = pp.extraer_texto(ruta)
    ruta = f._crear(tmp_path / "variante.pdf", texto.replace("Tasa de basuras", etiqueta))
    ext = pp.extraer_local(ruta)
    assert ext.tipo == pp.TipoDocumento.CATASTRO_IBI
    assert ext.campos["ibi_ejercicio"] == 450 and ext.campos["tasa_basura"] == 87.65
    pp.procesar_pdf(db, cliente["id"], ruta, ejercicio=2026)
    gastos = {g["categoria"]: g for g in db.listar("gastos", cliente["id"])}
    assert gastos["IBI"]["total"] == 450
    assert gastos["Tasa/basura"]["total"] == 87.65


@pytest.mark.parametrize("cantidad", [0, 1, 6, 8])
def test_bloque_excel_preserva_plantilla_y_formulas(db, cliente, cantidad):
    huella = hashlib.sha256(database.PLANTILLA_EXCEL.read_bytes()).hexdigest()
    codigos = [db.upsert("inmuebles", cliente["id"], {"direccion": f"Calle Prueba {i}"})
               for i in range(cantidad)]
    ruta = eg.exportar_cliente(db, cliente["id"], 2026)
    assert eg.verificar_estructura(ruta) == []
    wb = openpyxl.load_workbook(ruta)
    plantilla = openpyxl.load_workbook(database.PLANTILLA_EXCEL)
    try:
        assert wb.sheetnames == plantilla.sheetnames and len(wb.sheetnames) == 17
        ws = wb["00_RESUMEN"]
        original = plantilla["00_RESUMEN"]
        assert ws.max_column == original.max_column == 8
        assert ws["A29"].value == "GASTOS DEDUCIBLES DEL EJERCICIO POR INMUEBLE"
        assert "A29:H29" in set(map(str, ws.merged_cells.ranges))
        for atributo in ("font", "fill", "alignment"):
            assert copy(getattr(ws["A29"], atributo)) == copy(getattr(original["A18"], atributo))
        assert ws["A30"].value == "Categoría" and ws["B30"].value == "Total"
        assert copy(ws["A30"].font) == copy(original["A19"].font)
        visibles = codigos[:5 if cantidad > 6 else 6]
        for col, codigo in enumerate(visibles, 3):
            assert ws.cell(30, col).value.startswith(codigo + "\nCalle Prueba")
        for fila, categoria in enumerate(database.CATEGORIAS_GASTO, 31):
            assert ws.cell(fila, 1).value == categoria
            for col in range(2, 3 + len(visibles)):
                cell = ws.cell(fila, col)
                formula = cell.value
                assert cell.data_type == "f" and "$E$4" in formula
                assert "'08_GASTOS'!$R$2:$R$3000" in formula
                # Q (N*O*P de la plantilla) daría 0 con O/P vacíos: se usa N con O/P al 100 % por defecto
                assert "'08_GASTOS'!$Q$" not in formula
                for letra_rango in "NOP":
                    assert f"'08_GASTOS'!${letra_rango}$2:${letra_rango}$3000" in formula
                assert "'08_GASTOS'!$H$2:$H$3000" in formula
                assert f"$A{fila}" in formula
                assert '=""' in formula  # N×O×P solo si R está vacía; un R=0 sigue siendo final
                assert cell.number_format == '#,##0.00 [$€-es-ES]'
                if col >= 3:
                    assert "'08_GASTOS'!$C$2:$C$3000" in formula
                    assert f"LEFT({cell.column_letter}$30," in formula
            if cantidad > 6:
                assert ws["H30"].value == "Resto"
                assert ws[f"H{fila}"].value == f"=B{fila}-SUM(C{fila}:G{fila})"
        fila_total = 31 + len(database.CATEGORIAS_GASTO)
        assert ws.cell(fila_total, 1).value == "TOTAL"
        assert ws.cell(fila_total, 2).value == f"=SUM(B31:B{fila_total - 1})"
        assert "08_GASTOS (columna U)" in ws.cell(fila_total + 2, 1).value
        for hoja in plantilla:
            for fila in hoja:
                for celda in fila:
                    if celda.data_type == "f":
                        assert wb[hoja.title][celda.coordinate].value == celda.value
            assert wb[hoja.title].max_column == hoja.max_column
        # Celdas no editables del resumen original, incluidos textos y estilos, intactas.
        editables = set(database.RESUMEN_CELDAS.values())
        for fila in original:
            for celda in fila:
                if celda.coordinate not in editables:
                    assert ws[celda.coordinate].value == celda.value
                    assert ws[celda.coordinate]._style == celda._style
        # El bloque no crea registros al importar; el importador solo lee el esquema real.
        db.set_config(f"excel_firma|{cliente['id']}", "")
        eg.import_excel_to_db(db, cliente["id"])
        assert db.listar("gastos", cliente["id"]) == []
        assert len(db.listar("inmuebles", cliente["id"])) == cantidad
    finally:
        wb.close()
        plantilla.close()
    assert hashlib.sha256(database.PLANTILLA_EXCEL.read_bytes()).hexdigest() == huella


def test_desglose_rechaza_celdas_ocupadas_y_verificador_no_ignora_otras_combinaciones(db, cliente):
    wb = openpyxl.load_workbook(database.PLANTILLA_EXCEL)
    try:
        wb["00_RESUMEN"]["A31"] = "Celda reservada"
        with pytest.raises(ValueError, match="debe estar vacía"):
            eg._resumen_gastos(wb["00_RESUMEN"], [])
        assert wb["00_RESUMEN"]["A31"].value == "Celda reservada"
    finally:
        wb.close()
    ruta = eg.exportar_cliente(db, cliente["id"], 2026)
    wb = openpyxl.load_workbook(ruta)
    try:
        wb["00_RESUMEN"].merge_cells("A47:B47")
        wb.save(ruta)
    finally:
        wb.close()
    assert "00_RESUMEN: celdas combinadas diferentes" in eg.verificar_estructura(ruta)


@pytest.mark.parametrize("detalle,categoria,concepto", [
    ("Suministro de gas natural. Consumo: 250 kWh.", "Suministros", "Gas"),
    ("Suministro de agua. Consumo: 12 m3.", "Suministros", "Agua"),
    ("Suministro de electricidad. Consumo: 180 kWh.", "Suministros", "Luz"),
    ("Póliza de seguro de impago.", "Seguro", "Seguro de impago"),
    ("Póliza de responsabilidad civil.", "Seguro", "Seguro RC"),
    ("Póliza de hogar.", "Seguro", "Seguro del hogar"),
])
def test_categoria_y_concepto_se_infiere_del_servicio(tmp_path, detalle, categoria, concepto):
    ruta = f._crear(tmp_path / "entrada.pdf", f"""FACTURA O RECIBO
Proveedor: Empresa Sintética SA
{detalle}
Total factura: 100,00 euros
""")
    campos = pp.extraer_local(ruta).campos
    assert campos["categoria"] == categoria and campos["concepto"] == concepto


def test_concepto_generico_no_oculta_categoria_del_servicio(tmp_path):
    ruta = f._crear(tmp_path / "entrada.pdf", """FACTURA
Proveedor: Empresa Sintética SA
Concepto: Cuota mensual
Consumo: 180 kWh
Total factura: 100,00 euros
""")
    assert pp.extraer_local(ruta).campos["categoria"] == "Suministros"

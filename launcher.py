"""Entrada del ejecutable: instancia única y diagnóstico aislado del paquete."""
from __future__ import annotations
import json
import os
from pathlib import Path
import sys
import tempfile
import traceback


def diagnostico(salida: Path) -> None:
    # No usa el almacén real, incluso al probar un ejecutable recién construido.
    with tempfile.TemporaryDirectory(prefix="CAM-paquete-") as temporal:
        # Los runners Windows pueden redirigir TEMP mediante un junction.
        # Usamos el destino real de nuestro temporal, sin relajar la validación
        # de las carpetas de datos elegidas por el usuario.
        raiz = Path(temporal).resolve()
        os.environ["LOCALAPPDATA"] = str(raiz / "local")
        datos = raiz / "Datos ficticios"
        datos.mkdir()
        config = raiz / "local/ContabilidadAMedida/ubicacion.ini"
        config.parent.mkdir(parents=True)
        config.write_text("[Datos]\nCarpeta=" + str(datos) + "\n", encoding="utf-16")
        import app_paths
        import database
        import excel_generator
        import fitz
        import anthropic
        import keyring
        import flet
        import ui_actualizaciones
        import updater
        from version import VERSION
        from openpyxl import load_workbook
        assert app_paths.carpeta_datos() == datos
        for plantilla in database.PLANTILLAS.values():
            wb = load_workbook(plantilla, read_only=True)
            assert len(wb.sheetnames) >= 10
            wb.close()
        with database.Database(datos / "ficticia.db") as db:
            db.migrar_carteras()
            assert db.carpeta_base == datos
        cliente = app_paths.RECURSOS / "desktop/flet/flet.exe"
        assert cliente.is_file(), "No está incluido el cliente de escritorio de Flet"
        assert (app_paths.RECURSOS / "assets/icono.ico").is_file()
        assert (app_paths.RECURSOS / "packaging/aplicar_actualizacion.ps1").is_file()
        salida.write_text(json.dumps({"ok": True, "version": VERSION, "datos_reales": False}), encoding="utf-8")


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--diagnostico-paquete":
        try:
            diagnostico(Path(sys.argv[2]))
        except Exception as exc:
            try:
                Path(sys.argv[2]).write_text(json.dumps({"ok":False,"error":repr(exc),
                    "traceback":traceback.format_exc()}),encoding="utf-8")
            finally:
                # Un traceback sin capturar en el bootloader windowless abre
                # un diálogo modal y deja esperando al runner sin supervisión.
                sys.exit(1)
    else:
        if getattr(sys, 'frozen', False):
            os.environ['FLET_VIEW_PATH'] = str(Path(__file__).resolve().parent/'desktop/flet')
            os.chdir(Path(sys.executable).resolve().parent)
        import ctypes
        k32 = ctypes.WinDLL("kernel32", use_last_error=True)
        k32.CreateMutexW.argtypes = [ctypes.c_void_p, ctypes.c_bool, ctypes.c_wchar_p]
        k32.CreateMutexW.restype = ctypes.c_void_p
        mutex = k32.CreateMutexW(None, False, "Local\\ContabilidadAMedida")
        if not mutex:
            raise ctypes.WinError(ctypes.get_last_error())
        if ctypes.get_last_error() == 183:
            ctypes.windll.user32.MessageBoxW(None, "Contabilidad a Medida ya está abierta.", "Contabilidad a Medida", 0x40)
            sys.exit(0)
        import runpy
        # Conserva el arranque diferido original: errores del almacén se muestran en la ventana.
        runpy.run_module('main', run_name='__main__')

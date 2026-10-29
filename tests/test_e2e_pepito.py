"""Prueba E2E completa del expediente de Pepito Pérez mediante Controlador (sin UI)."""
import shutil
import sys
from datetime import date
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import excel_generator as eg
import fixtures_pdf as f
from main import Controlador


def _filas_hoja(ws):
    """Devuelve las filas de datos reales desde la fila 2."""
    out = []
    vacias = 0
    for fila in range(2, ws.max_row + 1):
        vals = (ws.cell(fila, 1).value, ws.cell(fila, 2).value, ws.cell(fila, 3).value)
        if any(v not in (None, "") for v in vals):
            out.append((fila, ws))
            vacias = 0
        else:
            vacias += 1
            if vacias >= 10:
                break
    return out


def test_e2e_pepito_flujo_completo(tmp_path):
    # Histórico anterior a 2024: conserva el ensayo trimestral y sus oráculos.
    # La agrupación anual moderna se comprueba en los tests fiscales y de ocupación.
    ctrl = Controlador(tmp_path / "contabilidad.db")
    ctrl.db.set_config("periodicidad_210_alquiler", "Trimestral")  # flujo histórico que presenta 3T
    ctrl.cambiar_carpeta_clientes(tmp_path / "Clientes")
    pdfs_dir = tmp_path / "pdfs_origen"
    pdfs_dir.mkdir(parents=True, exist_ok=True)

    try:
        # 1. Crear cliente Pepito Pérez (Francia, 2023) -> carpetas Soportes_PDF y _ARRASTRA_AQUI
        cliente = ctrl.crear_cliente("Pepito", "Pérez", "X1234567A", pais="Francia", ejercicio=2023)
        cid = cliente["id"]
        carpeta_cliente = Path(cliente["carpeta"])
        assert carpeta_cliente.name == "Pepito_Perez_X1234567A"
        assert (carpeta_cliente / "Soportes_PDF").is_dir()
        assert (carpeta_cliente / "_ARRASTRA_AQUI").is_dir()

        # 2. Escritura de compra 2019, 180.000 €, un solo titular, gananciales
        pdf_escritura = f.escritura_compra(
            pdfs_dir / "escritura_compra_2019.pdf",
            ref=f.REF,
            precio=180_000,
            fecha="2019-05-10",
            titulares=["Pepito Pérez"],
            gananciales=True,
        )
        r_esc = ctrl.procesar_pdf(cid, pdf_escritura, ejercicio=2023)
        assert r_esc["error_excel"] is None
        ruta_excel = r_esc["excel"]
        assert ruta_excel.is_file()

        wb = openpyxl.load_workbook(ruta_excel)
        try:
            ws_inm = wb["02_INMUEBLES"]
            assert ws_inm["A2"].value == "INM-001"
            assert ws_inm["F2"].value == f.REF
            assert ws_inm["H2"].value == 0.5
            assert ws_inm["J2"].value == 180_000
            assert ws_inm["Z2"].hyperlink is not None
            target_z2 = ws_inm["Z2"].hyperlink.target
            assert target_z2.startswith("Soportes_PDF/") and "\\" not in target_z2
            assert (carpeta_cliente / target_z2).is_file()
        finally:
            wb.close()

        # 3. IBI 2023 (VC 90.000, suelo 36.000, construcción 54.000, revisión 2019, cuota 450)
        pdf_ibi = f.recibo_ibi(
            pdfs_dir / "recibo_ibi_2023.pdf",
            ref=f.REF,
            ejercicio=2023,
            valor_catastral=90_000,
            valor_suelo=36_000,
            valor_construccion=54_000,
            anio_revision=2019,
            ibi=450,
        )
        def confirmar_ibi(ext):
            # La etiqueta genérica de la fixture exige confirmación humana de la revisión colectiva.
            assert ext.campos["anio_revision"] == 2019
            ext.revision_confirmada = True
            return ext

        r_ibi = ctrl.procesar_pdf(cid, pdf_ibi, ejercicio=2023, confirmar=confirmar_ibi)
        assert r_ibi["error_excel"] is None

        wb = openpyxl.load_workbook(ruta_excel)
        try:
            ws_cat = wb["03_CATASTRO_IBI"]
            assert ws_cat["B2"].value == 2023
            assert ws_cat["C2"].value == "INM-001"
            assert ws_cat["E2"].value == 90_000
            assert ws_cat["F2"].value == 36_000
            assert ws_cat["G2"].value == 54_000
            assert ws_cat["I2"].value == 2019
            assert ws_cat["M2"].value == 450

            ws_gas = wb["08_GASTOS"]
            assert ws_gas["A2"].value == "GAS-001"
            assert ws_gas["H2"].value == "IBI"
            assert ws_gas["N2"].value == 450
        finally:
            wb.close()

        # 4. Contrato 01/07/2023–30/06/2024, 900 €/mes, inquilina "Ana López"
        pdf_contrato = f.contrato_alquiler(
            pdfs_dir / "contrato_alquiler_2023.pdf",
            ref=f.REF,
            fecha_inicio="2023-07-01",
            fecha_fin_prevista="2024-06-30",
            renta_inicial=900,
            inquilinos=[{"nombre": "Ana López", "documento": "12345678Z", "pais": "España"}],
        )
        r_ctr = ctrl.procesar_pdf(cid, pdf_contrato, ejercicio=2023)
        assert r_ctr["error_excel"] is None

        wb = openpyxl.load_workbook(ruta_excel)
        try:
            ws_ctr = wb["04_CONTRATOS"]
            assert ws_ctr["A2"].value == "CTR-001"
            assert ws_ctr["B2"].value == "INM-001"
            assert ws_ctr["H2"].value == 900

            ws_inq = wb["05_INQUILINOS"]
            assert ws_inq["A2"].value == "INQ-001"
            assert ws_inq["D2"].value == "Ana López"

            ws_210a = wb["09_210_ALQUILER"]
            periodos_2023 = {
                ws_210a[f"E{fila}"].value: ws_210a[f"R{fila}"].value
                for fila, _ in _filas_hoja(ws_210a)
                if ws_210a[f"B{fila}"].value == 2023
            }
            assert set(periodos_2023.keys()) == {"3T", "4T"}

            ws_210i = wb["10_210_IMPUTACION"]
            dias_2023 = [
                ws_210i[f"D{fila}"].value
                for fila, _ in _filas_hoja(ws_210i)
                if ws_210i[f"B{fila}"].value == 2023
            ]
            assert dias_2023 == [181]
        finally:
            wb.close()

        # 5. 210 del 3T presentado (NRC ficticio) -> 09 3T "Presentado", 4T "Pendiente"; 14 refleja ambos
        pdf_210_3t = f.modelo_210(
            pdfs_dir / "modelo_210_3T_2023.pdf",
            ref=f.REF,
            ejercicio=2023,
            periodo="3T",
            tipo_renta="01",
            base=884.96,
            cuota=168.14,
            fecha_presentacion="2023-10-15",
            nrc="NRC-PEPITO-3T-2023",
        )
        r_210 = ctrl.procesar_pdf(cid, pdf_210_3t, ejercicio=2023)
        assert r_210["error_excel"] is None

        cuota_4t_antes = next(
            r["cuota"]
            for r in ctrl.db.listar("m210_alquiler", cid, ejercicio=2023)
            if r["periodo"] == "4T"
        )

        wb = openpyxl.load_workbook(ruta_excel)
        try:
            ws_210a = wb["09_210_ALQUILER"]
            filas_210_2023 = {
                ws_210a[f"E{fila}"].value: (
                    ws_210a[f"A{fila}"].value,
                    ws_210a[f"R{fila}"].value,
                    ws_210a[f"Q{fila}"].value,
                )
                for fila, _ in _filas_hoja(ws_210a)
                if ws_210a[f"B{fila}"].value == 2023
            }
            cod_3t, est_3t, nrc_3t = filas_210_2023["3T"]
            cod_4t, est_4t, _ = filas_210_2023["4T"]
            assert est_3t == "Presentado" and nrc_3t == "NRC-PEPITO-3T-2023"
            assert est_4t == "Pendiente"

            ws_obl = wb["14_OBLIGACIONES"]
            obl_por_modelo = {
                ws_obl[f"J{fila}"].value: ws_obl[f"E{fila}"].value
                for fila, _ in _filas_hoja(ws_obl)
                if ws_obl[f"B{fila}"].value == 2023 and ws_obl[f"C{fila}"].value == "210 Alquiler"
            }
            assert obl_por_modelo[cod_3t] == "Presentado"
            assert obl_por_modelo[cod_4t] == "Pendiente"

            # 6. 11_PATRIMONIO_714 observaciones contienen "Tiene 1 piso," y "NO está obligado"
            ws_pat = wb["11_PATRIMONIO_714"]
            obs_pat_2023 = [
                ws_pat[f"T{fila}"].value
                for fila, _ in _filas_hoja(ws_pat)
                if ws_pat[f"B{fila}"].value == 2023
            ]
            assert len(obs_pat_2023) == 1
            assert "Tiene 1 piso," in obs_pat_2023[0]
            assert "NO está obligado" in obs_pat_2023[0]
        finally:
            wb.close()

        # 7. Añade en 08_GASTOS una fila SIN ID (Comunidad, 2023-11-15, 300 €, INM-001), guarda y procesa valor de referencia
        wb = openpyxl.load_workbook(ruta_excel)
        ws_gas = wb["08_GASTOS"]
        ws_gas["A3"] = None
        ws_gas["B3"] = 2023
        ws_gas["C3"] = "INM-001"
        ws_gas["E3"] = date(2023, 11, 15)
        ws_gas["H3"] = "Comunidad"
        ws_gas["N3"] = 300
        wb.save(ruta_excel)
        wb.close()

        pdf_val_ref = f.valor_referencia(
            pdfs_dir / "valor_referencia_2023.pdf",
            ref=f.REF,
            ejercicio=2023,
            valor=210_000,
        )
        r_ref = ctrl.procesar_pdf(cid, pdf_val_ref, ejercicio=2023)
        assert r_ref["error_excel"] is None

        cuota_4t_despues = next(
            r["cuota"]
            for r in ctrl.db.listar("m210_alquiler", cid, ejercicio=2023)
            if r["periodo"] == "4T"
        )
        assert cuota_4t_despues < cuota_4t_antes
        assert cuota_4t_despues == pytest.approx(cuota_4t_antes - 150 * 0.19, abs=0.02)

        wb = openpyxl.load_workbook(ruta_excel)
        try:
            ws_gas = wb["08_GASTOS"]
            gastos_excel = [
                (ws_gas[f"A{fila}"].value, ws_gas[f"H{fila}"].value, ws_gas[f"N{fila}"].value)
                for fila, _ in _filas_hoja(ws_gas)
            ]
            gasto_comunidad = [g for g in gastos_excel if g[1] == "Comunidad"]
            assert len(gasto_comunidad) == 1
            assert str(gasto_comunidad[0][0]).startswith("GAS-")
            assert gasto_comunidad[0][2] == 300

            docs_antes = len(_filas_hoja(wb["13_DOCUMENTACION"]))
        finally:
            wb.close()

        # 8. Repite el PDF de la escritura -> duplicado=True, 13_DOCUMENTACION no crece
        r_dup = ctrl.procesar_pdf(cid, pdf_escritura, ejercicio=2023)
        assert r_dup["resultado"].duplicado is True
        wb = openpyxl.load_workbook(ruta_excel)
        try:
            docs_despues = len(_filas_hoja(wb["13_DOCUMENTACION"]))
            assert docs_despues == docs_antes
        finally:
            wb.close()

        # 9. verificar_estructura(excel) == []
        assert eg.verificar_estructura(ruta_excel) == []

        # 10. procesar_bandeja: PDF válido en _ARRASTRA_AQUI se procesa y borra; PDF corrupto va a _errores
        bandeja = carpeta_cliente / "_ARRASTRA_AQUI"
        pdf_res = f.certificado_residencia(
            pdfs_dir / "certificado_residencia_2023.pdf",
            pais="Francia",
            ejercicio=2023,
        )
        en_bandeja_ok = bandeja / "certificado_residencia_2023.pdf"
        shutil.copy2(pdf_res, en_bandeja_ok)

        en_bandeja_corrupto = bandeja / "corrupto.pdf"
        en_bandeja_corrupto.write_bytes(b"NO ES UN PDF VALIDO")

        res_bandeja = ctrl.procesar_bandeja(cid, ejercicio=2023)
        assert len(res_bandeja) == 2
        por_nombre = {r["archivo"]: r for r in res_bandeja}
        assert por_nombre["certificado_residencia_2023.pdf"]["ok"] is True
        assert not en_bandeja_ok.exists()

        assert por_nombre["corrupto.pdf"]["ok"] is False
        assert not en_bandeja_corrupto.exists()
        assert (bandeja / "_errores" / "corrupto.pdf").is_file()
        assert (bandeja / "_errores" / "corrupto.txt").is_file()
        assert eg.verificar_estructura(ruta_excel) == []
    finally:
        ctrl.cerrar()


def test_controlador_metodos_auxiliares_y_excel_bloqueado(tmp_path, monkeypatch):
    ctrl = Controlador(tmp_path / "auxiliar.db")
    ctrl.cambiar_carpeta_clientes(tmp_path / "Clientes")
    try:
        c = ctrl.crear_cliente("María", "Gómez", "Y9999999Z", pais="Alemania", ejercicio=2026)
        cid = c["id"]
        pdf_esc = f.escritura_compra(tmp_path / "esc.pdf", gananciales=False)
        ext = ctrl.extraer(pdf_esc)
        assert ext.campos["precio_escritura"] == 180_000

        ctrl.procesar_pdf(cid, pdf_esc, ejercicio=2026)
        r_gan = ctrl.set_gananciales(cid, "INM-001", True, ejercicio=2026)
        assert ctrl.db.get("inmuebles", cid, "INM-001")["pct_titularidad"] == 0.5
        assert r_gan["error_excel"] is None

        r_gas = ctrl.alta_gasto(cid, 2026, "INM-001", "Seguro", "2026-03-10", 200, proveedor="Mapfre")
        assert r_gas["gasto"].startswith("GAS-")
        assert ctrl.db.get("gastos", cid, r_gas["gasto"])["origen"] == "app"

        r_ven = ctrl.marcar_vendido(cid, "INM-001", "2026-09-01", 260_000, ejercicio=2026)
        assert ctrl.db.get("inmuebles", cid, "INM-001")["estado"] == "Vendido"
        assert r_ven["error_excel"] is None

        # Simular ExcelBloqueadoError: procesar_pdf no relanza, devuelve mensaje en error_excel
        def fallo_permiso(*_args, **_kwargs):
            raise PermissionError("Bloqueado por Excel")

        monkeypatch.setattr(eg.os, "replace", fallo_permiso)
        pdf_ibi = f.recibo_ibi(tmp_path / "ibi.pdf")
        r_bloq = ctrl.procesar_pdf(cid, pdf_ibi, ejercicio=2026)
        assert r_bloq["error_excel"] is not None
        assert "Cierra el Excel" in r_bloq["error_excel"]
    finally:
        ctrl.cerrar()

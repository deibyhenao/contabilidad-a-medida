"""Regresiones de la auditoría de 56d61b5, sin red ni expedientes reales."""
from types import SimpleNamespace as NS
import threading
import math
import openpyxl
import pytest
import ia_extractor as ia
import pdf_processor as pp
import excel_generator as eg
from controlador import Controlador
from test_enrutado_ia import pdf, _simular, _ext
from test_catastro_ejercicio_revision import expediente


@pytest.mark.parametrize('mensaje',[
    'La clave API no es válida.', 'Sin conexión con el servicio de IA.',
    'Límite de uso de la API alcanzado.', 'Saldo insuficiente.',
    'No se pudo leer el documento.',
])
def test_error_operativo_no_llama_a_sonnet(pdf,monkeypatch,mensaje):
    llamadas=_simular(monkeypatch,{'claude-haiku-5-5':ia.ErrorIA(mensaje),
                                 'claude-sonnet-5-5':_ext(pp.TipoDocumento.CATASTRO_IBI,.9)})
    ext=pp.extraer(pdf,ia.AjustesIA())
    assert llamadas==['claude-haiku-5-5'] and ext.motor=='local'


@pytest.mark.parametrize('valor',[None,[],{'tipo':'sin-esquema'}])
def test_esquema_invalido_si_admite_respaldo(valor):
    with pytest.raises(ia.ErrorIA) as error:
        ia.a_extraccion(valor,'ficticio.pdf')
    assert error.value.admite_respaldo


def test_respaldo_respeta_esfuerzo_y_se_limita_a_dos_lecturas(pdf,monkeypatch):
    llamadas=[]
    def extraer(ruta,ajustes,**kw):
        llamadas.append((ajustes.modelo,ajustes.esfuerzo,ajustes.respaldo,ajustes.permitir_respaldo_servidor))
        return _ext(pp.TipoDocumento.DESCONOCIDO,.1)
    monkeypatch.setattr(ia,'extraer',extraer)
    pp.extraer(pdf,ia.AjustesIA(esfuerzo='medium'))
    assert llamadas==[('claude-haiku-5-5','medium','claude-sonnet-5-5',True),
                      ('claude-sonnet-5-5','medium','',False)]


@pytest.mark.parametrize('modelo', ['claude-haiku-5-5','claude-sonnet-5-5'])
def test_coste_pdf_coincide_con_chat_incluida_cache(modelo):
    import asistente
    uso=NS(input_tokens=100,output_tokens=10,cache_creation_input_tokens=1000,
           cache_read_input_tokens=2000,cache_creation=NS(ephemeral_1h_input_tokens=400))
    assert ia.coste_estimado(modelo,uso)==pytest.approx(asistente.coste_chat(modelo,uso),abs=1e-8)


def test_haiku_largo_cuenta_cache_en_umbral():
    uso=NS(input_tokens=1,output_tokens=20,cache_read_input_tokens=100000)
    assert ia.coste_estimado('claude-haiku-5-5',uso)==pytest.approx((.5+20*2.5+100000*.05)/1e6,abs=1e-8)


@pytest.mark.parametrize('incluye_estado',[False,True])
def test_presentada_no_se_traslada_a_destino_provisional(expediente,monkeypatch,incluye_estado):
    db,cid,inm,_,_=expediente
    otro=db.upsert('inmuebles',cid,{'direccion':'Ficticia B','fecha_adquisicion':'2024-01-01'})
    db.upsert('catastro',cid,{'inmueble':otro,'ejercicio':2026,'valor_catastral':90000,'revision_pendiente':1},clave=f'CAT|{otro}|2026')
    codigo=db.upsert('m210_imputacion',cid,{'inmueble':inm,'ejercicio':2025,'estado':'Presentado','cuota':123.45})
    ctrl=object.__new__(Controlador);ctrl.db=db;ctrl._lock=threading.RLock()
    monkeypatch.setattr(ctrl,'_importar',lambda *a:None)
    monkeypatch.setattr(ctrl,'_recalcular_y_exportar',lambda *a:{})
    anterior=db.get('m210_imputacion',cid,codigo)
    datos={'inmueble':otro,'ejercicio':2026}
    if incluye_estado:datos['estado']='Presentado'
    with pytest.raises(ValueError,match='provisional'):
        ctrl.editar(cid,'m210_imputacion',codigo,datos,2026)
    assert db.get('m210_imputacion',cid,codigo)==anterior
    ctrl.editar(cid,'m210_imputacion',codigo,{'nrc':'NRC-FICTICIO','estado':'Presentado'},2025)
    assert db.get('m210_imputacion',cid,codigo)['nrc']=='NRC-FICTICIO'
    assert db.get('m210_imputacion',cid,codigo)['cuota']==123.45


@pytest.mark.parametrize('valor',[True,False,float('nan'),float('inf'),-float('inf')])
def test_importe_no_numerico_finito_no_se_exime_de_formula(valor):
    from database import TABLAS
    t=TABLAS['m210_imputacion'];letras={c.campo:c.letra for c in t.columnas_excel}
    wb=openpyxl.Workbook();ws=wb.active;ws.title=t.hoja
    ws[f"{letras['estado']}2"]='Presentado'
    celda=f"{letras['cuota_formula']}2";ws[celda]=valor
    assert celda not in eg._celdas_imputacion_congeladas(ws)
    wb.close()


def test_sonnet_de_respaldo_no_delega_a_otro_modelo(pdf,monkeypatch):
    from test_ia_extractor import ClienteFalso,RESPUESTA_ESCRITURA
    falso=ClienteFalso(RESPUESTA_ESCRITURA)
    monkeypatch.setattr(ia,'_cliente',lambda:falso)
    ia.leer_con_claude(pdf,ia.AjustesIA(modelo='claude-sonnet-5-5',permitir_respaldo_servidor=False))
    assert falso.ruta=='normal'
    assert 'fallbacks' not in falso.peticion and 'betas' not in falso.peticion


@pytest.mark.parametrize('delta,cuantas',[(0,1),(-.0001,2)])
def test_umbral_no_escala_al_igualarlo(pdf,monkeypatch,delta,cuantas):
    # El umbral bajó a 0,5 tras la prueba real (Sonnet no mejoró una lectura de Haiku a 0,70)
    llamadas=_simular(monkeypatch,{'claude-haiku-5-5':_ext(pp.TipoDocumento.CATASTRO_IBI,ia.UMBRAL_RESPALDO+delta),
                                 'claude-sonnet-5-5':_ext(pp.TipoDocumento.CATASTRO_IBI,.9)})
    pp.extraer(pdf,ia.AjustesIA())
    assert len(llamadas)==cuantas

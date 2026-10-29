"""Decisiones del asesor sobre un 303 negativo: arrastre, devolución, revisión y conservación."""
import json
from pathlib import Path

import openpyxl
import pytest

import database
import fiscal_residentes as fr
from controlador import Controlador


@pytest.fixture
def caso303(tmp_path, monkeypatch):
    c=Controlador(tmp_path/'base.db')
    c.cambiar_carpeta_clientes(tmp_path/'Datos')
    cli=c.db.crear_cliente('FICTICIO','Saldo303','QASALDO303',tipo_cliente='Autónomo')
    cid=cli['id']
    c.db.upsert('actividades',cid,{'regimen_iva':'General','regimen_irpf':'Directa simplificada'})
    codigo=c.db.upsert('m303',cid,{'ejercicio':2024,'periodo':'4T','estado':'Presentado',
        'resultado':-500,'a_ingresar':-500,'a_compensar':500,'nrc':'JUSTIFICANTE_FICTICIO'},clave='303|2024|4T')
    c.db.upsert('facturas_emitidas',cid,{'fecha':'2025-02-01','base':1000,'cuota_iva':210})
    monkeypatch.setattr(c,'_importar',lambda *a:None)
    monkeypatch.setattr(c,'_exportar_informe',lambda cid,ej,r:r)
    try:yield c,cid,codigo
    finally:c.cerrar()


@pytest.mark.parametrize('destino,saldo,esperado,pendiente',[
    ('compensar',500,0,False),('devolver',0,210,False),('pendiente',0,210,True)])
def test_decision_303_cambia_solo_arrastres_y_conserva_presentado(caso303,destino,saldo,esperado,pendiente):
    c,cid,codigo=caso303
    antes=c.db.get('m303',cid,codigo)
    c.confirmar_saldo_303(cid,codigo,2025,destino,saldo,'Justificante contrastado',fr.firma_saldo_303(antes),confirmado=True)
    assert c.db.get('m303',cid,codigo)==antes
    siguiente=c.db.get_por_clave('m303',cid,'303|2025|1T')
    assert siguiente['a_ingresar']==esperado
    assert fr.saldo_303(c.db,cid,antes)['pendiente'] is pendiente
    assert bool(fr.pendientes_saldo_303(c.db,cid,2025,'1T')) is pendiente
    if pendiente:
        assert 'PROVISIONAL' in siguiente['observaciones']


def test_cambiar_compensar_devolver_y_pendiente_no_duplica_ni_cambia_justificante(caso303):
    c,cid,codigo=caso303
    antes=c.db.get('m303',cid,codigo);firma=fr.firma_saldo_303(antes)
    for destino,saldo,cuota in [('compensar',500,0),('devolver',0,210),('compensar',700,0),('pendiente',0,210)]:
        c.confirmar_saldo_303(cid,codigo,2025,destino,saldo,'Historial contrastado',firma,confirmado=True)
        assert c.db.get_por_clave('m303',cid,'303|2025|1T')['a_ingresar']==cuota
        assert c.db.get('m303',cid,codigo)==antes
    assert len(c.db.listar('m303',cid,ejercicio=2025))==4


@pytest.mark.parametrize('destino,saldo,confirmado',[
    ('compensar',500,False),('compensar',-1,True),('compensar','50%',True),
    ('compensar',True,True),('compensar',float('nan'),True),('compensar',499,True),('inventado',500,True)])
def test_no_confirmar_o_importe_invalido_no_modifica_nada(caso303,destino,saldo,confirmado):
    c,cid,codigo=caso303
    antes=c.db.get('m303',cid,codigo)
    with pytest.raises(ValueError):
        c.confirmar_saldo_303(cid,codigo,2025,destino,saldo,'Justificante contrastado',fr.firma_saldo_303(antes),confirmado=confirmado)
    assert c.db.get('m303',cid,codigo)==antes
    assert c.db.get_config(f'saldo303|{cid}|{codigo}') is None


def test_devolucion_intermedia_no_admitida_en_motor_trimestral_general(caso303):
    c,cid,codigo=caso303
    c.db.upsert('m303',cid,{'codigo':codigo,'periodo':'2T'})
    fila=c.db.get('m303',cid,codigo)
    with pytest.raises(ValueError,match='4T'):
        c.confirmar_saldo_303(cid,codigo,2025,'devolver',0,'Justificante contrastado',fr.firma_saldo_303(fila),confirmado=True)


def test_justificante_cambiado_invalida_decision_y_bloquea_presentacion_dependiente(caso303):
    c,cid,codigo=caso303
    fila=c.db.get('m303',cid,codigo)
    c.confirmar_saldo_303(cid,codigo,2025,'compensar',500,'Justificante contrastado',fr.firma_saldo_303(fila),confirmado=True)
    c.db.upsert('m303',cid,{'codigo':codigo,'a_ingresar':-600})
    assert fr.saldo_303(c.db,cid,c.db.get('m303',cid,codigo))['pendiente']
    with pytest.raises(ValueError,match='cambió'):
        c.confirmar_saldo_303(cid,codigo,2025,'compensar',600,'Justificante contrastado',fr.firma_saldo_303(fila),confirmado=True)
    siguiente=c.db.get_por_clave('m303',cid,'303|2025|1T')
    with pytest.raises(ValueError,match='provisional'):
        c.marcar_presentado(cid,'m303',siguiente['codigo'],2025)


def test_303_rollback_si_falla_recalculo(caso303,monkeypatch):
    c,cid,codigo=caso303
    fila=c.db.get('m303',cid,codigo)
    monkeypatch.setattr(c,'_recalcular_bd',lambda *a:(_ for _ in ()).throw(RuntimeError('fallo provocado')))
    with pytest.raises(RuntimeError,match='fallo provocado'):
        c.confirmar_saldo_303(cid,codigo,2025,'compensar',500,'Justificante contrastado',fr.firma_saldo_303(fila),confirmado=True)
    assert c.db.get_config(f'saldo303|{cid}|{codigo}') is None


def test_confirmacion_303_excel_explica_historico_y_aplica_saldo_en_siguiente_periodo(caso303):
    import excel_generator as eg
    c,cid,codigo=caso303
    fila=c.db.get('m303',cid,codigo)
    c.confirmar_saldo_303(cid,codigo,2025,'devolver',0,'Justificante devolución 4T',fr.firma_saldo_303(fila),confirmado=True)
    ruta=eg.exportar_cliente(c.db,cid,2025)
    wb=openpyxl.load_workbook(ruta)
    try:
        tabla=database.TABLAS['m303'];ws=wb[tabla.hoja]
        col=next(c.letra for c in tabla.columnas_excel if c.campo=='a_compensar')
        assert ws[col+'2'].value==500
        assert 'devolver' in ws[col+'2'].comment.text
        assert '0.00 euros' in ws[col+'2'].comment.text
        col_ingreso=next(c.letra for c in tabla.columnas_excel if c.campo=='a_ingresar')
        assert ws[col_ingreso+'2'].value==0
        assert '-500.00' in ws[col_ingreso+'2'].comment.text
    finally:wb.close()
    eg.import_excel_to_db(c.db,cid)
    assert c.db.get('m303',cid,codigo)==fila
    assert c.db.get_por_clave('m303',cid,'303|2025|1T')['a_ingresar']==210


def test_303_negativo_presentado_no_resta_otros_impuestos_del_resumen(caso303):
    c,cid,codigo=caso303
    informe=fr.recalcular_residente(c.db,cid,2024)
    assert informe.trimestres[3].a_ingresar_303==0
    assert informe.total_a_ingresar>=0
    assert c.db.get('m303',cid,codigo)['a_ingresar']==-500


@pytest.mark.parametrize('valor',[True,'500',float('nan'),499])
def test_configuracion_saldo_corrupta_no_habilita_credito(caso303,valor):
    c,cid,codigo=caso303
    fila=c.db.get('m303',cid,codigo)
    c.db.set_config(f'saldo303|{cid}|{codigo}',json.dumps({'version':1,'firma':fr.firma_saldo_303(fila),
        'destino':'compensar','saldo':valor,'justificacion':'Nota aparentemente válida'}))
    assert fr.saldo_303(c.db,cid,fila)['pendiente']


def test_resultado_negativo_con_a_ingresar_cero_tambien_exige_saldo_minimo(caso303):
    c,cid,codigo=caso303
    c.db.upsert('m303',cid,{'codigo':codigo,'a_ingresar':0,'resultado':-500})
    fila=c.db.get('m303',cid,codigo)
    assert fr.minimo_saldo_303(fila)==500
    with pytest.raises(ValueError,match='inferior'):
        c.confirmar_saldo_303(cid,codigo,2025,'compensar',0,'Justificante contrastado',fr.firma_saldo_303(fila),confirmado=True)


def test_minimo_prioriza_importe_del_justificante_frente_a_estimacion_calculada():
    assert fr.minimo_saldo_303({'a_ingresar':-500,'resultado':-1000})==500

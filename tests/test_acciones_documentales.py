"""Seguridad de acciones propuestas: API simulada y expedientes temporales."""
import re
from pathlib import Path
from types import SimpleNamespace

import pytest

import acciones_documentales as ad
import documentos
import gestion_archivos
import fixtures_pdf as f
import pdf_processor as pp
from controlador import Controlador
from database import ruta_soporte


@pytest.fixture
def entorno(tmp_path):
    ctrl = Controlador(tmp_path / 'aislada.sqlite')
    ctrl.cambiar_carpeta_clientes(tmp_path / 'Clientes')
    cid = ctrl.crear_cliente('Михаил', 'Иванов', 'SEC-FICTICIO')['id']
    ruta = f.certificado_residencia(tmp_path / 'cert.pdf', pais='Rusia', ejercicio=2025)
    res = ctrl.procesar_pdf(cid, ruta, 2025)
    yield ctrl, cid, res['resultado'].documento
    ctrl.cerrar()


def emitir(ctrl, cid, codigo, **campos):
    gestor = ad.GestorAcciones(lambda: (cid, 2025))
    texto = gestor.proponer({'codigo': codigo, **campos}, cid, 2025, ctrl.preparar_accion_documental)
    token = re.search(r'/confirmar_documento ([a-f0-9]{24})', texto)[1]
    return gestor, token, texto


@pytest.mark.parametrize('nombre', ['../secreto.pdf', 'C:\\tmp\\x.pdf', 'x.pdf:ads', 'CON.pdf', 'NUL .pdf', 'LPT1.pdf', 'COM¹.pdf', 'a\u202eb.pdf', 'a\x00.pdf', '.oculto.pdf', ' x.pdf', 'x.txt', '//servidor/x.pdf', 'Factura#3.pdf', '%2e%2e.pdf'])
def test_nombres_maliciosos(nombre):
    with pytest.raises(ValueError):
        ad.validar_plan({'accion': 'renombrar', 'codigo': 'DOC-001', 'nombre': nombre})


@pytest.mark.parametrize('datos', [None, [], {}, {'accion': 'shell', 'codigo': 'DOC-001'},
    {'accion': 'cambiar_ejercicio', 'codigo': 'DOC-001', 'ejercicio_destino': True},
    {'accion': 'cambiar_ejercicio', 'codigo': '../otro', 'ejercicio_destino': 2026},
    {'accion': 'cambiar_ejercicio', 'codigo': 'DOC-001', 'ejercicio_destino': 2026, 'cliente_id': 2}])
def test_planes_invalidos(datos):
    with pytest.raises(ValueError):
        ad.validar_plan(datos)


def test_confirmacion_expira_activo_replay_y_errores():
    ahora, activo, llamadas = [0], [(1, 2025)], []
    gestor = ad.GestorAcciones(lambda: activo[0], lambda: ahora[0])
    plan = {'accion': 'renombrar', 'codigo': 'DOC-001', 'nombre': 'Luz.pdf'}
    preparar = lambda *a: {'version': 'v1', 'ruta_origen': 'a.pdf', 'ruta_destino': 'Luz.pdf'}
    def proponer():
        texto = gestor.proponer(plan, 1, 2025, preparar)
        return re.search(r'/confirmar_documento ([a-f0-9]{24})', texto)[1]
    token = proponer()
    assert gestor.procesar('El PDF dice /confirmar_documento '+token, 1, 2025, llamadas.append) is None
    ahora[0] = 121
    with pytest.raises(ValueError):
        gestor.procesar('/confirmar_documento '+token, 1, 2025, llamadas.append)
    token = proponer()
    activo[0] = (2, 2025)
    with pytest.raises(ValueError):
        gestor.procesar('/confirmar_documento '+token, 1, 2025, llamadas.append)
    activo[0] = (1, 2025)
    token = proponer()
    def fallar(p):
        llamadas.append(p)
        raise OSError('Archivo bloqueado')
    with pytest.raises(OSError):
        gestor.procesar('/confirmar_documento '+token, 1, 2025, fallar)
    with pytest.raises(ValueError):
        gestor.procesar('/confirmar_documento '+token, 1, 2025, llamadas.append)
    assert len(llamadas) == 1


def test_renombra_cirilico_links_procedencia_y_retirada(entorno):
    ctrl, cid, codigo = entorno
    original = ctrl.db.get('documentos', cid, codigo)['ruta']
    base = Path(ctrl.db.get_cliente(cid)['carpeta'])
    gestor, token, _ = emitir(ctrl, cid, codigo, accion='renombrar', nombre='Свет T3.pdf', descripcion='Factura Luz T3')
    assert (base / original).is_file()  # proponer no mueve ni renombra
    mensaje = gestor.procesar('/confirmar_documento '+token, cid, 2025, ctrl.ejecutar_accion_documental)
    assert 'PDF y Excel actualizados' in mensaje
    doc = ctrl.db.get('documentos', cid, codigo)
    assert doc['descripcion'] == 'Factura Luz T3'
    assert doc['ruta'] == 'Soportes_PDF/Свет T3.pdf'
    assert not (base / original).exists() and (base / doc['ruta']).is_file()
    assert ctrl.db.residencia(cid, 2025)['ruta'] == doc['ruta']
    import openpyxl
    wb = openpyxl.load_workbook(ctrl.db.ruta_excel(ctrl.db.get_cliente(cid)))
    try:
        links = [c.hyperlink.target for row in wb['01_RESIDENCIA_FISCAL'] for c in row if c.hyperlink]
        assert any('Свет' in x for x in links)
    finally:
        wb.close()
    ctrl.retirar_documento(cid, codigo, 2025)
    assert not ctrl.db.residencia(cid, 2025)
    ctrl.restaurar_documento(cid, codigo, 2025)
    assert ctrl.db.residencia(cid, 2025)['ruta'] == doc['ruta']


def test_cambia_ejercicio_pdf_bd_y_excel(entorno):
    ctrl, cid, codigo = entorno
    gestor, token, texto = emitir(ctrl, cid, codigo, accion='cambiar_ejercicio', ejercicio_destino=2026)
    assert '2025 → 2026' in texto
    gestor.procesar('/confirmar_documento '+token, cid, 2025, ctrl.ejecutar_accion_documental)
    doc = ctrl.db.get('documentos', cid, codigo)
    assert doc['ejercicio'] == 2026
    assert doc['ruta'] == f'Soportes_PDF/2026_{codigo}.pdf'
    assert not ctrl.db.residencia(cid, 2025)
    assert ctrl.db.residencia(cid, 2026)['ruta'] == doc['ruta']
    assert (Path(ctrl.db.get_cliente(cid)['carpeta']) / doc['ruta']).is_file()


def test_rechaza_confirmacion_obsoleta_y_cliente_ajeno(entorno):
    ctrl, cid, codigo = entorno
    gestor, token, _ = emitir(ctrl, cid, codigo, accion='renombrar', nombre='Nuevo.pdf')
    ctrl.db.upsert('documentos', cid, {'codigo': codigo, 'descripcion': 'Editado después'})
    with pytest.raises(ValueError, match='cambiado'):
        gestor.procesar('/confirmar_documento '+token, cid, 2025, ctrl.ejecutar_accion_documental)
    otro = ctrl.crear_cliente('Otro', 'Ficticio', 'OTRO-SEC')['id']
    with pytest.raises(ValueError):
        ctrl.preparar_accion_documental(otro, {'accion': 'renombrar', 'codigo': codigo, 'nombre': 'Nuevo.pdf'})


def test_rollback_de_calculo_no_pierde_pdf(entorno, monkeypatch):
    ctrl, cid, codigo = entorno
    base = Path(ctrl.db.get_cliente(cid)['carpeta'])
    original = ctrl.db.get('documentos', cid, codigo)['ruta']
    gestor, token, _ = emitir(ctrl, cid, codigo, accion='renombrar', nombre='Fallido.pdf')
    def fallar(*a):
        raise RuntimeError('fallo simulado')
    monkeypatch.setattr(ctrl, '_recalcular_bd', fallar)
    with pytest.raises(RuntimeError):
        gestor.procesar('/confirmar_documento '+token, cid, 2025, ctrl.ejecutar_accion_documental)
    assert ctrl.db.get('documentos', cid, codigo)['ruta'] == original
    assert (base / original).is_file() and not (base / 'Soportes_PDF/Fallido.pdf').exists()
    assert not ctrl.db.con.execute('SELECT 1 FROM documentos_archivos').fetchone()


def test_destino_existente_nunca_se_pisa(entorno):
    ctrl, cid, codigo = entorno
    destino = Path(ctrl.db.get_cliente(cid)['carpeta']) / 'Soportes_PDF/Existe.pdf'
    destino.write_bytes(b'NO MODIFICAR')
    with pytest.raises(ValueError, match='existe'):
        emitir(ctrl, cid, codigo, accion='renombrar', nombre='Existe.pdf')
    assert destino.read_bytes() == b'NO MODIFICAR'


@pytest.mark.parametrize('ruta', ['../fuera.pdf', 'Soportes_PDF/../../fuera.pdf', 'C:/fuera.pdf', 'Soportes_PDF/CON.pdf', 'Soportes_PDF/a\x01.pdf', 'Soportes_PDF/a.pdf:ads'])
def test_resolvedor_rutas_rechaza(tmp_path, ruta):
    with pytest.raises(ValueError):
        ruta_soporte(tmp_path, ruta)


def test_tool_use_solo_propone(monkeypatch):
    import asistente
    import ia_extractor as ia
    monkeypatch.setattr(ia, 'hay_clave', lambda: True)
    plan = {'accion': 'renombrar', 'codigo': 'DOC-001', 'nombre': 'Luz.pdf'}
    final = SimpleNamespace(stop_reason='tool_use', usage=None, content=[SimpleNamespace(type='tool_use', name='proponer_accion_documental', input=plan)])
    monkeypatch.setattr(asistente, '_llamar', lambda *a: final)
    conv = asistente.Conversacion(1)
    texto = conv.preguntar('Renombra DOC-001 a Luz.pdf', 'Contexto ficticio', ia.AjustesIA(activada=True))
    assert conv.acciones_propuestas == [plan]
    assert 'confirmación' in texto
    assert conv.mensajes[-1]['content'][0]['type'] == 'text'


@pytest.mark.parametrize('final', [None, SimpleNamespace(content=None), SimpleNamespace(content=[],stop_reason=None)])
def test_respuesta_api_invalida_no_colapsa(monkeypatch, final):
    import asistente
    import ia_extractor as ia
    monkeypatch.setattr(ia, 'hay_clave', lambda: True)
    monkeypatch.setattr(asistente, '_llamar', lambda *a: final)
    conv = asistente.Conversacion(1)
    with pytest.raises(asistente.ErrorAsistente, match='inválida'):
        conv.preguntar('Pregunta', 'Contexto ficticio', ia.AjustesIA(activada=True))
    assert conv.acciones_propuestas == [] and conv.mensajes == []


def test_sql_inyeccion_valores_e_identificadores(tmp_path):
    from database import Database
    db = Database(tmp_path / 'sql.sqlite')
    db.carpeta_clientes = tmp_path / 'Clientes'
    try:
        c = db.crear_cliente("Ficticio'); DROP TABLE clientes; --", 'Auditoria', 'SQL-001')
        assert db.get_cliente(c['id'])['nombre'].startswith('Ficticio')
        assert db.listar_clientes(busqueda="' OR 1=1 --") == []
        for metodo in (lambda: db.listar_pagina('gastos; DROP TABLE clientes', c['id']),
                       lambda: db.listar_pagina('gastos', c['id'], orden='codigo; DROP TABLE clientes'),
                       lambda: db.listar_pagina('gastos', c['id'], **{'x=? OR 1=1 --': 1})):
            with pytest.raises(KeyError):
                metodo()
        assert db.contar_clientes() == 1
    finally:
        db.cerrar()


def test_recuperacion_renombrado_tras_bloqueo_windows(entorno, monkeypatch):
    ctrl, cid, codigo = entorno
    base = Path(ctrl.db.get_cliente(cid)['carpeta'])
    original = ctrl.db.get('documentos', cid, codigo)['ruta']
    gestor, token, _ = emitir(ctrl, cid, codigo, accion='renombrar', nombre='Recuperado.pdf')
    recuperar = documentos.recuperar_archivos
    def bloquear(db, cliente):
        if db.con.execute('SELECT 1 FROM documentos_archivos WHERE cliente_id=?', (cliente,)).fetchone():
            raise PermissionError('Bloqueo simulado de Windows')
        return recuperar(db, cliente)
    monkeypatch.setattr(documentos, 'recuperar_archivos', bloquear)
    texto = gestor.procesar('/confirmar_documento '+token, cid, 2025, ctrl.ejecutar_accion_documental)
    assert 'pendiente' in texto and 'PDF y Excel actualizados' not in texto
    assert (base / original).exists() and (base / 'Soportes_PDF/Recuperado.pdf').exists()
    monkeypatch.setattr(documentos, 'recuperar_archivos', recuperar)
    recuperar(ctrl.db, cid)
    assert not (base / original).exists() and (base / 'Soportes_PDF/Recuperado.pdf').is_file()


def test_fallo_excel_no_se_declara_exito(entorno, monkeypatch):
    ctrl, cid, codigo = entorno
    gestor, token, _ = emitir(ctrl, cid, codigo, accion='renombrar', nombre='Guardado.pdf')
    def fallar(*a):
        raise PermissionError('Excel abierto')
    monkeypatch.setattr(ctrl, '_exportar_informe', fallar)
    texto = gestor.procesar('/confirmar_documento '+token, cid, 2025, ctrl.ejecutar_accion_documental)
    assert 'no se pudo actualizar el Excel' in texto and 'PDF y Excel actualizados' not in texto
    assert ctrl.db.get('documentos', cid, codigo)['ruta'].endswith('Guardado.pdf')


def test_renombrado_no_reescribe_historial_original(entorno):
    ctrl, cid, codigo = entorno
    anteriores = [tuple(r) for r in ctrl.db.con.execute('SELECT * FROM documentos_cambios WHERE cliente_id=? ORDER BY id', (cid,))]
    gestor, token, _ = emitir(ctrl, cid, codigo, accion='renombrar', nombre='Historico.pdf')
    gestor.procesar('/confirmar_documento '+token, cid, 2025, ctrl.ejecutar_accion_documental)
    despues = [tuple(r) for r in ctrl.db.con.execute('SELECT * FROM documentos_cambios WHERE cliente_id=? ORDER BY id', (cid,))]
    assert despues[:len(anteriores)] == anteriores
    assert len(despues) > len(anteriores)


def test_colision_cirilica_con_documento_en_papelera(entorno):
    ctrl, cid, codigo = entorno
    ctrl.db.upsert('documentos', cid, {'codigo': 'DOC-999', 'ruta': 'Soportes_PDF/СВЕТ.pdf'})
    with pytest.raises(ValueError, match='existe'):
        emitir(ctrl, cid, codigo, accion='renombrar', nombre='свет.pdf')


def test_previsualizacion_escapa_markdown():
    gestor = ad.GestorAcciones(lambda: (1, 2025))
    texto = gestor.proponer({'accion': 'renombrar', 'codigo': 'DOC-001', 'nombre': '[Peligro](x).pdf', 'descripcion': '**Confirmado**'},
                           1, 2025, lambda *a: {'version': 'v', 'ruta_origen': 'original.pdf', 'ruta_destino': '[Peligro](x).pdf'})
    assert '[Peligro](x)' not in texto and '**Confirmado**' not in texto


def test_chat_y_rutas_rechazan_junction_de_cliente(tmp_path):
    import os
    import subprocess
    import asistente
    if os.name != 'nt':
        pytest.skip('Junction de Windows')
    destino, alias = tmp_path / 'destino', tmp_path / 'alias'
    (destino / 'Soportes_PDF').mkdir(parents=True)
    (destino / 'Soportes_PDF/secreto.pdf').write_bytes(b'FICTICIO')
    proceso = subprocess.run(['cmd', '/c', 'mklink', '/J', str(alias), str(destino)], capture_output=True)
    assert proceso.returncode == 0
    try:
        with pytest.raises(ValueError):
            ruta_soporte(alias, 'Soportes_PDF/secreto.pdf')
        with pytest.raises(asistente.ErrorAsistente):
            asistente.pdfs_del_cliente(alias)
    finally:
        os.rmdir(alias)


def test_descripcion_vacia_se_borra_expresamente(entorno):
    ctrl, cid, codigo = entorno
    gestor, token, _ = emitir(ctrl, cid, codigo, accion='renombrar', nombre='Sin descripcion.pdf', descripcion='')
    gestor.procesar('/confirmar_documento '+token, cid, 2025, ctrl.ejecutar_accion_documental)
    assert ctrl.db.get('documentos', cid, codigo)['descripcion'] is None


def test_colision_concurrente_identica_no_borra_archivo_ajeno(entorno, monkeypatch):
    ctrl, cid, codigo = entorno
    base = Path(ctrl.db.get_cliente(cid)['carpeta'])
    origen = base / ctrl.db.get('documentos', cid, codigo)['ruta']
    destino = base / 'Soportes_PDF/Concurrente.pdf'
    gestor, token, _ = emitir(ctrl, cid, codigo, accion='renombrar', nombre='Concurrente.pdf')
    preparar = gestion_archivos.preparar
    llamadas = [0]
    def crear_destino_ajeno(*args):
        llamadas[0] += 1
        if llamadas[0] == 2:  # después de la protección, antes de copiar
            destino.write_bytes(origen.read_bytes())
        return preparar(*args)
    monkeypatch.setattr(gestion_archivos, 'preparar', crear_destino_ajeno)
    with pytest.raises(ValueError, match='existe'):
        gestor.procesar('/confirmar_documento '+token, cid, 2025, ctrl.ejecutar_accion_documental)
    assert origen.is_file() and destino.is_file()
    assert ctrl.db.get('documentos', cid, codigo)['ruta'] == origen.relative_to(base).as_posix()


def test_rollback_restauracion_no_borra_copia_ajena(entorno, monkeypatch):
    ctrl, cid, codigo = entorno
    base = Path(ctrl.db.get_cliente(cid)['carpeta'])
    ruta_original = ctrl.db.get('documentos', cid, codigo)['ruta']
    ctrl.retirar_documento(cid, codigo, 2025)
    meta = documentos._documento(ctrl.db, cid, codigo)
    papelera = base / meta['ruta_papelera']
    destino = base / ruta_original
    copiar = documentos._copiar_seguro
    def copiar_con_destino_ajeno(origen, objetivo, huella):
        objetivo.write_bytes(origen.read_bytes())
        return copiar(origen, objetivo, huella)
    def fallar(*a):
        raise RuntimeError('Fallo de recálculo simulado')
    monkeypatch.setattr(documentos, '_copiar_seguro', copiar_con_destino_ajeno)
    monkeypatch.setattr(ctrl, '_recalcular_bd', fallar)
    with pytest.raises(RuntimeError):
        ctrl.restaurar_documento(cid, codigo, 2025)
    assert destino.is_file() and papelera.is_file()
    assert documentos._documento(ctrl.db, cid, codigo)['en_papelera']


def test_etiqueta_hipervinculo_excel_es_texto_nunca_formula(tmp_path):
    import openpyxl
    import excel_generator as eg
    wb = openpyxl.Workbook()
    eg._link(wb.active['A1'], 'Soportes_PDF/=1+1.pdf')
    wb.save(tmp_path / 'etiqueta.xlsx')
    wb.close()
    verificado = openpyxl.load_workbook(tmp_path / 'etiqueta.xlsx', data_only=False)
    try:
        celda = verificado.active['A1']
        assert celda.value == '=1+1.pdf' and celda.data_type == 's'
        assert celda.hyperlink.target == 'Soportes_PDF/=1+1.pdf'
    finally:
        verificado.close()

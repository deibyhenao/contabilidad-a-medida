"""Lectores locales acotados. Nunca ejecutan fórmulas, macros ni vínculos.

Los límites generan ErrorDocumento: no se devuelve una extracción parcial.
"""
from __future__ import annotations
import csv
import io
import re
from contextvars import ContextVar
from itertools import islice
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile

FORMATOS = frozenset({'.pdf', '.xlsx', '.xls', '.docx', '.csv'})
MAX_ARCHIVO = 20 * 1024 * 1024
MAX_ARCHIVO_PDF = 32 * 1024 * 1024
MAX_DESCOMPRIMIDO = 80 * 1024 * 1024
MAX_ENTRADAS_ZIP = 2000
MAX_RATIO_ZIP = 200
MAX_FILAS = 100_000
MAX_CELDAS = 300_000
MAX_TEXTO = 2_000_000
MAX_PAGINAS = 1000
MAX_HOJAS = 100

class ErrorDocumento(ValueError):
    """Archivo ilegible, inseguro o superior a los límites locales."""

@dataclass
class ResultadoDocumento:
    texto: str
    escaneado: bool = False
    avisos: list[str] = field(default_factory=list)

_VISTA = ContextVar('limite_vista_documental', default=None)

class _VistaCompleta(Exception):
    def __init__(self, texto):
        self.texto = texto

class _Texto:
    def __init__(self):
        self.partes = []
        self.longitud = 0
        self.celdas = 0
    def agregar(self, valor):
        vista = _VISTA.get()
        if vista is not None and self.longitud + len(valor) + 1 > vista:
            disponible = max(0, vista - self.longitud - (1 if self.partes else 0))
            if disponible:
                self.partes.append(valor[:disponible])
            raise _VistaCompleta(self)
        self.longitud += len(valor) + 1
        if self.longitud > MAX_TEXTO:
            raise ErrorDocumento('El texto supera el límite permitido.')
        self.partes.append(valor)
    def fila(self, valores, *, numero=None):
        valores = list(valores)
        self.celdas += len(valores)
        if self.celdas > MAX_CELDAS:
            raise ErrorDocumento('El documento supera el límite de celdas.')
        valores = [_valor(v) for v in valores]
        if any(valores):
            prefijo = f'[Fila {numero}] ' if numero is not None else ''
            self.agregar(prefijo + '\t'.join(valores))
    def terminar(self, **kwargs):
        return ResultadoDocumento('\n'.join(self.partes), **kwargs)

def _valor(valor):
    if valor is None:
        return ''
    if isinstance(valor, (date, datetime)):
        return valor.isoformat()
    return str(valor).replace('\r\n', '\n').replace('\r', '\n')

def _revisar_ooxml(ruta, extension):
    obligatoria = 'xl/workbook.xml' if extension == '.xlsx' else 'word/document.xml'
    with ZipFile(ruta) as archivo:
        entradas = archivo.infolist()
        nombres = [e.filename for e in entradas]
        if len(entradas) > MAX_ENTRADAS_ZIP:
            raise ErrorDocumento('Demasiadas entradas comprimidas.')
        if len(set(nombres)) != len(nombres) or obligatoria not in nombres:
            raise ErrorDocumento('El archivo no corresponde al formato Office indicado.')
        total = 0
        for entrada in entradas:
            nombre = entrada.filename.replace('\\', '/')
            if nombre.startswith('/') or '..' in nombre.split('/') or entrada.flag_bits & 1:
                raise ErrorDocumento('Contenedor Office inseguro o cifrado.')
            impresion_pasiva = extension == '.xlsx' and re.fullmatch(
                r'xl/printerSettings/printerSettings[0-9]+\.bin', nombre) is not None
            if ('vbaproject' in nombre.lower() or
                    (nombre.lower().endswith('.bin') and not impresion_pasiva)):
                raise ErrorDocumento('No se admiten macros ni objetos binarios incrustados.')
            total += entrada.file_size
            if total > MAX_DESCOMPRIMIDO or entrada.file_size > max(1, entrada.compress_size) * MAX_RATIO_ZIP:
                raise ErrorDocumento('El documento supera los límites de descompresión.')
            if nombre.lower().endswith(('.xml', '.rels')):
                compacto = archivo.read(entrada).replace(b'\x00', b'').upper()
                if b'<!DOCTYPE' in compacto or b'<!ENTITY' in compacto:
                    raise ErrorDocumento('No se admiten entidades ni DTD en Office.')
                if nombre == '[Content_Types].xml' and b'MACROENABLED' in compacto:
                    raise ErrorDocumento('El documento declara macros.')

def _xlsx(ruta):
    from openpyxl import load_workbook
    _revisar_ooxml(ruta, '.xlsx')
    texto = _Texto()
    libro = load_workbook(ruta, read_only=True, data_only=True, keep_links=False)
    try:
        if len(libro.worksheets) > MAX_HOJAS:
            raise ErrorDocumento('El libro supera el límite de hojas.')
        filas = 0
        for hoja in libro.worksheets:
            texto.agregar(f'[Hoja: {hoja.title}]')
            hoja.reset_dimensions()
            for numero, fila in enumerate(hoja.iter_rows(), 1):
                filas += 1
                if filas > MAX_FILAS:
                    raise ErrorDocumento('El libro supera el límite total de filas.')
                texto.fila((c.value for c in fila), numero=numero)
    finally:
        libro.close()
    return texto.terminar(avisos=['Las fórmulas no se ejecutan: sus valores guardados pueden estar ausentes o desactualizados.'])

def _xls(ruta):
    try:
        import xlrd
    except ImportError as exc:
        raise ErrorDocumento('La lectura de XLS requiere instalar xlrd.') from exc
    with ruta.open('rb') as entrada:
        if entrada.read(8) != b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
            raise ErrorDocumento('El archivo no es un XLS binario válido.')
    texto = _Texto()
    libro = xlrd.open_workbook(str(ruta), on_demand=True)
    try:
        if libro.nsheets > MAX_HOJAS:
            raise ErrorDocumento('El libro supera el límite de hojas.')
        filas = 0
        for indice in range(libro.nsheets):
            hoja = libro.sheet_by_index(indice)
            filas += hoja.nrows
            if filas > MAX_FILAS or texto.celdas + hoja.nrows * hoja.ncols > MAX_CELDAS:
                raise ErrorDocumento('El libro supera el límite de filas o celdas.')
            texto.agregar(f'[Hoja: {hoja.name}]')
            for numero in range(hoja.nrows):
                texto.fila((xlrd.xldate_as_datetime(c.value, libro.datemode) if c.ctype == xlrd.XL_CELL_DATE else c.value for c in hoja.row(numero)), numero=numero + 1)
            libro.unload_sheet(indice)
    finally:
        libro.release_resources()
    return texto.terminar(avisos=['No se ejecutan fórmulas ni macros; los valores guardados pueden estar desactualizados.'])

def _docx(ruta):
    _revisar_ooxml(ruta, '.docx')
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise ErrorDocumento('La lectura de DOCX requiere instalar python-docx.') from exc
    documento = Document(str(ruta))
    texto = _Texto()
    filas = 0
    for elemento in documento.element.body.iterchildren():
        if elemento.tag.endswith('}p'):
            texto.agregar(Paragraph(elemento, documento).text)
        elif elemento.tag.endswith('}tbl'):
            for fila in Table(elemento, documento).rows:
                filas += 1
                if filas > MAX_FILAS:
                    raise ErrorDocumento('El documento supera el límite de filas.')
                texto.fila(c.text for c in fila.cells)
    return texto.terminar(avisos=['Se lee el cuerpo y sus tablas; imágenes, cabeceras y pies requieren revisión.'])

def _csv(ruta):
    datos = ruta.read_bytes()
    avisos = []
    if datos.startswith((b'\xff\xfe', b'\xfe\xff')):
        contenido = datos.decode('utf-16')
    else:
        try:
            contenido = datos.decode('utf-8-sig')
        except UnicodeDecodeError:
            contenido = datos.decode('cp1252')
            avisos.append('CSV interpretado como Windows-1252: comprueba los caracteres.')
    if '\x00' in contenido:
        raise ErrorDocumento('El CSV contiene datos binarios.')
    try:
        dialecto = csv.Sniffer().sniff(contenido[:65536], delimiters=';,\t|')
    except csv.Error:
        # Sniffer falla con ciertos campos multilínea. Puntuar registros CSV
        # completos evita confundir las comas decimales con separadores.
        candidatos = []
        for separador in ';,\t|':
            filas = list(islice(csv.reader(io.StringIO(contenido), delimiter=separador), 20))
            cantidades = [len(f) for f in filas if f]
            if cantidades and min(cantidades) > 1 and len(set(cantidades)) == 1:
                candidatos.append((cantidades[0], separador))
        separador = max(candidatos, default=(0, ','))[1]
        dialecto = type('CSVDetectado', (csv.excel,), {'delimiter': separador})
        if not candidatos:
            avisos.append('No se detectó un delimitador uniforme; se ha usado coma.')
    texto = _Texto()
    for numero, fila in enumerate(csv.reader(io.StringIO(contenido, newline=''), dialecto, strict=True), 1):
        if numero > MAX_FILAS:
            raise ErrorDocumento('El CSV supera el límite de filas.')
        texto.fila(fila)
    return texto.terminar(avisos=avisos)

def _pdf(ruta):
    import fitz
    texto = _Texto()
    escaneado = False
    avisos = []
    with fitz.open(str(ruta)) as documento:
        if documento.needs_pass:
            raise ErrorDocumento('El PDF está protegido por contraseña.')
        if documento.page_count > MAX_PAGINAS:
            raise ErrorDocumento('El PDF supera el límite de páginas.')
        for numero, pagina in enumerate(documento, 1):
            texto.agregar(f'[Página {numero}]')
            contenido = pagina.get_text('text', sort=True).strip()
            if contenido:
                texto.agregar(contenido)
            else:
                escaneado = True
            try:
                tablas = pagina.find_tables().tables if hasattr(pagina, 'find_tables') else []
                for tabla in tablas:
                    texto.agregar('[Tabla]')
                    for fila in tabla.extract():
                        texto.fila(fila)
            except (ErrorDocumento, _VistaCompleta):
                raise
            except Exception:
                aviso = 'No se pudieron estructurar todas las tablas del PDF; se conserva el texto disponible.'
                if aviso not in avisos:
                    avisos.append(aviso)
    if escaneado:
        avisos.append('Hay páginas sin texto extraíble: necesitan OCR o revisión visual.')
    return texto.terminar(escaneado=escaneado, avisos=avisos)

def extraer_documento(ruta: Path, *, max_caracteres: int | None = None) -> ResultadoDocumento:
    """Lee un soporte sin escribirlo. El llamador debe validar su autorización/ruta."""
    if max_caracteres is not None and (isinstance(max_caracteres, bool) or not isinstance(max_caracteres, int) or not 1 <= max_caracteres <= MAX_TEXTO):
        raise ErrorDocumento('El límite de vista previa debe ser un entero entre 1 y MAX_TEXTO.')
    ruta = Path(ruta)
    extension = ruta.suffix.lower()
    if extension not in FORMATOS:
        raise ErrorDocumento('Formato no admitido. Usa PDF, XLSX, XLS, DOCX o CSV.')
    token_vista = _VISTA.set(max_caracteres)
    try:
        if not ruta.is_file():
            raise ErrorDocumento('El documento no existe o no es un archivo.')
        limite_archivo = MAX_ARCHIVO_PDF if extension == '.pdf' else MAX_ARCHIVO
        if ruta.stat().st_size > limite_archivo:
            raise ErrorDocumento('El documento supera el límite de tamaño de archivo.')
        with ruta.open('rb') as entrada:
            cabecera = entrada.read(8)
        if extension == '.pdf' and not cabecera.startswith(b'%PDF-'):
            raise ErrorDocumento('El archivo no es un PDF válido.')
        if extension in {'.xlsx', '.docx'} and not cabecera.startswith(b'PK\x03\x04'):
            raise ErrorDocumento('El archivo no corresponde al formato Office indicado.')
        return {'.pdf': _pdf, '.xlsx': _xlsx, '.xls': _xls, '.docx': _docx, '.csv': _csv}[extension](ruta)
    except _VistaCompleta as vista:
        return vista.texto.terminar(avisos=['Vista previa parcial: se ha detenido la lectura al alcanzar el límite de caracteres.'])
    except ErrorDocumento:
        raise
    except Exception as exc:
        raise ErrorDocumento('No se pudo interpretar el documento: está dañado, bloqueado o no corresponde a su extensión.') from exc
    finally:
        _VISTA.reset(token_vista)

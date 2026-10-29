"""Oráculos independientes y conservación de datos al actualizar 1.0.6 a 1.0.7."""
from datetime import date
from pathlib import Path

import openpyxl
import pytest
import json

from controlador import Controlador
import database
import excel_generator as eg
import fiscal_residentes as fr
import tax_calculator as tc


@pytest.mark.parametrize('ejercicio', [2024, 2025, 2026])
def test_imputacion_anual_no_supera_porcentaje_legal(ejercicio):
    inm = {'codigo': 'INM-001', 'fecha_adquisicion': '2020-01-01', 'pct_titularidad': 1}
    cat = {'valor_catastral': 100000, 'pct_imputacion': .02, 'revision_pendiente': 0}
    r = tc.calcular_210_imputacion(inm, [], cat, tc.resolver_residencia({'pais': 'Rusia'}), ejercicio)
    assert (r.renta_imputada, r.cuota) == (2000, 480)
    assert r.dias == (366 if ejercicio == 2024 else 365)


def test_imputacion_bisiesta_parcial_conserva_criterio365():
    inm = {'codigo': 'INM-001', 'fecha_adquisicion': '2024-03-01', 'pct_titularidad': .5}
    cat = {'valor_catastral': 100000, 'pct_imputacion': .02, 'revision_pendiente': 0}
    r = tc.calcular_210_imputacion(inm, [], cat, tc.resolver_residencia({'pais': 'Rusia'}), 2024)
    assert (r.dias, r.renta_imputada, r.cuota) == (306, 838.36, 201.21)


@pytest.mark.parametrize('ejercicio', [2023, 2024, 2025, 2026])
def test_excepcion_revision_colectiva2012_y_porcentaje_manual(ejercicio):
    assert tc.pct_imputacion(2012, ejercicio) == .011
    assert tc.pct_imputacion(2011, ejercicio) == .02
    assert tc.pct_imputacion(ejercicio + 1, ejercicio) == .02
    assert tc.pct_imputacion(2012, ejercicio, .02) == .02


@pytest.mark.parametrize('ejercicio,cuota,esperado', [
    (2025, 100, date(2026, 1, 20)), (2026, 100, date(2027, 4, 20)),
    (2026, 0, date(2027, 1, 20)), (2025, 0, date(2026, 1, 20)),
])
def test_plazo210_segun_resultado(ejercicio, cuota, esperado):
    assert tc.plazo_210_alquiler(ejercicio, 'Anual', cuota) == esperado


@pytest.mark.parametrize('modelo,dia', [('111', 20), ('115', 20), ('303', 30), ('130', 30)])
def test_plazo4t_no_es_comun_a_todos_los_modelos(modelo, dia):
    assert dict(tc.plazos_autonomo(modelo, 2025))['4T'] == date(2026, 1, dia)


def test_devolucion_no_usa_inicio_de_ventana_como_vencimiento():
    assert tc.plazo_210_alquiler(2026, 'Anual', -50) is None


def test_347_ultimo_dia_febrero_bisiesto():
    assert dict(tc.plazos_autonomo('347', 2023))['Anual'] == date(2024, 2, 29)


def test_303_rectificativa_reduce_base_y_cuota():
    recibidas = [{'trimestre': '1T', 'base': 1000, 'cuota_iva': 210},
                 {'trimestre': '1T', 'base': -200, 'cuota_iva': -42}]
    r = fr.calcular_303([], recibidas, '1T')
    assert (r.base_deducible, r.iva_deducible) == (800, 168)


@pytest.mark.parametrize('ejercicio,ingreso', [(2023, 3720), (2024, 3800), (2025, 3800)])
def test_130_porcentaje_dificil_justificacion_por_ejercicio(ejercicio, ingreso):
    r = fr.calcular_130([{'trimestre': '1T', 'base': 30000}],
                       [{'trimestre': '1T', 'base': 10000}], '1T', 0, True, ejercicio)
    assert r.a_ingresar == ingreso


def test_130_no_aplica_propuesta_erronea_de_sumar_retenciones_a_pagos():
    emitidas = [{'trimestre': '1T', 'base': 10000, 'retencion': 500}]
    r1 = fr.calcular_130(emitidas, [], '1T', 0, False, 2025)
    r2 = fr.calcular_130(emitidas, [], '2T', r1.a_ingresar, False, 2025)
    assert (r1.a_ingresar, r2.a_ingresar) == (1500, 0)


@pytest.fixture
def expediente(tmp_path):
    ctrl = Controlador(tmp_path/'ficticio.db')
    ctrl.cambiar_carpeta_clientes(tmp_path/'Datos externos')
    cid = ctrl.db.crear_cliente('FICTICIO', 'Fiscal106', 'QA-FISCAL106')['id']
    ctrl.db.guardar_residencia(cid, 2024, pais='Rusia')
    inm = ctrl.db.upsert('inmuebles', cid, {'fecha_adquisicion': '2020-01-01', 'pct_titularidad': 1,
                                          'precio_escritura': 200000, 'observaciones': 'Nota del asesor'})
    ctrl.db.upsert('catastro', cid, {'ejercicio': 2024, 'inmueble': inm, 'valor_catastral': 100000,
                                    'pct_imputacion': .02, 'revision_pendiente': 0})
    try:
        yield ctrl, cid, inm
    finally:
        ctrl.cerrar()


def test_libro106_previo_se_renueva_una_vez_sin_perder_edicion_manual(expediente):
    c, cid, inm = expediente
    ruta = c.sincronizar(cid, 2024)['excel']
    wb = openpyxl.load_workbook(ruta)
    wb['10_210_IMPUTACION']['H2'] = '=IF(A2="","",ROUND(E2*F2*G2*D2/365,2))'
    columna = next(col.letra for col in database.TABLAS['inmuebles'].columnas_excel
                   if col.campo == 'observaciones')
    wb['02_INMUEBLES'][columna+'2'] = 'Edición manual anterior al parche'
    wb.save(ruta); wb.close()
    c.db.set_config(f'excel_revision|{cid}', '')
    assert eg.necesita_actualizar_formato(c.db, cid)
    c.abrir_cliente(cid, 2024)
    assert c.db.get('inmuebles', cid, inm)['observaciones'] == 'Edición manual anterior al parche'
    assert len(c.db.listar_clientes()) == 1
    wb = openpyxl.load_workbook(ruta)
    assert wb['10_210_IMPUTACION']['D2'].value == 366
    assert wb['10_210_IMPUTACION']['H2'].value == '=IF(A2="","",ROUND(E2*F2*G2*MIN(D2,365)/365,2))'
    wb.close()
    assert eg.verificar_estructura(ruta) == []
    assert not eg.necesita_actualizar_formato(c.db, cid)
    firma = eg._firma_excel(ruta)
    assert c.abrir_cliente(cid, 2024)['excel'] is None
    assert eg._firma_excel(ruta) == firma
    assert not list(Path(ruta).parent.glob('*PENDIENTE*.xlsx'))


def test_trimestre_presentado_no_se_reescribe_ni_duplica_en_anual(expediente):
    c, cid, inm = expediente
    codigo = c.db.upsert('m210_alquiler', cid, {'inmueble': inm, 'ejercicio': 2024, 'periodo': '1T',
        'estado': 'Presentado', 'ingresos': 1000, 'base': 1000, 'cuota': 240, 'pagado': 240,
        'observaciones': 'Justificante manual'}, clave=f'210A|2024|{inm}|1T')
    antes = c.db.get('m210_alquiler', cid, codigo)
    inf = tc.recalcular_cliente(c.db, cid, 2024)
    assert c.db.get('m210_alquiler', cid, codigo) == antes
    assert not [m for m in c.db.listar('m210_alquiler', cid) if m['periodo'] == 'Anual']
    assert any(a.nivel == 'grave' and 'Desde 2024' in a.mensaje for a in inf.alertas)


def test_modelo_automatico_anterior_conserva_importes_al_quedar_no_procede(expediente):
    c, cid, inm = expediente
    codigo = c.db.upsert('m210_alquiler', cid, {'inmueble': inm, 'ejercicio': 2024, 'periodo': '1T',
        'estado': 'Pendiente', 'ingresos': 1234, 'cuota': 296.16}, clave=f'210A|2024|{inm}|1T')
    tc.recalcular_cliente(c.db, cid, 2024)
    previo = c.db.get('m210_alquiler', cid, codigo)
    assert (previo['estado'], previo['ingresos'], previo['cuota']) == ('No procede', 1234, 296.16)


@pytest.mark.parametrize('estado,metadatos', [
    ('En preparación', {}), ('Pendiente', {'pagado': 240}),
    ('Pendiente', {'nrc': 'NRC-FICTICIO'}), ('Pendiente', {'fecha_presentacion': '2024-04-15'}),
])
def test_historico_protegido_bloquea_dependientes_y_otros_inmuebles_siguen(expediente, estado, metadatos):
    c, cid, inm = expediente
    codigo = c.db.upsert('m210_alquiler', cid, {'inmueble': inm, 'ejercicio': 2024, 'periodo': '1T',
        'estado': estado, 'ingresos': 1000, 'cuota': 240, **metadatos}, clave=f'210A|2024|{inm}|1T')
    arr = c.db.upsert('arrastres_gastos', cid, {'inmueble': inm, 'ejercicio': 2024,
        'generado': 3000, 'pendiente': 3000, 'saldos': '{"2024": 3000}'}, clave=f'ARR|{inm}|2024')
    modelo_antes = c.db.get('m210_alquiler', cid, codigo)
    arr_antes = c.db.get('arrastres_gastos', cid, arr)
    otro = c.db.upsert('inmuebles', cid, {'fecha_adquisicion': '2020-01-01', 'pct_titularidad': 1})
    c.db.upsert('contratos', cid, {'inmueble': otro, 'fecha_inicio': '2025-01-01',
        'fecha_fin_real': '2025-12-31', 'renta_inicial': 1000, 'periodicidad': 'Mensual',
        'tipo': 'Temporada', 'estado': 'Finalizado'})
    inf = tc.recalcular_cliente(c.db, cid, 2025)
    assert c.db.get('m210_alquiler', cid, codigo) == modelo_antes
    assert c.db.get('arrastres_gastos', cid, arr) == arr_antes
    assert not c.db.listar('m210_alquiler', cid, inmueble=inm, ejercicio=2025)
    assert c.db.listar('m210_alquiler', cid, inmueble=otro, ejercicio=2025)
    assert inm in inf.alquileres_bloqueados
    vista = next(v for v in c.vista_inmuebles(cid, 2024) if v['codigo'] == inm)
    assert vista['_revision_210'] and vista['_desglose_210'][0]['codigo'] == codigo


def test_preferencia_anual_global_no_agrupa_ingresos2023(expediente):
    c, cid, inm = expediente
    c.db.set_config('periodicidad_210_alquiler', 'Anual')
    assert tc.periodicidad_210_alquiler(c.db, cid, inm, 2023) == 'Trimestral'


def test_desglose_trimestral_informativo_se_construye_y_no_oculta_conflicto():
    import ui_cartera as uc
    import ui_estilo as ui
    vista = {'_revision_210': 'Conciliar histórico', '_desglose_210': [
        {'codigo': '210-001', 'periodo': '1T', 'estado': 'En preparación', 'cuota': 240}],
        '_ocupacion_trimestral': [{'periodo': '1T', 'dias_alquilado': 90, 'dias_vacio': 0,
                                  'ingresos_devengados': 3000}]}
    assert uc.filas_desglose_210(vista, 'Anual')[0]['codigo'] == '210-001'
    assert uc.desglose_210(vista, 'Anual', 2025, ui.CLARO)


def test_error_de_carpeta_no_deja_cliente_parcial_en_bd(tmp_path, monkeypatch):
    db = database.Database(tmp_path/'ficticio.db')
    db.carpeta_clientes = tmp_path/'Datos'
    def fallar(_cliente):
        raise PermissionError('Carpeta ficticia bloqueada por Windows')
    monkeypatch.setattr(db, 'asegurar_carpetas', fallar)
    try:
        with pytest.raises(PermissionError):
            db.crear_cliente('FICTICIO', 'Permisos', 'QA-PERMISOS')
        assert db.listar_clientes() == []
        assert db.get_cliente_por_nif('QA-PERMISOS') is None
    finally:
        db.cerrar()


def test_fila_trimestral_manual_pendiente_no_convive_con_anual(expediente):
    c, cid, inm = expediente
    codigo = c.db.upsert('m210_alquiler', cid, {'inmueble': inm, 'ejercicio': 2024, 'periodo': '1T',
        'estado': 'Pendiente', 'ingresos': 3000, 'cuota': 720}, clave='fila-manual-importada')
    antes = c.db.get('m210_alquiler', cid, codigo)
    inf = tc.recalcular_cliente(c.db, cid, 2024)
    assert inm in inf.alquileres_bloqueados
    assert c.db.get('m210_alquiler', cid, codigo) == antes
    assert len(c.db.listar('m210_alquiler', cid)) == 1


def test_cambio_de_calendario_se_avisa_y_no_borra_importes(expediente):
    c, cid, inm = expediente
    codigo = c.db.upsert('m210_alquiler', cid, {'inmueble': inm, 'ejercicio': 2024, 'periodo': '1T',
        'estado': 'Pendiente', 'ingresos': 1234, 'cuota': 296.16}, clave=f'210A|2024|{inm}|1T')
    inf = tc.recalcular_cliente(c.db, cid, 2024)
    assert any(a.tipo == 'cambio de calendario 210' and a.nivel == 'grave' for a in inf.alertas)
    assert c.db.get('m210_alquiler', cid, codigo)['ingresos'] == 1234
    obligacion = c.db.listar('obligaciones', cid, id_modelo=codigo)[0]
    assert 'conservados' in obligacion['notas'] and 'Sin ingresos' not in obligacion['notas']


def test_filas_ampliadas_reciben_formula_aunque_primera_este_presentada(expediente, tmp_path, monkeypatch):
    import shutil
    c, cid, inm = expediente
    copia = tmp_path/'plantilla-corta-qa.xlsx'
    shutil.copyfile(database.PLANTILLA_EXCEL, copia)
    wb = openpyxl.load_workbook(copia)
    ws = wb['10_210_IMPUTACION']
    ws.delete_rows(3, ws.max_row)
    wb.save(copia); wb.close()
    # La clave real de la cartera se toma del catálogo, sin modificar la plantilla original.
    from carteras import NR
    monkeypatch.setitem(database.PLANTILLAS, NR, copia)
    for indice, ejercicio in enumerate((2023, 2024, 2025)):
        c.db.upsert('m210_imputacion', cid, {'inmueble': inm, 'ejercicio': ejercicio, 'dias': 365,
            'valor_catastral': 100000, 'pct_imputacion': .02, 'pct_titularidad': 1, 'tipo': .24,
            'estado': 'Presentado' if indice == 0 else 'Pendiente', 'renta_imputada': 999.99,
            'cuota': 239.99})
    ruta = eg.exportar_cliente(c.db, cid, 2024)
    wb = openpyxl.load_workbook(ruta)
    try:
        ws = wb['10_210_IMPUTACION']
        assert ws['H2'].value == 999.99 and ws['H2'].comment
        assert ws['H3'].value == '=IF(A3="","",ROUND(E3*F3*G3*MIN(D3,365)/365,2))'
        assert ws['H4'].value == '=IF(A4="","",ROUND(E4*F4*G4*MIN(D4,365)/365,2))'
        assert ws['J4'].value == '=IF(A4="","",ROUND(H4*I4,2))'
    finally:
        wb.close()


def test_conciliacion_explicita_permite_futuro_sin_reescribir_lo_presentado(expediente):
    c, cid, inm = expediente
    c.db.guardar_residencia(cid, 2024, pais='Alemania')
    c.db.guardar_residencia(cid, 2025, pais='Alemania')
    codigo = c.db.upsert('m210_alquiler', cid, {'inmueble': inm, 'ejercicio': 2024, 'periodo': '1T',
        'estado': 'Presentado', 'ingresos': 3000, 'gastos': 3000, 'base': 0, 'cuota': 0, 'pagado': 0},
        clave=f'210A|2024|{inm}|1T')
    arr = c.db.upsert('arrastres_gastos', cid, {'inmueble': inm, 'ejercicio': 2024,
        'generado': 3000, 'pendiente': 3000, 'saldos': '{"2024":3000}'}, clave=f'ARR|{inm}|2024')
    c.db.upsert('contratos', cid, {'inmueble': inm, 'fecha_inicio': '2025-01-01',
        'fecha_fin_real': '2025-12-31', 'renta_inicial': 1000, 'periodicidad': 'Mensual',
        'tipo': 'Temporada', 'estado': 'Finalizado'})
    antes = c.db.get('m210_alquiler', cid, codigo)
    saldo_antes = c.db.get('arrastres_gastos', cid, arr)
    firma = tc.firma_historial_210(c.db.listar('m210_alquiler', cid), inm, 2024)
    with pytest.raises(ValueError, match='Confirma'):
        c.conciliar_historial_210(cid, inm, 2024, 2025, {2024: 3000}, 'Justificante contrastado', firma)
    salida = c.conciliar_historial_210(cid, inm, 2024, 2025, {2024: 3000},
                                     'Justificante contrastado', firma, confirmado=True)
    assert not salida.get('error_excel')
    assert inm not in salida['informe'].alquileres_bloqueados
    anual = c.db.get_por_clave('m210_alquiler', cid, f'210A|2025|{inm}|Anual')
    assert (anual['ingresos'], anual['gastos'], anual['base'], anual['cuota']) == (12000, 3000, 9000, 1710)
    assert c.db.get('m210_alquiler', cid, codigo) == antes
    assert c.db.get('arrastres_gastos', cid, arr) == saldo_antes
    assert tc.conflictos_210_alquiler(c.db, cid, 2024)  # no crea un anual duplicado en el año presentado
    c.db.upsert('m210_alquiler', cid, {'codigo': codigo, 'observaciones': 'Cambio posterior a la revisión'})
    assert inm in tc.conflictos_210_alquiler(c.db, cid, 2025)


@pytest.mark.parametrize('saldo', [-1, float('nan'), float('inf'), True, '', 1e20, '50%'])
def test_conciliacion_no_admite_saldos_invalidos(expediente, saldo):
    c, cid, inm = expediente
    with pytest.raises(ValueError):
        c.conciliar_historial_210(cid, inm, 2024, 2025, {2024: saldo}, 'Justificante contrastado',
                                 'firma-ficticia', confirmado=True)
    assert not c.db.get_config(f'conciliacion_210|{cid}|{inm}|2024')


def test_firma_conciliacion_ignora_solo_metadatos_sqlite():
    fila = {'id': 1, 'actualizado': '2026-10-28 10:00:00', 'cliente_id': 10,
            'codigo': '210A-001', 'inmueble': 'INM-001', 'ejercicio': 2024,
            'cuota': 123, 'pagado': 123, 'estado': 'Presentado', 'observaciones': 'Verificado'}
    firma = tc.firma_historial_210([fila], 'INM-001', 2024)
    assert tc.firma_historial_210([{**fila, 'id': 99, 'actualizado': '2026-10-28 11:00:00'}],
                                'INM-001', 2024) == firma
    for campo, valor in [('cliente_id', 11), ('cuota', 124), ('pagado', 124),
                         ('estado', 'Pendiente'), ('observaciones', 'Corregido')]:
        assert tc.firma_historial_210([{**fila, campo: valor}], 'INM-001', 2024) != firma


@pytest.mark.parametrize('saldo', [True, '100', float('inf'), -1, 1e20])
def test_conciliacion_guardada_rechaza_importes_corruptos(expediente, saldo):
    c, cid, inm = expediente
    filas = [{'codigo': '210A-001', 'inmueble': inm, 'ejercicio': 2024}]
    c.db.set_config(f'conciliacion_210|{cid}|{inm}|2024', json.dumps({
        'version': 1, 'firma': tc.firma_historial_210(filas, inm, 2024), 'saldos': {'2024': saldo}}))
    assert tc.conciliaciones_210(c.db, cid, filas) == {}


def test_borrar_inmueble_limpia_marcas_sin_alcanzar_otro_codigo(expediente):
    c, cid, inm = expediente
    for marca in ('periodicidad_210', 'conciliacion_210', 'revision_calendario_210'):
        c.db.set_config(f'{marca}|{cid}|{inm}|2024', 'propia')
        c.db.set_config(f'{marca}|{cid}|{inm}0|2024', 'ajena')
    c.db.borrar('inmuebles', cid, inm)
    for marca in ('periodicidad_210', 'conciliacion_210', 'revision_calendario_210'):
        assert not c.db.get_config(f'{marca}|{cid}|{inm}|2024')
        assert c.db.get_config(f'{marca}|{cid}|{inm}0|2024') == 'ajena'


@pytest.mark.parametrize('pais', ['Deutschland', 'Nederland', 'The Netherlands', 'Österreich',
                                  'Sverige', 'Polska', 'Suomi', 'Eesti', 'Latvija', 'Lietuva',
                                  'Hrvatska', 'Slovensko', 'Slovenija', 'Danmark', 'Norge',
                                  'Magyarország', 'Éire', 'România', 'Ellada', 'EL'])
def test_residencia_acepta_nombres_extranjeros_del_mismo_pais(pais):
    residencia = tc.resolver_residencia({'pais': pais})
    assert (residencia.grupo, residencia.tipo, residencia.deduce_gastos) == (tc.GRUPO_UE, .19, True)


def test_conciliacion_fallida_revierte_pago_por_nrc(expediente, monkeypatch):
    c, cid, inm = expediente
    codigo = c.db.upsert('m210_alquiler', cid, {'inmueble': inm, 'ejercicio': 2024, 'periodo': '1T',
        'estado': 'Presentado', 'ingresos': 1000, 'base': 1000, 'cuota': 240,
        'nrc': 'A'*22}, clave=f'210A|2024|{inm}|1T')
    antes = c.db.get('m210_alquiler', cid, codigo)
    firma = tc.firma_historial_210([antes], inm, 2024)
    def fallar(*args):
        assert c.db.get('m210_alquiler', cid, codigo)['pagado'] == 240
        raise RuntimeError('Fallo ficticio después del pago')
    monkeypatch.setattr(c, '_recalcular_bd', fallar)
    with pytest.raises(RuntimeError, match='Fallo ficticio'):
        c.conciliar_historial_210(cid, inm, 2024, 2025, {2024: 0},
                                 'Justificante contrastado', firma, confirmado=True)
    assert c.db.get('m210_alquiler', cid, codigo) == antes
    assert not c.db.get_config(f'conciliacion_210|{cid}|{inm}|2024')


@pytest.mark.parametrize('cantidad', [0, 1, 6])
def test_resumen_excel_muestra_gastos_sin_inmueble_aunque_haya_pocos(expediente, cantidad):
    c, cid, inm = expediente
    if cantidad == 0:
        c.db.borrar('inmuebles', cid, inm)
    for i in range(1, cantidad):
        c.db.upsert('inmuebles', cid, {'direccion': f'Calle ficticia {i}'})
    c.db.upsert('gastos', cid, {'ejercicio': 2025, 'categoria': 'Comunidad', 'total': 321})
    ruta = eg.exportar_cliente(c.db, cid, 2025)
    wb = openpyxl.load_workbook(ruta)
    try:
        ws = wb['00_RESUMEN']
        assert ws['H30'].value == 'Resto / sin asignar'
        fila = next(i for i in range(31, 45) if ws.cell(i, 1).value == 'Comunidad')
        assert ws.cell(fila, 8).value == f'=B{fila}-SUM(C{fila}:G{fila})'
        assert eg.verificar_estructura(ruta) == []
    finally:
        wb.close()


def test_cambio_cartera_guarda_ediciones_antes_de_confirmar_traslado(expediente):
    import main
    import ui_estilo as ui
    from types import SimpleNamespace
    from test_documentos_ui import buscar
    c, cid, _ = expediente
    app = main.App.__new__(main.App)
    app.p, app.ctrl, app.cid, app.ejercicio = ui.CLARO, c, cid, 2025
    app.page = SimpleNamespace(update=lambda: None)
    app.cerrar_dialogo = lambda: None
    eventos, dialogo, trabajos = [], {}, []
    app._refrescar_clientes = lambda: eventos.append('refresco')
    app._confirmar_cambio_cartera = lambda *args: eventos.append(('confirmar', args))
    app._dialogo = lambda titulo, contenido, acciones, **kw: dialogo.update(contenido=contenido, acciones=acciones)
    app.en_segundo_plano = lambda titulo, fn, *args, **kw: trabajos.append((fn, kw['despues']))
    app.dialogo_editar_cliente()
    buscar(dialogo['contenido'], 'Nombre').value = 'Nombre corregido'
    buscar(dialogo['contenido'], 'Tipo de cliente').value = 'Autónomo'
    buscar(dialogo['contenido'], 'País de residencia 2025').value = 'Francia'
    dialogo['acciones'][-1].on_click(None)
    assert not eventos and c.db.get_cliente(cid)['nombre'] != 'Nombre corregido'
    guardar, despues = trabajos.pop()
    guardar()
    assert c.db.get_cliente(cid)['nombre'] == 'Nombre corregido'
    assert c.db.get_cliente(cid)['tipo_cliente'] == 'No residente'
    assert c.db.residencia(cid, 2025)['pais'] == 'Francia'
    assert not eventos
    despues()
    assert eventos == ['refresco', ('confirmar', (cid, 'Autónomo', 2025))]


def test_venta_exige_amortizacion_acreditada_y_actualiza_excel(expediente):
    c, cid, inm = expediente
    salida = c.marcar_vendido(cid, inm, '2025-06-01', 300000, 2025)
    ob = c.db.get_por_clave('obligaciones', cid, f'VENTA|{inm}')
    assert ob['estado'] == 'Revisar' and 'PROVISIONAL' in ob['notas']
    with pytest.raises(ValueError, match='amortización acumulada'):
        c.marcar_presentado(cid, 'obligaciones', ob['codigo'], 2025)
    firma = tc.firma_amortizacion_venta(c.db, cid, inm)
    with pytest.raises(ValueError, match='Confirma'):
        c.confirmar_amortizacion_venta(cid, inm, 2025, 5000, 'Historial de construcción revisado', firma)
    salida = c.confirmar_amortizacion_venta(cid, inm, 2025, 5000,
        'Historial de construcción revisado', firma, confirmado=True)
    assert not salida.get('error_excel')
    ob = c.db.get_por_clave('obligaciones', cid, f'VENTA|{inm}')
    assert ob['estado'] == 'Pendiente' and 'PROVISIONAL' not in ob['notas']
    assert 'ganancia 105.000,00 €' in ob['notas'] and 'cuota 19 % 19.950,00 €' in ob['notas']
    assert 'resultado 10.950,00 €' in ob['notas']
    assert tc.amortizacion_venta_confirmada(c.db, cid, inm)['importe'] == 5000
    wb = openpyxl.load_workbook(salida['excel'])
    try:
        col = next(x.letra for x in database.TABLAS['obligaciones'].columnas_excel if x.campo == 'notas')
        assert any(cell.value == ob['notas'] for cell in wb['14_OBLIGACIONES'][col])
    finally:
        wb.close()
    # Una reimportación intacta no invalida la confirmación.
    c.sincronizar(cid, 2025)
    assert tc.amortizacion_venta_confirmada(c.db, cid, inm)['importe'] == 5000
    # Una cifra de compra distinta sí obliga a volver a contrastar el cálculo.
    c.db.upsert('inmuebles', cid, {'codigo': inm, 'precio_escritura': 210000}, respetar_manuales=False)
    assert tc.amortizacion_venta_confirmada(c.db, cid, inm) is None


def test_excel_no_puede_dar_por_presentada_venta_sin_confirmar(expediente):
    c, cid, inm = expediente
    salida = c.marcar_vendido(cid, inm, '2025-06-01', 300000, 2025)
    wb = openpyxl.load_workbook(salida['excel'])
    hoja = wb['14_OBLIGACIONES']
    fila = next(i for i in range(2, hoja.max_row+1) if hoja[f'D{i}'].value == inm and
                'ganancia patrimonial venta' in str(hoja[f'K{i}'].value))
    columna_estado = next(x.letra for x in database.TABLAS['obligaciones'].columnas_excel if x.campo == 'estado')
    hoja[f'{columna_estado}{fila}'] = 'Presentado'
    wb.save(salida['excel']); wb.close()
    antes = c.db.listar('obligaciones', cid)
    with pytest.raises(ValueError, match='amortización'):
        eg.import_excel_to_db(c.db, cid)
    assert c.db.listar('obligaciones', cid) == antes


def test_venta_presentada_106_permanece_intacta_sin_confirmacion_nueva(expediente):
    c, cid, inm = expediente
    c.db.upsert('inmuebles', cid, {'codigo': inm, 'fecha_venta': '2025-06-01', 'precio_venta': 300000})
    codigo = c.db.upsert('obligaciones', cid, {'ejercicio': 2025, 'tipo': 'Otra', 'inmueble': inm,
        'estado': 'Presentado', 'pagado': 1234.56, 'notas': 'Justificante histórico inmutable'}, clave=f'VENTA|{inm}')
    antes = c.db.get('obligaciones', cid, codigo)
    informe = tc.recalcular_cliente(c.db, cid, 2025)
    assert c.db.get('obligaciones', cid, codigo) == antes
    assert not any(a.tipo == 'amortización venta' for a in informe.alertas)


def test_venta_amortizacion_acreditada_ya_incluye_titularidad():
    resultado = tc.calcular_venta({'fecha_venta': '2025-06-01', 'precio_venta': 300000,
        'precio_escritura': 200000, 'pct_titularidad': .5}, 5000)
    assert (resultado.ganancia, resultado.cuota, resultado.retencion, resultado.resultado) == (55000, 10450, 4500, 5950)


def test_dialogo_venta_valida_sin_perder_campos_y_captura_antes_del_worker():
    import main
    import ui_estilo as ui
    from types import SimpleNamespace
    from test_documentos_ui import buscar, recorrer
    main._cargar_modulos()
    app = main.App.__new__(main.App)
    app.p, app.cid, app.ejercicio = ui.CLARO, 1, 2025
    app.page = SimpleNamespace(update=lambda: None)
    app.ocupada = lambda: False
    cierres, trabajos, llamadas, dialogo = [], [], [], {}
    app.cerrar_dialogo = lambda: cierres.append(True)
    app.ctrl = SimpleNamespace(confirmar_amortizacion_venta=lambda *a, **kw: llamadas.append((a, kw)))
    app.en_segundo_plano = lambda texto, fn, *a, **kw: trabajos.append((fn, kw)) or True
    app._dialogo = lambda titulo, contenido, acciones, **kw: dialogo.update(contenido=contenido, acciones=acciones)
    app.dialogo_amortizacion_venta('INM-001', {'firma': 'firma-ficticia', 'acreditada': None})
    importe = buscar(dialogo['contenido'], 'Amortización acumulada acreditada (€)')
    nota = buscar(dialogo['contenido'], 'Historial y justificantes contrastados')
    confirmar = next(c for c in recorrer(dialogo['contenido']) if isinstance(c, main.ft.Checkbox))
    confirmar.value = True; importe.value = 'no válido'; nota.value = 'Historial contrastado'
    dialogo['acciones'][-1].on_click(None)
    assert not cierres and not trabajos and importe.value == 'no válido'
    importe.value = '5.000,00'
    dialogo['acciones'][-1].on_click(None)
    assert not cierres and len(trabajos) == 1  # El diálogo permanece hasta confirmar el guardado.
    importe.value = '999'; nota.value = 'Cambió el control después'
    trabajo, callbacks = trabajos.pop()
    trabajo()
    assert llamadas == [((1, 'INM-001', 2025, 5000.0, 'Historial contrastado', 'firma-ficticia'), {'confirmado': True})]


@pytest.mark.parametrize('valor', ['2025-01-02BASURA', '01/02/2025 99:99', '2025-02-30', '2025-01-02T99:00:00'])
def test_fechas_invalidas_no_se_corrigen_recortando_texto(valor):
    with pytest.raises(ValueError, match='Fecha no reconocida'):
        database.a_iso(valor)


@pytest.mark.parametrize('valor,esperado', [('1/2/2025', '2025-02-01'), ('2025-01-02T12:30:00Z', '2025-01-02'),
    ('2025-01-02 12:30:00', '2025-01-02'), ('01/02/2025 12:30', '2025-02-01')])
def test_fechas_y_marcas_de_tiempo_validas_se_conservan(valor, esperado):
    assert database.a_iso(valor) == esperado


@pytest.mark.parametrize('campo', ['precio_escritura', 'pct_titularidad'])
@pytest.mark.parametrize('valor', [True, False])
def test_booleanos_no_se_convierten_en_importes_o_porcentajes(campo, valor):
    columna = next(c for c in database.TABLAS['inmuebles'].columnas_bd if c.campo == campo)
    with pytest.raises(ValueError, match='valor lógico'):
        database._normalizar(columna, valor)


def test_porcentaje_con_tres_decimales_no_se_confunde_con_millares():
    assert database.a_num('1.500%') == .015
    assert database.a_num('1.500') == 1500  # Formato monetario español conservado.
    columna = next(c for c in database.TABLAS['inmuebles'].columnas_bd if c.campo == 'pct_titularidad')
    assert database._normalizar(columna, '1.500%') == .015


def test_campo_monetario_no_acepta_porcentaje_como_si_fueran_euros():
    columna = next(c for c in database.TABLAS['inmuebles'].columnas_bd if c.campo == 'precio_escritura')
    with pytest.raises(ValueError, match='euros'):
        database._normalizar(columna, '50%')


def test_prevision2027_no_se_presenta_como_calculo_definitivo(expediente):
    c, cid, inm = expediente
    c.db.upsert('catastro', cid, {'inmueble': inm, 'ejercicio': 2027, 'valor_catastral': 100000,
        'pct_imputacion': .02, 'revision_pendiente': 0})
    c.sincronizar(cid, 2027)
    fila = c.db.listar('m210_imputacion', cid, ejercicio=2027)[0]
    assert tc.AVISO_IMPUTACION_2027 in fila['observaciones']
    with pytest.raises(ValueError, match='2027'):
        c.marcar_presentado(cid, 'm210_imputacion', fila['codigo'], 2027)
    obligacion = c.db.listar('obligaciones', cid, id_modelo=fila['codigo'])[0]
    with pytest.raises(ValueError, match='2027'):
        c.marcar_presentado(cid, 'obligaciones', obligacion['codigo'], 2027)
    assert c.db.get('m210_imputacion', cid, fila['codigo'])['estado'] != 'Presentado'


@pytest.mark.parametrize('inicial,esperado', [
    (date(2026, 2, 28), date(2026, 3, 2)), (date(2027, 1, 30), date(2027, 2, 1)),
    (date(2026, 12, 6), date(2026, 12, 9)), (date(2024, 12, 6), date(2024, 12, 9)),
    (date(2027, 8, 15), date(2027, 8, 17)), (date(2025, 11, 9), date(2025, 11, 11)),
    (date(2026, 5, 15), date(2026, 5, 18)), (date(2026, 12, 31), date(2026, 12, 31)),
])
def test_vencimientos_electronicos_saltan_inhabiles_confirmados(inicial, esperado):
    from calendario_fiscal import siguiente_habil
    assert siguiente_habil(inicial, municipio='Madrid') == esperado


def test_347_y_4t303_no_vencen_en_sabado():
    assert dict(tc.plazos_autonomo('347', 2025))['Anual'] == date(2026, 3, 2)
    assert dict(tc.plazos_autonomo('303', 2026))['4T'] == date(2027, 2, 1)


def test_festivo_regional_no_se_extiende_a_todos_los_contribuyentes():
    from calendario_fiscal import siguiente_habil
    assert siguiente_habil(date(2025, 7, 25)) == date(2025, 7, 25)
    assert siguiente_habil(date(2025, 7, 25), municipio='Madrid') == date(2025, 7, 28)


def test_borrado_dependientes_revierte_si_falla_a_mitad(expediente, monkeypatch):
    c, cid, inm = expediente
    codigo = c.db.upsert('contratos', cid, {'inmueble': inm, 'fecha_inicio': '2025-01-01'})
    c.db.upsert('inquilinos', cid, {'contrato': codigo, 'nombre': 'Ficticio'})
    c.db.upsert('rentas', cid, {'contrato': codigo, 'importe': 1000})
    antes = {t: c.db.listar(t, cid) for t in ('contratos', 'inquilinos', 'rentas')}
    borrar = c.db.borrar
    def fallo(tabla, *args):
        if tabla == 'rentas':
            raise PermissionError('Fallo simulado durante la cascada')
        return borrar(tabla, *args)
    monkeypatch.setattr(c.db, 'borrar', fallo)
    with pytest.raises(PermissionError, match='cascada'):
        c.borrar_registro(cid, 'contratos', codigo, 2025)
    assert {t: c.db.listar(t, cid) for t in antes} == antes


def test_contrato_soporte_de_presentado_no_se_elimina(expediente):
    c, cid, inm = expediente
    ctr = c.db.upsert('contratos', cid, {'inmueble': inm, 'fecha_inicio': '2025-01-01'})
    c.db.upsert('m210_alquiler', cid, {'inmueble': inm, 'ejercicio': 2025, 'periodo': 'Anual',
        'estado': 'Presentado', 'cuota': 240}, clave=f'210A|2025|{inm}|Anual')
    with pytest.raises(ValueError, match='soporte histórico'):
        c.borrar_registro(cid, 'contratos', ctr, 2025)
    assert c.db.get('contratos', cid, ctr) is not None


def test_justificante303_a_ingresar_no_arrastra_credito_automatico(expediente):
    c, cid, _ = expediente
    fila = fr._guardar(c.db, 'm303', cid, '303|2025|1T', {'ejercicio': 2025, 'periodo': '1T',
        'a_ingresar': 0, 'a_compensar': 100}, 'Pendiente', {'importe': 210, 'csv': 'CSV-FICTICIO'})
    assert (fila['estado'], fila['a_ingresar'], fila['a_compensar']) == ('Presentado', 210, 0)
    resultado = fr.calcular_303([{'trimestre': '2T', 'base': 1000, 'cuota_iva': 210}], [], '2T', fila['a_compensar'])
    assert resultado.a_ingresar == 210


@pytest.mark.parametrize('modelo', ['303', '130', '111', '115'])
@pytest.mark.parametrize('nrc,pago_manual,esperado', [('N'*22, None, 210), ('CSV-FICTICIO', None, None),
                                                  ('N'*22, 0, 0), ('N'*22, 77, 77)])
def test_pago_residente_requiere_nrc_y_conserva_importe_manual(expediente, modelo, nrc, pago_manual, esperado):
    c, cid, _ = expediente
    tabla = {'303': 'm303', '130': 'm130'}.get(modelo, 'm111_115')
    campo = 'a_ingresar' if modelo in ('303', '130') else 'retenciones'
    codigo = c.db.upsert(tabla, cid, {'ejercicio': 2025, 'periodo': '1T', 'modelo': modelo,
        'estado': 'Presentado', campo: 210, 'nrc': nrc})
    if pago_manual is not None:
        c.db.upsert('obligaciones', cid, {'ejercicio': 2025, 'estado': 'Presentado', 'pagado': pago_manual},
                    clave=f'OBL|{modelo}|2025|1T')
    fila = c.db.get(tabla, cid, codigo)
    informe = fr.InformeResidente(2025, 'Autónomo', None, 'Exento', '', False)
    fr._calendario(c.db, cid, 2025, informe,
        {'1T': fila} if modelo == '303' else {}, {'1T': fila} if modelo == '130' else {},
        {('1T', modelo): fila} if modelo in ('111', '115') else {}, 'Autónomo')
    assert c.db.get_por_clave('obligaciones', cid, f'OBL|{modelo}|2025|1T')['pagado'] == esperado


def test_retenciones_inactivas_conservan_historia_sin_sumar_deuda(expediente):
    c, cid, _ = expediente
    residente = c.db.crear_cliente('FICTICIO', 'Retenciones', 'QARET107', tipo_cliente='Autónomo')['id']
    c.db.upsert('actividades', residente, {'regimen_iva': 'General', 'regimen_irpf': 'Directa simplificada'})
    codigo = c.db.upsert('m111_115', residente, {'ejercicio': 2025, 'periodo': '1T', 'modelo': '111',
        'perceptores': 1, 'base': 1000, 'retenciones': 150, 'estado': 'Pendiente'}, clave='RET|2025|1T|111')
    informe = fr.recalcular_residente(c.db, residente, 2025)
    fila = c.db.get('m111_115', residente, codigo)
    assert (fila['estado'], fila['base'], fila['retenciones']) == ('No procede', 1000, 150)
    assert informe.trimestres[0].retenciones_111 == 0


@pytest.mark.parametrize('nuevo_valor,nueva_construccion', [(110000, 70000), (100000, 60000)])
def test_ibi_posterior_no_pisa_valores_ni_revision_confirmada(expediente, tmp_path, nuevo_valor, nueva_construccion):
    import pdf_processor as pp
    import fixtures_pdf as fx
    c, cid, inm = expediente
    codigo = c.db.upsert('catastro', cid, {'inmueble': inm, 'ejercicio': 2024, 'valor_catastral': 100000,
        'valor_suelo': 40000, 'valor_construccion': 60000, 'anio_revision': 2012, 'pct_imputacion': .011,
        'revision_pendiente': 0}, clave=f'CAT|{inm}|2024')
    ruta = fx.recibo_ibi(tmp_path/'ibi-ficticio.pdf', ejercicio=2024, valor_catastral=nuevo_valor,
                         valor_suelo=40000, valor_construccion=nueva_construccion)
    extraccion = pp.extraer_local(ruta)
    extraccion.campos.pop('anio_revision', None)
    resultado = pp.aplicar(c.db, cid, extraccion, ruta, inmueble_destino=inm, ejercicio=2024)
    fila = c.db.get('catastro', cid, codigo)
    assert (fila['valor_catastral'], fila['valor_suelo'], fila['valor_construccion']) == (100000, 40000, 60000)
    assert (fila['anio_revision'], fila['pct_imputacion'], fila['revision_pendiente']) == (2012, .011, 0)
    if nuevo_valor != 100000:
        assert any('se conservan los valores' in texto for texto in resultado.mensajes)


def test_editor_generico_no_elude_confirmacion_de_venta(expediente):
    c, cid, inm = expediente
    c.marcar_vendido(cid, inm, '2025-06-01', 300000, 2025)
    fila = c.db.get_por_clave('obligaciones', cid, f'VENTA|{inm}')
    with pytest.raises(ValueError, match='amortización acumulada'):
        c.editar(cid, 'obligaciones', fila['codigo'], {'estado': 'Presentado'}, 2025)
    assert c.db.get('obligaciones', cid, fila['codigo'])['estado'] != 'Presentado'


@pytest.mark.parametrize('tabla', ['m210_imputacion', 'obligaciones'])
def test_editor_generico_no_elude_guardia_prevision2027(expediente, tabla):
    c, cid, inm = expediente
    c.sincronizar(cid, 2027)
    modelo = c.db.listar('m210_imputacion', cid, ejercicio=2027)[0]
    fila = modelo if tabla == 'm210_imputacion' else c.db.listar('obligaciones', cid, id_modelo=modelo['codigo'])[0]
    with pytest.raises(ValueError, match='2027'):
        c.editar(cid, tabla, fila['codigo'], {'estado': 'Presentado'}, 2027)
    assert c.db.get(tabla, cid, fila['codigo'])['estado'] != 'Presentado'


@pytest.mark.parametrize('ejercicio,tipo', [(2025, 'venta'), (2027, 'imputacion')])
def test_importador_preserva_presentados_anteriores_sin_guardias_nuevas(expediente, ejercicio, tipo):
    c, cid, inm = expediente
    if tipo == 'venta':
        c.db.upsert('inmuebles', cid, {'codigo': inm, 'fecha_venta': f'{ejercicio}-06-01', 'precio_venta': 300000})
        tabla = 'obligaciones'
        codigo = c.db.upsert(tabla, cid, {'ejercicio': ejercicio, 'inmueble': inm,
            'estado': 'Presentado', 'pagado': 123.45, 'notas': 'Justificante antiguo'}, clave=f'VENTA|{inm}')
    else:
        tabla = 'm210_imputacion'
        codigo = c.db.upsert(tabla, cid, {'ejercicio': ejercicio, 'inmueble': inm, 'estado': 'Presentado',
            'dias': 365, 'renta_imputada': 1234, 'cuota': 296.16, 'pagado': 296.16}, clave=f'210I|{ejercicio}|{inm}')
    antes = c.db.get(tabla, cid, codigo)
    eg.exportar_cliente(c.db, cid, ejercicio)
    eg.import_excel_to_db(c.db, cid)
    assert c.db.get(tabla, cid, codigo) == antes


def test_alias_territorio_no_clasifica_el_salvador_como_ue():
    assert tc.grupo_territorio('El Salvador') == tc.GRUPO_EXTRA
    assert tc.grupo_territorio('Venezuela') == tc.GRUPO_EXTRA


@pytest.mark.parametrize('cuota,estado,conflicto', [(240, 'Pendiente', True), (0, 'Pendiente', False),
    (-10, 'Pendiente', False), (240, 'Presentado', False)])
def test_anual_anterior2024_a_ingresar_se_revisa_sin_tocar_presentados(cuota, estado, conflicto):
    modelos = [{'codigo': '210A-001', 'clave': '210A|2023|INM-001|Anual', 'inmueble': 'INM-001',
        'ejercicio': 2023, 'periodo': 'Anual', 'cuota': cuota, 'estado': estado}]
    assert bool(tc._conflictos_modelos_210(modelos, 2023)) == conflicto


@pytest.mark.parametrize('valor', ['50%', True, False, float('nan'), float('inf')])
def test_importe_monetario_de_confirmacion_no_acepta_porcentajes_o_no_finitos(valor):
    with pytest.raises(ValueError):
        database.a_euros(valor)


@pytest.mark.parametrize('emitida', [False, True])
def test_numeracion_facturas_se_repite_por_ano_sin_pisar_106_ni_duplicar(expediente, tmp_path, emitida):
    import pdf_processor as pp
    import fixtures_pdf as fx
    c, _, _ = expediente
    cli = c.db.crear_cliente('FICTICIO', 'Facturas', 'QAFACT107', tipo_cliente='Autónomo')
    cid = cli['id']
    tabla = 'facturas_emitidas' if emitida else 'facturas_recibidas'
    prefijo = 'FE' if emitida else 'FR'
    nif_proveedor = cli['nif'] if emitida else 'B11111111'
    nif_receptor = 'B22222222' if emitida else cli['nif']
    tercero = nif_receptor if emitida else nif_proveedor
    legado = c.db.upsert(tabla, cid, {'numero': '1', 'fecha': '2024-02-01', 'base': 70,
        'cuota_iva': 14.7, 'nif_proveedor': nif_proveedor, 'nif_receptor': nif_receptor},
        clave=f'{prefijo}|1|{tercero}' if emitida else f'{prefijo}|{tercero}|1')
    for anio, base in [(2025, 100), (2024, 70)]:
        ruta = fx._crear(tmp_path/f'factura-{anio}.pdf', f'Factura ficticia 1 del ejercicio {anio}')
        campos = {'n_factura': '1', 'fecha': f'{anio}-02-01', 'base': base, 'iva': base*.21,
            'nif_proveedor': nif_proveedor, 'nif_receptor': nif_receptor, 'proveedor': 'Proveedor ficticio'}
        ext = pp.Extraccion(tipo=pp.TipoDocumento.FACTURA_GASTO, confianza=1,
            nombre_archivo=ruta.name, campos=campos)
        pp.aplicar(c.db, cid, ext, ruta, ejercicio=anio)
    filas = c.db.listar(tabla, cid)
    assert len(filas) == 2
    assert {f['ejercicio']: f['base'] for f in filas} == {2024: 70, 2025: 100}
    assert c.db.get(tabla, cid, legado)['fecha'] == '2024-02-01'


def test_rectificativa_sin_base_infiere_signo_retencion_y_base_correctos():
    import pdf_processor as pp
    tabla, datos, _ = pp.factura_residente({'nif': 'CLIENTE'}, {'fecha': '2025-02-01',
        'nif_proveedor': 'PROVEEDOR', 'nif_receptor': 'CLIENTE', 'n_factura': 'AB1',
        'total': -106, 'iva': -21, 'retencion': 15}, 'Factura rectificativa', 'ficticia')
    assert tabla == 'facturas_recibidas'
    assert (datos['base'], datos['cuota_iva'], datos['retencion']) == (-100, -21, -15)


def test_nif_del_bloque_cliente_no_se_atribuye_al_proveedor(tmp_path):
    import fixtures_pdf as fx
    import pdf_processor as pp
    ruta = fx._crear(tmp_path/'factura-ficticia.pdf', 'FACTURA\nProveedor: Suministros Ficticios\n'
        'Cliente: Negocio Ficticio\nNIF: B22222222\nFactura numero: 1\nFecha: 01/02/2025\n'
        'Base imponible: 100\nIVA: 21\nTotal: 121')
    ext = pp.extraer_local(ruta)
    assert not ext.campos.get('nif_proveedor')
    assert ext.campos['nif_receptor'] == 'B22222222'
    tabla, _, _ = pp.factura_residente({'nif': 'B22222222', 'nombre': 'Negocio Ficticio'},
                                      ext.campos, ext.texto, 'qa')
    assert tabla == 'facturas_recibidas'


def test_revisar_mismo_desglose_de_gastos_no_lo_detecta_como_factura_duplicada(expediente, tmp_path):
    import ia_extractor as ia
    import pdf_processor as pp
    from test_multiformato_integracion import respuesta
    c, cid, inm = expediente
    ruta = tmp_path/'gastos-ficticios.csv'
    ruta.write_text('factura;importe\n1;100', encoding='utf-8')
    campos = {'ejercicio': '2025', 'periodo_desde': '2025-01-01', 'periodo_hasta': '2025-01-31',
              'total': '100', 'n_factura': '1', 'proveedor': 'Proveedor ficticio'}
    fila = {'tabla': 'gastos', 'evidencia': 'fila 2', 'campos': [
        {'campo': k, 'valor': v, 'evidencia': 'fila 2'} for k, v in campos.items()]}
    ext = ia.a_extraccion(respuesta(registros=[fila]), ruta.name)
    ext.campos['inmueble'] = inm
    primero = pp.aplicar(c.db, cid, ext, ruta, ejercicio=2025)
    with c.db.transaccion():
        pp._aplicar(c.db, cid, ext, ruta, ejercicio=2025, documento_existente=primero.documento)
    assert len(c.db.listar('gastos', cid)) == 1


def test_presentacion_y_obligacion_se_revierten_juntas_si_falla_recalculo(expediente, monkeypatch):
    c, cid, inm = expediente
    c.db.upsert('catastro', cid, {'inmueble': inm, 'ejercicio': 2025, 'valor_catastral': 100000,
        'pct_imputacion': .02, 'revision_pendiente': 0}, clave=f'CAT|{inm}|2025')
    c.sincronizar(cid, 2025)
    fila = c.db.listar('m210_imputacion', cid, ejercicio=2025)[0]
    obligaciones = c.db.listar('obligaciones', cid, id_modelo=fila['codigo'])
    antes = c.db.get('m210_imputacion', cid, fila['codigo'])
    monkeypatch.setattr(c, '_recalcular_bd', lambda *a, **kw: (_ for _ in ()).throw(RuntimeError('fallo simulado')))
    with pytest.raises(RuntimeError, match='fallo simulado'):
        c.marcar_presentado(cid, 'm210_imputacion', fila['codigo'], 2025, pagado=480)
    assert c.db.get('m210_imputacion', cid, fila['codigo']) == antes
    assert c.db.listar('obligaciones', cid, id_modelo=fila['codigo']) == obligaciones


def test_facturas_no_residente_con_numero_anual_no_pisan_el_ejercicio_anterior(expediente, tmp_path):
    import pdf_processor as pp
    import fixtures_pdf as fx
    c, cid, inm = expediente
    codigo = c.db.upsert('gastos', cid, {'inmueble': inm, 'ejercicio': 2024, 'fecha':'2024-01-15',
        'total':70, 'n_factura':'1', 'proveedor':'Proveedor ficticio', 'nif_proveedor':'B11111111'},
        clave=f'FAC|{inm}|B11111111|Proveedor ficticio|1')
    for anio, total in [(2025, 100), (2024, 70)]:
        ruta = fx._crear(tmp_path/f'gasto-{anio}.pdf', f'Factura ficticia del ejercicio {anio}')
        ext = pp.Extraccion(tipo=pp.TipoDocumento.FACTURA_GASTO, confianza=1, nombre_archivo=ruta.name,
            campos={'fecha':f'{anio}-01-15','total':total,'n_factura':'1','proveedor':'Proveedor ficticio',
                'nif_proveedor':'B11111111', 'nif_receptor':c.db.get_cliente(cid)['nif']})
        pp.aplicar(c.db, cid, ext, ruta, inmueble_destino=inm, ejercicio=anio)
    filas = c.db.listar('gastos',cid)
    assert len(filas) == 2
    assert {f['ejercicio']:f['total'] for f in filas} == {2024:70,2025:100}
    assert c.db.get('gastos',cid,codigo)['fecha'] == '2024-01-15'


@pytest.mark.parametrize('presentado',[False,True])
def test_abrir_ano_actual_actualiza_imputacion_antigua_sin_alquiler_y_conserva_presentado(expediente,presentado,monkeypatch):
    _fijar_ano_controlador(monkeypatch,2026)
    c,cid,inm=expediente
    c.db.upsert('catastro',cid,{'inmueble':inm,'ejercicio':2024,'valor_catastral':100000,
        'pct_imputacion':.02,'revision_pendiente':0},clave=f'CAT|{inm}|2024')
    codigo=c.db.upsert('m210_imputacion',cid,{'ejercicio':2024,'inmueble':inm,'dias':366,
        'valor_catastral':100000,'pct_imputacion':.02,'renta_imputada':2005.48,'cuota':481.32,
        'estado':'Presentado' if presentado else 'Pendiente'},clave=f'210I|2024|{inm}')
    antes=c.db.get('m210_imputacion',cid,codigo)
    eg.exportar_cliente(c.db,cid,2024)
    c.db.set_config(f'excel_version|{cid}','1.0.6')
    c.abrir_cliente(cid,2026)
    despues=c.db.get('m210_imputacion',cid,codigo)
    if presentado:
        assert despues==antes
    else:
        assert (despues['renta_imputada'],despues['cuota'])==(2000,480)
    assert not c.db.listar('m210_alquiler',cid)


def test_editar_titularidad_en_ano_actual_actualiza_imputacion_anterior_sin_alquiler(expediente,monkeypatch):
    _fijar_ano_controlador(monkeypatch,2026)
    c,cid,inm=expediente
    c.db.upsert('catastro',cid,{'inmueble':inm,'ejercicio':2024,'valor_catastral':100000,
        'pct_imputacion':.02,'revision_pendiente':0},clave=f'CAT|{inm}|2024')
    c.sincronizar(cid,2024)
    c.set_gananciales(cid,inm,True,2026)
    fila=c.db.listar('m210_imputacion',cid,ejercicio=2024)[0]
    assert (fila['pct_titularidad'],fila['renta_imputada'],fila['cuota'])==(.5,1000,240)



def _fijar_ano_controlador(monkeypatch,anio):
    import controlador
    class FechaFija(date):
        @classmethod
        def today(cls):
            return cls(anio,10,29)
    monkeypatch.setattr(controlador,'date',FechaFija)


@pytest.mark.parametrize('historico,activo,incluido',[(2022,2026,True),(2021,2026,False),(2021,2021,True)])
def test_ventana_renovacion_no_abre_historicos_ajenos_ni_anos_vacios(expediente,monkeypatch,historico,activo,incluido):
    _fijar_ano_controlador(monkeypatch,2026)
    c,cid,inm=expediente
    c.db.upsert('m210_imputacion',cid,{'ejercicio':historico,'inmueble':inm,'estado':'Pendiente'},
                clave=f'210I|{historico}|{inm}')
    otro=c.db.crear_cliente('FICTICIO','Ajeno','QA-AJENO107')['id']
    c.db.upsert('m210_imputacion',otro,{'ejercicio':2023,'inmueble':'INM-001'},clave='210I|2023|INM-001')
    llamados=[]
    def motor(db,cliente,ejercicios):
        assert cliente==cid
        llamados.extend(ejercicios)
        return {e:object() for e in ejercicios}
    monkeypatch.setattr(tc,'recalcular_ejercicios',motor)
    c._recalcular_bd(cid,{activo},activo)
    assert (historico in llamados)==incluido
    assert set(llamados)==({activo,historico} if incluido else {activo})



@pytest.mark.parametrize('grupo',[None,tc.GRUPO_UE,tc.GRUPO_EXTRA])
def test_recalculo_no_persiste_hipotesis_de_residencia_vacia_y_respeta_grupo_manual(expediente,grupo):
    c,cid,_=expediente
    codigo=c.db.upsert('residencias',cid,{'codigo':'2025','ejercicio':2025,'grupo':grupo},clave='RES|2025')
    tc.recalcular_cliente(c.db,cid,2025)
    fila=c.db.get('residencias',cid,codigo)
    assert fila['pais'] is None and fila['certificado'] is None and fila['ruta'] is None
    if grupo is None:
        assert fila['grupo'] is None and fila['gastos_deducibles'] is None
        assert tc.resolver_residencia(fila).asumida
    else:
        assert fila['grupo']==grupo
        assert not tc.resolver_residencia(fila).asumida
        assert fila['gastos_deducibles']==('Sí' if grupo==tc.GRUPO_UE else 'No')



@pytest.mark.parametrize('pais,grupo', [(' ',None),(None,' '),('None','nan'),('nan','invalido'),('null','desconocido')])
def test_residencia_vacia_o_grupo_invalido_sigue_provisional_sin_mutar_entrada(pais,grupo):
    fila={'pais':pais,'grupo':grupo};antes=dict(fila)
    r=tc.resolver_residencia(fila)
    assert r.asumida and r.tipo==.24 and not r.deduce_gastos
    assert fila==antes


@pytest.mark.parametrize('pais,grupo', [(' Francia ',None),(None,' UE/EEE '),(' no ',None)])
def test_espacios_no_invalidan_pais_ni_grupo_y_no_sigue_siendo_codigo_noruega(pais,grupo):
    r=tc.resolver_residencia({'pais':pais,'grupo':grupo})
    assert not r.asumida and r.tipo==.19 and r.deduce_gastos


def test_no_deducir_manual_con_espacios_y_minusculas_se_respeta():
    r=tc.resolver_residencia({'pais':'Francia','gastos_deducibles':' no '})
    assert r.tipo==.19 and not r.deduce_gastos

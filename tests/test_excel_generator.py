"""Contratos de intercambio Excel: fórmulas, estilos y conservación de datos."""
import hashlib
import sys
from copy import copy
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from openpyxl.formula.translate import Translator

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import database
import excel_generator as eg
from database import Database


@pytest.fixture
def db(tmp_path):
    d = Database(tmp_path / "test.db")
    d.carpeta_clientes = tmp_path / "Clientes"
    yield d
    d.cerrar()


@pytest.fixture
def cliente(db):
    return db.crear_cliente("Pepito", "Pérez", "X1234567A", nacionalidad="Francesa")


def test_estructura_resumen_estilos_y_plantilla_intacta(db, cliente):
    huella = hashlib.sha256(database.PLANTILLA_EXCEL.read_bytes()).hexdigest()
    cid = cliente["id"]
    db.guardar_residencia(cid, 2026, pais="Francia")
    db.upsert("inmuebles", cid, {"fecha_adquisicion": "2019-05-10", "precio_escritura": 180000,
                                "pct_titularidad": 0.5, "carpeta": "Soportes_PDF/compra.pdf"})
    ruta = eg.exportar_cliente(db, cid, 2026)
    assert ruta == db.ruta_excel(cliente)
    assert eg.verificar_estructura(ruta) == []
    wb = openpyxl.load_workbook(ruta)
    plantilla = openpyxl.load_workbook(database.PLANTILLA_EXCEL)
    try:
        assert len(wb.sheetnames) == 17
        assert wb["00_RESUMEN"]["B4"].value == "Pepito Pérez"
        assert wb["00_RESUMEN"]["B5"].value == "X1234567A"
        assert wb["00_RESUMEN"]["B12"].value.date() == date.today()
        assert wb["00_RESUMEN"]["E4"].value == 2026
        assert wb["00_RESUMEN"]["B11"].hyperlink.target == Path(cliente["carpeta"]).resolve().as_uri()
        assert wb["01_RESIDENCIA_FISCAL"]["A2"].value == 2026
        ws, original = wb["02_INMUEBLES"], plantilla["02_INMUEBLES"]
        assert ws["G2"].value.date() == date(2019, 5, 10)
        assert ws["H2"].value == 0.5
        assert ws["J2"].value == 180000
        assert ws["Z2"].value == "compra.pdf"
        assert ws["Z2"].hyperlink.target == "Soportes_PDF/compra.pdf"
        assert ws["Z2"].font.color.rgb.endswith("0000FF")
        assert ws["Z2"].font.underline == "single"
        assert ws["Z2"].font.name == original["Z2"].font.name
        assert copy(ws["Z2"].border) == copy(original["Z2"].border)
        for coord in ("G2", "H2", "J2", "Q2"):
            assert ws[coord]._style == original[coord]._style
        assert [[c.value for c in row] for row in wb["16_COMO_USAR"]] == [
            [c.value for c in row] for row in plantilla["16_COMO_USAR"]]
    finally:
        wb.close()
        plantilla.close()
    assert hashlib.sha256(database.PLANTILLA_EXCEL.read_bytes()).hexdigest() == huella


def test_ida_vuelta_todas_las_tablas_sin_perder_campos_solo_bd(db, cliente):
    cid = cliente["id"]
    for nombre, tabla in database.TABLAS.items():
        if not tabla.hoja:
            continue
        datos = {}
        for col in tabla.columnas_bd:
            if col.campo == "codigo":
                continue
            if col.tipo == "date":
                valor = date(2026, 2, 3)
            elif col.tipo == "eur":
                valor = 1234.56
            elif col.tipo == "pct":
                valor = 0.5
            elif col.tipo == "int":
                valor = 2026
            elif col.tipo == "link":
                valor = "Soportes_PDF/prueba.pdf"
            else:
                valor = "Prueba con tildes ñ " + col.campo
            datos[col.campo] = valor
        if nombre == "residencias":
            datos.update(codigo="2026", ejercicio=2026)
        db.upsert(nombre, cid, datos, clave=f"PRUEBA|{nombre}")
    antes = {n: db.listar(n, cid) for n, t in database.TABLAS.items() if t.hoja}
    ruta = eg.exportar_cliente(db, cid, 2026)
    db.set_config(f"excel_firma|{cid}", "")
    conteos = eg.import_excel_to_db(db, cid)
    assert conteos and all(n == 1 for n in conteos.values())
    for nombre, filas in antes.items():
        despues = db.get(nombre, cid, filas[0]["codigo"])
        assert {k: v for k, v in despues.items() if k != "actualizado"} == {
            k: v for k, v in filas[0].items() if k != "actualizado"}
    assert eg.verificar_estructura(ruta) == []


def test_gasto_manual_sin_id_sobrevive_y_descuenta_210(db, cliente):
    cid = cliente["id"]
    db.guardar_residencia(cid, 2026, pais="Francia")
    inm = db.upsert("inmuebles", cid, {"fecha_adquisicion": "2019-01-01", "pct_titularidad": 0.5})
    db.upsert("contratos", cid, {"inmueble": inm, "fecha_inicio": "2026-01-01",
                                "renta_inicial": 1000, "estado": "Vigente"})
    ruta = eg.exportar_cliente(db, cid, 2026)
    wb = openpyxl.load_workbook(ruta)
    ws = wb["08_GASTOS"]
    for coord, valor in {"B2": 2026, "C2": inm, "E2": date(2026, 8, 1),
                          "H2": "Comunidad", "N2": 300, "V2": "Introducido a mano"}.items():
        ws[coord] = valor
    wb.save(ruta)
    wb.close()
    eg.import_excel_to_db(db, cid)
    eg.import_excel_to_db(db, cid)  # reintentar antes de exportar tampoco duplica
    assert len(db.listar("gastos", cid)) == 1
    eg.sincronizar(db, cid, 2026)
    gasto = db.listar("gastos", cid)[0]
    assert gasto["origen"] == "excel" and gasto["observaciones"] == "Introducido a mano"
    assert gasto["codigo"].startswith("GAS-")
    modelo = db.get_por_clave("m210_alquiler", cid, f"210A|2026|{inm}|3T")
    assert modelo["gastos"] == 150
    wb = openpyxl.load_workbook(ruta)
    try:
        assert wb["08_GASTOS"]["A2"].value == gasto["codigo"]
        assert wb["08_GASTOS"]["V2"].value == "Introducido a mano"
    finally:
        wb.close()


def test_celda_vaciada_y_resumen_mandan_sobre_bd(db, cliente):
    cid = cliente["id"]
    cod = db.upsert("inmuebles", cid, {"direccion": "Calle Mayor", "observaciones": "Borrar", "precio_escritura": 100})
    db.actualizar_cliente(cid, email="anterior@example.invalid")
    ruta = eg.exportar_cliente(db, cid, 2026)
    wb = openpyxl.load_workbook(ruta)
    wb["02_INMUEBLES"]["AA2"] = None
    wb["02_INMUEBLES"]["J2"] = 0
    wb["00_RESUMEN"]["B8"] = None
    wb["00_RESUMEN"]["B10"] = "Ana"
    wb.save(ruta)
    wb.close()
    eg.import_excel_to_db(db, cid)
    fila = db.get("inmuebles", cid, cod)
    assert fila["observaciones"] is None and fila["precio_escritura"] == 0
    assert fila["direccion"] == "Calle Mayor"
    assert db.get_cliente(cid)["email"] is None
    assert db.get_cliente(cid)["responsable"] == "Ana"


def test_filas_borradas_solo_eliminan_gastos_de_excel(db, cliente):
    cid = cliente["id"]
    manual = db.upsert("gastos", cid, {"origen": "excel", "total": 10})
    pdf = db.upsert("gastos", cid, {"origen": "pdf", "total": 20})
    ruta = eg.exportar_cliente(db, cid, 2026)
    wb = openpyxl.load_workbook(ruta)
    ws = wb["08_GASTOS"]
    for fila in (2, 3):
        for col in database.TABLAS["gastos"].columnas_excel:
            if col.tipo != "formula":
                ws[f"{col.letra}{fila}"] = None
    wb.save(ruta)
    wb.close()
    eg.import_excel_to_db(db, cid)
    assert db.get("gastos", cid, manual) is None
    assert db.get("gastos", cid, pdf)["total"] == 20
    assert db.get_por_clave("notas", cid, f"BORRADA|gastos|{pdf}") is not None
    eg.exportar_cliente(db, cid, 2026)
    eg.import_excel_to_db(db, cid)
    assert len([n for n in db.listar("notas", cid) if n["clave"] == f"BORRADA|gastos|{pdf}"]) == 1


def test_excel_bloqueado_conserva_original_y_deja_copia(db, cliente, monkeypatch):
    cid = cliente["id"]
    ruta = eg.exportar_cliente(db, cid, 2026)
    anterior = ruta.read_bytes()
    db.upsert("inmuebles", cid, {"direccion": "Datos nuevos"})

    def bloqueado(*args):
        raise PermissionError("Archivo en uso")

    monkeypatch.setattr(eg.os, "replace", bloqueado)
    with pytest.raises(eg.ExcelBloqueadoError, match="Cierra el Excel.*pulsa Sincronizar") as error:
        eg.exportar_cliente(db, cid, 2026)
    pendiente = error.value.ruta_pendiente
    assert "_PENDIENTE_" in pendiente.name and pendiente.exists()
    assert ruta.read_bytes() == anterior
    wb = openpyxl.load_workbook(pendiente)
    try:
        assert wb["02_INMUEBLES"]["B2"].value == "Datos nuevos"
    finally:
        wb.close()
    assert not list(ruta.parent.glob(".nr_*.xlsx"))


def test_backups_limitados_a_diez(db, cliente):
    cid = cliente["id"]
    ruta = eg.exportar_cliente(db, cid, 2026)
    backups = ruta.parent / "_backups"
    backups.mkdir()
    for n in range(12):
        copia = backups / f"{ruta.stem}_20260101_120000_{n:06d}.xlsx"
        copia.write_bytes(b"prueba")
        eg.os.utime(copia, (n + 1, n + 1))
    eg.exportar_cliente(db, cid, 2026)
    assert len(list(backups.glob("*.xlsx"))) == 10
    assert not (backups / f"{ruta.stem}_20260101_120000_000000.xlsx").exists()
    assert any(p.name.startswith(ruta.stem) for p in backups.glob("*.xlsx"))


def test_filas_adicionales_copian_estilo_y_traducen_formula(db, cliente, monkeypatch):
    plantilla = openpyxl.load_workbook(database.PLANTILLA_EXCEL)
    ws = plantilla["02_INMUEBLES"]
    ultima = ws.max_row + 1
    formula = Translator(ws["Q2"].value, origin="Q2").translate_formula(f"Q{ultima}")
    estilo = ws["J2"]._style
    plantilla.close()
    iterar_original = db.iterar

    def iterar(tabla, cid, *args, **kwargs):
        if tabla == "inmuebles":
            yield from ({"codigo": f"INM-{n:03}", "precio_escritura": n} for n in range(1, ultima))
        else:
            yield from iterar_original(tabla, cid, *args, **kwargs)

    monkeypatch.setattr(db, "iterar", iterar)
    ruta = eg.exportar_cliente(db, cliente["id"], 2026)
    wb = openpyxl.load_workbook(ruta)
    try:
        assert wb["02_INMUEBLES"][f"Q{ultima}"].value == formula
        assert wb["02_INMUEBLES"][f"J{ultima}"]._style == estilo
    finally:
        wb.close()
    assert eg.verificar_estructura(ruta) == []


def test_verificador_detecta_formula_validacion_y_cabecera(db, cliente):
    ruta = eg.exportar_cliente(db, cliente["id"], 2026)
    wb = openpyxl.load_workbook(ruta)
    ws = wb["02_INMUEBLES"]
    ws["Q2"] = 0
    ws["A1"] = "Otra cabecera"
    ws.column_dimensions["B"].width = 3
    ws.data_validations.dataValidation.clear()
    wb.save(ruta)
    wb.close()
    diferencias = eg.verificar_estructura(ruta)
    assert any("Q2" in d for d in diferencias)
    assert any("fila 1" in d for d in diferencias)
    assert any("anchos" in d for d in diferencias)
    assert any("validaciones" in d for d in diferencias)


def test_ruta_opcional_inexistente_y_proteccion_plantilla(db, cliente, tmp_path):
    assert eg.import_excel_to_db(db, cliente["id"]) == {}
    destino = tmp_path / "otro" / "cliente.xlsx"
    assert eg.exportar_cliente(db, cliente["id"], 2026, destino) == destino
    assert eg.import_excel_to_db(db, cliente["id"], destino)["inmuebles"] == 0
    with pytest.raises(ValueError, match="plantilla"):
        eg.exportar_cliente(db, cliente["id"], 2026, database.PLANTILLA_EXCEL)


def test_import_excel_to_db_omite_apertura_si_firma_coincide(db, cliente, monkeypatch):
    cid = cliente["id"]
    ruta = eg.exportar_cliente(db, cid, 2026)
    firma = db.get_config(f"excel_firma|{cid}")
    assert firma == f"{ruta.stat().st_mtime_ns}|{ruta.stat().st_size}"

    load_original = eg.openpyxl.load_workbook
    aperturas = []

    def espiar_load(*args, **kwargs):
        aperturas.append(args)
        return load_original(*args, **kwargs)

    monkeypatch.setattr(eg.openpyxl, "load_workbook", espiar_load)
    assert eg.import_excel_to_db(db, cid) == {}
    assert aperturas == []

    # Si el usuario modifica el Excel, la firma cambia y sí se abre e importa
    wb = load_original(ruta)
    wb["02_INMUEBLES"]["B2"] = "Calle Modificada 10"
    wb.save(ruta)
    wb.close()

    conteos = eg.import_excel_to_db(db, cid)
    assert len(aperturas) == 1
    assert conteos["inmuebles"] == 1
    assert db.listar("inmuebles", cid)[0]["direccion"] == "Calle Modificada 10"


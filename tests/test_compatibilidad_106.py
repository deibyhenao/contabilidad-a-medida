"""Políticas de compatibilidad: libros existentes, calendario heredado y respaldo previo.

Son pruebas de integración sintéticas, NO sustituyen el ensayo de los instaladores públicos con Excel real.
"""
from pathlib import Path
import sqlite3

import openpyxl
import pytest

from controlador import Controlador
import database
import excel_generator as eg
import tax_calculator as tc


@pytest.fixture
def expediente(tmp_path):
    ctrl = Controlador(tmp_path / 'ficticio.db')
    ctrl.cambiar_carpeta_clientes(tmp_path / 'Datos externos')
    cid = ctrl.db.crear_cliente('FICTICIO', 'Compatibilidad', 'QA106NR01')['id']
    inm = ctrl.db.upsert('inmuebles', cid, {'direccion': 'Dirección manual conservada',
        'fecha_adquisicion': '2020-01-01', 'pct_titularidad': 1, 'observaciones': 'Nota manual conservada'})
    ctrl.db.upsert('contratos', cid, {'inmueble': inm, 'fecha_inicio': '2025-01-01',
        'fecha_fin_real': '2025-09-30', 'renta_inicial': 1000, 'tipo': 'Temporada', 'estado': 'Finalizado'})
    try:
        yield ctrl, cid, inm
    finally:
        ctrl.cerrar()


def test_trimestrales_pendientes_anteriores_conservan_calendario(expediente):
    ctrl, cid, inm = expediente
    for p in ('1T', '2T', '3T', '4T'):
        ctrl.db.upsert('m210_alquiler', cid, {'ejercicio': 2025, 'inmueble': inm,
            'periodo': p, 'estado': 'Pendiente'}, clave=f'210A|2025|{inm}|{p}')
    assert tc.periodicidad_210_alquiler(ctrl.db, cid, inm, 2025) == 'Trimestral'
    informe = tc.recalcular_cliente(ctrl.db, cid, 2025)
    assert all(ctrl.db.get('m210_alquiler', cid, c)['periodo'] != 'Anual' for c in informe.filas_210_alquiler)


def test_cliente_sin_modelos_anteriores_recibe_anual(expediente):
    ctrl, cid, inm = expediente
    assert tc.periodicidad_210_alquiler(ctrl.db, cid, inm, 2025) == 'Anual'


def test_libro_anterior_sin_edicion_se_renueva_al_abrir_y_no_se_repite(expediente):
    ctrl, cid, inm = expediente
    ruta = ctrl.sincronizar(cid, 2025)['excel']
    wb = openpyxl.load_workbook(ruta)
    ws = wb['00_RESUMEN']
    inicio = next(c.row for row in ws for c in row if str(c.value or '').startswith('OCUPACIÓN Y RENTAS'))
    for row in ws.iter_rows(min_row=inicio, max_row=inicio + 3, max_col=8):
        for c in row:
            c.value = None
    wb.save(ruta)
    wb.close()
    # Simula un libro exportado y reconocido por la aplicación anterior, no una edición pendiente.
    ctrl.db.set_config(f'excel_version|{cid}', 'anterior')
    ctrl.db.set_config(f'excel_firma|{cid}', eg._firma_excel(ruta))
    assert not eg.excel_modificado(ctrl.db, cid)
    salida = ctrl.abrir_cliente(cid, 2025)
    assert salida['excel'] == ruta
    wb = openpyxl.load_workbook(ruta)
    try:
        assert str(wb['00_RESUMEN'].cell(inicio, 1).value).startswith('OCUPACIÓN Y RENTAS')
    finally:
        wb.close()
    assert ctrl.db.get('inmuebles', cid, inm)['observaciones'] == 'Nota manual conservada'
    firma = eg._firma_excel(ruta)
    assert ctrl.abrir_cliente(cid, 2025)['excel'] is None
    assert eg._firma_excel(ruta) == firma
    assert len(ctrl.db.listar_clientes()) == 1


def test_respaldo_de_actualizacion_una_sola_vez_y_antes_de_tocar_datos(expediente):
    ctrl, cid, _ = expediente
    ruta = ctrl.db.db_path
    ctrl.db.set_config('version_datos_aplicacion', 'anterior')
    ctrl.cerrar()
    carpeta = ruta.parent / 'Backups_Seguridad'
    antes = set(carpeta.glob('*.db'))
    nuevo = Controlador(ruta)
    nuevo.cerrar()
    copias = set(carpeta.glob('*.db')) - antes
    assert len(copias) == 1
    with sqlite3.connect(next(iter(copias))) as con:
        assert con.execute('SELECT count(*) FROM clientes').fetchone()[0] == 1
        assert con.execute('SELECT valor FROM config WHERE clave=?', ('version_datos_aplicacion',)).fetchone()[0] == 'anterior'
    otra = Controlador(ruta)
    otra.cerrar()
    assert set(carpeta.glob('*.db')) - antes == copias


def test_no_recrea_base_vacia_si_hay_expedientes_anteriores(tmp_path):
    from app_paths import AlmacenNoDisponible
    ruta = tmp_path / 'contabilidad.db'
    (tmp_path / 'Clientes_No_Residentes' / 'FICTICIO').mkdir(parents=True)
    with pytest.raises(AlmacenNoDisponible, match='vacía'):
        Controlador(ruta)
    assert not ruta.exists()


def test_fallo_de_respaldo_no_marca_version_ni_modifica_clientes(expediente, monkeypatch):
    import migracion_actualizacion as ma
    ctrl, cid, _ = expediente
    ctrl.db.set_config('version_datos_aplicacion', 'anterior')
    monkeypatch.setattr(ma, 'crear_snapshot', lambda *a: (_ for _ in ()).throw(PermissionError('Prueba de respaldo')))
    with pytest.raises(PermissionError):
        ma.preparar_actualizacion(ctrl.db)
    assert ctrl.db.get_config('version_datos_aplicacion') == 'anterior'
    assert ctrl.db.get_cliente(cid)['nombre'] == 'FICTICIO'


def test_no_recrea_base_truncada_a_cero_si_hay_expedientes(tmp_path):
    from app_paths import AlmacenNoDisponible
    ruta = tmp_path / 'contabilidad.db'
    ruta.touch()
    (tmp_path / 'Clientes_No_Residentes' / 'FICTICIO').mkdir(parents=True)
    with pytest.raises(AlmacenNoDisponible, match='vacía'):
        Controlador(ruta)
    assert ruta.read_bytes() == b''


def test_excel_de_otro_cliente_no_pisa_el_expediente(expediente):
    import shutil
    ctrl, cid, inm = expediente
    ruta = ctrl.sincronizar(cid, 2025)['excel']
    otro = ctrl.db.crear_cliente('OTRO FICTICIO', 'No mezclar', 'QA106NR02')['id']
    ctrl.db.upsert('inmuebles', otro, {'direccion': 'NO IMPORTAR EN OTRO CLIENTE', 'codigo': inm})
    ajeno = eg.exportar_cliente(ctrl.db, otro, 2025)
    shutil.copyfile(ajeno, ruta)
    contenido = ruta.read_bytes()
    with pytest.raises(ValueError, match='otro cliente'):
        ctrl.abrir_cliente(cid, 2025)
    assert ctrl.db.get('inmuebles', cid, inm)['direccion'] == 'Dirección manual conservada'
    assert ruta.read_bytes() == contenido


def test_identidad_borrada_del_excel_no_se_importa(expediente):
    ctrl, cid, inm = expediente
    ruta = ctrl.sincronizar(cid, 2025)['excel']
    wb = openpyxl.load_workbook(ruta)
    wb['00_RESUMEN']['B5'] = None
    wb['02_INMUEBLES']['B2'] = 'NO IMPORTAR SIN IDENTIDAD'
    wb.save(ruta)
    wb.close()
    with pytest.raises(ValueError, match='Falta el NIF'):
        ctrl.abrir_cliente(cid, 2025)
    assert ctrl.db.get('inmuebles', cid, inm)['direccion'] == 'Dirección manual conservada'


@pytest.mark.parametrize('contenido', [b'', b'NO ES UN EXCEL', b'PK\x03\x04TRUNCADO'])
def test_excel_corrupto_no_se_sobrescribe_ni_modifica_bd(expediente, contenido):
    from zipfile import BadZipFile
    ctrl, cid, inm = expediente
    ruta = ctrl.sincronizar(cid, 2025)['excel']
    ruta.write_bytes(contenido)
    with pytest.raises((BadZipFile, ValueError, OSError)):
        ctrl.abrir_cliente(cid, 2025)
    assert ruta.read_bytes() == contenido
    assert ctrl.db.get('inmuebles', cid, inm)['observaciones'] == 'Nota manual conservada'


def test_excel_bloqueado_no_confirma_formato_nuevo_y_reintenta(expediente, monkeypatch):
    ctrl, cid, inm = expediente
    ruta = ctrl.sincronizar(cid, 2025)['excel']
    ctrl.db.set_config(f'excel_version|{cid}', 'anterior')
    anterior = ruta.read_bytes()
    reemplazar = eg.os.replace

    def bloqueado(origen, destino):
        if Path(destino) == ruta:
            raise PermissionError('Bloqueo ficticio de Excel')
        return reemplazar(origen, destino)

    with monkeypatch.context() as contexto:
        contexto.setattr(eg.os, 'replace', bloqueado)
        resultado = ctrl.abrir_cliente(cid, 2025)
    assert resultado['error_excel']
    assert 'PENDIENTE' in resultado['excel'].name
    assert resultado['excel'].exists()
    assert ruta.read_bytes() == anterior
    assert ctrl.db.get_config(f'excel_version|{cid}') == 'anterior'
    assert eg.necesita_actualizar_formato(ctrl.db, cid)
    salida = ctrl.abrir_cliente(cid, 2025)
    assert salida['error_excel'] is None
    assert not eg.necesita_actualizar_formato(ctrl.db, cid)
    assert ctrl.db.get('inmuebles', cid, inm)['observaciones'] == 'Nota manual conservada'


def test_copia_auxiliar_no_da_por_actualizado_el_excel_canonico(expediente, tmp_path):
    ctrl, cid, _ = expediente
    ctrl.sincronizar(cid, 2025)
    ctrl.db.set_config(f'excel_version|{cid}', 'anterior')
    eg.exportar_cliente(ctrl.db, cid, 2025, tmp_path / 'copia.xlsx')
    assert ctrl.db.get_config(f'excel_version|{cid}') == 'anterior'
    assert eg.necesita_actualizar_formato(ctrl.db, cid)


@pytest.mark.parametrize('tamano', [None, 0])
def test_base_ausente_o_truncada_con_solo_respaldos_no_se_recrea(tmp_path, tamano):
    from app_paths import AlmacenNoDisponible
    ruta = tmp_path / 'contabilidad.db'
    if tamano is not None:
        ruta.touch()
    carpeta = tmp_path / 'Backups_Seguridad'
    carpeta.mkdir()
    # Basta la señal de un respaldo previo; no se restaura automáticamente.
    (carpeta / 'backup_pre_update_FICTICIO.db').write_bytes(b'COPIA FICTICIA')
    with pytest.raises(AlmacenNoDisponible):
        Controlador(ruta)
    if tamano is None:
        assert not ruta.exists()
    else:
        assert ruta.read_bytes() == b''


def test_marca_de_version_atomica_si_falla_despues_de_crear_respaldo(expediente, monkeypatch):
    import migracion_actualizacion as ma
    ctrl, cid, _ = expediente
    ctrl.db.set_config('version_datos_aplicacion', 'anterior')
    original = ctrl.db.set_config

    def interrumpido(clave, valor):
        if clave == 'version_datos_aplicacion':
            raise OSError('Interrupción ficticia tras respaldar')
        return original(clave, valor)

    with monkeypatch.context() as contexto:
        contexto.setattr(ctrl.db, 'set_config', interrumpido)
        with pytest.raises(OSError, match='Interrupción'):
            ma.preparar_actualizacion(ctrl.db)
    assert ctrl.db.get_config('version_datos_aplicacion') == 'anterior'
    assert ctrl.db.get_config('respaldo_actualizacion|' + ma.VERSION) is None
    copias = list((ctrl.db.db_path.parent / 'Backups_Seguridad').glob('*.db'))
    assert copias
    with sqlite3.connect(copias[-1]) as con:
        assert con.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
        assert con.execute('SELECT count(*) FROM clientes').fetchone()[0] == 1
    assert ma.preparar_actualizacion(ctrl.db).is_file()
    assert ma.preparar_actualizacion(ctrl.db) is None


def test_edicion_manual_en_libro_antiguo_se_importa_antes_de_renovarlo(expediente):
    ctrl, cid, inm = expediente
    ruta = ctrl.sincronizar(cid, 2025)['excel']
    ctrl.db.set_config(f'excel_version|{cid}', 'anterior')
    columna = next(c.letra for c in database.TABLAS['contratos'].columnas_excel if c.campo == 'renta_inicial')
    wb = openpyxl.load_workbook(ruta)
    wb['04_CONTRATOS'][columna + '2'] = 1100
    wb.save(ruta)
    wb.close()
    salida = ctrl.abrir_cliente(cid, 2025)
    assert salida['error_excel'] is None
    assert salida['excel'] == ruta
    assert ctrl.db.listar('contratos', cid)[0]['renta_inicial'] == 1100
    modelo = ctrl.db.get_por_clave('m210_alquiler', cid, f'210A|2025|{inm}|Anual')
    assert modelo['ingresos'] == 9900
    assert modelo['cuota'] == 2376
    assert not list(ruta.parent.glob('*PENDIENTE*.xlsx'))
    wb = openpyxl.load_workbook(ruta)
    try:
        assert wb['04_CONTRATOS'][columna + '2'].value == 1100
        assert any(str(c.value or '').startswith('OCUPACIÓN Y RENTAS') for row in wb['00_RESUMEN'] for c in row)
    finally:
        wb.close()


def test_preferencia_global_no_cambia_calendario_ya_trabajado(expediente):
    ctrl, cid, inm = expediente
    ctrl.db.upsert('m210_alquiler', cid, {'ejercicio': 2025, 'inmueble': inm,
        'periodo': '1T', 'estado': 'Pendiente'}, clave=f'210A|2025|{inm}|1T')
    ctrl.db.set_config('periodicidad_210_alquiler', 'Anual')
    assert tc.periodicidad_210_alquiler(ctrl.db, cid, inm, 2025) == 'Trimestral'
    # La elección expresa para ESTE inmueble/año sí tiene prioridad.
    ctrl.db.set_config(f'periodicidad_210|{cid}|{inm}|2025', 'Anual')
    assert tc.periodicidad_210_alquiler(ctrl.db, cid, inm, 2025) == 'Anual'


def test_segundo_conflicto_no_sobrescribe_copia_pendiente_anterior(expediente):
    import os
    ctrl, cid, _ = expediente
    ruta = ctrl.sincronizar(cid, 2025)['excel']

    def provocar_conflicto():
        ctrl._firmas[cid] = eg._firma_excel(ruta)
        st = ruta.stat()
        os.utime(ruta, ns=(st.st_atime_ns, st.st_mtime_ns + 1_000_000_000))
        return ctrl._exportar_informe(cid, 2025, {'informe': None})['excel']

    primera = provocar_conflicto()
    wb = openpyxl.load_workbook(primera)
    wb['02_INMUEBLES']['AA2'] = 'EDICIÓN MANUAL DE LA COPIA: CONSERVAR'
    wb.save(primera)
    wb.close()
    contenido = primera.read_bytes()
    segunda = provocar_conflicto()
    assert segunda != primera
    assert primera.read_bytes() == contenido
    assert segunda.is_file()


def test_codigo_duplicado_en_excel_no_sobrescribe_la_primera_fila(expediente):
    ctrl, cid, inm = expediente
    ruta = ctrl.sincronizar(cid, 2025)['excel']
    wb = openpyxl.load_workbook(ruta)
    wb['02_INMUEBLES']['A3'] = inm
    wb['02_INMUEBLES']['B3'] = 'COPIA PEGADA CON ID DUPLICADO'
    wb.save(ruta)
    wb.close()
    with pytest.raises(ValueError, match='duplicado'):
        ctrl.abrir_cliente(cid, 2025)
    assert ctrl.db.get('inmuebles', cid, inm)['direccion'] == 'Dirección manual conservada'


def test_fin_de_contrato_invertido_en_excel_se_rechaza_sin_perder_datos(expediente):
    from datetime import date
    ctrl, cid, _ = expediente
    ruta = ctrl.sincronizar(cid, 2025)['excel']
    contrato = ctrl.db.listar('contratos', cid)[0]
    wb = openpyxl.load_workbook(ruta)
    wb['04_CONTRATOS']['E2'] = date(2024, 12, 31)
    wb['02_INMUEBLES']['B2'] = 'NO CONFIRMAR SI FALLA EL CONTRATO'
    wb.save(ruta)
    wb.close()
    with pytest.raises(ValueError, match='anterior al inicio'):
        ctrl.abrir_cliente(cid, 2025)
    assert ctrl.db.get('contratos', cid, contrato['codigo']) == contrato
    assert ctrl.db.listar('inmuebles', cid)[0]['direccion'] == 'Dirección manual conservada'


@pytest.mark.parametrize('valor', ['2025,9', 2025.5, 'NaN', 'Inf', float('-inf')])
def test_un_ejercicio_invalido_no_se_trunca_a_otro_anio(valor):
    columna = next(c for c in database.TABLAS['m210_alquiler'].columnas_bd if c.campo == 'ejercicio')
    with pytest.raises(ValueError):
        database._normalizar(columna, valor)


@pytest.mark.parametrize('valor', [2025, 2025.0, '2025', '2025,0'])
def test_un_entero_valido_se_sigue_aceptando(valor):
    columna = next(c for c in database.TABLAS['m210_alquiler'].columnas_bd if c.campo == 'ejercicio')
    assert database._normalizar(columna, valor) == 2025

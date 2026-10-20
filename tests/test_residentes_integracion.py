"""Residentes de punta a punta: Excel propio, facturas en PDF, justificantes y chat por cartera."""
import hashlib
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))
import asistente as asis  # noqa: E402
import database  # noqa: E402
import excel_generator  # noqa: E402
import fixtures_pdf as fx  # noqa: E402
from controlador import Controlador  # noqa: E402

SHA_PLANTILLA_NR = "19c0c4ae5298"


@pytest.fixture
def ctrl(tmp_path):
    c = Controlador(tmp_path / "t.db")
    c.cambiar_carpeta_clientes(tmp_path)
    yield c
    c.cerrar()


@pytest.fixture
def luis(ctrl):
    return ctrl.crear_cliente("Luis", "Gil", "12345678Z", tipo_cliente="Autónomo",
                              regimen_iva="General", regimen_irpf="Directa simplificada")["id"]


def _pdf(ruta: Path, texto: str) -> Path:
    return fx._crear(ruta, texto)


def test_excel_residente_con_su_plantilla_y_ediciones_manuales(ctrl, luis):
    ctrl.editar(luis, "facturas_emitidas", None, {"fecha": "2025-02-10", "numero": "1", "receptor": "Cliente SL",
                                                   "nif_receptor": "B11111111", "base": 1000, "tipo_iva": 0.21,
                                                   "cuota_iva": 210}, 2025)
    cliente = ctrl.db.get_cliente(luis)
    ruta = ctrl.db.ruta_excel(cliente)
    assert ruta.name.startswith("RES_") and ruta.exists()
    assert excel_generator.verificar_estructura(ruta, database.PLANTILLA_EXCEL_RES) == []
    wb = openpyxl.load_workbook(ruta)
    assert wb["02_FACTURAS_EMITIDAS"]["I2"].value == 1000 and wb["02_FACTURAS_EMITIDAS"]["C2"].value == "1T"
    assert wb["02_FACTURAS_EMITIDAS"]["N2"].value.startswith("=IF(")      # total con fórmula intacta
    fila_303 = [r for r in wb["04_MODELO_303"].iter_rows(min_row=2, values_only=True) if r[2] == "1T"][0]
    assert fila_303[4] == 210 and wb["00_RESUMEN"]["B6"].value == "Autónomo"
    # Edición manual: cambia la base y añade una factura escrita a mano
    ws = wb["02_FACTURAS_EMITIDAS"]
    ws["I2"], ws["K2"] = 2000, 420
    for col, valor in zip("BDEFGIJK", (2025, "2025-03-01", "2", "Otro SL", "B22222222", 500, 0.21, 105)):
        ws[f"{col}3"] = valor
    wb.save(ruta)
    wb.close()
    ctrl.sincronizar(luis, 2025)
    emitidas = ctrl.db.listar("facturas_emitidas", luis, ejercicio=2025)
    assert sorted(float(f["base"]) for f in emitidas) == [500, 2000]
    assert any(f["origen"] == "excel" for f in emitidas)
    assert ctrl.db.get_por_clave("m303", luis, "303|2025|1T")["iva_devengado"] == 525
    sha = hashlib.sha256(database.PLANTILLA_EXCEL.read_bytes()).hexdigest()
    assert sha.startswith(SHA_PLANTILLA_NR)


def test_factura_emitida_recibida_y_alquiler_de_local(ctrl, luis, tmp_path):
    emitida = _pdf(tmp_path / "f1.pdf", """FACTURA
Proveedor: Luis Gil
NIF: 12345678Z
Cliente: Talleres Norte SL
NIF cliente: B87654321
Nº de factura: 2025-014
Fecha de emisión: 12/03/2025
Concepto: Desarrollo de software a medida
Base imponible: 1.500,00 euros
IVA (21%): 315,00 euros
Retención IRPF (15%): 225,00 euros
Total factura: 1.590,00 euros""")
    recibida = _pdf(tmp_path / "f2.pdf", """FACTURA
Proveedor: Gestoría López SL
CIF proveedor: B11223344
Cliente: Luis Gil
NIF cliente: 12345678Z
Nº de factura: G-77
Fecha de emisión: 05/04/2025
Concepto: Honorarios de asesoría fiscal del primer trimestre
Base imponible: 100,00 euros
IVA (21%): 21,00 euros
Retención (15%): 15,00 euros
Total factura: 106,00 euros""")
    local = _pdf(tmp_path / "f3.pdf", """FACTURA
Proveedor: Inmuebles Centro SL
CIF proveedor: B99887766
Cliente: Luis Gil
NIF cliente: 12345678Z
Nº de factura: A-04
Fecha de emisión: 02/04/2025
Concepto: Alquiler del local de la calle Mayor, abril
Base imponible: 800,00 euros
IVA (21%): 168,00 euros
Retención (19%): 152,00 euros
Total factura: 816,00 euros""")
    for ruta in (emitida, recibida, local):
        ctrl.procesar_pdf(luis, ruta, 2025)
    em = ctrl.db.listar("facturas_emitidas", luis)
    assert len(em) == 1 and em[0]["nif_receptor"] == "B87654321" and em[0]["retencion"] == 225
    assert (em[0]["tipo_iva"], em[0]["tipo_retencion"], em[0]["trimestre"]) == (0.21, 0.15, "1T")
    rc = {f["numero"]: f for f in ctrl.db.listar("facturas_recibidas", luis)}
    assert rc["G-77"]["categoria"] == "Servicios profesionales" and rc["A-04"]["categoria"] == "Alquiler local"
    assert ctrl.db.get_por_clave("m111_115", luis, "RET|2025|2T|111")["retenciones"] == 15
    assert ctrl.db.get_por_clave("m111_115", luis, "RET|2025|2T|115")["retenciones"] == 152
    assert ctrl.db.get_por_clave("m130", luis, "130|2025|1T")["retenciones"] == 225
    soportes = Path(ctrl.db.get_cliente(luis)["carpeta"]) / "Soportes_PDF"
    assert len(list(soportes.glob("*.pdf"))) == 3
    ctrl.procesar_pdf(luis, emitida, 2025)  # el mismo PDF otra vez no duplica
    assert len(ctrl.db.listar("facturas_emitidas", luis)) == 1


def test_justificante_303_marca_el_trimestre_presentado(ctrl, luis, tmp_path):
    ctrl.editar(luis, "facturas_emitidas", None, {"fecha": "2025-01-15", "numero": "9", "nif_receptor": "B1",
                                                   "base": 1000, "tipo_iva": 0.21, "cuota_iva": 210}, 2025)
    ctrl.procesar_pdf(luis, fx.modelo_presentado(tmp_path / "303.pdf", modelo="303", ejercicio=2025, periodo="1T",
                                                 importe=210, fecha_presentacion="2025-04-18", csv="CSV-303-1T"), 2025)
    fila = ctrl.db.get_por_clave("m303", luis, "303|2025|1T")
    assert (fila["estado"], fila["nrc"], fila["a_ingresar"]) == ("Presentado", "CSV-303-1T", 210)
    ctrl.editar(luis, "facturas_emitidas", None, {"fecha": "2025-02-15", "numero": "10", "nif_receptor": "B1",
                                                   "base": 5000, "tipo_iva": 0.21, "cuota_iva": 1050}, 2025)
    assert ctrl.db.get_por_clave("m303", luis, "303|2025|1T")["a_ingresar"] == 210  # congelado


def test_036_en_no_residente_no_cambia_su_cartera(ctrl, tmp_path):
    cid = ctrl.crear_cliente("Ana", "Ruiz", "Y1111111A", pais="Francia", ejercicio=2026)["id"]
    res = ctrl.procesar_pdf(cid, fx.modelo_036(tmp_path / "036.pdf"), 2026)["resultado"]
    cliente = ctrl.db.get_cliente(cid)
    assert cliente["tipo_cliente"] == "No residente" and cliente["cartera"] == "NR"
    assert any("cámbialo a Residentes" in m for m in res.mensajes)


def test_chat_del_residente_lee_solo_su_expediente(ctrl, luis, tmp_path):
    ctrl.editar(luis, "facturas_recibidas", None, {"fecha": "2025-05-20", "numero": "L-1", "proveedor": "Luz SA",
                                                    "nif_proveedor": "A1", "categoria": "Suministros", "base": 50,
                                                    "tipo_iva": 0.21, "cuota_iva": 10.5}, 2025)
    soportes = Path(ctrl.db.get_cliente(luis)["carpeta"]) / "Soportes_PDF"
    _pdf(soportes / "2025-05-20_FACTURA_luz.pdf", "Factura de luz mayo 2025 total 60,50 euros")
    otra = ctrl.crear_cliente("Eva", "Sanz", "87654321X", tipo_cliente="Pyme")
    _pdf(Path(otra["carpeta"]) / "Soportes_PDF" / "secreto.pdf", "DATOS DE EVA")
    texto, n = asis.construir_contexto(ctrl.lector, luis, 2025)
    assert "cartera Residentes" in texto and "Resumen trimestral" in texto and "facturas_recibidas" in texto
    assert "Factura de luz mayo 2025" in texto and "DATOS DE EVA" not in texto and n == 1
    assert "gastos_irpf" in texto and "m303 (04_MODELO_303)" in texto
    mal = Path(ctrl.db.get_cliente(luis)["carpeta"])
    destino = mal.parent.parent / "Clientes_No_Residentes" / mal.name
    mal.rename(destino)
    ctrl.db.actualizar_cliente(luis, carpeta=str(destino))
    with pytest.raises(asis.ErrorAsistente, match="no mezclar"):
        asis.construir_contexto(ctrl.lector, luis, 2025)

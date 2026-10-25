"""Casos independientes para el cierre A7; todos usan expedientes temporales."""
import json
from datetime import date
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

import database
import gastos as registro
import tax_calculator as tax
from controlador import Controlador
from excel_generator import _resumen_gastos
from ia_extractor import convertir_valor


@pytest.fixture
def expediente(tmp_path):
    with database.Database(tmp_path / "prueba.sqlite") as db:
        db.carpeta_clientes = tmp_path / "clientes"
        cid = db.crear_cliente("Prueba", "A7", "PRUEBA-A7")["id"]
        db.guardar_residencia(cid, 2020, pais="Francia")
        inm = db.upsert("inmuebles", cid, {"fecha_adquisicion": "2020-01-01", "pct_titularidad": 1})
        db.upsert("contratos", cid, {"inmueble": inm, "fecha_inicio": "2020-01-01",
                                      "renta_inicial": 100, "periodicidad": "Mensual", "estado": "Vigente"})
        db.set_config("periodicidad_210_alquiler", "Anual")
        yield db, cid, inm


@pytest.mark.parametrize("valor", ["150 %", "150", 150, 1.5, -0.01, "-1 %", float("nan"), float("inf")])
def test_porcentajes_invalidos_en_bd_controlador_ia(expediente, valor):
    db, cid, _ = expediente
    assert Controlador.__new__(Controlador).validar("inmuebles", {"pct_titularidad": valor})
    with pytest.raises(ValueError):
        db.upsert("inmuebles", cid, {"pct_titularidad": valor})
    with pytest.raises((ValueError, TypeError)):
        convertir_valor("pct_titularidad", valor)


@pytest.mark.parametrize("valor,esperado", [("100 %", 1), (100, 1), (50, .5), ("50 %", .5),
                                            ("0,5", .5), ("1,5 %", .015), (0, 0)])
def test_porcentajes_validos(expediente, valor, esperado):
    db, cid, _ = expediente
    cod = db.upsert("inmuebles", cid, {"pct_titularidad": valor})
    assert db.get("inmuebles", cid, cod)["pct_titularidad"] == esperado


def test_motor_rechaza_titularidad_preexistente_invalida():
    with pytest.raises(ValueError):
        tax.pct_titularidad({"pct_titularidad": 1.5, "gananciales_forzado": "Sí"})


@pytest.mark.parametrize("valor,esperado", [(1.005, 1.01), (-1.005, -1.01), (2.675, 2.68)])
def test_redondeo_monetario_explicito(valor, esperado):
    assert registro.redondear_eur(valor) == esperado
    assert tax.r2(valor) == esperado


def test_resumen_app_redondea_cada_linea_y_respeta_manual():
    assert sum(registro.importe_registrado({"total": .04, "pct_inmueble": .1}) for _ in range(2)) == 0
    assert registro.importe_registrado({"total": 999, "importe_deducible": 0}) == 0
    assert registro.importe_registrado({"base": 12.5}) == 12.5
    assert registro.importe_registrado({"total": 0, "base": 12.5}) == 0


def test_ejercicio_explicito_no_aparece_en_dos_anios(expediente):
    from main import App
    db, cid, _ = expediente
    cod = db.upsert("gastos", cid, {"ejercicio": 2025, "fecha": "2026-01-01", "total": 20})
    app = SimpleNamespace(ctrl=SimpleNamespace(lector=db), cid=cid, ejercicio=2025,
                          _importe_registrado=registro.importe_registrado)
    assert [g["codigo"] for g in App._gastos_del_ejercicio(app)] == [cod]
    app.ejercicio = 2026
    assert App._gastos_del_ejercicio(app) == []
    inferido = db.upsert("gastos", cid, {"fecha": "2026-02-03", "total": 10})
    assert db.get("gastos", cid, inferido)["ejercicio"] == 2026
    assert registro.ejercicio_registrado({}) is None


def test_edicion_parcial_no_cambia_ejercicio_explicito(expediente):
    db, cid, _ = expediente
    cod = db.upsert("gastos", cid, {"ejercicio": 2025, "fecha": "2026-01-01", "total": 20})
    db.upsert("gastos", cid, {"codigo": cod, "fecha": "2026-02-01"})
    assert db.get("gastos", cid, cod)["ejercicio"] == 2025


def test_formulario_no_expone_desglose_interno():
    from controlador import campos_editables
    assert "detalle_limite" not in {c.campo for c in campos_editables("m210_alquiler")}


def test_formula_excel_redondea_por_fila_y_usa_mismo_ejercicio():
    wb = Workbook()
    ws = wb.active
    _resumen_gastos(ws, [])
    formula = ws["B31"].value
    assert "*ROUND(" in formula and "$B$2:$B$3000=$E$4" in formula
    assert "('08_GASTOS'!$N$2:$N$3000=\"\")*'08_GASTOS'!$L$2:$L$3000" in formula
    assert "IF(" not in formula  # comprobado en Excel 16: el IF matricial exige otro modo de cálculo
    wb.close()


@pytest.mark.parametrize("valor", [1.01, 2, 50, 150, -0.1])
def test_excel_rechaza_fraccion_superior_a_uno(expediente, tmp_path, valor):
    from excel_generator import import_excel_to_db
    db, cid, _ = expediente
    wb = Workbook()
    wb.remove(wb.active)
    for tabla in database.TABLAS.values():
        if tabla.hoja and tabla.hoja not in wb.sheetnames:
            wb.create_sheet(tabla.hoja)
    tabla = database.TABLAS["inmuebles"]
    col = next(c for c in tabla.columnas_bd if c.campo == "pct_titularidad")
    ws = wb[tabla.hoja]
    ws["A2"] = "INVALIDO"
    ws[f"{col.letra}2"] = valor
    ws[f"{col.letra}2"].number_format = "0%"
    ruta = tmp_path / "porcentaje-invalido.xlsx"
    wb.save(ruta)
    wb.close()
    with pytest.raises(ValueError, match="porcentaje"):
        import_excel_to_db(db, cid, ruta)
    assert db.get("inmuebles", cid, "INVALIDO") is None


def liquidar(ej=2026, importe=1500, previos=None, titulo=1, grupo="Francia", otros=0, periodicidad="Anual"):
    inm = {"codigo": "INM", "pct_titularidad": titulo}
    contratos = [{"codigo": "CTR", "fecha_inicio": f"{ej}-01-01", "renta_inicial": 100,
                  "periodicidad": "Mensual"}]
    gs = [{"ejercicio": ej, "categoria": "Intereses/financiación", "total": importe, "pct_alquiler": 1}]
    if otros:
        gs.append({"ejercicio": ej, "categoria": "Comunidad", "total": otros, "pct_alquiler": 1})
    return tax.liquidar_210_alquiler(inm, contratos, {}, gs, None,
                                      tax.resolver_residencia({"pais": grupo}), ej, periodicidad, previos)


def test_limite_genera_exceso_sin_limitar_comunidad():
    r = liquidar(otros=200)
    assert r.resultados[0].ingresos == 1200
    assert r.resultados[0].gastos == 1400
    assert r.resultados[0].cuota == 0
    assert r.generado == 300 and r.saldos == {2026: 300}


def test_intereses_y_reparaciones_comparten_limite():
    r = tax.liquidar_210_alquiler({"codigo": "INM"}, [{"codigo": "CTR", "fecha_inicio": "2026-01-01",
        "renta_inicial": 100}], {}, [{"ejercicio": 2026, "categoria": categoria, "importe_deducible": 800}
        for categoria in tax.CATEGORIAS_LIMITADAS], None, tax.resolver_residencia({"pais": "Francia"}), 2026, "Anual")
    assert r.generado == 400 and r.resultados[0].gastos == 1200


def test_saldos_anteriores_primero_y_mas_antiguo_antes():
    r = liquidar(importe=400, previos={2022: 500, 2024: 900})
    assert r.aplicaciones == {2022: 500, 2024: 700}
    assert r.generado == 400 and r.saldos == {2024: 200, 2026: 400}


def test_cuarto_anio_valido_quinto_caduca():
    r = liquidar(importe=0, previos={2021: 700, 2022: 500})
    assert r.caducados == {2021: 700}
    assert r.aplicaciones == {2022: 500}
    assert r.resultados[0].cuota == 133


def test_titularidad_se_aplica_una_vez_al_saldo():
    r = liquidar(titulo=.5)
    assert r.resultados[0].ingresos == 600
    assert r.resultados[0].gastos == 600 and r.generado == 150
    siguiente = liquidar(ej=2027, importe=0, titulo=.5, previos=r.saldos)
    assert siguiente.resultados[0].gastos == 150
    assert siguiente.resultados[0].cuota == 85.5


def test_extraue_no_deduce_ni_genera_pero_conserva_saldo_vigente():
    r = liquidar(grupo="Reino Unido", previos={2024: 200})
    assert r.resultados[0].gastos == 0 and r.resultados[0].cuota == 288
    assert r.generado == 0 and r.aplicaciones == {} and r.saldos == {2024: 200}


def test_exceso_de_un_trimestre_no_compensa_otro_del_mismo_anio():
    r = tax.liquidar_210_alquiler({"codigo": "INM"}, [{"codigo": "CTR", "fecha_inicio": "2026-01-01",
        "renta_inicial": 100}], {}, [{"ejercicio": 2026, "fecha": "2026-01-01",
        "categoria": "Intereses/financiación", "importe_deducible": 900}], None,
        tax.resolver_residencia({"pais": "Francia"}), 2026)
    assert [x.gastos for x in r.resultados] == [300, 0, 0, 0]
    assert r.generado == 600


def test_ingreso_cero_con_gasto_explicitamente_alquiler_genera_exceso():
    r = tax.liquidar_210_alquiler({"codigo": "INM"}, [], {}, [
        {"ejercicio": 2026, "categoria": "Reparación/conservación", "total": 80, "pct_alquiler": 1}],
        None, tax.resolver_residencia({"pais": "Francia"}), 2026, "Anual")
    assert r.resultados == [] and r.saldos == {2026: 80}


def agregar_gasto(db, cid, inm, ej=2025, total=1500):
    return db.upsert("gastos", cid, {"inmueble": inm, "ejercicio": ej, "total": total,
                     "categoria": "Intereses/financiación", "pct_alquiler": 1})


def saldo(db, cid, inm, ej):
    return db.get_por_clave("arrastres_gastos", cid, f"ARR|{inm}|{ej}")


def modelo(db, cid, inm, ej):
    return db.get_por_clave("m210_alquiler", cid, f"210A|{ej}|{inm}|Anual")


def test_abrir_futuro_primero_y_recalcular_no_duplica(expediente):
    db, cid, inm = expediente
    agregar_gasto(db, cid, inm)
    tax.recalcular_cliente(db, cid, 2026)
    assert saldo(db, cid, inm, 2025)["generado"] == 300
    assert saldo(db, cid, inm, 2026)["aplicado"] == 300
    assert modelo(db, cid, inm, 2026)["cuota"] == 171
    assert db.listar("m210_alquiler", cid, ejercicio=2025) == []
    tax.recalcular_cliente(db, cid, 2026)
    assert saldo(db, cid, inm, 2026)["aplicado"] == 300
    assert modelo(db, cid, inm, 2026)["cuota"] == 171


def test_editar_o_borrar_gasto_pasado_actualiza_futuro(expediente):
    db, cid, inm = expediente
    cod = agregar_gasto(db, cid, inm)
    tax.recalcular_cliente(db, cid, 2026)
    db.upsert("gastos", cid, {"codigo": cod, "total": 1300})
    informes = tax.recalcular_ejercicios(db, cid, [2025])
    assert 2026 in informes
    assert modelo(db, cid, inm, 2026)["cuota"] == 209
    db.borrar("gastos", cid, cod)
    tax.recalcular_ejercicios(db, cid, [2025])
    assert saldo(db, cid, inm, 2025)["pendiente"] == 0
    assert modelo(db, cid, inm, 2026)["cuota"] == 228


def test_presentado_congela_importes_y_consumo_una_sola_vez(expediente):
    db, cid, inm = expediente
    cod = agregar_gasto(db, cid, inm)
    tax.recalcular_ejercicios(db, cid, [2025, 2026])
    fila = modelo(db, cid, inm, 2025)
    db.upsert("m210_alquiler", cid, {"codigo": fila["codigo"], "estado": "Presentado"})
    congelada = db.get("m210_alquiler", cid, fila["codigo"])
    db.upsert("gastos", cid, {"codigo": cod, "total": 9000})
    tax.recalcular_ejercicios(db, cid, [2025])
    assert db.get("m210_alquiler", cid, fila["codigo"]) == congelada
    assert saldo(db, cid, inm, 2025)["generado"] == 300
    assert modelo(db, cid, inm, 2026)["cuota"] == 171


def test_presentado_antiguo_sin_desglose_requiere_revision(expediente):
    db, cid, inm = expediente
    agregar_gasto(db, cid, inm)
    cod = db.upsert("m210_alquiler", cid, {"ejercicio": 2025, "inmueble": inm,
        "periodo": "Anual", "estado": "Presentado", "ingresos": 1200, "gastos": 1500,
        "base": 0, "tipo": .19, "cuota": 0}, clave=f"210A|2025|{inm}|Anual")
    congelada = db.get("m210_alquiler", cid, cod)
    tax.recalcular_cliente(db, cid, 2026)
    assert db.get("m210_alquiler", cid, cod) == congelada
    assert saldo(db, cid, inm, 2026)["estado"] == "Revisar"
    assert modelo(db, cid, inm, 2026)["cuota"] == 228


def test_cambio_tenencia_no_borra_arrastre_de_presentado(expediente):
    db, cid, inm = expediente
    agregar_gasto(db, cid, inm)
    tax.recalcular_ejercicios(db, cid, [2025, 2026])
    fila = modelo(db, cid, inm, 2025)
    db.upsert("m210_alquiler", cid, {"codigo": fila["codigo"], "estado": "Presentado"})
    db.upsert("inmuebles", cid, {"codigo": inm, "fecha_adquisicion": "2026-01-01"})
    tax.recalcular_ejercicios(db, cid, [2025])
    assert saldo(db, cid, inm, 2025)["generado"] == 300
    assert modelo(db, cid, inm, 2026)["cuota"] == 171


def test_saldos_no_cruzan_inmuebles_ni_clientes(expediente):
    db, cid, inm = expediente
    agregar_gasto(db, cid, inm)
    otro = db.upsert("inmuebles", cid, {})
    db.upsert("contratos", cid, {"inmueble": otro, "fecha_inicio": "2020-01-01", "renta_inicial": 100})
    cid2 = db.crear_cliente("Otra", "Prueba", "OTRA-A7")["id"]
    db.guardar_residencia(cid2, 2020, pais="Francia")
    inm2 = db.upsert("inmuebles", cid2, {"codigo": inm})
    db.upsert("contratos", cid2, {"inmueble": inm2, "fecha_inicio": "2020-01-01", "renta_inicial": 100})
    tax.recalcular_cliente(db, cid, 2026)
    tax.recalcular_cliente(db, cid2, 2026)
    assert modelo(db, cid, otro, 2026)["cuota"] == 228
    assert modelo(db, cid2, inm2, 2026)["cuota"] == 228


def test_gasto_sin_anclaje_no_se_deduce_repetidamente(expediente):
    db, cid, inm = expediente
    cod = db.upsert("gastos", cid, {"inmueble": inm, "categoria": "Intereses/financiación",
                                    "total": 1500, "pct_alquiler": 1})
    tax.recalcular_ejercicios(db, cid, [2025, 2026])
    assert saldo(db, cid, inm, 2025)["generado"] == 0
    assert saldo(db, cid, inm, 2026)["generado"] == 0
    assert modelo(db, cid, inm, 2026)["cuota"] == 228
    db.upsert("gastos", cid, {"codigo": cod, "ejercicio": 2025})
    tax.recalcular_ejercicios(db, cid, [2025])
    assert saldo(db, cid, inm, 2026)["aplicado"] == 300
    assert modelo(db, cid, inm, 2026)["cuota"] == 171


def test_periodo_invertido_no_depende_del_orden_de_apertura(expediente):
    db, cid, inm = expediente
    db.upsert("gastos", cid, {"inmueble": inm, "ejercicio": 2026, "periodo_desde": "2026-12-31",
        "periodo_hasta": "2025-01-01", "categoria": "Intereses/financiación", "total": 3000,
        "pct_alquiler": 1})
    tax.recalcular_cliente(db, cid, 2026)
    directo = saldo(db, cid, inm, 2026)["generado"]
    tax.recalcular_ejercicios(db, cid, [2025, 2026])
    assert directo == saldo(db, cid, inm, 2026)["generado"] == 600


def test_categoria_no_elude_limite_por_mayusculas(expediente):
    db, cid, inm = expediente
    cod = agregar_gasto(db, cid, inm)
    db.upsert("gastos", cid, {"codigo": cod, "categoria": "INTERESES/FINANCIACIÓN"})
    assert db.get("gastos", cid, cod)["categoria"] == "Intereses/financiación"
    tax.recalcular_cliente(db, cid, 2025)
    assert saldo(db, cid, inm, 2025)["generado"] == 300
    puro = tax.liquidar_210_alquiler({"codigo": "INM"}, [{"codigo": "CTR", "fecha_inicio": "2026-01-01",
        "renta_inicial": 100}], {}, [{"ejercicio": 2026, "categoria": "INTERESES/FINANCIACIÓN",
        "total": 1500, "pct_alquiler": 1}], None,
        tax.resolver_residencia({"pais": "Francia"}), 2026, "Anual")
    assert puro.generado == 300


def test_controller_recalcula_anios_anterior_nuevo_y_borrado(expediente):
    import threading
    db, cid, inm = expediente
    cod = agregar_gasto(db, cid, inm)
    ctrl = Controlador.__new__(Controlador)
    ctrl.db = db
    ctrl._lock = threading.RLock()
    ctrl._importar = lambda cid: None
    ctrl._recalcular_y_exportar = lambda cid, anios, activo: {"anios": anios}
    resultado = ctrl.editar(cid, "gastos", cod, {"ejercicio": 2027}, 2026)
    assert resultado["anios"] == {2025, 2026, 2027}
    resultado = ctrl.borrar_registro(cid, "gastos", cod, 2026)
    assert resultado["anios"] == {2026, 2027}


def test_presentado_no_admite_cambio_silencioso_de_periodicidad(expediente):
    db, cid, inm = expediente
    tax.recalcular_cliente(db, cid, 2025)
    fila = modelo(db, cid, inm, 2025)
    db.upsert("m210_alquiler", cid, {"codigo": fila["codigo"], "estado": "Presentado"})
    db.set_config("periodicidad_210_alquiler", "Trimestral")
    with pytest.raises(ValueError, match="periodicidad"):
        tax.recalcular_cliente(db, cid, 2026)

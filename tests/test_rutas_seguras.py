"""Rutas y copias de Excel: solo datos ficticios y raíces temporales."""
import os
from pathlib import Path
import subprocess

import openpyxl
import pytest

from rutas_seguras import sanitize_safe_path
import excel_generator as eg
import pdf_processor as pp


@pytest.mark.parametrize("ruta", ["../Otro/soporte.pdf", "sub/../../Otro.pdf",
    r"sub\..\Otro.pdf", "factura.pdf:secreto", "CON.pdf", "NUL.xlsx", "a. ", "a\x00.pdf"])
def test_rechaza_escape_y_nombres_windows(tmp_path, ruta):
    with pytest.raises(ValueError):
        sanitize_safe_path(tmp_path, ruta)


def test_aisla_cliente_y_conserva_cirilico(tmp_path):
    cliente = tmp_path / "Михаил_Иванов"
    cliente.mkdir()
    assert sanitize_safe_path(cliente, "Soportes_PDF/Москва.pdf") == cliente / "Soportes_PDF/Москва.pdf"
    with pytest.raises(ValueError):
        sanitize_safe_path(cliente, tmp_path / "Otro" / "privado.pdf")


def _junction(alias, destino):
    if os.name != "nt":
        alias.symlink_to(destino, target_is_directory=True)
    else:
        resultado = subprocess.run(["cmd", "/c", "mklink", "/J", str(alias), str(destino)],
                                   capture_output=True)
        assert resultado.returncode == 0


def test_alias_raiz_no_archiva_en_otro_expediente(tmp_path):
    otro, alias = tmp_path / "Otro", tmp_path / "Alias"
    otro.mkdir()
    fuente = tmp_path / "factura.pdf"
    fuente.write_bytes(b"PDF ficticio")
    _junction(alias, otro)
    try:
        with pytest.raises(ValueError):
            pp.archivar_pdf(fuente, alias, pp.TipoDocumento.FACTURA_GASTO, None)
        assert not list(otro.iterdir())
    finally:
        if os.name == "nt":
            os.rmdir(alias)
        else:
            alias.unlink()


def test_backups_junction_no_escribe_ni_reemplaza_excel(tmp_path):
    cliente, otro = tmp_path / "Cliente", tmp_path / "Otro"
    cliente.mkdir()
    otro.mkdir()
    destino = cliente / "expediente.xlsx"
    destino.write_bytes(b"original")
    enlace = cliente / "_backups"
    _junction(enlace, otro)
    libro = openpyxl.Workbook()
    try:
        with pytest.raises(ValueError):
            eg._guardar_atomico(libro, destino, cliente)
        assert destino.read_bytes() == b"original"
        assert not list(otro.iterdir())
    finally:
        libro.close()
        if os.name == "nt":
            os.rmdir(enlace)
        else:
            enlace.unlink()


def test_rotacion_solo_copias_del_libro_no_borra_archivos_usuario(tmp_path):
    cliente = tmp_path / "Cliente"
    cliente.mkdir()
    backups = cliente / "_backups"
    backups.mkdir()
    ajeno = backups / "presupuesto_manual.xlsx"
    ajeno.write_bytes(b"datos manuales")
    otra_copia = backups / "otro_20260101_120000_000001.xlsx"
    otra_copia.write_bytes(b"otro expediente")
    for n in range(12):
        (backups / f"expediente_20260101_120000_{n:06d}.xlsx").write_bytes(b"copia")
    destino = cliente / "expediente.xlsx"
    destino.write_bytes(b"anterior")
    libro = openpyxl.Workbook()
    try:
        eg._guardar_atomico(libro, destino, cliente)
    finally:
        libro.close()
    assert ajeno.read_bytes() == b"datos manuales"
    assert otra_copia.read_bytes() == b"otro expediente"
    assert len(list(backups.glob("expediente_*.xlsx"))) == 10
    assert destino.read_bytes().startswith(b"PK")


@pytest.mark.parametrize("valor", ["https://ejemplo.invalid/robar.pdf", "../Otro/privado.pdf",
    r"\\servidor\datos\privado.pdf", "C:/Windows/privado.pdf", "Soportes_PDF/CON.pdf"])
def test_excel_no_activa_vinculos_fuera_del_expediente(tmp_path, valor):
    libro = openpyxl.Workbook()
    try:
        celda = libro.active["A1"]
        eg._link(celda, valor, tmp_path)
        assert celda.hyperlink is None
        assert celda.data_type == "s"
        assert celda.value == valor
    finally:
        libro.close()


def test_backup_excel_ruta_larga_conserva_libro_y_diez_copias(tmp_path):
    # Un destino que cabe en Windows puede dejar de caber al añadir el sello.
    cliente = tmp_path
    while len(str(cliente)) < 190:
        cliente = cliente / ("directorio_" + "x" * min(24, 190-len(str(cliente))))
    cliente.mkdir(parents=True)
    destino = cliente / "EXPEDIENTE_CON_NOMBRE_LARGO_DE_CLIENTE.xlsx"
    libro = openpyxl.Workbook()
    try:
        libro.active['B1'] = 12345.67
        for n in range(12):
            libro.active['A1'] = n
            eg._guardar_atomico(libro, destino, cliente)
        copias = list((cliente/'_backups').glob('*.xlsx'))
        assert len(copias) == 10
        assert all(p.name.startswith('CAM_') for p in copias)
        assert all(len(str(p)) < 260 for p in copias)
        with destino.open('rb') as archivo:
            guardado = openpyxl.load_workbook(archivo)
            assert guardado.active['B1'].value == 12345.67
            guardado.close()
    finally:
        libro.close()


def test_excel_bloqueado_en_ruta_larga_conserva_copia_pendiente(tmp_path, monkeypatch):
    cliente = tmp_path
    while len(str(cliente)) < 190:
        cliente = cliente / ("directorio_" + "x" * min(24, 190-len(str(cliente))))
    cliente.mkdir(parents=True)
    destino = cliente / "EXPEDIENTE_CON_NOMBRE_LARGO_DE_CLIENTE.xlsx"
    destino.write_bytes(b'original abierto en Excel')
    libro = openpyxl.Workbook()
    libro.active['A1'] = 'Corrección manual'
    def bloqueado(*args):
        raise PermissionError('Excel abierto')
    monkeypatch.setattr(eg.os, 'replace', bloqueado)
    try:
        with pytest.raises(eg.ExcelBloqueadoError) as error:
            eg._guardar_atomico(libro, destino, cliente)
        pendientes = list(cliente.glob('PENDIENTE_*.xlsx'))
        assert len(pendientes) == 1 and len(str(pendientes[0])) < 260
        assert destino.read_bytes() == b'original abierto en Excel'
        assert pendientes[0].read_bytes().startswith(b'PK')
    finally:
        libro.close()

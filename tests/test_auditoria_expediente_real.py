"""Fallos hallados al probar un expediente real de no residente (datos de prueba ficticios con la misma forma):
prórroga legal de la vivienda, declaraciones presentadas que no cuadran, Excel de una imputación presentada,
cobros tras una rescisión y nombres de archivo con letras cirílicas (también en NFD)."""
import sys
import unicodedata
from datetime import date
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))
import database  # noqa: E402
import excel_generator as eg  # noqa: E402
import tax_calculator as tc  # noqa: E402
from controlador import Controlador  # noqa: E402


@pytest.fixture
def ctrl(tmp_path):
    c = Controlador(tmp_path / "t.db")
    c.cambiar_carpeta_clientes(tmp_path)
    yield c
    c.cerrar()


def _expediente(ctrl):
    cid = ctrl.crear_cliente("Prueba", "Extra UE", "X0000001T", pais="Rusia", ejercicio=2025)["id"]
    db = ctrl.db
    for ej in (2024, 2025, 2026):
        db.guardar_residencia(cid, ej, pais="Rusia", grupo=tc.grupo_territorio("Rusia"), certificado="Sí")
    inm = db.upsert("inmuebles", cid, {"direccion": "Calle Ficticia 1", "ref_catastral": "0000000AA0000A0001AA",
                                       "fecha_adquisicion": "2024-06-04", "pct_titularidad": 1,
                                       "precio_escritura": 300000})
    db.upsert("catastro", cid, {"ejercicio": 2024, "inmueble": inm, "valor_catastral": 100000,
                                "valor_construccion": 40000, "anio_revision": 2001})
    temporada = db.upsert("contratos", cid, {"inmueble": inm, "fecha_inicio": "2024-09-16",
                                             "fecha_fin_prevista": "2025-08-16", "fecha_fin_real": "2024-12-23",
                                             "estado": "Rescindido", "tipo": "Temporada", "renta_inicial": 2000,
                                             "periodicidad": "Mensual"})
    vivienda = db.upsert("contratos", cid, {"inmueble": inm, "fecha_inicio": "2025-06-01",
                                            "fecha_fin_prevista": "2026-05-31", "estado": "Vigente",
                                            "tipo": "Vivienda", "renta_inicial": 2400, "periodicidad": "Mensual"})
    return cid, inm, temporada, vivienda


def test_vivienda_se_prorroga_por_ley_y_temporada_no():
    viv = {"fecha_inicio": "2025-06-01", "fecha_fin_prevista": "2026-05-31", "estado": "Vigente", "tipo": "Vivienda"}
    assert tc.fin_con_prorroga(viv) == date(2030, 5, 31)
    assert tc.intervalo_contrato(viv) == (date(2025, 6, 1), date(2030, 5, 31))
    assert tc.fin_con_prorroga({**viv, "fecha_inicio": "2016-01-01", "fecha_fin_prevista": "2017-01-01"}) \
        == date(2018, 12, 31)                                              # LAU 2013-2019: 3 años
    assert tc.fin_con_prorroga({**viv, "tipo": "Temporada"}) is None
    assert tc.fin_con_prorroga({**viv, "fecha_fin_real": "2025-12-31"}) is None
    assert tc.fin_con_prorroga({**viv, "estado": "Rescindido"}) is None


def test_solo_un_contrato_vigente_se_prorroga_por_ley():
    viv = {"fecha_inicio": "2021-03-01", "fecha_fin_prevista": "2024-02-28", "tipo": "Vivienda"}
    for estado in database.ESTADOS_CONTRATO:
        if estado != "Vigente":
            assert tc.fin_con_prorroga({**viv, "estado": estado}) is None, estado
    assert tc.intervalo_contrato({**viv, "estado": "Prorrogado"})[1] == date(9999, 12, 31)   # sigue indefinido
    assert tc.intervalo_contrato({**viv, "estado": "Rescindido anticipadamente"})[1] == date(2024, 2, 28)
    assert tc.intervalo_contrato({**viv, "estado": "Vigente"})[1] == date(2026, 2, 28)


def test_aviso_de_prorroga_solo_mientras_dura(ctrl):
    cid, _inm, _t, _v = _expediente(ctrl)
    assert not [a for a in tc.recalcular_cliente(ctrl.db, cid, 2031).alertas if a.tipo.startswith("prórroga")]


def test_editar_imputacion_ya_presentada_no_se_bloquea(ctrl):
    cid, inm, _t, _v = _expediente(ctrl)
    codigo = ctrl.db.upsert("m210_imputacion", cid, {"ejercicio": 2024, "inmueble": inm, "renta_imputada": 387.04,
                                                     "cuota": 92.89, "estado": "Presentado"},
                            clave=f"210I|2024|{inm}")
    ctrl.editar(cid, "m210_imputacion", codigo, {"estado": "Presentado", "nrc": "NRC0001"}, 2024)
    assert ctrl.db.get("m210_imputacion", cid, codigo)["nrc"] == "NRC0001"


def test_prorroga_sigue_alquilado_y_avisa(ctrl):
    ctrl.db.set_config("periodicidad_210_alquiler", "Trimestral")  # este caso contrasta 2T y 4T
    cid, inm, _t, _v = _expediente(ctrl)
    inf = tc.recalcular_cliente(ctrl.db, cid, 2026)
    filas = {f["periodo"]: f for f in ctrl.db.listar("m210_alquiler", cid, ejercicio=2026)}
    assert filas["2T"]["ingresos"] == 7200 and filas["4T"]["ingresos"] == 7200   # junio y siguientes siguen alquilados
    imp = ctrl.db.get_por_clave("m210_imputacion", cid, f"210I|2026|{inm}")
    assert imp is None or not imp.get("dias")                                  # ningún día vacío en 2026
    assert any(a.tipo.startswith("prórroga") and "31/05/2030" in a.mensaje for a in inf.alertas)


def test_presentado_que_no_cuadra_avisa(ctrl):
    cid, inm, _t, _v = _expediente(ctrl)
    db = ctrl.db
    db.upsert("m210_alquiler", cid, {"ejercicio": 2025, "inmueble": inm, "periodo": "Anual", "base": 14620,
                                     "cuota": 3508.8, "pagado": 3508.8, "estado": "Presentado"},
              clave=f"210A|2025|{inm}|Anual")
    db.upsert("m210_imputacion", cid, {"ejercicio": 2024, "inmueble": inm, "dias": 104, "renta_imputada": 387.04,
                                       "cuota": 92.89, "pagado": 92.89, "estado": "Presentado"},
              clave=f"210I|2024|{inm}")
    alq = tc.recalcular_cliente(db, cid, 2025).alertas
    assert any("14.620,00 €" in a.mensaje and "16.800,00 €" in a.mensaje for a in alq)   # 7 × 2.400 devengados
    imp = tc.recalcular_cliente(db, cid, 2024).alertas
    assert any(a.tipo == "presentado 210 imputación" and "387,04 €" in a.mensaje for a in imp)
    # Lo presentado sigue congelado
    assert db.get_por_clave("m210_imputacion", cid, f"210I|2024|{inm}")["renta_imputada"] == 387.04


def test_base_presentada_que_no_cuadra_con_ingresos_menos_gastos(ctrl):
    """Residente UE: la base puede ser menor que los ingresos por los gastos; lo que no cuadre con ellos, avisa."""
    cid = ctrl.crear_cliente("Prueba", "Union Europea", "Y0000002T", pais="Alemania", ejercicio=2025)["id"]
    db = ctrl.db
    db.guardar_residencia(cid, 2025, pais="Alemania", grupo=tc.grupo_territorio("Alemania"), certificado="Sí")
    inm = db.upsert("inmuebles", cid, {"direccion": "Calle Ficticia 2", "fecha_adquisicion": "2020-01-01",
                                       "pct_titularidad": 1, "precio_escritura": 200000})
    db.upsert("contratos", cid, {"inmueble": inm, "fecha_inicio": "2024-01-01", "estado": "Vigente",
                                 "tipo": "Vivienda", "renta_inicial": 1000, "periodicidad": "Mensual"})
    for per, base in (("1T", 2400), ("2T", 2000)):      # 1T cuadra con 3.000 − 600; 2T no
        db.upsert("m210_alquiler", cid, {"ejercicio": 2025, "inmueble": inm, "periodo": per, "ingresos": 3000,
                                         "gastos": 600, "base": base, "estado": "Presentado"},
                  clave=f"210A|2025|{inm}|{per}")
    avisos = [a for a in tc.recalcular_cliente(db, cid, 2025).alertas if a.tipo.startswith(("base", "presentado"))]
    assert [a.tipo for a in avisos] == ["base 210 2T"] and "2.400,00 €" in avisos[0].mensaje


def test_excel_de_imputacion_presentada_muestra_lo_presentado(ctrl):
    cid, inm, _t, _v = _expediente(ctrl)
    ctrl.db.upsert("m210_imputacion", cid, {"ejercicio": 2024, "inmueble": inm, "dias": 112,
                                            "valor_catastral": 100000, "pct_imputacion": 0.02, "pct_titularidad": 1,
                                            "tipo": 0.24, "renta_imputada": 387.04, "cuota": 92.89,
                                            "estado": "Presentado"}, clave=f"210I|2024|{inm}")
    ruta = ctrl.sincronizar(cid, 2024)["excel"]
    ws = openpyxl.load_workbook(ruta)["10_210_IMPUTACION"]
    fila = next(r for r in range(2, 30) if ws[f"B{r}"].value == 2024)
    assert ws[f"H{fila}"].value == 387.04 and ws[f"J{fila}"].value == 92.89       # no la fórmula (= 613,70 €)
    assert "declaración presentada" in ws[f"H{fila}"].comment.text and eg.verificar_estructura(ruta) == []


def test_cobros_de_un_contrato_rescindido_no_quedan_pendientes(ctrl):
    cid, inm, temporada, _v = _expediente(ctrl)
    db = ctrl.db
    db.vaciar_campo("contratos", cid, temporada, "fecha_fin_real")
    db.upsert("contratos", cid, {"codigo": temporada, "estado": "Vigente"})
    tc.recalcular_cliente(db, cid, 2025)                     # sin rescindir: genera enero-agosto de 2025
    assert [c for c in db.listar("cobros", cid, ejercicio=2025) if c.get("contrato") == temporada]
    db.upsert("contratos", cid, {"codigo": temporada, "fecha_fin_real": "2024-12-23", "estado": "Rescindido"})
    tc.recalcular_cliente(db, cid, 2025)
    assert not [c for c in db.listar("cobros", cid, ejercicio=2025) if c.get("contrato") == temporada]


def test_nombres_conservan_el_cirilico_tambien_en_nfd():
    nfd = unicodedata.normalize("NFD", "Майорка_аренда_2025")
    assert database.componente_ruta(nfd) == "Майорка_аренда_2025"
    assert database.componente_ruta("Sede Electrónica Peñón") == "Sede_Electronica_Penon"


def test_resumen_marca_provisional_la_imputacion_sin_revision_acreditada(ctrl):
    cid, inm, _t, _v = _expediente(ctrl)
    cat = ctrl.db.catastro(cid, inm, 2025)["codigo"]
    ctrl.db.vaciar_campo("catastro", cid, cat, "anio_revision")
    ctrl.db.upsert("catastro", cid, {"codigo": cat, "revision_pendiente": 1})
    inf = tc.recalcular_cliente(ctrl.db, cid, 2025)
    fila = next(e for e in inf.estados if "210 Imputación 2025" in e)
    assert "cuota" in fila and "PROVISIONAL: revisión catastral sin acreditar" in fila
    imp = ctrl.db.get_por_clave("m210_imputacion", cid, f"210I|2025|{inm}")
    assert imp["estado"] == "Falta documentación"


def test_rescision_guarda_un_estado_valido_del_desplegable(ctrl):
    cid, _inm, temporada, _v = _expediente(ctrl)
    ctrl.db.upsert("contratos", cid, {"codigo": temporada, "estado": "Rescindido"})   # valor de versiones antiguas
    assert ctrl.db.get("contratos", cid, temporada)["estado"] == "Rescindido anticipadamente"
    assert "Rescindido anticipadamente" in database.ESTADOS_CONTRATO

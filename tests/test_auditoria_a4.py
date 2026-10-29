"""Regresiones A4; auditoría estática de Codex, ejecución pendiente de Claude."""
from datetime import date
from types import SimpleNamespace

import openpyxl
import pytest

import database
import ia_extractor as ia
import main
import tax_calculator as tc
import ui_estilo as ui


@pytest.fixture
def ctrl(tmp_path):
    c = main.Controlador(tmp_path / "a4.db")
    c.cambiar_carpeta_clientes(tmp_path / "Clientes")
    yield c
    c.cerrar()


@pytest.mark.parametrize("valor,esperado", [("0.125", 0.125), ("-0.125", -0.125),
                                            ("1.234", 1234), ("180.000,50", 180000.5)])
def test_a4_decimales_y_miles(valor, esperado):
    assert database.a_num(valor) == esperado


def test_a4_porcentaje_ia_con_tres_decimales(ctrl):
    cid = ctrl.crear_cliente("Cliente", "A4", "A4-PCT")["id"]
    valor = ia.convertir_valor("pct_titularidad", "0.125")
    codigo = ctrl.db.upsert("inmuebles", cid, {"pct_titularidad": valor})
    assert valor == 0.125
    assert ctrl.db.get("inmuebles", cid, codigo)["pct_titularidad"] == 0.125


@pytest.mark.parametrize("estado", ["Disponible", "No procede", "Falta"])
def test_a4_pendientes_conserva_documento_completado(ctrl, estado):
    cid = ctrl.crear_cliente("Cliente", "A4", "A4-DOC")["id"]
    inm = ctrl.db.upsert("inmuebles", cid, {"fecha_adquisicion": "2020-01-01"})
    tc.revisar_pendientes(ctrl.db, cid, 2026)
    clave = f"FALTA|2026|{inm}|ibi"
    doc = ctrl.db.get_por_clave("documentos", cid, clave)
    ctrl.db.upsert("documentos", cid, {"codigo": doc["codigo"], "estado": estado,
                                       "ruta": "Soportes_PDF/manual.pdf", "observaciones": "Conservar"})
    tc.revisar_pendientes(ctrl.db, cid, 2026)  # antes podía volver a marcarlo como Falta
    assert ctrl.db.get("documentos", cid, doc["codigo"])["estado"] == estado
    ctrl.db.upsert("catastro", cid, {"ejercicio": 2026, "inmueble": inm, "valor_catastral": 100000})
    tc.revisar_pendientes(ctrl.db, cid, 2026)  # antes podía borrarlo
    conservado = ctrl.db.get("documentos", cid, doc["codigo"])
    assert conservado["ruta"] == "Soportes_PDF/manual.pdf"
    assert conservado["observaciones"] == "Conservar"


def test_a4_faltas_anuales_no_colisionan(ctrl):
    cid = ctrl.crear_cliente("Cliente", "A4", "A4-ANIOS")["id"]
    tc.revisar_pendientes(ctrl.db, cid, 2025)
    tc.revisar_pendientes(ctrl.db, cid, 2026)
    anterior = ctrl.db.get_por_clave("documentos", cid, "FALTA|2025|-|residencia")
    actual = ctrl.db.get_por_clave("documentos", cid, "FALTA|2026|-|residencia")
    assert anterior["codigo"] != actual["codigo"]
    ctrl.db.guardar_residencia(cid, 2026, certificado="Sí")
    tc.revisar_pendientes(ctrl.db, cid, 2026)
    assert ctrl.db.get("documentos", cid, anterior["codigo"])
    assert ctrl.db.get("documentos", cid, actual["codigo"]) is None


def test_a4_vaciar_solo_un_campo_y_conservar_formula(ctrl):
    cid = ctrl.crear_cliente("Cliente", "A4", "A4-EDIT")["id"]
    codigo = ctrl.editar(cid, "inmuebles", None, {"fecha_adquisicion": "2020-01-01",
                         "notario": "Anterior", "direccion": "Conservar", "precio_escritura": 100000}, 2026)["codigo"]
    res = ctrl.editar(cid, "inmuebles", codigo, {"notario": ""}, 2026)
    fila = ctrl.db.get("inmuebles", cid, codigo)
    assert fila["notario"] is None and fila["direccion"] == "Conservar"
    wb = openpyxl.load_workbook(res["excel"])
    try:
        assert wb["02_INMUEBLES"]["Q2"].data_type == "f"
        assert wb["02_INMUEBLES"]["B2"].value == "Conservar"
    finally:
        wb.close()


@pytest.mark.parametrize("tabla", ["m210_alquiler", "m210_imputacion", "patrimonio_714", "its_718"])
def test_a4_presentacion_resuelve_el_prefijo_exacto(ctrl, tabla):
    cid = ctrl.crear_cliente("Cliente", "A4", "A4-MODELO")["id"]
    codigo = database.TABLAS[tabla].prefijo + "-999"
    ctrl.db.upsert(tabla, cid, {"codigo": codigo, "ejercicio": 2026, "estado": "Pendiente"})
    obl = ctrl.db.upsert("obligaciones", cid, {"ejercicio": 2026, "id_modelo": codigo, "estado": "Pendiente"})
    ctrl.marcar_presentado(cid, "obligaciones", obl, 2026, date(2026, 10, 17), 20)
    assert ctrl.db.get(tabla, cid, codigo)["estado"] == "Presentado"
    for otra in ("m210_alquiler", "m210_imputacion", "patrimonio_714", "its_718"):
        if otra != tabla:
            assert ctrl.db.get(otra, cid, codigo) is None


def test_a4_errores_ia_no_exponen_texto_del_proveedor():
    import anthropic
    import httpx
    secreto = "sk-ant-ficticia-a4"
    response = httpx.Response(400, request=httpx.Request("POST", "https://api.example.test"))
    error = anthropic.BadRequestError(secreto, response=response, body={"error": secreto})
    assert secreto not in ia._mensaje_error(error)
    assert secreto not in ia._mensaje_error(RuntimeError(secreto))


def test_a4_keyring_no_expone_excepcion(monkeypatch):
    def fallar(*_args):
        raise RuntimeError("sk-ant-ficticia-a4")
    monkeypatch.setattr(ia, "_keyring", lambda: SimpleNamespace(set_password=fallar))
    with pytest.raises(ia.ErrorIA) as captura:
        ia.guardar_clave("sk-ant-ficticia-a4")
    assert "sk-ant-ficticia-a4" not in str(captura.value)
    assert captura.value.__suppress_context__


def test_a4_tarea_fija_cliente_y_no_aplica_informe_a_otro():
    app = object.__new__(main.App)
    cola, llamadas, resultados = [], [], []
    app.cid, app.ejercicio = 1, 2026
    app.cerrando, app._turno = False, __import__("threading").Lock()  # admisión de tareas (A4-P03)
    app.page = SimpleNamespace(run_thread=cola.append)
    app.ctrl = SimpleNamespace(set_gananciales=lambda *args: llamadas.append(args) or {"informe": "cliente 1"})
    app.set_ocupado = lambda *_args: None
    app.avisar = lambda *_args, **_kw: None
    app._pintar_principal = lambda: None
    app._aplicar_resultado = resultados.append
    app._toggle_gananciales("INM-001", True)
    app.cid, app.ejercicio = 2, 2025
    cola.pop()()
    assert llamadas == [(1, "INM-001", True, 2026)]
    assert resultados == []


def test_a4_cierre_antes_de_registrar_dialogo_cancela():
    app = object.__new__(main.App)
    app.p, app.cerrando = ui.CLARO, True
    with pytest.raises(main.CanceladoPorUsuario):
        app.revisar_extraccion(None)
    assert app._evento_dialogo is None


def test_a4_editor_cliente_importa_antes_y_vacia_residencia(ctrl):
    cliente = ctrl.crear_cliente("Cliente", "A4", "A4-CLIENTE", pais="Francia", ejercicio=2026)
    cid = cliente["id"]
    ctrl.db.actualizar_cliente(cid, email="anterior@example.test", telefono="111")
    ctrl.db.guardar_residencia(cid, 2026, certificado="Sí")
    ruta = ctrl.sincronizar(cid, 2026)["excel"]
    wb = openpyxl.load_workbook(ruta)
    wb["00_RESUMEN"]["B9"] = "222"  # cambio externo todavía sin importar
    wb.save(ruta)
    wb.close()
    app = object.__new__(main.App)
    app.p, app.cid, app.ejercicio, app.ctrl = ui.CLARO, cid, 2026, ctrl
    app.page = SimpleNamespace(pop_dialog=lambda: None)
    app.en_segundo_plano = lambda _texto, trabajo, *_args, **_kw: trabajo()
    app._refrescar_clientes = lambda: None
    dialogo = {}
    app._dialogo = lambda _titulo, contenido, acciones, **_kw: dialogo.update(contenido=contenido, acciones=acciones)
    app.dialogo_editar_cliente()
    contenido = dialogo["contenido"]
    campos = {c.content.label: c.content for c in contenido.controls[0].controls}
    campos["Email"].value = "nuevo@example.test"
    for container in contenido.controls[4].controls:  # tras el desplegable de tipo y la cabecera
        container.content.value = ""  # vaciar país y certificado, conservando los demás campos
    dialogo["acciones"][-1].on_click(None)
    cliente = ctrl.db.get_cliente(cid)
    assert cliente["email"] == "nuevo@example.test"
    assert cliente["telefono"] == "222"
    residencia = ctrl.db.get("residencias", cid, "2026")
    assert residencia["pais"] is None and residencia["certificado"] is None
    wb = openpyxl.load_workbook(ruta)
    try:
        assert wb["00_RESUMEN"]["B8"].value == "nuevo@example.test"
        assert wb["00_RESUMEN"]["B9"].value == "222"
    finally:
        wb.close()


def test_a4_pendientes_autonomo_se_actualizan_en_el_mismo_ciclo(ctrl):
    cid = ctrl.crear_cliente("Cliente", "A4", "A4-AUT", tipo_cliente="Autónomo")["id"]
    ctrl.db.upsert("actividades", cid, {"fecha_alta": "2026-01-01", "modelos": "303"})
    informe = ctrl.sincronizar(cid, 2026)["informe"]
    assert any("Modelo 303" in p.detalle for p in informe.pendientes)
    # Desde la separación de carteras el 303 del autónomo se calcula y su obligación es OBL|303|…
    obligacion = ctrl.db.get_por_clave("obligaciones", cid, "OBL|303|2026|1T")
    ctrl.db.upsert("obligaciones", cid, {"codigo": obligacion["codigo"], "estado": "Presentado"})
    informe = ctrl.abrir_cliente(cid, 2026)["informe"]
    assert all(p.codigo != obligacion["codigo"] for p in informe.pendientes)

"""Valores visibles, costes de compra y Excel; expedientes siempre ficticios."""
from types import SimpleNamespace
from unittest.mock import Mock
from pathlib import Path
import hashlib
import inspect
import openpyxl
import pytest
import flet as ft
import database
import main
import ui_estilo as ui
import valores_inmuebles as vi
import tax_calculator as tc
import excel_generator as eg
from controlador import Controlador
from test_documentos_ui import recorrer

DATOS={'precio_escritura':200000,'itp_iva_ajd':20000,'notaria':1200,'registro':650,
       'agencia':3500,'abogado_gestoria':900,'otros_gastos':250,'valor_referencia':230000}


def test_tres_valores_no_se_mezclan_y_no_cambia_patrimonio():
    cat={'valor_catastral':150000,'valor_referencia':240000}
    antes=tc.valor_patrimonial(DATOS,cat)
    r=vi.resumen(DATOS,cat)
    assert (r['catastral'],r['adquisicion'],r['referencia'])==(150000,226500,240000)
    assert r['gastos']==26500 and r['referencia_compra']==230000
    assert tc.valor_patrimonial(DATOS,cat)==antes
    # Sin fecha de compra desde 2022 la referencia no cuenta en Patrimonio (AEAT): manda la adquisición con gastos
    assert antes.valor==226500


@pytest.mark.parametrize('datos,valor,parcial', [({},None,False),({'precio_escritura':0},0,False),
    ({'notaria':100},100,True),({'precio_escritura':200000,'agencia':0},200000,False)])
def test_sin_dato_cero_y_total_parcial(datos,valor,parcial):
    r=vi.resumen(datos)
    assert r['adquisicion']==valor and r['parcial']==parcial


@pytest.fixture
def app():
    a=main.App.__new__(main.App)
    a.p,a.cid,a.ejercicio=ui.CLARO,7,2026
    a.ctrl=SimpleNamespace(lector=SimpleNamespace(get=Mock(return_value={**DATOS,'codigo':'INM-001','direccion':'Calle ficticia'})),editar=Mock())
    a._dialogo,a.avisar,a.ocupada=Mock(),Mock(),Mock(return_value=False)
    a.page=SimpleNamespace(update=Mock(),pop_dialog=Mock())
    a.cerrar_dialogo=Mock()
    def ejecutar(texto, trabajo, *args, **kwargs):
        trabajo()
        if kwargs.get('despues'):kwargs['despues']()
        return True
    a.en_segundo_plano=Mock(side_effect=ejecutar)
    return a


def test_editor_total_en_vivo_no_suma_referencia_y_conserva_otras_partidas(app):
    app.dialogo_adquisicion('INM-001')
    contenido,acciones=app._dialogo.call_args.args[1:3]
    campos=[c for c in recorrer(contenido) if isinstance(c,ft.TextField)]
    assert len(campos)==7
    agencia=next(c for c in campos if c.label=='Agencia / inmobiliaria (€)')
    agencia.value='4.500,00'
    agencia.on_change(SimpleNamespace(control=agencia))
    textos=[c.value for c in recorrer(contenido) if isinstance(c,ft.Text)]
    assert ui.eur(227500) in textos
    acciones[-1].on_click(None)
    acciones[-1].on_click(None)
    app.ctrl.editar.assert_called_once_with(7,'inmuebles','INM-001',{'agencia':'4.500,00'},2026)


@pytest.mark.parametrize('cambio',['cliente','ejercicio','importe'])
def test_editor_no_guarda_destino_cambiado_o_importe_invalido(app,cambio):
    app.dialogo_adquisicion('INM-001')
    contenido,acciones=app._dialogo.call_args.args[1:3]
    campo=next(c for c in recorrer(contenido) if isinstance(c,ft.TextField))
    campo.value='nan' if cambio=='importe' else '123'
    if cambio=='cliente':app.cid=8
    if cambio=='ejercicio':app.ejercicio=2025
    acciones[-1].on_click(None)
    app.ctrl.editar.assert_not_called()
    app.cerrar_dialogo.assert_not_called()


def test_tarjeta_conserva_escritura_comprobado_referencias_y_desglose(app):
    cat={'valor_catastral':150000,'valor_referencia':240000}
    fila={**DATOS,'codigo':'INM-001','_catastro':cat,'_vp':tc.valor_patrimonial({**DATOS,'valor_comprobado':250000},cat),
          '_tit':.5,'_computado':125000}
    tarjeta=app._tarjeta_inmueble(fila)
    textos=' '.join(str(c.value) for c in recorrer(tarjeta) if isinstance(c,ft.Text))
    for esperado in ('Valor catastral','Valor de adquisición','Valor de referencia','Valor comprobado',
                     'Precio de escritura:',ui.eur(226500),'Referencia en adquisición','Referencia Catastro'):
        assert esperado in textos


def test_excel_resumen_formulas_y_reimportacion_preservan_costes(tmp_path):
    ctrl=Controlador(tmp_path/'aislada.sqlite')
    ctrl.cambiar_carpeta_clientes(tmp_path/'Clientes')
    huella=hashlib.sha256(database.PLANTILLA_EXCEL.read_bytes()).hexdigest()
    try:
        cid=ctrl.crear_cliente('Valores','Ficticios','VAL-001')['id']
        cod=ctrl.db.upsert('inmuebles',cid,{**DATOS,'direccion':'Dirección ficticia','fecha_adquisicion':'2025-02-10','pct_titularidad':.5})
        ctrl.db.upsert('catastro',cid,{'inmueble':cod,'ejercicio':2026,'valor_catastral':150000,'valor_referencia':240000})
        res=ctrl.sincronizar(cid,2026)
        wb=openpyxl.load_workbook(res['excel'])
        assert len(wb.sheetnames)==17
        assert wb['02_INMUEBLES']['Q2'].value=='=IF(A2="","",SUM(J2:P2))'
        ws=wb['00_RESUMEN']
        fila=next(i for i in range(1,ws.max_row+1) if str(ws.cell(i,1).value or '').startswith(cod+'\n'))
        assert ws.cell(fila,2).value==150000
        assert ws.cell(fila,3).value=='=IF(COUNT(\'02_INMUEBLES\'!J2:P2)=0,"",\'02_INMUEBLES\'!Q2)'
        assert ws.cell(fila,8).value==240000
        assert wb['02_INMUEBLES']['R2'].value==230000
        assert wb['02_INMUEBLES']['Q1'].comment and wb['02_INMUEBLES']['N1'].comment
        wb['02_INMUEBLES']['N2']=4500
        wb.save(res['excel']);wb.close()
        ctrl.sincronizar(cid,2026)
        inm=ctrl.db.get('inmuebles',cid,cod)
        assert tc.valor_adquisicion_total(inm)==227500
        for c,v in DATOS.items():
            assert inm[c]==(4500 if c=='agencia' else v)
        assert ctrl.db.catastro(cid,cod,2026)['valor_catastral']==150000
        assert hashlib.sha256(database.PLANTILLA_EXCEL.read_bytes()).hexdigest()==huella
    finally:ctrl.cerrar()


@pytest.mark.parametrize('cat,compra,esperado', [(None,None,None),(0,None,0),(0,123,0),(123,456,123),(None,456,456)])
def test_referencia_cero_no_se_oculta(cat,compra,esperado):
    assert vi.resumen({'valor_referencia':compra},{'valor_referencia':cat})['referencia']==esperado


def test_gastos_sin_datos_no_se_muestran_como_cero():
    assert vi.resumen({'precio_escritura':100})['gastos'] is None
    assert vi.resumen({'agencia':0})['gastos']==0


def test_fallo_de_guardado_conserva_dialogo_e_importes(app):
    import threading
    app.cerrando=False
    app._turno=threading.Lock()
    app.set_ocupado,app._pintar_principal,app._aplicar_resultado=Mock(),Mock(),Mock()
    app.page.run_thread=lambda f:f()
    app.en_segundo_plano=lambda *a,**kw:main.App.en_segundo_plano(app,*a,**kw)
    app.ctrl.editar.side_effect=OSError('Fallo ficticio')
    app.dialogo_adquisicion('INM-001')
    contenido,acciones=app._dialogo.call_args.args[1:3]
    campo=next(c for c in recorrer(contenido) if isinstance(c,ft.TextField))
    campo.value='220000'
    acciones[-1].on_click(None)
    app.cerrar_dialogo.assert_not_called()
    assert campo.value=='220000' and not campo.read_only
    assert not acciones[-1].disabled
    assert not app._turno.locked()

"""Gestoría sintética: 100 clientes, 20 residentes fiscales rusos, PDFs y correcciones.

Sin red ni datos reales. Oráculos explícitos e independientes del motor fiscal.
"""
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from types import SimpleNamespace
import hashlib
import json
import time
import tracemalloc

import openpyxl
import pytest

import database
import excel_generator
import fiscal_residentes
import fixtures_pdf as fx
import ia_extractor as ia
import pdf_processor as pp
import tax_calculator as tc
from controlador import Controlador

EJ = 2025


def euros(valor):
    return float(Decimal(str(valor)).quantize(Decimal('.01'), rounding=ROUND_HALF_UP))


@pytest.fixture
def db(tmp_path):
    with database.Database(tmp_path / 'sintetica.sqlite') as d:
        d.carpeta_base = tmp_path / 'Gestoria'
        yield d


def crear(d, i, ruso=False, tipo='No residente'):
    return d.crear_cliente('Александр' if ruso else 'Cliente', 'Смирнов_Москва' if ruso else f'Ficticio_{i:03}',
                          f'MASIVO{i:08}', nacionalidad='Rusia' if ruso else 'Francia', tipo_cliente=tipo)['id']


def poblar_inmuebles(d, cid, pais):
    d.guardar_residencia(cid, EJ, pais=pais, certificado='Sí', gastos_deducibles='Sí')
    for i in range(4):
        cod = f'INM-{i+1:03}'
        datos = dict(codigo=cod, direccion=f'Piso sintético {i}', fecha_adquisicion='2020-01-01',
                     precio_escritura=200000, pct_titularidad=1, estado='Activo')
        if i == 3:
            datos.update(fecha_venta='2025-07-01', precio_venta=250000, estado='Vendido')
        d.upsert('inmuebles', cid, datos)
        d.upsert('catastro', cid, dict(inmueble=cod, ejercicio=EJ, valor_catastral=100000,
                                      valor_suelo=40000, valor_construccion=60000, anio_revision=2000))
        if i < 2:
            ini, fin = ('2025-01-01', '2025-03-31') if i == 0 else ('2025-07-01', '2025-09-30')
            d.upsert('contratos', cid, dict(codigo=f'CTR-{i+1:03}', inmueble=cod, fecha_inicio=ini,
                                          fecha_fin_real=fin, renta_inicial=1000, periodicidad='Mensual', tipo='Temporada'))
            d.upsert('gastos', cid, dict(ejercicio=EJ, inmueble=cod, fecha=ini, categoria='IBI', total=900,
                                       periodo_desde=ini, periodo_hasta=fin, pct_inmueble=1, pct_alquiler=1))


def test_100_clientes_20_rusos_multinmueble_y_paginacion(db, record_property):
    inicio = time.perf_counter()
    tracemalloc.start()
    try:
        with db.transaccion():
            rusos = [crear(db, i, ruso=True) for i in range(20)]
            europeos = [crear(db, i) for i in range(20, 60)]
            residentes = [crear(db, i, tipo='Autónomo' if i < 80 else 'Pyme') for i in range(60, 100)]
            for cid in rusos:
                poblar_inmuebles(db, cid, 'Rusia')
                tc.recalcular_cliente(db, cid, EJ)
            poblar_inmuebles(db, europeos[0], 'Francia')
            tc.recalcular_cliente(db, europeos[0], EJ)
        assert db.contar_clientes() == 100
        assert db.con.execute('SELECT count(*) FROM residencias WHERE pais=?', ('Rusia',)).fetchone()[0] == 20
        vistos = [c['id'] for offset in range(0, 100, 17) for c in db.listar_clientes(limite=17, desplazamiento=offset)]
        assert len(vistos) == len(set(vistos)) == 100
        assert len(db.listar_clientes(busqueda='александр смирнов')) == 20
        for cid in rusos:
            assert len(db.listar('inmuebles', cid)) == 4
            modelos = [f for f in db.listar('m210_alquiler', cid, ejercicio=EJ) if f['ingresos']]
            assert len(modelos) == 2
            for f in modelos:
                assert (f['ingresos'], f['gastos'], f['base'], f['tipo'], f['cuota']) == (3000, 0, 3000, .24, 720)
            imp = {r['inmueble']: r for r in db.listar('m210_imputacion', cid, ejercicio=EJ)}
            assert imp['INM-003']['dias'] == 365
            # El motor termina la tenencia el día anterior a la venta.
            assert imp['INM-004']['dias'] == (date(2025, 7, 1) - date(2025, 1, 1)).days
            assert db.residencia(cid, EJ)['gastos_deducibles'] == 'No'
            assert not any(c['id'] == cid for c in db.listar_clientes(cartera='RES'))
        modelos_ue = [r for r in db.listar('m210_alquiler', europeos[0], ejercicio=EJ) if r['ingresos']]
        assert {f['periodo'] for f in modelos_ue} == {'Anual'}
        for f in modelos_ue:
            dias = 90 if f['inmueble'] == 'INM-001' else 92  # duración real de cada contrato, independiente del calendario
            amort = euros(200000 * .6 * .03 * dias / 365)
            gastos = euros(900 + amort)
            assert f['tipo'] == .19
            assert f['gastos'] == pytest.approx(gastos, abs=.01)
            assert f['cuota'] == euros((3000 - gastos) * .19)
        assert db.contar_clientes(cartera='RES') == len(residentes) == 40
        peak = tracemalloc.get_traced_memory()[1]
        record_property('pico_python_bytes', peak)
        record_property('duracion_segundos', time.perf_counter() - inicio)
        assert peak < 128 * 1024 * 1024
    finally:
        tracemalloc.stop()


@pytest.mark.parametrize('tipo', ['Autónomo', 'Pyme'])
def test_200_facturas_pdf_por_cliente_y_excel(db, tmp_path, tipo, record_property):
    cid = crear(db, 101 if tipo == 'Autónomo' else 102, ruso=True, tipo=tipo)
    cliente = db.get_cliente(cid)
    d = Path(cliente['carpeta'])
    inicio = time.perf_counter()
    with db.transaccion():
        db.upsert('actividades', cid, dict(regimen_iva='General', regimen_irpf='Directa normal', modelos='303,130'))
        for t in range(4):
            for n in range(50):
                mes = 1 + 3*t
                pdf = fx.factura_gasto(tmp_path / f'{tipo}_{t}_{n}.pdf', fecha=f'2025-{mes:02}-10',
                                       base=100, iva=21, total=121, n_factura=f'F-{t}-{n}',
                                       concepto='Gasolina' if n % 2 else 'Suministro electricidad')
                ext = pp.extraer_local(pdf)
                assert ext.tipo == pp.TipoDocumento.FACTURA_GASTO
                pp.aplicar(db, cid, ext, pdf, ejercicio=EJ)
        fiscal_residentes.recalcular_residente(db, cid, EJ)
    facturas = db.listar('facturas_recibidas', cid, ejercicio=EJ)
    assert len(facturas) == 200
    for t in ('1T', '2T', '3T', '4T'):
        fs = [f for f in facturas if f['trimestre'] == t]
        assert len(fs) == 50
        assert sum(f['base'] for f in fs) == 5000
        assert sum(f['cuota_iva'] for f in fs) == 1050
        modelo = db.get_por_clave('m303', cid, f'303|{EJ}|{t}')
        assert modelo['iva_deducible'] == 1050
    ruta = excel_generator.exportar_cliente(db, cid, EJ)
    assert 'Москва' in str(ruta)
    wb = openpyxl.load_workbook(ruta)
    try:
        ws = wb['03_FACTURAS_RECIBIDAS']
        assert ws['J201'].value == 100
        assert ws['P201'].data_type == 'f'
        assert ws['S2'].hyperlink.target.startswith('Soportes_PDF/')
    finally:
        wb.close()
    assert len(list((d / 'Soportes_PDF').glob('*.pdf'))) == 200
    assert all((d / f['ruta']).is_file() for f in facturas)
    record_property('duracion_segundos', time.perf_counter() - inicio)


def test_override_persiste_por_campo_y_no_cruza_clientes(db):
    c1, c2 = crear(db, 120), crear(db, 121)
    for cid in (c1, c2):
        db.upsert('inmuebles', cid, dict(codigo='INM-001', precio_escritura=100000))
    db.proteger_campos_inmueble(c1, 'INM-001', ['precio_escritura'])
    db.upsert('inmuebles', c1, dict(codigo='INM-001', precio_escritura=900000, municipio='Москва'))
    db.upsert('inmuebles', c2, dict(codigo='INM-001', precio_escritura=900000))
    assert db.get('inmuebles', c1, 'INM-001')['precio_escritura'] == 100000
    assert db.get('inmuebles', c1, 'INM-001')['municipio'] == 'Москва'
    assert db.get('inmuebles', c2, 'INM-001')['precio_escritura'] == 900000
    with database.Database(db.db_path) as otro:
        assert otro.get('inmuebles', c1, 'INM-001')['manual_override'] == 1
        assert otro.campos_manuales_inmueble(c1, 'INM-001') == {'precio_escritura'}


@pytest.mark.parametrize('respuesta', [{}, None, 'TIMEOUT', 'AUTH'])
def test_api_fallida_no_impide_alta_manual(db, tmp_path, monkeypatch, respuesta):
    import anthropic
    import httpx

    def llamar(**kw):
        if respuesta == 'TIMEOUT':
            raise TimeoutError('Simulado')
        if respuesta == 'AUTH':
            raise anthropic.AuthenticationError('Simulado', response=httpx.Response(401, request=httpx.Request('POST', 'https://invalid.test')), body={})
        return SimpleNamespace(stop_reason='end_turn', content=[SimpleNamespace(type='text', text=json.dumps(respuesta))])
    cliente = SimpleNamespace(messages=SimpleNamespace(create=llamar))
    monkeypatch.setattr(ia, 'obtener_clave', lambda: 'clave-ficticia-sin-red')
    monkeypatch.setattr(ia, '_cliente', lambda: cliente)
    pdf = fx.escritura_compra(tmp_path / 'respaldo.pdf')
    ext = pp.extraer(pdf, ia.AjustesIA())
    assert ext.motor == 'local' and ext.campos['precio_escritura'] == 180000
    cid = crear(db, 130)
    codigo = db.upsert('inmuebles', cid, dict(precio_escritura=180000))
    db.proteger_campos_inmueble(cid, codigo, ['precio_escritura'])
    assert db.get('inmuebles', cid, codigo)['manual_override'] == 1


def test_ia_manual_pdf_chat_excel_y_patrimonio_reales(tmp_path):
    import chat_correcciones as cc

    ctrl = Controlador(tmp_path / 'chat.sqlite')
    ctrl.cambiar_carpeta_clientes(tmp_path / 'Gestoria')
    try:
        cid = ctrl.db.crear_cliente('Михаил', 'Иванов', 'CHAT0001')['id']
        ctrl.db.guardar_residencia(cid, EJ, pais='Rusia')
        pdf1 = fx.escritura_compra(tmp_path / 'primera.pdf', precio=180000)
        ext = pp.Extraccion(tipo=pp.TipoDocumento.ESCRITURA_COMPRA, confianza=.95, motor='ia',
                            campos=dict(ref_catastral=fx.REF, precio_escritura=180000,
                                        pct_titularidad=1, fecha_adquisicion='2020-01-01'), nombre_archivo=pdf1.name)
        resultado = pp.aplicar(ctrl.db, cid, ext, pdf1, ejercicio=EJ)
        codigo = resultado.codigos['inmueble']
        ctrl.editar(cid, 'inmuebles', codigo, {'precio_escritura':250000}, EJ)
        assert ctrl.db.get('inmuebles', cid, codigo)['manual_override'] == 1
        pdf2 = fx.escritura_compra(tmp_path / 'segunda.pdf', precio=900000)
        ext2 = pp.extraer_local(pdf2)
        pp.aplicar(ctrl.db, cid, ext2, pdf2, inmueble_destino=codigo, ejercicio=EJ)
        assert ctrl.db.get('inmuebles', cid, codigo)['precio_escritura'] == 250000
        assert ctrl.db.get('inmuebles', cid, codigo)['municipio'] == 'Alicante'
        activo = [(cid, EJ)]
        g = cc.GestorCorrecciones(lambda: activo[0])
        p = g.proponer(cid, EJ, codigo, 'valor_adquisicion', '350000,00')
        r = g.confirmar(p.token, cid, EJ, lambda x: ctrl.aplicar_correccion_chat(x, cid, EJ))
        assert 'guardada' in r.mensaje.lower()
        assert ctrl.db.get('inmuebles', cid, codigo)['precio_escritura'] == 350000
        assert ctrl.db.get_por_clave('patrimonio_714', cid, f'714|{EJ}')['valor_bienes'] == 350000
        with pytest.raises(cc.ErrorCorreccion):
            g.confirmar(p.token, cid, EJ, lambda x: pytest.fail('No repetir guardado'))
        ruta = ctrl.db.ruta_excel(ctrl.db.get_cliente(cid))
        wb = openpyxl.load_workbook(ruta)
        try:
            assert wb['02_INMUEBLES']['J2'].value == 350000
            wb['02_INMUEBLES']['J2'] = 375000
            wb.save(ruta)
        finally:
            wb.close()
        ctrl.sincronizar(cid, EJ)
        ctrl.db.upsert('inmuebles', cid, dict(codigo=codigo, precio_escritura=999999))
        assert ctrl.db.get('inmuebles', cid, codigo)['precio_escritura'] == 375000
        assert ctrl.db.get_por_clave('patrimonio_714', cid, f'714|{EJ}')['valor_bienes'] == 375000
        ctrl.db.upsert('patrimonio_714', cid, {'codigo':ctrl.db.get_por_clave('patrimonio_714', cid, f'714|{EJ}')['codigo'],
                                            'estado':'Presentado'})
        p2 = g.proponer(cid, EJ, codigo, 'valor_adquisicion', '400000')
        with pytest.raises(cc.ErrorCorreccion, match='presentadas'):
            g.confirmar(p2.token, cid, EJ, lambda x: ctrl.aplicar_correccion_chat(x, cid, EJ))
        assert ctrl.db.get('inmuebles', cid, codigo)['precio_escritura'] == 375000
    finally:
        ctrl.cerrar()


def test_ruta_pdf_larga_preserva_contenido_y_duplicado(tmp_path):
    padre = tmp_path
    # Ruta total cercana al límite clásico; los nombres de componente siguen siendo válidos.
    while len(str(padre / 'Soportes_PDF')) < 195:
        padre = padre / 'abcdefghijklmnop'
    padre.mkdir(parents=True)
    pdf = fx.factura_gasto(tmp_path / ('nombre_factura_muy_largo_' * 3 + '.pdf'))
    rel, sha = pp.archivar_pdf(pdf, padre, pp.TipoDocumento.FACTURA_GASTO, date(2025, 1, 1))
    salida = padre / rel
    assert len(str(salida)) < 248
    assert hashlib.sha256(salida.read_bytes()).hexdigest() == sha
    assert pp.archivar_pdf(pdf, padre, pp.TipoDocumento.FACTURA_GASTO, date(2025, 1, 1)) == (rel, sha)


def test_nombres_legados_hex_no_impiden_deduplicar(tmp_path):
    carpeta = tmp_path / 'cliente'
    soportes = carpeta / 'Soportes_PDF'
    soportes.mkdir(parents=True)
    previo = fx.factura_gasto(soportes / 'factura_0000000000000000.pdf')
    rel, _ = pp.archivar_pdf(previo, carpeta, pp.TipoDocumento.FACTURA_GASTO, date(2025, 1, 1))
    assert rel == 'Soportes_PDF/factura_0000000000000000.pdf'
    assert len(list(soportes.glob('*.pdf'))) == 1


def test_comandos_ui_no_envian_token_a_claude(monkeypatch):
    import re
    import threading
    import asistente
    import chat_correcciones as cc
    import main

    a = main.App.__new__(main.App)
    a.cid, a.ejercicio, a.cerrando, a.informe = 1, EJ, False, None
    a._turno, a.conversaciones = threading.Lock(), {}
    a._chat_ocupado, a._chat_en_curso = False, None
    a._pintar_principal = a._bajar_al_final = lambda: None
    a._actualizar_chat = lambda *args, **kw: None
    a.avisar = lambda *args: pytest.fail('No debe fallar el comando')
    a.page = SimpleNamespace(run_thread=lambda fn: fn())
    llamados = []
    def guardar(p, cid, ej, al_guardar):
        llamados.append((cid, ej, p.codigo, p.valor))
        al_guardar({'informe': 'recalculado'})
        return cc.ReciboGuardado(p.id, cid, ej, p.codigo, p.campo, p.valor, 1)
    a.ctrl = SimpleNamespace(aplicar_correccion_chat=guardar, _lock=threading.RLock())
    monkeypatch.setattr(asistente.Conversacion, 'preguntar', lambda *a, **k: pytest.fail('No enviar comandos a API'))
    a._enviar_pregunta('/corregir INM-001 valor_adquisicion 350000,00')
    conv = a.conversaciones[1]
    assert not conv.mensajes and not llamados
    token = re.search(r'/confirmar ([a-f0-9]{16})', conv.turnos[-1].texto)[1]
    a._enviar_pregunta('/confirmar ' + token)
    assert len(llamados) == 1
    assert not conv.mensajes
    assert 'guardada' in conv.turnos[-1].texto.lower()
    assert not a._turno.locked() and not a._chat_ocupado


def test_acciones_manuales_pueden_cambiar_campos_protegidos(tmp_path):
    ctrl = Controlador(tmp_path / 'manuales.sqlite')
    ctrl.cambiar_carpeta_clientes(tmp_path / 'Gestoria')
    try:
        cid = ctrl.db.crear_cliente('Ana', 'Ficticia', 'MANUAL01')['id']
        codigo = ctrl.db.upsert('inmuebles', cid, {'pct_titularidad':.75, 'estado':'Activo', 'precio_escritura':100000})
        ctrl.db.proteger_campos_inmueble(cid, codigo, ['pct_titularidad','estado'])
        ctrl.set_gananciales(cid, codigo, True, EJ)
        assert ctrl.db.get('inmuebles', cid, codigo)['pct_titularidad'] == .5
        ctrl.set_gananciales(cid, codigo, False, EJ)
        assert ctrl.db.get('inmuebles', cid, codigo)['pct_titularidad'] == .75
        ctrl.marcar_vendido(cid, codigo, '2025-07-01', 150000, EJ)
        ctrl.db.upsert('inmuebles', cid, {'codigo':codigo, 'estado':'Activo', 'precio_venta':999999})
        fila = ctrl.db.get('inmuebles', cid, codigo)
        assert fila['estado'] == 'Vendido' and fila['precio_venta'] == 150000
    finally:
        ctrl.cerrar()

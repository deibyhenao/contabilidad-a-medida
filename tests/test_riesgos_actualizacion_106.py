"""Reproducciones independientes de los avisos, exclusivamente con datos ficticios."""
from pathlib import Path
import sys
from types import SimpleNamespace as NS

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import openpyxl
import pytest
import database
import excel_generator as eg
import ia_extractor
from controlador import Controlador


@pytest.fixture
def expediente(tmp_path, monkeypatch):
    monkeypatch.setattr(database, 'PROYECTO_DIR', tmp_path)
    monkeypatch.setattr(database, 'DB_PATH_DEFECTO', tmp_path/'predeterminada-ficticia.db')
    monkeypatch.setattr(ia_extractor, 'obtener_clave', lambda: None)
    monkeypatch.setattr(ia_extractor, '_cliente', lambda: (_ for _ in ()).throw(AssertionError('API prohibida')))
    ctrl = Controlador(tmp_path/'aislada.db')
    ctrl.cambiar_carpeta_clientes(tmp_path/'Datos')
    c = ctrl.crear_cliente('FICTICIO', 'Actualización', 'QA106-RIESGO', pais='Rusia', ejercicio=2025)
    cid = c['id']
    inm = ctrl.db.upsert('inmuebles', cid, {'direccion': 'Prueba de conservación',
        'fecha_adquisicion': '2020-01-01', 'pct_titularidad': 1, 'observaciones': 'ANTERIOR'})
    ctrl.db.upsert('contratos', cid, {'inmueble': inm, 'fecha_inicio': '2025-01-01',
        'fecha_fin_real': '2025-09-30', 'renta_inicial': 1000, 'estado': 'Finalizado', 'tipo': 'Temporada'})
    ctrl.sincronizar(cid, 2025)
    try:
        yield ctrl, cid, inm
    finally:
        ctrl.cerrar()


def test_renombrar_ficha_no_genera_conflicto_ni_excel_huerfano(expediente):
    ctrl, cid, _ = expediente
    original = ctrl.db.ruta_excel(ctrl.db.get_cliente(cid))
    ctrl._importar(cid)
    ctrl.db.actualizar_cliente(cid, apellidos='Nombre corregido')
    salida = ctrl._recalcular_y_exportar(cid, {2025}, 2025)
    assert not salida.get('error_excel')
    canonico = ctrl.db.ruta_excel(ctrl.db.get_cliente(cid))
    assert canonico.exists()
    assert canonico == original or not original.exists()
    assert not list(canonico.parent.glob('*PENDIENTE*.xlsx'))


def test_carpeta_movida_no_se_recrea_silenciosamente(expediente):
    ctrl, cid, _ = expediente
    cliente = ctrl.db.get_cliente(cid)
    original = Path(cliente['carpeta'])
    trasladada = original.with_name(original.name+'_movida')
    original.rename(trasladada)
    with pytest.raises((FileNotFoundError, RuntimeError), match='carpeta|expediente'):
        ctrl.abrir_cliente(cid, 2025)
    assert not original.exists()
    assert list(trasladada.glob('*.xlsx'))
    assert ctrl.db.get_cliente(cid)['carpeta'] == cliente['carpeta']


def test_guardado_externo_durante_serializacion_conserva_excel_del_usuario(expediente, monkeypatch):
    ctrl, cid, inm = expediente
    ruta = ctrl.db.ruta_excel(ctrl.db.get_cliente(cid))
    ctrl._importar(cid)
    ctrl.db.upsert('inmuebles', cid, {'codigo': inm, 'observaciones': 'CAMBIO DE APP'})
    guardar = openpyxl.workbook.workbook.Workbook.save
    ejecutado = False

    def guardar_y_simular_usuario(wb, destino):
        nonlocal ejecutado
        guardar(wb, destino)
        if not ejecutado and Path(destino).parent == ruta.parent:
            ejecutado = True
            externo = openpyxl.load_workbook(ruta)
            externo['02_INMUEBLES']['AA2'] = 'CAMBIO EXTERNO A CONSERVAR'
            guardar(externo, ruta)
            externo.close()

    monkeypatch.setattr(openpyxl.workbook.workbook.Workbook, 'save', guardar_y_simular_usuario)
    salida = ctrl.commit_and_sync_excel(cid, 2025)
    wb = openpyxl.load_workbook(ruta)
    try:
        assert wb['02_INMUEBLES']['AA2'].value == 'CAMBIO EXTERNO A CONSERVAR'
    finally:
        wb.close()
    assert salida.get('error_excel')
    assert salida['excel'] != ruta and salida['excel'].exists()
    assert ctrl.db.get('inmuebles', cid, inm)['observaciones'] == 'CAMBIO DE APP'


def test_borrar_gasto_en_excel_tras_bloqueo_no_lo_resucita(expediente, monkeypatch):
    ctrl, cid, inm = expediente
    ruta = ctrl.db.ruta_excel(ctrl.db.get_cliente(cid))
    wb = openpyxl.load_workbook(ruta)
    for celda, valor in {'B2': 2025, 'C2': inm, 'H2': 'Comunidad', 'N2': 100,
                         'V2': 'GASTO MANUAL A BORRAR'}.items():
        wb['08_GASTOS'][celda] = valor
    wb.save(ruta)
    wb.close()
    ctrl.sincronizar(cid, 2025)
    assert len(ctrl.db.listar('gastos', cid)) == 1
    reemplazar = eg.os.replace

    def bloqueado(origen, destino):
        if Path(destino) == ruta:
            raise PermissionError('Bloqueo simulado')
        return reemplazar(origen, destino)

    with monkeypatch.context() as contexto:
        contexto.setattr(eg.os, 'replace', bloqueado)
        assert ctrl.sincronizar(cid, 2025).get('error_excel')
    wb = openpyxl.load_workbook(ruta)
    wb['08_GASTOS'].delete_rows(2)
    wb.save(ruta)
    wb.close()
    salida = ctrl.sincronizar(cid, 2025)
    assert not salida.get('error_excel')
    assert ctrl.db.listar('gastos', cid) == []


def _controles(control):
    yield control
    for hijo in getattr(control, 'controls', None) or []:
        yield from _controles(hijo)
    contenido = getattr(control, 'content', None)
    if contenido is not None and not isinstance(contenido, str):
        yield from _controles(contenido)


def test_error_recalculando_ficha_revierte_datos_del_cliente(expediente, monkeypatch):
    import main
    import ui_estilo as ui
    main._cargar_modulos()
    ctrl, cid, _ = expediente
    antes = ctrl.db.get_cliente(cid)
    app = main.App.__new__(main.App)
    app.p, app.ctrl, app.cid, app.ejercicio = ui.CLARO, ctrl, cid, 2025
    app.page = NS(update=lambda: None)
    app.cerrar_dialogo = lambda: None
    app._refrescar_clientes = lambda: None
    dialogo, trabajos = {}, []
    app._dialogo = lambda titulo, contenido, acciones, **kw: dialogo.update(contenido=contenido, acciones=acciones)
    app.en_segundo_plano = lambda titulo, fn, *a: trabajos.append(fn)
    app.dialogo_editar_cliente()
    nombre = next(c for c in _controles(dialogo['contenido']) if getattr(c, 'label', None) == 'Nombre')
    nombre.value = 'NO CONFIRMAR SI FALLA'
    dialogo['acciones'][-1].on_click(None)
    monkeypatch.setattr(ctrl, '_recalcular_bd', lambda *a, **k: (_ for _ in ()).throw(RuntimeError('Fallo fiscal simulado')))
    with pytest.raises(RuntimeError, match='simulado'):
        trabajos.pop()()
    assert ctrl.db.get_cliente(cid) == antes


def test_firma_detecta_cambio_con_igual_tamano_y_fecha(tmp_path):
    import os
    ruta = tmp_path / 'firma.xlsx'
    ruta.write_bytes(b'AAAA')
    previo = ruta.stat()
    firma = eg._firma_excel(ruta)
    ruta.write_bytes(b'BBBB')
    os.utime(ruta, ns=(previo.st_atime_ns, previo.st_mtime_ns))
    assert eg._firma_excel(ruta) != firma


def test_metadatos_de_excel_no_permiten_salir_de_la_carpeta(expediente):
    ctrl, cid, _ = expediente
    for nombre in ('../ajeno.xlsx', 'C:/ajeno.xlsx', 'sub/ajeno.xlsx'):
        ctrl.db.set_config(f'excel_nombre|{cid}', nombre)
        with pytest.raises(ValueError):
            ctrl.db.ruta_excel(ctrl.db.get_cliente(cid))


def test_error_fiscal_revierte_tambien_residencia(expediente, monkeypatch):
    ctrl, cid, _ = expediente
    cliente = ctrl.db.get_cliente(cid)
    residencias = ctrl.db.listar('residencias', cid)
    monkeypatch.setattr(ctrl, '_recalcular_bd', lambda *a, **k: (_ for _ in ()).throw(RuntimeError('Fallo simulado')))
    with pytest.raises(RuntimeError, match='simulado'):
        ctrl.editar_datos_cliente(cid, 2025, {'nombre': 'CAMBIO'}, {'pais': 'Francia'}, cliente['tipo_cliente'])
    assert ctrl.db.get_cliente(cid) == cliente
    assert ctrl.db.listar('residencias', cid) == residencias


def test_copia_previa_de_version_sobrevive_rotacion(tmp_path):
    import hashlib
    ruta = tmp_path / 'principal.xlsx'
    wb = openpyxl.Workbook()
    wb.active['A1'] = 'EDICION ANTERIOR CONSERVADA'
    wb.save(ruta)
    original = ruta.read_bytes()
    for numero in range(12):
        wb.active['A1'] = numero
        eg._guardar_atomico(wb, ruta, tmp_path, guardar_version=numero == 0)
    wb.close()
    copia = tmp_path / '_backups' / f'ANTES_{eg.VERSION}_{hashlib.sha256(original).hexdigest()[:12]}.xlsx'
    assert copia.read_bytes() == original
    assert len(list((tmp_path / '_backups').glob('principal_*.xlsx'))) == 10


def test_borrado_en_excel_no_elimina_cambio_concurrente_de_app(expediente):
    ctrl, cid, inm = expediente
    codigo = ctrl.db.upsert('gastos', cid, {'origen': 'excel', 'ejercicio': 2025,
        'inmueble': inm, 'categoria': 'Comunidad', 'total': 100})
    ruta = ctrl.sincronizar(cid, 2025)['excel']
    ctrl.db.upsert('gastos', cid, {'codigo': codigo, 'total': 200})
    ctrl.db.set_config(f'excel_pendiente|{cid}', '1')
    wb = openpyxl.load_workbook(ruta)
    wb['08_GASTOS'].delete_rows(2)
    wb.save(ruta)
    wb.close()
    original = ruta.read_bytes()
    with pytest.raises(eg.ExcelConflictoImportacion, match='se eliminó'):
        ctrl.sincronizar(cid, 2025)
    assert ctrl.db.get('gastos', cid, codigo)['total'] == 200
    assert ruta.read_bytes() == original

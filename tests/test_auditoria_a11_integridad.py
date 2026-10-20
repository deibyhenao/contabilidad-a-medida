"""A11: importación de declaraciones congeladas y selección tras eliminar."""
import threading
from types import SimpleNamespace as NS

import openpyxl
import pytest

import database
import excel_generator as eg
import main
import ui_estilo as ui
from carteras import RES


@pytest.fixture
def db(tmp_path):
    with database.Database(tmp_path / "a11.sqlite") as conexion:
        conexion.carpeta_base = tmp_path
        yield conexion


def _libro(db, cid, ruta):
    """Libro mínimo con todas las hojas y columnas reales del importador."""
    cartera = db.get_cliente(cid)["cartera"]
    wb = openpyxl.Workbook()
    wb.active.title = "00_RESUMEN"
    for tabla, hoja in database.hojas_cartera(cartera):
        ws = wb.create_sheet(hoja)
        for col in eg._columnas(tabla):
            ws[f"{col.letra}1"] = col.campo
        for numero, fila in enumerate(db.listar(tabla.nombre, cid), tabla.fila_inicio):
            for col in eg._columnas(tabla):
                ws[f"{col.letra}{numero}"] = eg._valor_excel(col, fila.get(col.campo))
    wb.save(ruta)
    return wb


@pytest.mark.parametrize("tipo,tabla,importe", [
    ("No residente", "m210_alquiler", "cuota"),
    ("Autónomo", "m303", "a_ingresar"),
    ("Autónomo", "m130", "a_ingresar"),
    ("Autónomo", "m111_115", "retenciones"),
])
@pytest.mark.parametrize("edicion", ["importe_estado", "vaciar", "borrar"])
def test_presentado_excel_conserva_fila_completa(db, tmp_path, tipo, tabla, importe, edicion):
    cliente = db.crear_cliente("Cliente", "Ficticio", "A11-PRESENTADO", tipo_cliente=tipo)
    cid = cliente["id"]
    codigo = db.upsert(tabla, cid, {"ejercicio": 2025, "estado": "Presentado", importe: 210,
                                     "nrc": "NRC-FICTICIO", "fecha_presentacion": "2025-04-15"})
    antes = db.get(tabla, cid, codigo)
    ruta = tmp_path / "manual.xlsx"
    wb = _libro(db, cid, ruta)
    definicion = database.TABLAS[tabla]
    ws = wb[definicion.hoja_en(cliente["cartera"])]
    columnas = {col.campo: col.letra for col in eg._columnas(definicion)}
    if edicion == "borrar":
        ws.delete_rows(2)
    elif edicion == "vaciar":
        for campo in (importe, "nrc", "fecha_presentacion"):
            ws[f"{columnas[campo]}2"] = None
    else:
        ws[f"{columnas[importe]}2"] = 999
        ws[f"{columnas['estado']}2"] = "Pendiente"
    wb.save(ruta)
    wb.close()
    eg.import_excel_to_db(db, cid, ruta)
    assert db.get(tabla, cid, codigo) == antes
    assert db.listar("notas", cid) == []


def test_pendiente_puede_marcarse_presentado_desde_excel(db, tmp_path):
    cliente = db.crear_cliente("Cliente", "Ficticio", "A11-PENDIENTE", tipo_cliente="Autónomo")
    cid = cliente["id"]
    codigo = db.upsert("m303", cid, {"ejercicio": 2025, "estado": "Pendiente", "a_ingresar": 210})
    ruta = tmp_path / "manual.xlsx"
    wb = _libro(db, cid, ruta)
    ws = wb["04_MODELO_303"]
    ws["L2"], ws["M2"], ws["N2"] = "Presentado", "2025-04-15", "NRC-FICTICIO"
    wb.save(ruta)
    wb.close()
    eg.import_excel_to_db(db, cid, ruta)
    fila = db.get("m303", cid, codigo)
    assert (fila["estado"], fila["fecha_presentacion"], fila["nrc"], fila["a_ingresar"]) == (
        "Presentado", "2025-04-15", "NRC-FICTICIO", 210)


def test_eliminar_ultimo_residente_no_abre_no_residente(db, tmp_path):
    nr = db.crear_cliente("Ana", "Ficticia", "A11-NR")
    residente = db.crear_cliente("Luis", "Ficticio", "A11-RES", tipo_cliente="Autónomo")
    app = main.App.__new__(main.App)
    app.p, app.cid, app.ejercicio, app.cartera = ui.CLARO, residente["id"], 2025, RES
    app.cerrando, app._turno = False, threading.Lock()
    trabajos, acciones, abiertos, errores = [], [], [], []
    app.page = NS(run_thread=trabajos.append, pop_dialog=lambda: None)

    def eliminar(cid):
        db.con.execute("DELETE FROM clientes WHERE id=?", (cid,))
        db.con.commit()
        return tmp_path / "_Eliminados" / "ficticio"

    app.ctrl = NS(lector=db, eliminar_cliente=eliminar)
    app.avisar = lambda mensaje, error=False: errores.append(mensaje) if error else None
    app.set_ocupado = lambda *args: None
    app._refrescar_clientes = lambda: None
    app._pintar_principal = lambda: None
    app.abrir_cliente = abiertos.append
    app._dialogo = lambda titulo, contenido, botones, **kw: acciones.extend(botones)
    app.confirmar_eliminar_cliente()
    acciones[-1].on_click(None)
    trabajos.pop()()
    assert not errores
    assert app.cid is None and app.cartera == RES and abiertos == []
    assert db.get_cliente(nr["id"]) is not None

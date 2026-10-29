"""Fechas reales, periodos parciales y exportación automática del mismo guardado."""
import calendar
from datetime import date
from pathlib import Path

import openpyxl
import pytest
import database
import excel_generator as eg
import ia_extractor
from controlador import Controlador


@pytest.fixture
def expediente(tmp_path, monkeypatch):
    monkeypatch.setattr(database, 'PROYECTO_DIR', tmp_path)
    monkeypatch.setattr(ia_extractor, 'obtener_clave', lambda: None)
    c = Controlador(tmp_path / 'ficticio.db')
    c.cambiar_carpeta_clientes(tmp_path / 'Datos')
    cliente = c.crear_cliente('FICTICIO', 'PARCIAL', 'QA-PARCIAL', pais='Rusia', ejercicio=2025)
    cid = cliente['id']
    inm = c.db.upsert('inmuebles', cid, {'direccion': 'Inmueble ficticio', 'estado': 'Activo',
        'fecha_adquisicion': '2020-01-01', 'pct_titularidad': 1, 'precio_escritura': 200000})
    c.db.upsert('catastro', cid, {'inmueble': inm, 'ejercicio': 2025, 'valor_catastral': 100000,
        'valor_suelo': 40000, 'valor_construccion': 60000, 'pct_imputacion': .02})
    ctr = c.db.upsert('contratos', cid, {'inmueble': inm, 'fecha_inicio': '2025-01-01',
        'fecha_fin_prevista': '2027-12-31', 'fecha_fin_real': '2025-12-31', 'renta_inicial': 1000,
        'periodicidad': 'Mensual', 'tipo': 'Vivienda', 'estado': 'Vigente'})
    c.sincronizar(cid, 2025)
    c.establecer_periodicidad_210(cid, inm, 2025, 'Anual')
    try:
        yield c, cid, inm, ctr
    finally:
        c.cerrar()


@pytest.mark.parametrize('fin,dias,ingresos,periodos,cuota,imputacion', [
    ('2025-09-30', 273, 9000, 3, 2160, 120.99),
    ('2025-06-30', 181, 6000, 2, 1440, 241.97),
    ('2025-05-15', 135, 4483.87, 2, 1076.13, 302.46),
])
def test_fin_real_prevalece_y_actualiza_excel_en_el_mismo_guardado(
        expediente, fin, dias, ingresos, periodos, cuota, imputacion):
    c, cid, inm, ctr = expediente
    salida = c.editar(cid, 'contratos', ctr, {'fecha_fin_real': fin}, 2025)
    assert not salida.get('error_excel')
    modelos = [f for f in c.db.listar('m210_alquiler', cid, ejercicio=2025) if f['estado'] != 'No procede']
    assert len(modelos) == 1 and modelos[0]['periodo'] == 'Anual'
    vista = c.vista_inmuebles(cid, 2025)[0]
    assert len([p for p in vista['_ocupacion_trimestral'] if p['dias_alquilado']]) == periodos
    assert round(sum(f['cuota'] for f in modelos), 2) == cuota
    assert c.db.listar('m210_imputacion', cid, ejercicio=2025)[0]['cuota'] == imputacion
    wb = openpyxl.load_workbook(salida['excel'])
    try:
        ws = wb['00_RESUMEN']
        fila = next(x.row for row in ws for x in row if str(x.value or '').startswith('OCUPACIÓN Y RENTAS')) + 3
        assert [ws.cell(fila, col).value for col in (3, 4, 5, 6)] == [365, dias, 365 - dias, ingresos]
        letra = next(col.letra for col in database.TABLAS['contratos'].columnas_excel if col.campo == 'fecha_fin_real')
        assert str(wb['04_CONTRATOS'][letra + '2'].value).startswith(fin)
        assert wb['10_210_IMPUTACION']['H2'].value == '=IF(A2="","",ROUND(E2*F2*G2*MIN(D2,365)/365,2))'
        assert wb['10_210_IMPUTACION']['J2'].value == '=IF(A2="","",ROUND(H2*I2,2))'
    finally:
        wb.close()
    assert eg.verificar_estructura(salida['excel']) == []
    assert not list(Path(salida['excel']).parent.glob('*PENDIENTE*.xlsx'))


def test_fecha_invalida_en_excel_no_amplia_el_alquiler(expediente):
    c, cid, _, ctr = expediente
    previo = c.db.get('contratos', cid, ctr)
    ruta = c.db.ruta_excel(c.db.get_cliente(cid))
    letra = next(col.letra for col in database.TABLAS['contratos'].columnas_excel if col.campo == 'fecha_fin_real')
    wb = openpyxl.load_workbook(ruta)
    wb['04_CONTRATOS'][letra + '2'] = '31/02/2025'
    wb.save(ruta)
    wb.close()
    contenido = ruta.read_bytes()
    with pytest.raises(ValueError):
        c.abrir_cliente(cid, 2025)
    assert c.db.get('contratos', cid, ctr) == previo and ruta.read_bytes() == contenido


def test_recortar_alquiler_conserva_cobro_manual_posterior(expediente):
    c, cid, _, ctr = expediente
    julio = next(f for f in c.db.listar('cobros', cid, ejercicio=2025) if f['periodo_desde'] == '2025-07-01')
    c.editar(cid, 'cobros', julio['codigo'], {'importe_cobrado': 200, 'fecha_cobro': '2025-07-15'}, 2025)
    salida = c.editar(cid, 'contratos', ctr, {'fecha_fin_real': '2025-05-15'}, 2025)
    actual = c.db.get('cobros', cid, julio['codigo'])
    assert actual['importe_cobrado'] == 200 and actual['importe_contractual'] == 0
    assert actual['fecha_cobro'] == '2025-07-15'
    activos = [f['periodo'] for f in c.db.listar('m210_alquiler', cid, ejercicio=2025) if f['estado'] != 'No procede']
    assert set(activos) == {'Anual'}
    assert [p['periodo'] for p in c.vista_inmuebles(cid, 2025)[0]['_ocupacion_trimestral']
            if p['dias_alquilado']] == ['1T', '2T']
    assert not salida.get('error_excel')


def test_presentada_no_recibe_formula_nueva(expediente):
    c, cid, inm, _ = expediente
    c.db.upsert('m210_imputacion', cid, {'ejercicio': 2025, 'inmueble': inm, 'estado': 'Presentado',
        'renta_imputada': 999.99, 'cuota': 239.99, 'pagado': 239.99, 'nrc': 'NRC-FICTICIO'})
    ruta = eg.exportar_cliente(c.db, cid, 2025)
    wb = openpyxl.load_workbook(ruta)
    try:
        assert wb['10_210_IMPUTACION']['H2'].value == 999.99
        assert wb['10_210_IMPUTACION']['J2'].value == 239.99
        assert wb['10_210_IMPUTACION']['H2'].comment and wb['10_210_IMPUTACION']['J2'].comment
    finally:
        wb.close()
    assert eg.verificar_estructura(ruta) == []

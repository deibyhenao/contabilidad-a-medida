"""Soportes sintéticos; nunca se leen expedientes ni se llama a una API."""
import sys
from zipfile import ZipFile, ZIP_DEFLATED

import fitz
import pytest
from openpyxl import Workbook
import lectores_documentales as lectores


def test_csv_europeo_bom_cirilico_y_campos_multilinea(tmp_path):
    ruta = tmp_path / 'Михаил.csv'
    ruta.write_text('nombre;importe;nota\nМосква;1.234,56;"línea 1\nlínea 2"\n', encoding='utf-8-sig')
    resultado = lectores.extraer_documento(ruta)
    assert 'Москва\t1.234,56\tlínea 1\nlínea 2' in resultado.texto
    assert not resultado.escaneado


@pytest.mark.parametrize('codificacion', ['utf-16', 'cp1252'])
def test_csv_codificaciones_windows(tmp_path, codificacion):
    ruta = tmp_path / 'recibo.csv'
    ruta.write_bytes('concepto;importe\nNotaría;123,45\n'.encode(codificacion))
    resultado = lectores.extraer_documento(ruta)
    assert 'Notaría\t123,45' in resultado.texto
    if codificacion == 'cp1252':
        assert any('1252' in aviso for aviso in resultado.avisos)


def test_xlsx_todas_hojas_formulas_sin_ejecutar_y_cierre(tmp_path):
    ruta = tmp_path / 'Смета.xlsx'
    libro = Workbook()
    libro.active.title = 'Oculta'
    libro.active.append(['Москва', 123.45, '=1+2'])
    libro.create_sheet('Detalle').append(['Notaría', 90])
    libro.save(ruta)
    libro.close()
    resultado = lectores.extraer_documento(ruta)
    assert '[Hoja: Oculta]' in resultado.texto
    assert '[Hoja: Detalle]' in resultado.texto
    assert 'Москва\t123.45' in resultado.texto
    assert '=1+2' not in resultado.texto
    assert resultado.avisos
    ruta.rename(tmp_path / 'renombrado.xlsx')


def test_xlsx_dimensiones_erroneas_no_ocultan_datos(tmp_path):
    ruta = tmp_path / 'original.xlsx'
    libro = Workbook()
    libro.active['A1'] = 'principio'
    libro.active['B3'] = 'final'
    libro.save(ruta)
    with ZipFile(ruta) as entrada:
        contenidos = {n: entrada.read(n) for n in entrada.namelist()}
    contenidos['xl/worksheets/sheet1.xml'] = contenidos['xl/worksheets/sheet1.xml'].replace(b'A1:B3', b'A1:A1')
    with ZipFile(ruta, 'w', ZIP_DEFLATED) as salida:
        for nombre, contenido in contenidos.items():
            salida.writestr(nombre, contenido)
    assert 'final' in lectores.extraer_documento(ruta).texto


def test_pdf_texto_y_pagina_escaneada(tmp_path):
    ruta = tmp_path / 'escritura.pdf'
    with fitz.open() as pdf:
        pdf.new_page().insert_text((40, 50), 'Compra 2025: 123456 euros')
        pdf.new_page()
        pdf.save(ruta)
    resultado = lectores.extraer_documento(ruta)
    assert 'Compra 2025' in resultado.texto
    assert resultado.escaneado
    assert any('OCR' in aviso for aviso in resultado.avisos)
    ruta.rename(tmp_path / 'cerrado.pdf')


def test_pdf_cifrado_rechazado(tmp_path):
    ruta = tmp_path / 'privado.pdf'
    with fitz.open() as pdf:
        pdf.new_page()
        pdf.save(ruta, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw='secreto', owner_pw='propietario')
    with pytest.raises(lectores.ErrorDocumento, match='contraseña'):
        lectores.extraer_documento(ruta)


@pytest.mark.parametrize('extension', ['pdf', 'xlsx', 'docx'])
def test_extension_falsa(tmp_path, extension):
    ruta = tmp_path / f'mal.{extension}'
    ruta.write_text('no es el formato declarado', encoding='utf-8')
    with pytest.raises(lectores.ErrorDocumento):
        lectores.extraer_documento(ruta)


@pytest.mark.parametrize('contenido,nombre', [
    (b'A' * 100_000, 'word/document.xml'),
    (b'<!DOCTYPE x [<!ENTITY evil "x">]><x/>', 'word/document.xml'),
    (b'x', 'word/vbaProject.bin'),
    (b'x', '../exterior.xml'),
], ids=['bomba', 'entidades', 'macro', 'traversal'])
def test_ooxml_bombas_entidades_macros_y_traversal(tmp_path, contenido, nombre):
    ruta = tmp_path / 'mal.docx'
    with ZipFile(ruta, 'w', ZIP_DEFLATED) as zip_doc:
        if nombre != 'word/document.xml':
            zip_doc.writestr('word/document.xml', '<document/>')
        zip_doc.writestr(nombre, contenido)
    with pytest.raises(lectores.ErrorDocumento):
        lectores.extraer_documento(ruta)


@pytest.mark.parametrize('limite', ['MAX_ARCHIVO', 'MAX_TEXTO', 'MAX_CELDAS', 'MAX_FILAS'])
def test_limites_no_devuelven_resultado_parcial(tmp_path, monkeypatch, limite):
    ruta = tmp_path / 'mucho.csv'
    ruta.write_text('a;b\nc;d\ne;f\n', encoding='utf-8')
    monkeypatch.setattr(lectores, limite, 1)
    with pytest.raises(lectores.ErrorDocumento, match='límite'):
        lectores.extraer_documento(ruta)


def test_xlsx_limite_celdas_cierra_libro(tmp_path, monkeypatch):
    ruta = tmp_path / 'limite.xlsx'
    libro = Workbook()
    libro.active.append([1, 2, 3])
    libro.save(ruta)
    monkeypatch.setattr(lectores, 'MAX_CELDAS', 2)
    with pytest.raises(lectores.ErrorDocumento, match='celdas'):
        lectores.extraer_documento(ruta)
    ruta.unlink()


def test_docx_orden_parrafos_tablas(tmp_path):
    docx = pytest.importorskip('docx')
    ruta = tmp_path / 'Москва.docx'
    documento = docx.Document()
    documento.add_paragraph('Antes Москва')
    tabla = documento.add_table(rows=1, cols=2)
    tabla.cell(0, 0).text = 'Valor'
    tabla.cell(0, 1).text = '123,45'
    documento.add_paragraph('Después')
    documento.save(ruta)
    resultado = lectores.extraer_documento(ruta)
    assert resultado.texto.index('Antes') < resultado.texto.index('Valor\t123,45') < resultado.texto.index('Después')


def test_dependencia_xls_ausente_es_error_explicito(tmp_path, monkeypatch):
    ruta = tmp_path / 'antiguo.xls'
    ruta.write_bytes(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1')
    monkeypatch.setitem(sys.modules, 'xlrd', None)
    with pytest.raises(lectores.ErrorDocumento, match='xlrd'):
        lectores.extraer_documento(ruta)


def test_formato_no_admitido_y_archivo_ausente(tmp_path):
    with pytest.raises(lectores.ErrorDocumento, match='Formato'):
        lectores.extraer_documento(tmp_path / 'macro.xlsm')
    with pytest.raises(lectores.ErrorDocumento, match='no existe'):
        lectores.extraer_documento(tmp_path / 'ausente.csv')



def test_vista_previa_parcial_explicita_y_aislamiento(tmp_path):
    ruta = tmp_path / 'vista.csv'
    ruta.write_text('nombre;importe\nConcepto largo;12345\nSegundo;42\n', encoding='utf-8')
    vista = lectores.extraer_documento(ruta, max_caracteres=20)
    assert len(vista.texto) <= 20
    assert any('parcial' in aviso for aviso in vista.avisos)
    assert 'Segundo' in lectores.extraer_documento(ruta).texto


def test_csv_malformado_no_se_acepta_como_completo(tmp_path):
    ruta = tmp_path / 'incompleto.csv'
    ruta.write_text('a;b\n1;"sin cierre', encoding='utf-8')
    with pytest.raises(lectores.ErrorDocumento):
        lectores.extraer_documento(ruta)


def test_xls_todas_hojas_fechas_y_liberacion(tmp_path, monkeypatch):
    from types import SimpleNamespace
    from datetime import datetime
    ruta = tmp_path / 'legado.xls'
    ruta.write_bytes(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1')
    liberadas = []
    hoja = SimpleNamespace(name='Москва', nrows=1, ncols=2,
                           row=lambda n: [SimpleNamespace(ctype=1, value='Compra'),
                                          SimpleNamespace(ctype=3, value=45000)])
    libro = SimpleNamespace(nsheets=2, datemode=0,
                            sheet_by_index=lambda n: hoja,
                            unload_sheet=lambda n: liberadas.append(n),
                            release_resources=lambda: liberadas.append('cerrado'))
    def abrir(nombre, *, on_demand):
        assert nombre == str(ruta)
        assert on_demand is True
        return libro
    modulo = SimpleNamespace(open_workbook=abrir, XL_CELL_DATE=3,
                             xldate_as_datetime=lambda v, modo: datetime(2025, 1, 2))
    monkeypatch.setitem(sys.modules, 'xlrd', modulo)
    resultado = lectores.extraer_documento(ruta)
    assert resultado.texto.count('[Hoja: Москва]') == 2
    assert '2025-01-02T00:00:00' in resultado.texto
    assert liberadas == [0, 1, 'cerrado']


def test_xls_limite_libera_recursos(tmp_path, monkeypatch):
    from types import SimpleNamespace
    ruta = tmp_path / 'limite.xls'
    ruta.write_bytes(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1')
    liberadas = []
    libro = SimpleNamespace(nsheets=999,
                            release_resources=lambda: liberadas.append(True))
    monkeypatch.setitem(sys.modules, 'xlrd', SimpleNamespace(open_workbook=lambda *a, **k: libro))
    with pytest.raises(lectores.ErrorDocumento, match='hojas'):
        lectores.extraer_documento(ruta)
    assert liberadas == [True]


def test_xlsx_procedencia_filas_con_huecos_y_columnas_vacias(tmp_path):
    ruta = tmp_path / 'filas.xlsx'
    libro = Workbook()
    libro.active.title = 'Rentas'
    libro.active['A2'] = 'enero'
    libro.active['C2'] = 1200
    libro.active['B10'] = 'IBI'
    libro.save(ruta)
    resultado = lectores.extraer_documento(ruta)
    assert '[Hoja: Rentas]' in resultado.texto
    assert '[Fila 2] enero\t\t1200' in resultado.texto
    assert '[Fila 10] \tIBI' in resultado.texto
    assert '[Fila 3]' not in resultado.texto


def test_pdf_tiene_limite_independiente_de_office(tmp_path, monkeypatch):
    ruta = tmp_path / 'limite.pdf'
    with fitz.open() as pdf:
        pdf.new_page()
        pdf.save(ruta)
    monkeypatch.setattr(lectores, 'MAX_ARCHIVO', 1)
    assert lectores.extraer_documento(ruta).escaneado
    monkeypatch.setattr(lectores, 'MAX_ARCHIVO_PDF', 1)
    with pytest.raises(lectores.ErrorDocumento, match='tamaño'):
        lectores.extraer_documento(ruta)


def test_xlsx_configuracion_impresora_pasiva_admitida(tmp_path):
    ruta = tmp_path / 'impresion.xlsx'
    libro = Workbook()
    libro.active['A1'] = 'Renta anual'
    libro.active['B1'] = 1200
    libro.save(ruta)
    with ZipFile(ruta, 'a') as contenedor:
        contenedor.writestr('xl/printerSettings/printerSettings1.bin', b'Configuracion pasiva sintetica')
    lectura = lectores.extraer_documento(ruta)
    assert 'Renta anual\t1200' in lectura.texto


@pytest.mark.parametrize('nombre', [
    'xl/embeddings/oleObject1.bin',
    'xl/printerSettings/vbaProject.bin',
    'xl/printerSettings/printerSettingsA.bin',
    'xl/printerSettings/printerSettings1.bin.exe.bin',
    'word/printerSettings/printerSettings1.bin',
])
def test_xlsx_binarios_fuera_de_allowlist_siguen_prohibidos(tmp_path, nombre):
    ruta = tmp_path / 'binario.xlsx'
    libro = Workbook()
    libro.active['A1'] = 'Datos'
    libro.save(ruta)
    with ZipFile(ruta, 'a') as contenedor:
        contenedor.writestr(nombre, b'Binario sintetico')
    with pytest.raises(lectores.ErrorDocumento, match='macros|binarios'):
        lectores.extraer_documento(ruta)

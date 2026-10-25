"""Pago derivado de un NRC: nunca sustituye pagos expresos ni recalcula lo presentado."""
from datetime import date
import math

import pytest

from controlador import Controlador
import pdf_processor as pp
import tax_calculator as tc
import fixtures_pdf as fx

NRC = "1234567890123ABCDEFGHI"


@pytest.mark.parametrize("cambios", [
    {"pagado": 0}, {"pagado": 37.5}, {"estado": "Pendiente"}, {"nrc": None}, {"nrc": ""},
    {"nrc": "NRC0001"}, {"nrc": NRC + "X"}, {"nrc": "Ａ" * 22}, {"nrc": "ß" * 11},
    {"cuota": 0}, {"cuota": -10}, {"cuota": math.inf}, {"cuota": math.nan}, {"cuota": True},
])
def test_no_se_infiere_pago_sin_evidencia_suficiente_o_si_es_manual(cambios):
    assert tc.pago_inferido_por_nrc({"estado": "Presentado", "nrc": NRC, "cuota": 120, **cambios}) is None


def test_nrc_copiado_con_mayusculas_y_espacios_externos():
    assert tc.pago_inferido_por_nrc({"estado": "Presentado", "nrc": " " + NRC.lower() + " ",
                                    "cuota": 120.125}) == 120.13


@pytest.fixture
def caso(tmp_path):
    ctrl = Controlador(tmp_path / "prueba.db")
    ctrl.cambiar_carpeta_clientes(tmp_path / "Clientes")
    cid = ctrl.crear_cliente("Prueba", "Pago", "12345678Z", pais="Rusia", ejercicio=2025)["id"]
    inmueble = ctrl.db.upsert("inmuebles", cid, {"direccion": "Calle Ficticia 1",
        "fecha_adquisicion": "2024-01-01", "precio_escritura": 100000, "pct_titularidad": 1})
    yield ctrl, cid, inmueble
    ctrl.cerrar()


@pytest.mark.parametrize("renta,tabla,campo,clave", [
    ("35", "m210_alquiler", "210_alquiler", "210A|2025|{inmueble}|Anual"),
    ("02", "m210_imputacion", "210_imputacion", "210I|2025|{inmueble}"),
])
def test_extraccion_completa_pago_y_no_pisa_manual(caso, tmp_path, renta, tabla, campo, clave):
    ctrl, cid, inmueble = caso
    ext = pp.Extraccion(tipo=pp.TipoDocumento.MODELO_210, confianza=.95, nombre_archivo="modelo.pdf",
        campos={"tipo_renta": renta, "periodo": "0A", "ejercicio": 2025, "nrc": NRC, "base": 500, "cuota": 120})
    ruta = fx._crear(tmp_path / "modelo.pdf", "Modelo ficticio")
    salida = pp.aplicar(ctrl.db, cid, ext, ruta, inmueble_destino=inmueble, ejercicio=2025)
    fila = ctrl.db.get(tabla, cid, salida.codigos[campo])
    assert fila["pagado"] == 120 and fila["cuota"] == 120
    assert tc.NOTA_PAGO_NRC in fila["observaciones"]
    ctrl.db.upsert(tabla, cid, {"codigo": fila["codigo"], "pagado": 0})
    ext.campos["pagado"] = 999  # Otra lectura no puede pisar lo escrito por el asesor.
    corregido = fx._crear(tmp_path / "corregido.pdf", "Otra copia ficticia del modelo corregido")
    salida = pp.aplicar(ctrl.db, cid, ext, corregido, inmueble_destino=inmueble, ejercicio=2025)
    assert ctrl.db.get(tabla, cid, salida.codigos[campo])["pagado"] == 0


@pytest.mark.parametrize("tabla,tipo", [("m210_alquiler", "210 Alquiler"), ("m210_imputacion", "210 Imputación")])
@pytest.mark.parametrize("pago_obligacion", [None, 0, 37.5])
def test_obligacion_presentada_completa_vacio_y_respeta_pago_propio(caso, tabla, tipo, pago_obligacion):
    ctrl, cid, inmueble = caso
    modelo = ctrl.db.upsert(tabla, cid, {"ejercicio": 2025, "inmueble": inmueble, "estado": "Presentado",
        "cuota": 120, "nrc": NRC, "observaciones": "Nota del asesor"})
    obligacion = ctrl.db.upsert("obligaciones", cid, {"ejercicio": 2025, "tipo": tipo, "inmueble": inmueble,
        "estado": "Presentado", "id_modelo": modelo, "pagado": pago_obligacion, "notas": "Conservar nota"}, clave="PRUEBA")
    for _ in range(2):
        tc._obligacion(ctrl.db, cid, 2025, tipo, inmueble, "PRUEBA", date(2026, 1, 20), modelo, "Presentado")
    fila = ctrl.db.get(tabla, cid, modelo)
    oblig = ctrl.db.get("obligaciones", cid, obligacion)
    assert fila["pagado"] == 120 and fila["observaciones"].startswith("Nota del asesor")
    assert fila["observaciones"].count(tc.NOTA_PAGO_NRC) == 1
    assert oblig["pagado"] == (120 if pago_obligacion is None else pago_obligacion)
    assert oblig["notas"].startswith("Conservar nota")


def test_sin_nrc_no_se_rellena_pago_ni_nota(caso):
    ctrl, cid, inmueble = caso
    codigo = ctrl.db.upsert("m210_alquiler", cid, {"estado": "Presentado", "cuota": 120, "inmueble": inmueble})
    fila = tc.completar_pago_210_nrc(ctrl.db, cid, "m210_alquiler", codigo)
    assert fila["pagado"] is None and fila["observaciones"] is None


def test_modelos_distintos_del_210_no_admiten_esta_inferencia(caso):
    ctrl, cid, _ = caso
    with pytest.raises(ValueError, match="solo se completa"):
        tc.completar_pago_210_nrc(ctrl.db, cid, "patrimonio_714", "714-001")


def test_venta_con_nrc_conserva_pago_y_notas_al_reimportar(caso, tmp_path):
    ctrl, cid, inmueble = caso
    ext = pp.Extraccion(tipo=pp.TipoDocumento.MODELO_210, confianza=.95, nombre_archivo="venta.pdf",
        campos={"tipo_renta": "28", "ejercicio": 2025, "nrc": NRC, "base": 500, "cuota": 120})
    a = fx._crear(tmp_path / "venta.pdf", "Venta ficticia presentada")
    resultado = pp.aplicar(ctrl.db, cid, ext, a, inmueble_destino=inmueble, ejercicio=2025)
    codigo = resultado.codigos["venta"]
    fila = ctrl.db.get("obligaciones", cid, codigo)
    assert fila["pagado"] == 120 and tc.NOTA_PAGO_NRC in fila["notas"]
    ctrl.db.upsert("obligaciones", cid, {"codigo": codigo, "pagado": 0, "notas": fila["notas"] + "; Nota manual"})
    b = fx._crear(tmp_path / "venta2.pdf", "Otra copia de venta ficticia")
    pp.aplicar(ctrl.db, cid, ext, b, inmueble_destino=inmueble, ejercicio=2025)
    fila = ctrl.db.get("obligaciones", cid, codigo)
    assert fila["pagado"] == 0 and "Nota manual" in fila["notas"]
    assert fila["notas"].count(tc.NOTA_PAGO_NRC) == 1


@pytest.mark.parametrize("tabla,clave,campo,tipo", [
    ("m210_alquiler", "210A|2025|{inmueble}|Anual", "base", "210 Alquiler"),
    ("m210_imputacion", "210I|2025|{inmueble}", "renta_imputada", "210 Imputación"),
])
def test_recalculo_completo_conserva_presentado_y_completa_pago(caso, tabla, clave, campo, tipo):
    ctrl, cid, inmueble = caso
    clave = clave.format(inmueble=inmueble)
    modelo = ctrl.db.upsert(tabla, cid, {"ejercicio": 2025, "inmueble": inmueble, "periodo": "Anual",
        "estado": "Presentado", campo: 500, "cuota": 120, "nrc": NRC, "observaciones": "Nota manual"}, clave=clave)
    for _ in range(2):
        tc.recalcular_cliente(ctrl.db, cid, 2025)
    fila = ctrl.db.get(tabla, cid, modelo)
    oblig = ctrl.db.get_por_clave("obligaciones", cid, "OBL|" + clave)
    assert fila[campo] == 500 and fila["cuota"] == 120 and fila["estado"] == "Presentado"
    assert fila["pagado"] == 120 and oblig["pagado"] == 120
    assert "Nota manual" in fila["observaciones"] and fila["observaciones"].count(tc.NOTA_PAGO_NRC) == 1


def test_pago_parcial_explicito_de_documento_no_se_reinterpreta(caso, tmp_path):
    ctrl, cid, inmueble = caso
    ext = pp.Extraccion(tipo=pp.TipoDocumento.MODELO_210, confianza=.95, nombre_archivo="parcial.pdf",
        campos={"tipo_renta": "35", "ejercicio": 2025, "periodo": "0A", "nrc": NRC,
                "base": 500, "cuota": 120, "pagado": 40})
    ruta = fx._crear(tmp_path / "parcial.pdf", "Ingreso parcial ficticio")
    resultado = pp.aplicar(ctrl.db, cid, ext, ruta, inmueble_destino=inmueble, ejercicio=2025)
    fila = ctrl.db.get("m210_alquiler", cid, resultado.codigos["210_alquiler"])
    assert fila["pagado"] == 40 and tc.NOTA_PAGO_NRC not in (fila["observaciones"] or "")


def test_resumen_aclara_total_pagado_y_cuota_provisional_sin_tocar_formulas(caso):
    from openpyxl import load_workbook
    import database
    import excel_generator as eg
    ctrl, cid, inmueble = caso
    ctrl.db.upsert("m210_imputacion", cid, {"ejercicio": 2025, "inmueble": inmueble,
        "cuota": 45.8, "estado": "Falta documentación"})
    salida = eg.exportar_cliente(ctrl.db, cid, 2025)
    wb = load_workbook(salida)
    plantilla = load_workbook(database.PLANTILLAS['NR'])
    try:
        for celda in ("H7", "E10", "E12", "E14", "H10"):
            assert wb["00_RESUMEN"][celda].value == plantilla["00_RESUMEN"][celda].value
            assert wb["00_RESUMEN"][celda].comment
        assert "PROVISIONAL" in wb["00_RESUMEN"]["E14"].comment.text
        assert "45,80" in wb["00_RESUMEN"]["E14"].comment.text
        assert wb["00_RESUMEN"].column_dimensions['F'].width == plantilla["00_RESUMEN"].column_dimensions['F'].width
        from openpyxl.comments import Comment
        wb["00_RESUMEN"]["H7"].comment = Comment("Nota de plantilla", "Autor ficticio", width=500, height=200)
        eg._aclaraciones_resumen(wb, ctrl.db, cid, 2025)
        eg._aclaraciones_resumen(wb, ctrl.db, cid, 2025)
        comentario = wb["00_RESUMEN"]["H7"].comment
        assert comentario.author == "Autor ficticio" and comentario.width == 500 and comentario.height == 200
        assert comentario.text.startswith("Nota de plantilla") and comentario.text.count("Total de impuestos pagados") == 1
    finally:
        wb.close(); plantilla.close()
